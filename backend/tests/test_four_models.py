"""Owner's four-model routing: independent checks and two mandatory frontier approvals. Since the
owner's decision of 2026-09-30 (one accountable final reviewer), new runs have one approval; this file
pins the earlier policy, which jobs priced before that decision keep (tests/test_single_reviewer.py
covers the new one)."""

import io
import json

import pytest
from docx import Document

from app.ai import costs, orchestration
from app.ai.orchestration import DRAFTING_TASKS, ROUTINE_TASKS, STEPS, WORK_ROLES, AIRunner, Revision, current_engine
from app.ai.styles import writing_brief
from app.core.config import Settings
from app.proposals.ai import ProposalRunner
from tests.fake_models import FakeModels
from tests.test_ai_coverage import _judgment, _passage
from tests.test_api import STUDENT, _submit, enable_score


@pytest.fixture(autouse=True)
def earlier_workflow(monkeypatch):
    """These tests cover the multi-provider algorithm that engines priced before the Gemini workflow
    (owner decision 2026-10-07) keep running on, and that GEMINI_WORKFLOW=false brings back."""
    monkeypatch.setenv("GEMINI_WORKFLOW", "false")

BRIEF = writing_brief("PRESERVE_VOICE", "STANDARD")


@pytest.fixture(autouse=True)
def dual_approval_engine(monkeypatch):
    """The engine of jobs priced before one final reviewer: two approvals and Opus's guidance."""
    monkeypatch.setenv("SINGLE_REVIEWER", "false")
    monkeypatch.setenv("FRONTIER_GUIDANCE", "true")
    from app.core.config import get_settings

    get_settings.cache_clear()
    yield
    get_settings.cache_clear()


@pytest.fixture
def reviewers(monkeypatch):
    models = FakeModels()
    settings = Settings(_env_file=None, openai_api_key="sk-test", anthropic_api_key="sk-test",
                        model_prices={"fake:" + m: (0, 0, 0) for m in ("gpt-6-luna", "gpt-6-sol", "claude-sonnet-5-5", "claude-opus-5-5")})
    refs = []

    def route(ref, settings):
        refs.append(ref)
        return models, ref.partition(":")[2]

    monkeypatch.setattr(orchestration, "provider_for", route)
    return AIRunner(settings, lambda c: None, lambda: 0, 5), models, refs


def test_every_new_task_has_the_requested_model():
    runner = AIRunner(Settings(_env_file=None), lambda c: None, lambda: 0, 5)
    assert runner.model_for("analyse") == "openai:gpt-6-luna"
    assert runner.model_for("analyse_peer") == "anthropic:claude-sonnet-5-5"
    for task in ROUTINE_TASKS:
        assert runner.model_for(task) == "openai:gpt-6-luna"
    for task in DRAFTING_TASKS:
        assert runner.model_for(task) == "anthropic:claude-sonnet-5-5"
    assert runner.model_for("analyse_after") == "openai:gpt-6-luna"
    assert runner.model_for("analyse_after_peer") == "anthropic:claude-sonnet-5-5"
    works = {t for t, s in STEPS.items() if s.role in WORK_ROLES}  # the works algorithm has its own roles (test_works_flows)
    for task in set(STEPS) - ROUTINE_TASKS - DRAFTING_TASKS - works - {"analyse", "analyse_peer", "analyse_after", "analyse_after_peer"}:
        assert runner.model_for(task) == ("openai:gpt-6-sol" if STEPS[task].role == "lead" else "anthropic:claude-opus-5-5")


def test_a_frozen_four_model_engine_survives_configuration_changes():
    settings = Settings(_env_file=None)
    frozen = current_engine(settings)
    changed = settings.model_copy(update={"routine_model": "other", "drafting_model": "other", "ai_check_model": "other", "ai_check_peer_model": "other"})
    runner = AIRunner(changed, lambda c: None, lambda: 0, 5, engine=frozen)
    assert runner.model_for("analyse") == "openai:gpt-6-luna"
    assert runner.model_for("analyse_peer") == runner.model_for("refine") == "anthropic:claude-sonnet-5-5"
    assert runner.model_for("review") == "openai:gpt-6-sol"
    assert runner.model_for("review_peer") == "anthropic:claude-opus-5-5"


def test_checkers_are_independent_and_disagreement_is_preserved(reviewers):
    runner, models, refs = reviewers
    models.overrides["analyse"] = lambda p: {"blocks": [{**_judgment(b["id"]), "explanation": "PRIVATE FIRST REVIEW"} for b in p["blocks"]]}
    models.overrides["analyse_peer"] = lambda p: {"blocks": [{**_judgment(b["id"]), "riskBand": "high"} for b in p["blocks"]]}
    answers, seen = runner.analyse([_passage("b1")], [], {})
    assert seen == {"b1"}
    assert answers["b1"].score == pytest.approx((0.1 + 0.85) / 2)
    assert answers["b1"].bands == ["low", "high"]
    payloads = [json.loads(r[len(t):]) for t, r in zip(models.tasks, models.requests, strict=True)]
    assert payloads[0] == payloads[1] and "PRIVATE FIRST REVIEW" not in models.requests[1]
    assert refs == ["openai:gpt-6-luna", "anthropic:claude-sonnet-5-5"]


def test_a_missing_peer_assessment_prevents_complete_coverage(reviewers):
    runner, models, _ = reviewers
    models.overrides["analyse_peer"] = lambda p: {"blocks": []}
    _, seen = runner.analyse([_passage("b1")], [], {})
    assert seen == set()


@pytest.mark.parametrize("rejecting", ["review", "review_peer"])
def test_either_frontier_can_veto_a_rewrite(reviewers, rejecting):
    runner, models, refs = reviewers
    models.overrides[rejecting] = lambda p: {"results": [{"id": r["id"], "grade": "REPAIR", "issues": [], "note": "Not approved", "riskBand": "low"} for r in p["pairs"]]}
    result = runner.review([Revision("b1", "Original claim.", "Changed claim.", [])], {}, BRIEF)
    assert "b1" in result.issues
    assert refs == ["openai:gpt-6-sol", "anthropic:claude-opus-5-5"]


def test_omitted_opus_review_is_not_permission(reviewers):
    runner, models, _ = reviewers
    models.overrides["review_peer"] = lambda p: {"results": []}
    result = runner.review([Revision("b1", "Original claim.", "Changed claim.", [])], {}, BRIEF)
    assert result.issues == {"b1": ["NOT_REVIEWED"]}


def test_both_approvals_are_required_after_a_repair(reviewers):
    runner, models, refs = reviewers
    revision = Revision("b1", "Original claim.", "Changed claim.", [])
    assert runner.review([revision], {}, BRIEF).issues == {}
    models.overrides["review_peer"] = lambda p: {"results": []}
    result = runner.review([Revision("b1", revision.original, "Repaired claim.", [])], {}, BRIEF)
    assert result.issues == {"b1": ["NOT_REVIEWED"]}
    assert refs == ["openai:gpt-6-sol", "anthropic:claude-opus-5-5"] * 2


def test_a_missing_opus_proposal_review_cannot_pass(reviewers):
    runner, models, _ = reviewers
    models.overrides["p_review_peer"] = lambda p: {"results": []}
    proposal = ProposalRunner(runner.settings, lambda c: None, lambda: 0, 5)
    assert proposal.grade([{"key": "problem", "text": ["A claim."], "_words": "A claim."}], {}) == {}


def test_sonnet_55_has_a_verified_price():
    assert costs.price_for("anthropic", "claude-sonnet-5-5") == (2.0, 10.0, 0.20)


def test_rejected_rewrites_are_not_exported(client):
    client.models.overrides["review_peer"] = lambda p: {"results": []}
    job_id, job = _submit(client, {"writing": "REFINE", "academic": False})
    assert job["status"] == "COMPLETED" and job["outcome"] == "PARTIAL"
    assert job["refinement"]["refinedBlocks"] == 0
    assert all(c["kept"] for c in job["refinement"]["changes"])
    from tests.conftest import fixture_bytes

    output = client.get(f"/api/jobs/{job_id}/outputs/paper", headers=STUDENT).content
    assert [p.text for p in Document(io.BytesIO(output)).paragraphs] == [p.text for p in Document(io.BytesIO(fixture_bytes("simple_essay.docx"))).paragraphs]


def test_disagreement_lowers_confidence_and_is_visible_in_the_result(client):
    enable_score()
    client.models.overrides["analyse_peer"] = lambda p: {"blocks": [{**_judgment(b["id"]), "riskBand": "high"} for b in p["blocks"]]}
    _, job = _submit(client, {"writing": "AI_CHECK", "academic": False})
    assert job["analysis"]["coverageComplete"] is True
    assert job["analysis"]["disagreementBlocks"] and job["analysis"]["confidence"] == "LOW"


@pytest.mark.parametrize("reviewer", ["p_plan_review", "p_plan_review_peer"])
def test_either_frontier_can_block_a_proposal_plan_before_release(client, reviewer):
    from tests.test_proposals import _create, _run

    project = _create(client)
    client.models.overrides[reviewer] = lambda p: {"approved": False, "issues": ["Plan is not aligned."]}
    available = client.get("/api/wallet", headers=STUDENT).json()["available"]
    job = _run(client, project["id"], "PLAN")
    assert job["status"] == "FAILED" and job["failure"]["code"] == "DOCUMENT_NOT_APPROVED"
    assert job["outputs"] == [] and job["billing"]["charged"] == 0
    assert client.get("/api/wallet", headers=STUDENT).json()["available"] == available


def test_estimate_handoffs_remain_unique_after_200_calls(client, monkeypatch):
    from datetime import timedelta

    from app.jobs import pipeline
    from app.jobs.models import EstimateRun, Job, JobStatus, ModelCall, ServiceSelection, Stage, utcnow
    from app.runtime import get_runtime

    rt = get_runtime()
    call = ModelCall(stage=Stage.ANALYSING, provider="fake", model="m", prompt_version="p", input_tokens=0, output_tokens=0, cached_tokens=0, latency_ms=0, cost_usd=0)
    job = Job(id="job_handoff", status=JobStatus.DRAFT, owner_uid="u", owner_email="u@example.com", expires_at=utcnow() + timedelta(days=1),
              estimate=EstimateRun(id="estimate_long", status="RUNNING", selection=ServiceSelection(writing="REFINE"), source_sha256="s", budget_usd=1, fee_cap=0),
              model_calls=[call] * 200)
    rt.store.create(job)
    names = []
    monkeypatch.setattr(rt.queue, "enqueue", lambda jid, task_name, **kwargs: names.append(task_name))

    def handoff(ctx):
        raise pipeline.StageContinues

    monkeypatch.setattr(pipeline, "run_estimate", handoff)
    pipeline._run_estimate_task(rt, job.id)
    pipeline._run_estimate_task(rt, job.id)
    assert len(set(names)) == 2
    assert len(rt.store.get(job.id).model_calls) == 200
