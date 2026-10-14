"""Named code calculations get attribution, never fabricated empirical evidence."""

import io
from datetime import timedelta
from types import SimpleNamespace

import pytest
from docx import Document

from app.jobs.models import utcnow
from app.proposals import calculation_references as methods
from app.proposals import evidence, export, pipeline, service
from app.proposals.models import ChapterDocument, ChapterSection, Project, ProposalInputs, StoredChapterState
from app.proposals.rulebook import DEFAULT


def chapter(text="The formula published by Krejcie and Morgan will be used.", method="KREJCIE_MORGAN"):
    return ChapterDocument(number=3, title="Methodology", plan_version=1, words=len(text.split()), cited=[],
                           sections=[ChapterSection(key="sampling", number="3.5", heading="Sample Size", paragraphs=[text])],
                           calculation_source=methods.source(method))


def project():
    return Project(id="prj_synthetic", owner_uid="synthetic", owner_email="synthetic@example.com", rulebook=DEFAULT, citation="APA7",
                   inputs=ProposalInputs(topic="Synthetic topic for a deterministic export", level="MASTERS"),
                   chapters=[StoredChapterState(number=n) for n in (1, 2, 3)], expires_at=utcnow() + timedelta(days=30))


@pytest.mark.parametrize("style", ["APA6", "APA7", "HARVARD"])
@pytest.mark.parametrize("method,name,year", [("KREJCIE_MORGAN", "Krejcie and Morgan", "1970"), ("COCHRAN", "Cochran", "1977"), ("YAMANE", "Yamane", "1967")])
def test_calculated_method_has_a_registered_reference_and_year(style, method, name, year):
    record = methods.source(method)
    text = f"The formula published by {name} will be used."
    assert methods.render(text, record, evidence.Citer({}, style)) == f"The formula published by {name} ({year}) will be used."
    entry = evidence.reference(record, style)
    assert year in entry and record.title.split(" (")[0] in entry
    assert record.metadata != "MODEL" and not record.retraction_checked


def test_method_rendering_does_not_invent_a_mention_or_duplicate_a_year():
    record = methods.source("KREJCIE_MORGAN")
    for text in ("The calculation will use the student's supplied assumptions.", "Krejcie and Morgan (1970) provide the formula."):
        assert methods.render(text, record, evidence.Citer({}, "APA7")) == text
    assert methods.source("SATURATION") is None and methods.source("STATED") is None
    assert methods.source("COCHRAN") is not methods.source("COCHRAN")


def test_preview_and_word_export_share_the_same_calculation_reference(monkeypatch):
    doc = chapter()
    p = project()
    monkeypatch.setattr(service, "_owned", lambda *a: p)
    monkeypatch.setattr(service, "_chapter_doc", lambda rt, stored, version=None: doc if stored.number == 3 else None)
    monkeypatch.setattr(service, "_citer", lambda *a: evidence.Citer({}, "APA7"))
    view = service.chapter(SimpleNamespace(), SimpleNamespace(), p.id, 3)
    assert "Krejcie and Morgan (1970)" in view.sections[0].paragraphs[0]
    assert len(view.references) == 1 and "607–610" in view.references[0]
    text = "\n".join(p.text for p in Document(io.BytesIO(export.build(p, {3: doc}, {}, draft=True))).paragraphs)
    assert view.sections[0].paragraphs[0] in text and view.references[0] in text
    assert "REFERENCES" in text


def test_historical_and_legacy_chapters_do_not_gain_new_method_references(monkeypatch):
    legacy = chapter().model_copy(update={"calculation_source": None})
    current = chapter()
    p = project()
    monkeypatch.setattr(service, "_owned", lambda *a: p)
    monkeypatch.setattr(service, "_chapter_doc", lambda rt, stored, version=None: legacy if version else current)
    monkeypatch.setattr(service, "_citer", lambda *a: evidence.Citer({}, "APA7"))
    view = service.chapter(SimpleNamespace(), SimpleNamespace(), p.id, 3, version=1)
    assert view.references == [] and "(1970)" not in view.sections[0].paragraphs[0]
    text = "\n".join(p.text for p in Document(io.BytesIO(export.build(p, {3: legacy}, {}, draft=True))).paragraphs)
    assert "Determining sample size" not in text


def test_other_section_revisions_keep_the_sampling_reference_and_removal_drops_it():
    base = chapter()
    unrelated = chapter().model_copy(update={"sections": [ChapterSection(key="design", number="3.1", heading="Design", paragraphs=["A cross-sectional design will be used."])], "calculation_source": None})
    kept, missing = pipeline._merge(base, unrelated, {"design"})
    assert kept.calculation_source == base.calculation_source and not missing
    revised = chapter("Sample size will be determined with the supervisor.").model_copy(update={"calculation_source": None})
    removed, _ = pipeline._merge(base, revised, {"sampling"})
    assert removed.calculation_source is None


def test_document_only_freezes_method_metadata_when_the_approved_sampling_text_names_it():
    from app.proposals.ai import SectionText, Table
    from tests.test_scoped_hardening_20261013 import _chapter_three_input

    input_ = _chapter_three_input(3)
    input_.calculation_source = methods.source("KREJCIE_MORGAN")
    items = {i["key"]: i for i in pipeline._section_items(input_, {}, {})}
    named = pipeline._document(input_, {"sampling": SectionText(key="sampling", paragraphs=chapter().sections[0].paragraphs, table=Table(caption="", rows=[]))}, items, {})
    assert named.calculation_source == input_.calculation_source and named.cited == []
    unnamed = pipeline._document(input_, {"sampling": SectionText(key="sampling", paragraphs=["The supervisor will confirm the calculation."], table=Table(caption="", rows=[]))}, items, {})
    assert unnamed.calculation_source is None
