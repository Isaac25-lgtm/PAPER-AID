"""The algorithm revision of 2026-10-09 (owner, with Codex's plan): each change with the case that shows it.
Synthetic inputs and stand-in models only; nothing here makes a paid call."""

import threading
import time

import pytest

from app.ai import orchestration
from app.ai.orchestration import AIRunner
from app.core.errors import PermanentStageError
from tests.test_ai import ScriptedProvider, real_settings
from tests.test_api import STUDENT

PASSAGE = [{"id": "b1", "text": "A short paragraph of about twenty words that needs a quick look from the lead model before pricing.", "signals": []}]

# --- A call waits for the budget before it takes a slot of the shared limit ------------------------------------------


class _Priced(ScriptedProvider):
    name = "openai"


def _gated(monkeypatch, heartbeat=None, budget=5.0):
    scripted = _Priced({"analyse": [{"blocks": []}, {"blocks": []}]})
    monkeypatch.setattr(orchestration, "provider_for", lambda ref, settings: (scripted, "gpt-6-sol"))
    events: list[str] = []

    def gate(model, searching):
        events.append("taken")
        return (lambda: events.append("released")), 0

    runner = AIRunner(real_settings(), lambda c: None, lambda: 0.0, budget, heartbeat=heartbeat, gate=gate)
    return runner, scripted, events


def test_a_call_waiting_for_the_budget_holds_no_slot(monkeypatch):
    runner, scripted, events = _gated(monkeypatch)
    runner._inflight = 4.99  # a call alongside holds nearly the whole cap at its estimate; nothing is really spent
    def call():
        try:
            runner.analyse(PASSAGE, [], {})
        except PermanentStageError:
            pass  # the second assessor has no price in this stand-in: only the first call matters here

    thread = threading.Thread(target=call)
    thread.start()
    time.sleep(0.3)
    assert events == [] and scripted.calls == []  # it waits for the other call to be recorded, with no slot taken
    with runner._inflight_lock:  # the other call is recorded at its real cost
        runner._inflight = 0.0
        runner._inflight_lock.notify_all()
    thread.join(timeout=10)
    assert events[:2] == ["taken", "released"] and scripted.calls[:1] == ["analyse"]


def test_a_job_stopped_during_the_budget_wait_sends_nothing_and_frees_its_slot(monkeypatch):
    stopped = {"now": False}

    def heartbeat():
        if stopped["now"]:
            raise PermanentStageError("JOB_STOPPED", "Stopped.", "test: stopped while waiting")

    runner, scripted, events = _gated(monkeypatch, heartbeat=heartbeat)
    runner._inflight = 4.99
    failure: list[BaseException] = []

    def call():
        try:
            runner.analyse(PASSAGE, [], {})
        except PermanentStageError as exc:
            failure.append(exc)

    thread = threading.Thread(target=call)
    thread.start()
    time.sleep(0.3)
    stopped["now"] = True
    with runner._inflight_lock:
        runner._inflight = 0.0
        runner._inflight_lock.notify_all()
    thread.join(timeout=10)
    assert failure and failure[0].code == "JOB_STOPPED"
    assert scripted.calls == [] and events == ["taken", "released"]  # checked again after the wait: nothing was sent


def test_part_of_the_cap_can_be_kept_for_the_final_review(monkeypatch):
    runner, scripted, _ = _gated(monkeypatch, budget=4.0)
    runner.hold_back(3.9995, frozenset({"review"}))
    assert runner._cap("review") == 4.0 and runner._cap("analyse") == pytest.approx(0.0005)
    with pytest.raises(PermanentStageError) as err:  # another step may not spend what is kept for the review
        runner.analyse(PASSAGE, [], {})
    assert err.value.code == "BUDGET_EXCEEDED" and scripted.calls == []


# --- Workflow 2: frozen into the engine; topic planning has its own stage ---------------------------------------------


def test_the_workflow_is_frozen_into_each_engine(monkeypatch):
    from app.ai.orchestration import current_engine
    from app.core.config import Settings

    monkeypatch.setenv("VERTEX_PROJECT", "paperaid")
    old = current_engine(Settings())
    assert old.workflow == 1 and old.prompts["w_needs"] == "w-needs-v1" and old.vertex_stages["w_needs"] == "planner"
    monkeypatch.setenv("WORKFLOW", "2")
    new = current_engine(Settings())
    assert new.workflow == 2 and new.prompts["w_needs"] == "w-needs-v2" and new.prompts["p_needs"] == "p-needs-v2"
    assert new.vertex_stages["w_needs"] == "topics" and new.vertex_thinking["topics"] == "MEDIUM"
    assert new.vertex_thinking["planner"] == "HIGH"  # plans keep their own level


# --- The index: a key, a pool of candidates, and what a search came to -------------------------------------------------


@pytest.mark.parametrize("response, outcome", [
    (None, "UNAVAILABLE"),
    ((429, b"", "", {"x-ratelimit-remaining-usd": "0"}), "ALLOWANCE_USED"),
    ((429, b"", "", {"x-ratelimit-remaining-usd": "0.4"}), "RATE_LIMITED"),
    ((403, b"", "", {}), "ACCESS_DENIED"),
    ((503, b"", "", {}), "UNAVAILABLE"),
    ((200, b"not json", "application/json", {}), "INVALID"),
    ((200, b'{"results": []}', "application/json", {}), "NO_RESULTS"),
])
def test_an_index_search_says_what_it_came_to(monkeypatch, response, outcome):
    from app.analysis import fetch

    monkeypatch.setattr(fetch, "_request", lambda url, redirects=3: response)
    assert fetch.openalex_find("malaria vaccine uptake", 2015, 50) == (outcome, [])


def test_the_index_key_goes_only_to_the_index_and_no_redirect_is_followed(monkeypatch):
    from app.analysis import fetch

    seen = []
    monkeypatch.setattr(fetch, "_request", lambda url, redirects=3: seen.append((url, redirects)) or (200, b'{"results": []}', "", {}))
    monkeypatch.setattr(fetch, "_openalex_key", "")
    fetch.set_openalex_key(" test-key ")
    fetch.openalex_find("malaria", 2015, 50)
    fetch.openalex_retracted("10.1000/x")
    assert all(url.startswith("https://api.openalex.org/") and url.endswith("api_key=test-key") and redirects == 0 for url, redirects in seen) and len(seen) == 2
    monkeypatch.setattr(fetch, "_fetch", lambda url: seen.append((url, None)) or (404, b"", ""))
    fetch.crossref_lookup("10.1000/x")
    assert "test-key" not in seen[-1][0]  # another service never sees it


def test_candidates_are_ranked_for_the_need_not_only_by_the_index():
    from app.analysis import research

    def work(n, title, abstract, **more):
        return {"doi": f"10.1/{n}", "title": title, "abstract": abstract, **more}

    works = [
        work(1, "Vaccine hesitancy in European adults", "A survey of hesitancy among adults in Europe."),
        work(2, "Malaria vaccine uptake among caregivers in Uganda", "Uptake of the malaria vaccine among caregivers in Mukono, Uganda."),
        work(3, "Malaria vaccine uptake among caregivers in Uganda", "A retracted copy.", retracted="yes"),
        {"doi": "10.1/2", "title": "Malaria vaccine uptake among caregivers in Uganda (repeat)", "abstract": "The same work again."},
        work(5, "Caregivers and vaccines", "Caregivers in Uganda and uptake of vaccines."),
    ]
    ranked = research.rank_works(works, "What determines malaria vaccine uptake among caregivers in Uganda?", "malaria vaccine uptake caregivers", 3)
    assert [w["doi"] for w in ranked] == ["10.1/2", "10.1/5", "10.1/1"]  # the local study first; the retracted work and the repeat left out


# --- Workflow 2: how each topic is answered ---------------------------------------------------------------------------


@pytest.fixture
def v2(tmp_path, monkeypatch):
    from tests.conftest import _client

    monkeypatch.setenv("WORKFLOW", "2")
    monkeypatch.setenv("WORKS_ENABLED", '["CONCEPT_NOTE","COURSEWORK","FUNDING_PROPOSAL"]')
    yield from _client(tmp_path, monkeypatch, "cost")


def _need(n, category, essential=False, covered=(), broader=""):
    return {"id": f"n{n}", "need": f"Evidence need number {n} about maternal health", "category": category, "essential": essential,
            "query": f"maternal health topic {n}", "broader": broader, "coveredBy": list(covered)}


def _plan(client, needs):
    from app.runtime import get_runtime
    from tests.test_works import _coursework, _run

    client.models.overrides["w_needs"] = lambda payload: {"needs": needs}
    work = _coursework(client)
    _, job = _run(client, work["id"], "PLAN")
    return work, job, get_runtime().store.get(job["id"])


def test_a_scholarly_topic_is_read_from_a_ranked_pool_and_never_searched_on_the_web_when_the_index_answers(v2):
    work, job, stored = _plan(v2, [_need(1, "STUDY"), _need(2, "METHOD")])
    assert job["status"] == "COMPLETED", job.get("failure")
    assert "w_search" not in v2.models.tasks and v2.models.tasks.count("w_extract") == 2
    assert [(t.category, t.index, t.web, t.verified) for t in stored.topics] == [("STUDY", "FOUND", "NOT_USED", 1), ("METHOD", "FOUND", "NOT_USED", 1)]


def test_the_index_is_asked_for_a_pool_and_the_reader_gets_only_the_best_few(v2, monkeypatch):
    from app.analysis import fetch
    from tests.fake_models import WORK

    asked = []

    def find(query, from_year, limit):
        asked.append(limit)
        return "FOUND", [{**WORK, "doi": f"10.1/{n}", "title": f"{WORK['title']} {n}"} for n in range(limit)]

    monkeypatch.setattr(fetch, "openalex_find", find)
    read = []
    real = v2.models.default
    v2.models.overrides["w_extract"] = lambda payload: read.append(len(payload["works"])) or real("w_extract", payload)
    _, job, stored = _plan(v2, [_need(1, "STUDY")])
    assert job["status"] == "COMPLETED" and asked == [50] and read == [8]
    assert (stored.topics[0].candidates, stored.topics[0].kept) == (50, 8)


def test_an_official_figure_goes_to_the_web_and_the_reason_is_recorded(v2):
    _, job, stored = _plan(v2, [_need(1, "STATISTIC"), _need(2, "POLICY")])
    assert job["status"] == "COMPLETED", job.get("failure")
    assert v2.models.tasks.count("w_search") == 2 and "w_extract" not in v2.models.tasks
    assert [(t.index, t.web, t.reason) for t in stored.topics] == [("NOT_TRIED", "FOUND", "OFFICIAL_SOURCE")] * 2


def test_an_optional_scholarly_topic_the_index_cannot_answer_is_not_sent_to_the_web(v2):
    v2.models.works = []  # the index has nothing
    _, job, stored = _plan(v2, [_need(1, "STUDY", broader="maternal health"), _need(2, "STATISTIC")])
    assert job["status"] == "COMPLETED", job.get("failure")
    assert v2.models.tasks.count("w_search") == 1  # only the official figure
    assert v2.models.searched[:2] == ["maternal health topic 1", "maternal health"]  # the broader query was tried once
    assert (stored.topics[0].index, stored.topics[0].web) == ("NO_RESULTS", "NOT_USED")


def test_an_essential_topic_the_index_cannot_answer_is_searched_on_the_web(v2):
    v2.models.works = []
    _, job, stored = _plan(v2, [_need(1, "STUDY", essential=True)])
    assert job["status"] == "COMPLETED", job.get("failure")
    assert (stored.topics[0].index, stored.topics[0].web, stored.topics[0].reason, stored.topics[0].essential) == ("NO_RESULTS", "FOUND", "INDEX_EMPTY", True)


def test_an_index_that_cannot_be_asked_is_not_mistaken_for_no_evidence(v2, monkeypatch):
    from app.analysis import fetch

    monkeypatch.setattr(fetch, "openalex_find", lambda query, from_year, limit: ("ALLOWANCE_USED", []))
    _, job, stored = _plan(v2, [_need(1, "STUDY")])  # optional, yet the web is tried: the index never answered
    assert job["status"] == "COMPLETED", job.get("failure")
    assert (stored.topics[0].index, stored.topics[0].web, stored.topics[0].reason) == ("ALLOWANCE_USED", "FOUND", "INDEX_FAILED")


def test_a_step_stops_without_charge_when_an_essential_topic_has_no_evidence(v2):
    v2.models.works = []
    v2.models.opened = []  # the web search opens nothing either
    _, job, stored = _plan(v2, [_need(1, "STUDY", essential=True), _need(2, "STUDY")])
    assert job["status"] == "FAILED" and job["failure"]["code"] == "EVIDENCE_MISSING"
    assert "Evidence need number 1" in job["failure"]["userMessage"] and "Nothing was charged" in job["failure"]["userMessage"]
    assert job["paymentStatus"] in ("REFUNDED", "NOT_REQUIRED")
    assert "w_plan" not in v2.models.tasks  # stopped in research, before anything was written


def test_saved_evidence_is_not_searched_for_again(v2):
    from tests.test_works import _run, _work

    work, job, _ = _plan(v2, [_need(1, "STUDY")])
    assert job["status"] == "COMPLETED", job.get("failure")
    work = _work(v2, work["id"])
    work = v2.post(f"/api/works/{work['id']}/plan/approve", headers=STUDENT, json={"baseVersion": work["planVersion"]}).json()
    library = []

    def needs(payload):
        library.extend(e["id"] for e in payload["library"])
        return {"needs": [_need(1, "STUDY", covered=library[:1]), _need(2, "STUDY", covered=["Enot0real"]), _need(3, "STATISTIC")]}

    v2.models.overrides["w_needs"] = needs
    before = (v2.models.tasks.count("w_extract"), v2.models.tasks.count("w_search"))
    from app.runtime import get_runtime

    _, drafted = _run(v2, work["id"], "DRAFT")
    assert drafted["status"] == "COMPLETED", drafted.get("failure")
    assert library and (v2.models.tasks.count("w_extract"), v2.models.tasks.count("w_search")) == (before[0] + 1, before[1] + 1)
    topics = get_runtime().store.get(drafted["id"]).topics
    assert [(t.category, t.covered) for t in topics] == [("STUDY", False), ("STATISTIC", False), ("STUDY", True)]  # an unknown id is not coverage


# --- Workflow 2: the final editor reviews, corrects and decides --------------------------------------------------------


def _draft(client, edit=None, resolve=None, flag=None):
    """A coursework drafted on workflow 2: (work, the draft's job, the stored document)."""
    from app.runtime import get_runtime
    from app.works import service
    from tests.test_works import _coursework, _run, _work

    work = _coursework(client)
    _, planned = _run(client, work["id"], "PLAN")
    assert planned["status"] == "COMPLETED", planned.get("failure")
    work = _work(client, work["id"])
    work = client.post(f"/api/works/{work['id']}/plan/approve", headers=STUDENT, json={"baseVersion": work["planVersion"]}).json()
    for task, answer in (("w_edit", edit), ("w_resolve", resolve), ("w_flag", flag)):
        if answer is not None:
            client.models.overrides[task] = answer
    _, job = _run(client, work["id"], "DRAFT")
    rt = get_runtime()
    return work, job, service._doc(rt, rt.store.get_work(work["id"]))


def _verdicts(payload, **more):
    from tests import fake_works

    return {**fake_works.answer("w_edit", payload), **more}


def _correct(payload, text, kind="WORDING", action="REPLACE", paragraph=None, section=None):
    target = payload["editable"][0]
    return {"section": section or target["key"], "paragraph": paragraph or target["paragraphs"][0]["id"], "action": action, "text": text, "kind": kind, "reason": "Test."}


def test_a_sound_draft_is_reviewed_once_by_the_final_editor_and_by_nothing_else(v2):
    _, job, document = _draft(v2)
    assert job["status"] == "COMPLETED", job.get("failure")
    tasks = v2.models.tasks
    assert tasks.count("w_edit") == 1 and not {"w_final", "w_repair", "w_flag", "w_resolve", "w_plan_review"} & set(tasks)
    approval = document.approval
    assert approval.reviewer == "vertex:gemini-3.1-pro-preview" and approval.corrections == 0 and not approval.resolved
    assert approval.reviewed_sha256 == approval.accepted_sha256  # nothing changed between what it read and what was delivered


def test_the_editor_is_given_the_text_it_may_correct_with_ids_tokens_evidence_and_every_rule(v2):
    seen = {}

    def edit(payload):
        seen.update(payload)
        return _verdicts(payload)

    _, job, _ = _draft(v2, edit=edit)
    assert job["status"] == "COMPLETED", job.get("failure")
    section = seen["editable"][0]
    assert section["paragraphs"][0]["id"] == "p1" and section["wordsNow"] > 0 and "words" in section
    assert all(e["token"].startswith("\u27e6E") and e["sourceWords"] for e in seen["evidence"]) and seen["evidence"]
    assert any(r.get("sections") for r in seen["rules"]) and any("sections" not in r for r in seen["rules"])  # section rules and document rules
    assert {"document", "coverage", "numberTokens", "findings", "length"} <= seen.keys()


def test_a_correction_of_punctuation_and_case_only_is_applied_without_a_support_check(v2):
    def edit(payload):
        first = payload["editable"][0]["paragraphs"][0]
        return _verdicts(payload, corrections=[_correct(payload, first["text"].rstrip(".") + " !")])

    from app.proposals.pipeline import changes_words

    _, job, document = _draft(v2, edit=edit)
    assert job["status"] == "COMPLETED", job.get("failure")
    assert document.sections[0].paragraphs[0].endswith(" !") and document.approval.corrections == 1
    assert "w_flag" not in v2.models.tasks and "w_resolve" not in v2.models.tasks  # code can see no word, figure or citation changed
    assert not changes_words("Uptake fell, slightly ⟦E1a2b3c⟧.", "uptake fell slightly ⟦E1a2b3c⟧") and changes_words("Uptake fell ⟦E1a2b3c⟧.", "Uptake fell ⟦E9z8y7x⟧.")
    assert changes_words("The source reports an association.", "The intervention causes the outcome.") and changes_words("It rose by 12%.", "It rose by 21%.")


def test_a_change_of_substance_labelled_wording_is_still_checked_against_its_evidence(v2):
    """Codex audit of fd74ff3, finding 3: the support check skipped whatever the editor called WORDING."""
    checked = []

    def edit(payload):
        return _verdicts(payload, corrections=[_correct(payload, "The intervention causes the outcome in every district.", kind="WORDING")])

    def flag(payload):
        checked.extend(payload["corrections"])
        return {"results": [{"id": c["id"], "supported": False, "problem": "The source reports an association."} for c in payload["corrections"] if "causes" in c["after"]]
                           + [{"id": c["id"], "supported": True, "problem": ""} for c in payload["corrections"] if "causes" not in c["after"]]}

    def resolve(payload):
        flagged = payload["flagged"][0]
        return _verdicts(payload, corrections=[_correct(payload, "The source reports an association between the two.", kind="WORDING", paragraph=flagged["paragraph"])])

    _, job, document = _draft(v2, edit=edit, flag=flag, resolve=resolve)
    assert job["status"] == "COMPLETED", job.get("failure")
    assert [c["after"] for c in checked] == ["The intervention causes the outcome in every district.", "The source reports an association between the two."]
    assert "causes the outcome" not in " ".join(p for sec in document.sections for p in sec.paragraphs) and document.approval.flagged == 1


def test_an_unsupported_claim_written_on_the_last_look_stops_the_work(v2):
    """Finding 3: the resolution's own corrections were never checked. There is no further look, so the work stops."""
    def edit(payload):
        return _verdicts(payload, corrections=[_correct(payload, "Distance causes low uptake everywhere.", kind="CLAIM")])

    def resolve(payload):
        return _verdicts(payload, corrections=[_correct(payload, "Distance is the only cause of low uptake.", kind="CLAIM", paragraph=payload["flagged"][0]["paragraph"])])

    _, job, document = _draft(v2, edit=edit, resolve=resolve,
                              flag=lambda payload: {"results": [{"id": c["id"], "supported": False, "problem": "Not what the source says."} for c in payload["corrections"]]})
    assert job["status"] == "FAILED" and job["failure"]["code"] == "DOCUMENT_NOT_APPROVED" and document is None
    assert v2.models.tasks.count("w_resolve") == 1 and v2.models.tasks.count("w_flag") == 2 and "w_repair" not in v2.models.tasks


def test_a_correction_refused_on_the_last_look_is_never_read_as_approval(v2):
    """Finding 1: a correction to an unknown section was recorded under no section and dropped, so the editor's
    decision stood for text that was never made."""
    def bad(payload):
        return _verdicts(payload, corrections=[{"section": "no-such-section", "paragraph": "p1", "action": "REPLACE", "text": "A replacement that cannot be placed.",
                                                "kind": "CLAIM", "reason": "Test."}])

    seen = {}

    def resolve(payload):
        seen.update(payload)
        return bad(payload)

    _, job, document = _draft(v2, edit=bad, resolve=resolve)
    assert job["status"] == "FAILED" and job["failure"]["code"] == "DOCUMENT_NOT_APPROVED" and document is None
    assert [c["section"] for c in seen["checks"]] == [""] and "could not be applied" in seen["checks"][0]["issues"][0]  # told once, for the whole document


def test_a_heading_two_sections_share_names_neither(v2):
    """Finding 6: a duplicate heading was silently given to the last section that had it."""
    from app.runtime import get_runtime
    from app.works import service
    from tests.test_works import _coursework, _run, _work

    work = _coursework(v2)
    _, planned = _run(v2, work["id"], "PLAN")
    assert planned["status"] == "COMPLETED", planned.get("failure")
    work = _work(v2, work["id"])
    plan = work["plan"]
    for section in plan["sections"][:2]:
        section["heading"] = "Summary"
    saved = v2.post(f"/api/works/{work['id']}/plan", headers=STUDENT, json={"plan": plan, "baseVersion": work["planVersion"]})
    assert saved.status_code == 200, saved.json()
    approved = v2.post(f"/api/works/{work['id']}/plan/approve", headers=STUDENT, json={"baseVersion": saved.json()["planVersion"], "acknowledge": ["PLAN_OBJECTIONS"]})
    assert approved.status_code == 200, approved.json()
    seen = {}

    def edit(payload):
        return _verdicts(payload, corrections=[_correct(payload, "A replacement addressed only by a heading two sections share.", section="Summary")])

    def resolve(payload):
        seen.update(payload)
        return _verdicts(payload)

    v2.models.overrides.update(w_edit=edit, w_resolve=resolve)
    _, job = _run(v2, work["id"], "DRAFT")
    assert job["status"] == "COMPLETED", job.get("failure")
    rt = get_runtime()
    document = service._doc(rt, rt.store.get_work(work["id"]))
    assert "A replacement addressed only" not in " ".join(p for sec in document.sections for p in sec.paragraphs)
    assert document.approval.refused == 1 and any("could not be applied" in i for c in seen["checks"] for i in c["issues"])


def test_a_wording_correction_is_applied_by_paragraph_and_supported_needs_no_second_look(v2):
    marker = "This sentence was sharpened by the final editor."

    def edit(payload):
        first = payload["editable"][0]["paragraphs"][0]
        return _verdicts(payload, corrections=[_correct(payload, first["text"] + " " + marker)])

    work, job, document = _draft(v2, edit=edit)
    assert job["status"] == "COMPLETED", job.get("failure")
    assert marker in document.sections[0].paragraphs[0] and sum(marker in p for s in document.sections for p in s.paragraphs) == 1
    assert v2.models.tasks.count("w_flag") == 1 and "w_resolve" not in v2.models.tasks  # checked (it added words), supported, no second look
    assert document.approval.corrections == 1 and document.approval.substantive == 0
    assert document.approval.reviewed_sha256 != document.approval.accepted_sha256
    shown = v2.get(f"/api/works/{work['id']}/document", headers=STUDENT).json()
    assert "approval" not in shown and marker in " ".join(p for sec in shown["sections"] for p in sec["paragraphs"])  # the student sees the text, never the record


def test_a_corrected_claim_is_checked_and_a_supported_one_stands(v2):
    checked = []

    def edit(payload):
        first = payload["editable"][0]["paragraphs"][0]
        return _verdicts(payload, corrections=[_correct(payload, first["text"] + " The association is modest.", kind="CLAIM")])

    def flag(payload):
        checked.extend(payload["corrections"])
        return {"results": [{"id": c["id"], "supported": True, "problem": ""} for c in payload["corrections"]]}

    _, job, document = _draft(v2, edit=edit, flag=flag)
    assert job["status"] == "COMPLETED", job.get("failure")
    assert len(checked) == 1 and checked[0]["before"] and checked[0]["after"].endswith("The association is modest.")
    assert "w_resolve" not in v2.models.tasks and document.approval.substantive == 1 and document.approval.flagged == 0


def test_a_flagged_claim_goes_back_to_the_editor_once_and_the_check_never_rewrites(v2):
    asked = {}

    def edit(payload):
        first = payload["editable"][0]["paragraphs"][0]
        return _verdicts(payload, corrections=[_correct(payload, first["text"] + " Distance causes low uptake everywhere.", kind="CLAIM")])

    def resolve(payload):
        asked.update(payload)
        flagged = payload["flagged"][0]
        target = next(p for e in payload["editable"] for p in e["paragraphs"] if p["id"] == flagged["paragraph"])
        return _verdicts(payload, corrections=[_correct(payload, target["text"].replace(" Distance causes low uptake everywhere.", " Distance is associated with lower uptake."),
                                                        kind="CLAIM", paragraph=flagged["paragraph"])])

    def flag(payload):  # the first claim is not what the source says; the narrowed one is
        return {"results": [{"id": c["id"], "supported": "causes low uptake everywhere" not in c["after"], "problem": "The source shows an association in one district."}
                            for c in payload["corrections"]]}

    _, job, document = _draft(v2, edit=edit, resolve=resolve, flag=flag)
    assert job["status"] == "COMPLETED", job.get("failure")
    text = " ".join(p for s in document.sections for p in s.paragraphs)
    assert "is associated with lower uptake" in text and "causes low uptake everywhere" not in text
    assert v2.models.tasks.count("w_resolve") == 1 and v2.models.tasks.count("w_flag") == 2  # one resolution, and what it wrote is checked too
    assert asked["flagged"][0]["problem"].startswith("The source shows") and "findings" not in asked
    assert document.approval.flagged == 1 and document.approval.resolved and "w_repair" not in v2.models.tasks


def test_wording_the_editor_cannot_trace_is_sent_back_and_never_released_if_it_stays(v2):
    untraceable = " A recent survey found that 87 percent of clinics lack supplies (Smith, 2020)."
    checks = {}

    def edit(payload):
        first = payload["editable"][0]["paragraphs"][0]
        return _verdicts(payload, corrections=[_correct(payload, first["text"] + untraceable)])

    def resolve(payload):
        checks.update({"checks": payload["checks"]})
        return _verdicts(payload)  # it leaves the sentence in

    _, job, document = _draft(v2, edit=edit, resolve=resolve)
    assert "cannot be traced" in str(checks["checks"])  # code found it and told the editor which paragraph
    assert job["status"] == "FAILED" and job["failure"]["code"] == "DOCUMENT_NOT_APPROVED" and document is None
    assert job["paymentStatus"] in ("REFUNDED", "NOT_REQUIRED")


def test_a_correction_that_names_its_section_by_heading_is_applied_to_that_section(v2):
    """Real-model trial 2026-10-09: on its second look the editor named sections by the headings it had been given,
    every such correction was refused, and three of six drafts stopped. Sections are now given by key, and a heading
    still names its section."""
    seen = {}

    def edit(payload):
        target = payload["editable"][0]
        seen["editable"], seen["length"] = target, payload["length"]
        return _verdicts(payload, corrections=[_correct(payload, "The essay first sets out the question and the argument it will make about maternal health.", section=f"  {target['heading'].upper()} ")])

    _, job, document = _draft(v2, edit=edit)
    assert job["status"] == "COMPLETED", job.get("failure")
    section = next(s for s in document.sections if s.key == seen["editable"]["key"])
    assert section.paragraphs[0] == "The essay first sets out the question and the argument it will make about maternal health."
    assert document.approval.refused == 0 and document.approval.corrections == 1
    assert seen["length"] | {"words": 0} == {"words": 0, "limit": 1500, "aimFor": 1455, "minimum": 1276}  # the real limit, a working target under it, how short is too short


def test_what_the_editor_is_sent_back_names_sections_by_key(v2):
    seen = {}

    def edit(payload):
        return _verdicts(payload, corrections=[_correct(payload, "Malaria kills 73% of children in the district, as Smith (2019) showed.", kind="CLAIM", action="INSERT_AFTER")])

    def resolve(payload):
        seen.update(payload)
        target = payload["editable"][0]
        return _verdicts(payload, corrections=[{"section": target["key"], "paragraph": "p2", "action": "DELETE", "text": "", "kind": "CLAIM", "reason": "Test."}])

    _, job, document = _draft(v2, edit=edit, resolve=resolve, flag=lambda p: {"results": [{"id": c["id"], "supported": False, "problem": "No source gives this."} for c in p["corrections"]]})
    assert job["status"] == "COMPLETED", job.get("failure")
    key = seen["editable"][0]["key"]
    assert [(f["section"], f["paragraph"]) for f in seen["flagged"]] == [(key, "p2")] and seen["flagged"][0]["heading"]
    assert [c["section"] for c in seen["checks"]] == [key] and "73%" not in " ".join(p for s in document.sections for p in s.paragraphs)


def test_a_draft_well_under_its_limit_is_developed_by_the_writer_before_the_editor_reads_it(v2):
    """Real-model trial 2026-10-09: drafts of 1,244 words for 1,500 went to the editor, which corrects and does not
    write, and stopped as well under their limit."""
    from tests import fake_works

    order = []

    def short_draft(payload):
        answer = fake_works.draft(payload)
        for s in answer["sections"]:
            words = " ".join(s["paragraphs"]).split()
            s["paragraphs"] = [" ".join(words[: int(len(words) * 0.7)]).rstrip(".,;") + "."]
        return answer

    def repair(payload):
        order.append("writer")
        out = []
        for s in payload["sections"]:
            assert any("well short" in i for i in s["issues"])
            out.append({"key": s["key"], "paragraphs": fake_works._text({"key": s["key"], "heading": s["heading"], "words": s["words"], "evidence": s["evidence"]}, []),
                        "table": {"caption": "", "rows": []}})
        return {"sections": out}

    v2.models.overrides["w_draft"] = short_draft
    v2.models.overrides["w_repair"] = repair
    _, job, document = _draft(v2, edit=lambda payload: order.append("editor") or _verdicts(payload))
    assert job["status"] == "COMPLETED", job.get("failure")
    assert order == ["writer", "editor"] and 1275 <= document.words <= 1500  # one development, then the one review


def test_a_correction_that_names_an_unknown_place_is_refused_never_guessed_at():
    from app.proposals.ai import Correction, SectionText, Table
    from app.proposals.pipeline import apply_corrections as _apply_corrections

    def section(key, *paragraphs):
        return SectionText(key=key, paragraphs=list(paragraphs), table=Table(caption="T", rows=[["a"]]))

    current = {"intro": section("intro", "One.", "Two.", "Three."), "body": section("body", "Kept.")}

    def c(section_key, paragraph, action, text="", kind="WORDING"):
        return Correction(section=section_key, paragraph=paragraph, action=action, text=text, kind=kind)

    applied, refused = _apply_corrections(current, [
        c("intro", "p2", "REPLACE", "Two,  rewritten."), c("intro", "p2", "REPLACE", "Again."), c("intro", "p1", "INSERT_AFTER", "New."),
        c("intro", "p3", "DELETE"), c("intro", "p9", "REPLACE", "Nowhere."), c("body", "p1", "REPLACE", "Not allowed."),
        c("intro", "p1", "REMOVE_TABLE"), c("intro", "p1", "REPLACE", ""),
    ], ["intro"])
    assert current["intro"].paragraphs == ["One.", "New.", "Two, rewritten."] and current["intro"].table.rows == []
    assert current["body"].paragraphs == ["Kept."] and current["body"].table.rows == [["a"]]  # a section this step may not change
    assert [(a["action"], a["paragraph"]) for a in applied] == [("REPLACE", "p3"), ("INSERT_AFTER", "p2"), ("DELETE", ""), ("REMOVE_TABLE", "p1")]
    assert len(refused["intro"]) == 3 and len(refused[""]) == 1  # corrected twice, no such paragraph, no text; and the other section


def test_a_blocker_the_editor_cannot_settle_stops_the_draft_without_any_repair_round(v2):
    def edit(payload):
        answer = _verdicts(payload)
        answer["coverage"] = [{**c, "answered": False, "where": ""} for c in answer["coverage"]]
        answer["blockers"] = [{"where": "", "issue": "The question asks for a comparison the evidence cannot support."}]
        return answer

    _, job, document = _draft(v2, edit=edit)
    assert job["status"] == "FAILED" and job["failure"]["code"] == "DOCUMENT_NOT_READY" and document is None
    assert v2.models.tasks.count("w_edit") == 1 and "w_repair" not in v2.models.tasks and "w_final" not in v2.models.tasks


def test_when_the_final_editor_cannot_answer_the_draft_waits_and_nothing_else_approves(v2, monkeypatch):
    from app.core.errors import RetryableStageError
    from app.runtime import get_runtime

    get_runtime().settings.review_wait_sec = 1
    calls = {"n": 0}

    def edit(payload):
        calls["n"] += 1
        if calls["n"] <= 2:
            raise RetryableStageError("PROVIDER_UNAVAILABLE", "Unavailable.", "test: the premium model is down")
        return _verdicts(payload)

    _, job, document = _draft(v2, edit=edit)
    assert job["status"] == "COMPLETED", job.get("failure")
    stored = get_runtime().store.get(job["id"])
    assert calls["n"] == 3 and stored.capacity_waits == 2 and stored.attempts == 0  # two waits, no retry of the stage used up
    assert [t.code for t in stored.timings if t.outcome == "WAITING"] == ["REVIEW", "REVIEW"]
    assert v2.models.tasks.count("w_draft") == 1 and "w_final" not in v2.models.tasks  # the draft was kept; no stand-in approved
    assert document.approval.reviewer == "vertex:gemini-3.1-pro-preview"


def test_a_coursework_plan_is_decided_by_code_and_says_so(v2):
    from app.runtime import get_runtime
    from tests.test_works import _coursework, _run

    work = _coursework(v2)
    _, job = _run(v2, work["id"], "PLAN")
    assert job["status"] == "COMPLETED", job.get("failure")
    assert "w_plan_review" not in v2.models.tasks and v2.models.tasks.count("w_plan") == 1
    review = get_runtime().store.get_work(work["id"]).plan_review
    assert (review.outcome, review.reason) == ("APPROVED", "CODE_CHECKS")  # not recorded as a reviewer's approval


def test_the_final_review_s_share_of_the_cap_is_kept_from_every_other_step(monkeypatch):
    from app.ai.orchestration import EDITOR_TASKS

    runner, _, _ = _gated(monkeypatch, budget=10.0)
    monkeypatch.setattr(AIRunner, "worst_cost", lambda self, task, chars: 0.6)
    runner.keep_for_editor(12_000)  # one part: the review, its resolution and the support check after each
    assert runner._held == pytest.approx(4 * 0.6) and runner._held_for == EDITOR_TASKS
    assert runner._cap("w_edit") == 10.0 and runner._cap("w_search") == pytest.approx(10 - 2.4)
    runner, _, _ = _gated(monkeypatch, budget=40.0)
    runner.keep_for_editor(60_000)  # a long work is reviewed in parts, and each part's review and resolution is kept
    assert runner._held == pytest.approx(5 * 2 * 0.6 + 2 * 0.6)
    runner, _, _ = _gated(monkeypatch, budget=4.0)
    monkeypatch.setattr(AIRunner, "worst_cost", lambda self, task, chars: 5.0)
    runner.keep_for_editor(12_000)
    assert runner._held == 2.0  # never more than half the cap


def test_workflow_1_drafts_are_reviewed_as_they_were_priced(tmp_path, monkeypatch):
    from tests.conftest import _client

    monkeypatch.setenv("WORKS_ENABLED", '["CONCEPT_NOTE","COURSEWORK","FUNDING_PROPOSAL"]')
    for client in _client(tmp_path, monkeypatch, "cost"):
        _, job, document = _draft(client)
        assert job["status"] == "COMPLETED", job.get("failure")
        assert "w_final" in client.models.tasks and "w_plan_review" in client.models.tasks and "w_edit" not in client.models.tasks
        assert document.approval is None


def test_asking_for_changes_is_approved_by_the_final_editor_too(v2):
    from app.runtime import get_runtime
    from app.works import service
    from tests.test_works import _run, _work

    work, job, document = _draft(v2)
    assert job["status"] == "COMPLETED", job.get("failure")
    work = _work(v2, work["id"])
    key = document.sections[0].key
    asked = v2.post(f"/api/works/{work['id']}/requests", headers=STUDENT, json={"instruction": "Say more about rural clinics.", "sections": [key]})
    assert asked.status_code == 200, asked.json()
    before = v2.models.tasks.count("w_edit")
    seen = {}
    v2.models.overrides["w_edit"] = lambda payload: seen.update(payload) or _verdicts(payload)
    _, revised = _run(v2, work["id"], "REVISE")
    assert revised["status"] == "COMPLETED", revised.get("failure")
    assert v2.models.tasks.count("w_edit") == before + 1 and "w_final" not in v2.models.tasks
    assert [e["key"] for e in seen["editable"]] == [key]  # only the section the student asked about may be corrected
    rt = get_runtime()
    latest = service._doc(rt, rt.store.get_work(work["id"]))
    assert latest.approval is not None and latest.revised == [key]


def test_the_built_word_file_is_read_back_before_a_draft_can_complete(v2, monkeypatch):
    from app.works import export

    _, job, document = _draft(v2)
    assert job["status"] == "COMPLETED", job.get("failure")  # the real file carries every heading and paragraph
    real = export.build

    def without_a_paragraph(*args, **kwargs):
        import io

        from docx import Document

        built = Document(io.BytesIO(real(*args, **kwargs)))
        body = [p for p in built.paragraphs if len(p.text.split()) > 8][0]
        body._element.getparent().remove(body._element)
        out = io.BytesIO()
        built.save(out)
        return out.getvalue()

    monkeypatch.setattr(export, "build", without_a_paragraph)
    _, job, document = _draft(v2)
    assert job["status"] == "FAILED" and job["failure"]["code"] == "EXPORT_FAILED" and document is None


# --- Workflow 2: proposal chapters -------------------------------------------------------------------------------------


def _chapter_one(client, **answers):
    """A proposal planned and approved, so Chapter One is written: (project id, the chapter's job, the stored chapter)."""
    from app.proposals import service
    from app.runtime import get_runtime
    from tests.test_api import wait
    from tests.test_proposals import _create, _run

    pid = _create(client)["id"]
    planned = _run(client, pid, "PLAN")
    assert planned["status"] == "COMPLETED", planned.get("failure")
    client.models.overrides.update(answers)
    project = client.get(f"/api/projects/{pid}", headers=STUDENT).json()
    saved = client.post(f"/api/projects/{pid}/plan", headers=STUDENT, json={"plan": project["plan"], "baseVersion": project["planVersion"]}).json()
    acknowledge = ["SAMPLING"] if saved["plan"].get("samplingAssumed") else []
    approved = client.post(f"/api/projects/{pid}/plan/approve", headers=STUDENT, json={"baseVersion": saved["planVersion"], "acknowledge": acknowledge})
    assert approved.status_code == 200, approved.json()
    job = wait(client, approved.json()["activeJob"])
    rt = get_runtime()
    return pid, job, service._chapter_doc(rt, rt.store.get_project(pid).chapter(1))


def _passes(payload, corrections=None):
    """Every section passed, with `corrections` ({section key: [correction]})."""
    return {"results": [{"key": s["key"], "grade": "PASS", "issues": [], "note": "", "corrections": (corrections or {}).get(s["key"], [])} for s in payload["sections"]]}


def _change(section, text, kind="WORDING", action="REPLACE", paragraph=0):
    return {"paragraph": section["paragraphs"][paragraph]["id"], "action": action, "text": text, "kind": kind, "reason": "Test."}


def _section(payload, key):
    return next((s for s in payload["sections"] if s["key"] == key), None)


def test_a_sound_chapter_is_reviewed_once_by_the_final_editor_and_by_no_other(v2):
    _, job, chapter = _chapter_one(v2)
    assert job["status"] == "COMPLETED", job.get("failure")
    tasks = v2.models.tasks
    assert "p_edit" in tasks and not {"p_review", "p_review_peer", "p_fix", "p_flag", "p_resolve"} & set(tasks)
    approval = chapter.approval
    assert approval.reviewer == "vertex:gemini-3.1-pro-preview" and approval.corrections == 0 and not approval.resolved
    assert approval.reviewed_sha256 == approval.accepted_sha256 and all(s.reviewed for s in chapter.sections)
    shown = v2.get(f"/api/projects/{_}/chapters/1", headers=STUDENT).json()
    assert "approval" not in shown  # the record is for admins and audits, never the student's page


def test_a_chapter_correction_is_applied_by_paragraph_and_the_plan_s_statements_stay_word_for_word(v2):
    from app.proposals.models import ProposalPlan
    from app.proposals.pipeline import plan_statements

    def edit(payload):
        corrections = {}
        if (background := _section(payload, "background")):
            corrections["background"] = [_change(background, "The study will follow the plan the student approved.", paragraph=1)]
        if (objectives := _section(payload, "objectives")):
            corrections["objectives"] = [_change(objectives, "To do something else entirely.", kind="CLAIM", paragraph=len(objectives["paragraphs"]) - 1)]
        return _passes(payload, corrections)

    pid, job, chapter = _chapter_one(v2, p_edit=edit)
    assert job["status"] == "COMPLETED", job.get("failure")
    sections = {s.key: s for s in chapter.sections}
    assert sections["background"].paragraphs[1] == "The study will follow the plan the student approved."
    plan = ProposalPlan.model_validate(v2.get(f"/api/projects/{pid}", headers=STUDENT).json()["plan"])
    statements = plan_statements(plan, "objectives", sections["objectives"].number)
    assert sections["objectives"].paragraphs[-len(statements):] == statements  # put back by code, whatever the editor changed
    # the kept correction changed words, so it is checked (and supported); the undone one is not in the text: nothing to check
    assert v2.models.tasks.count("p_flag") == 1 and "p_resolve" not in v2.models.tasks
    assert chapter.approval.corrections == 1 and chapter.approval.reviewed_sha256 != chapter.approval.accepted_sha256


def test_a_flagged_chapter_claim_goes_back_to_the_editor_once_for_that_section_only(v2):
    seen = {}

    def edit(payload):
        background = _section(payload, "background")
        if background is None:
            return _passes(payload)
        token = f"⟦{background['evidence'][0]['id']}⟧"
        return _passes(payload, {"background": [_change(background, f"Distance causes lower uptake among every caregiver in the country. {token}", kind="CLAIM")]})

    def flag(payload):
        seen.setdefault("flag", payload)
        return {"results": [{"id": c["id"], "supported": "causes" not in c["after"], "problem": "The source shows an association among rural caregivers, not a cause for all."}
                            for c in payload["corrections"]]}

    def resolve(payload):
        seen["resolve"] = payload
        background = _section(payload, "background")
        token = f"⟦{background['evidence'][0]['id']}⟧"
        return _passes(payload, {"background": [_change(background, f"Distance was associated with lower uptake among rural caregivers. {token}", kind="CLAIM")]})

    _, job, chapter = _chapter_one(v2, p_edit=edit, p_flag=flag, p_resolve=resolve)
    assert job["status"] == "COMPLETED", job.get("failure")
    tasks = v2.models.tasks
    assert tasks.count("p_flag") == 2 and tasks.count("p_resolve") == 1 and "p_fix" not in tasks and "p_review" not in tasks  # the last look is checked too
    correction = seen["flag"]["corrections"][0]
    assert correction["after"].startswith("Distance causes") and correction["evidence"][0]["sourceWords"]  # checked against the source's own words
    assert [s["key"] for s in seen["resolve"]["sections"]] == ["background"]
    assert seen["resolve"]["sections"][0]["flagged"] == [{"paragraph": "p1", "problem": "The source shows an association among rural caregivers, not a cause for all."}]
    background = next(s for s in chapter.sections if s.key == "background")
    assert background.paragraphs[0].startswith("Distance was associated with lower uptake among rural caregivers.")
    approval = chapter.approval
    assert (approval.flagged, approval.resolved, approval.substantive) == (1, True, 2)


def test_chapter_wording_code_cannot_trace_after_the_editor_s_last_look_is_never_released(v2):
    def edit(payload):
        background = _section(payload, "background")
        if background is None:
            return _passes(payload)
        return _passes(payload, {"background": [_change(background, "Malaria kills 73% of children in the district, as Smith (2019) showed.", kind="CLAIM", action="INSERT_AFTER")]})

    seen = {}

    def resolve(payload):
        seen.update(payload)
        return _passes(payload)  # the editor leaves it

    pid, job, chapter = _chapter_one(v2, p_edit=edit, p_resolve=resolve)
    assert job["status"] == "FAILED" and job["failure"]["code"] == "DOCUMENT_NOT_APPROVED" and job["billing"]["charged"] == 0
    assert chapter is None and "p_fix" not in v2.models.tasks
    checks = seen["sections"][0]["paperaidChecks"]
    assert [s["key"] for s in seen["sections"]] == ["background"] and any(c.startswith("Paragraph p2 carries") for c in checks)


def test_a_section_the_editor_does_not_pass_is_not_delivered(v2):
    def edit(payload):
        answer = _passes(payload)
        for result in answer["results"]:
            if result["key"] == "problem":
                result.update(grade="REPAIR", issues=["The size of the problem in the study area is not given, and only the student can give it."])
        return answer

    _, job, chapter = _chapter_one(v2, p_edit=edit)
    assert job["status"] == "FAILED" and job["failure"]["code"] == "DOCUMENT_NOT_APPROVED" and chapter is None
    assert v2.models.tasks.count("p_resolve") == 0 and "p_fix" not in v2.models.tasks  # nothing flagged or found: no second look, no writer


def test_when_the_final_editor_cannot_answer_a_chapter_waits_with_its_draft_kept(v2):
    from app.core.errors import RetryableStageError
    from app.runtime import get_runtime

    get_runtime().settings.review_wait_sec = 1
    calls = {"n": 0}

    def edit(payload):
        calls["n"] += 1
        if calls["n"] <= 2:
            raise RetryableStageError("VERTEX_MODEL_UNAVAILABLE", "Unavailable.", "test: the premium model is down")
        return _passes(payload)

    _, job, chapter = _chapter_one(v2, p_edit=edit)
    assert job["status"] == "COMPLETED", job.get("failure")
    stored = get_runtime().store.get(job["id"])
    assert stored.capacity_waits == 2 and stored.attempts == 0 and [t.code for t in stored.timings if t.outcome == "WAITING"] == ["REVIEW", "REVIEW"]
    assert [t.outcome for t in stored.timings if t.stage == "DRAFTING"] == ["DONE"]  # written once: the waits kept the draft
    assert "p_review" not in v2.models.tasks and chapter.approval is not None


def test_a_chapter_revision_is_approved_by_the_final_editor_and_its_rewrite_is_paid_for_once(v2):
    from app.core.errors import RetryableStageError
    from app.proposals import service
    from app.runtime import get_runtime
    from tests.test_api import wait

    pid, job, _ = _chapter_one(v2)
    assert job["status"] == "COMPLETED", job.get("failure")
    added = v2.post(f"/api/projects/{pid}/chapters/1/request", headers=STUDENT, json={"instruction": "Show the size of the problem.", "sections": ["problem"]})
    assert added.status_code == 200, added.json()
    get_runtime().settings.review_wait_sec = 1
    calls = {"n": 0}
    seen = {}

    def edit(payload):
        calls["n"] += 1
        if calls["n"] == 1:
            raise RetryableStageError("PROVIDER_UNAVAILABLE", "Unavailable.", "test: the premium model is down")
        seen.update(payload)
        return _passes(payload)

    v2.models.overrides["p_edit"] = edit
    quoted = v2.post(f"/api/projects/{pid}/steps", headers=STUDENT, json={"step": "REVISE_1", "comments": [added.json()["feedback"][0]["id"]]}).json()
    v2.post(f"/api/projects/{pid}/steps/{quoted['job']['id']}/submit", headers=STUDENT, json={"quoteId": quoted["quote"]["id"]})
    revised = wait(v2, quoted["job"]["id"])
    assert revised["status"] == "COMPLETED", revised.get("failure")
    assert v2.models.tasks.count("p_fix") == 1 and "p_review" not in v2.models.tasks  # the rewrite was kept through the wait
    assert [s["key"] for s in seen["sections"]] == ["problem"]
    rt = get_runtime()
    latest = service._chapter_doc(rt, rt.store.get_project(pid).chapter(1))
    assert latest.revised == ["problem"] and latest.approval is not None
    assert any("supervisor's comment" in p for s in latest.sections if s.key == "problem" for p in s.paragraphs)


def test_a_proposal_s_scholarly_topic_is_read_from_the_index_and_not_searched_on_the_web(v2):
    from app.runtime import get_runtime
    from tests.test_proposals import _create, _run

    v2.models.overrides["p_needs"] = lambda payload: {"needs": [_need(1, "STUDY"), _need(2, "STATISTIC")]}
    job = _run(v2, _create(v2)["id"], "PLAN")
    assert job["status"] == "COMPLETED", job.get("failure")
    topics = get_runtime().store.get(job["id"]).topics
    assert [(t.category, t.index, t.web) for t in topics] == [("STUDY", "FOUND", "NOT_USED"), ("STATISTIC", "NOT_TRIED", "FOUND")]
    assert (v2.models.tasks.count("p_extract"), v2.models.tasks.count("p_search")) == (1, 1)


def test_workflow_1_chapters_are_reviewed_as_they_were_priced(client):
    _, job, chapter = _chapter_one(client)
    assert job["status"] == "COMPLETED", job.get("failure")
    assert "p_review" in client.models.tasks and "p_edit" not in client.models.tasks and chapter.approval is None


def test_a_chapter_s_final_review_keeps_its_share_of_the_cap(v2, monkeypatch):
    held = []
    keep = orchestration.AIRunner.keep_for_editor
    monkeypatch.setattr(orchestration.AIRunner, "keep_for_editor", lambda self, chars, service="w": (held.append((service, chars)), keep(self, chars, service))[1])
    _, job, _ = _chapter_one(v2)
    assert job["status"] == "COMPLETED", job.get("failure")
    assert held and {service for service, _ in held} == {"p"} and all(chars > 0 for _, chars in held)  # every stage of the chapter, never the plan


# --- Codex's audit of fd74ff3: chapters --------------------------------------------------------------------------------


def test_a_chapter_section_passed_with_issues_listed_is_asked_again_and_never_approved_as_it_stands(v2):
    """Finding 2: grade PASS with "The central claim is unsupported." among its issues was read as approval."""
    seen = {}

    def contradictory(payload):
        answer = _passes(payload)
        for result in answer["results"]:
            if result["key"] == "problem":
                result["issues"] = ["The central claim is unsupported."]
        return answer

    def resolve(payload):
        seen.update(payload)
        return contradictory(payload)

    _, job, chapter = _chapter_one(v2, p_edit=contradictory, p_resolve=resolve)
    assert job["status"] == "FAILED" and job["failure"]["code"] == "DOCUMENT_NOT_APPROVED" and chapter is None
    assert [s["key"] for s in seen["sections"]] == ["problem"] and "also listed issues" in seen["sections"][0]["paperaidChecks"][0]


def test_an_unsupported_chapter_claim_written_on_the_last_look_is_not_delivered(v2):
    def edit(payload, text="Distance causes lower uptake among every caregiver in the country."):
        background = _section(payload, "background")
        if background is None:
            return _passes(payload)
        return _passes(payload, {"background": [_change(background, f"{text} ⟦{background['evidence'][0]['id']}⟧", kind="WORDING")]})

    _, job, chapter = _chapter_one(v2, p_edit=edit, p_resolve=lambda payload: edit(payload, "Distance is the only cause of lower uptake."),
                                   p_flag=lambda payload: {"results": [{"id": c["id"], "supported": False, "problem": "The source shows an association."} for c in payload["corrections"]]})
    assert job["status"] == "FAILED" and job["failure"]["code"] == "DOCUMENT_NOT_APPROVED" and chapter is None
    assert v2.models.tasks.count("p_flag") == 2 and v2.models.tasks.count("p_resolve") == 1 and "p_fix" not in v2.models.tasks


# --- Codex's audit of fd74ff3: research --------------------------------------------------------------------------------


def _runner_v(workflow):
    from types import SimpleNamespace

    return SimpleNamespace(_engine=SimpleNamespace(workflow=workflow), budget_reached=False)


def _saved(item_id, statement, scope="Uganda", title="A study", need="An earlier need"):
    from app.proposals.models import EvidenceItem, EvidenceSource

    return EvidenceItem(id=item_id, chapter=0, need=need, statement=statement, passage=statement, scope=scope, access="ABSTRACT", verified=True, support="SUPPORTED",
                        source=EvidenceSource(url="https://example.org/" + item_id, title=title, year="2020"), retrieved_on="2026-10-09")


def _asked(n, need, query, essential=False, covered=(), category="STUDY"):
    from app.proposals.ai import Need

    return Need.model_validate({"id": f"n{n}", "need": need, "category": category, "essential": essential, "query": query, "broader": "", "coveredBy": list(covered)})


def test_an_essential_topic_is_never_the_one_the_search_allowance_cuts_off():
    """Finding 5: with one optional need before an essential one and room for one, the essential need vanished and
    the check that it was answered passed."""
    from app.proposals.pipeline import essential_answered, topics_to_search

    optional = _asked(1, "Background on vaccination campaigns", "vaccination campaigns background")
    essential = _asked(2, "Uptake of the malaria vaccine among caregivers in Mukono", "malaria vaccine uptake caregivers Mukono", essential=True)
    search, covered, left = topics_to_search([optional, essential], {}, 1, _runner_v(2))
    assert search == [essential] and covered == [] and left == [optional]
    many = [_asked(n, f"Essential matter number {n} in Uganda", f"essential matter {n}", essential=True) for n in range(1, 6)]
    search, _, left = topics_to_search(many, {}, 2, _runner_v(2))
    assert len(search) == 4 and left == many[4:]  # every essential need up to twice the allowance; the rest reported, never forgotten
    with pytest.raises(PermanentStageError) as stopped:
        essential_answered(_runner_v(2), [], [], [], left)
    assert stopped.value.code == "EVIDENCE_MISSING"
    assert topics_to_search([optional, essential], {}, 1, _runner_v(1)) == ([optional], [], [])  # workflow 1 as it was priced


def test_saved_evidence_about_something_else_does_not_cover_a_need():
    """Finding 5: any usable evidence id counted as coverage, so a general vaccination study could stand in for an
    essential question about local malaria-vaccine uptake and its search was skipped."""
    from app.analysis import research
    from app.proposals.pipeline import topics_to_search

    library = {"Egeneral": _saved("Egeneral", "Childhood vaccination coverage rose across Europe after school campaigns.", scope="Europe", title="Vaccination in Europe"),
               "Elocal": _saved("Elocal", "Uptake of the malaria vaccine among caregivers in Mukono was lower where clinics were far.", scope="Mukono, Uganda")}
    need = "Uptake of the malaria vaccine among caregivers in Mukono"
    query = "malaria vaccine uptake caregivers Mukono"
    wrong = _asked(1, need, query, essential=True, covered=["Egeneral"])
    right = _asked(2, need, query, essential=True, covered=["Elocal"])
    unknown = _asked(3, need, query, covered=["Enot-there"])
    search, covered, _ = topics_to_search([wrong, right, unknown], library, 5, _runner_v(2))
    assert covered == [right] and search == [wrong, unknown]
    assert not research.about(need, query, "Childhood vaccination coverage rose across Europe.") and research.about(need, query, library["Elocal"].statement)


def test_a_need_whose_found_quotation_fails_the_support_check_is_asked_on_the_other_route(v2):
    """Finding 7: the web was skipped because the index's quotation was found, then the support check rejected the
    finding, and the essential need failed without the route it was allowed."""
    from app.runtime import get_runtime

    def verify(payload):  # the index's finding does not support its claim; the web's does
        return {"results": [{"id": c["id"], "support": "SUPPORTED" if c["sources"][0]["access"] != "ABSTRACT" else "CONTRADICTED", "note": ""} for c in payload["claims"]]}

    v2.models.overrides["w_verify"] = verify
    _, job, stored = _plan(v2, [_need(1, "STUDY", essential=True), _need(2, "STUDY")])
    assert job["status"] == "COMPLETED", job.get("failure")
    first, second = get_runtime().store.get(job["id"]).topics[:2]
    assert (first.index, first.web, first.reason) == ("FOUND", "FOUND", "NOT_SUPPORTED")  # the essential need: the web, once
    assert (second.index, second.web) == ("FOUND", "NOT_USED")  # an optional need is not sent to the web for it
    assert v2.models.tasks.count("w_search") == 1 and v2.models.tasks.count("w_verify") == 2


def test_a_method_is_looked_for_in_any_year_and_a_study_s_second_search_has_no_date_floor(v2, monkeypatch):
    """Finding 8: the index was always asked for the last fifteen years, and workflow 2 no longer sends an empty
    search to the web, where foundational sources used to come from."""
    from app.analysis import fetch
    from tests.fake_models import WORK

    asked = []

    def find(query, from_year, limit):
        asked.append((query, from_year))
        return ("FOUND", [WORK]) if from_year is None else ("NO_RESULTS", [])

    monkeypatch.setattr(fetch, "openalex_find", find)
    _, job, _ = _plan(v2, [_need(1, "METHOD"), _need(2, "STUDY")])
    assert job["status"] == "COMPLETED", job.get("failure")
    by_query = {}
    for query, year in asked:
        by_query.setdefault(query, []).append(year)
    assert by_query["maternal health topic 1"] == [None]  # a method's own source, however old
    window, anytime = by_query["maternal health topic 2"]
    assert window is not None and anytime is None  # a study: recent first, then any year


# --- Codex's audit of fd74ff3: request size, waiting, reporting --------------------------------------------------------


def test_every_request_to_the_editor_fits_the_bound_and_each_paragraph_is_offered_once(v2):
    """Finding 4: the document was measured before the text the editor may correct, the evidence and the findings
    were added, and a section split between parts was sent whole with each."""
    from app.ai import orchestration
    from tests import fake_works

    def long_draft(payload):
        answer = fake_works.draft(payload)
        for section in answer["sections"]:
            body = section["paragraphs"]
            section["paragraphs"] = [f"{p} (paragraph {n} of this section)" for n in range(1, 13) for p in body]
        return answer

    sizes, offered = [], []

    def edit(payload):
        sizes.append(orchestration.payload_words(payload))
        offered.extend((e["key"], p["id"]) for e in payload["editable"] for p in e["paragraphs"])
        return _verdicts(payload)

    v2.models.overrides["w_draft"] = long_draft
    v2.models.overrides["w_compress"] = lambda payload: {"sections": [{"key": s["key"], "paragraphs": s["text"], "table": s["table"]} for s in payload["sections"]]}
    _, job, _ = _draft(v2, edit=edit)
    assert len(sizes) > 1 and max(sizes) <= orchestration.FINAL_PART_WORDS, sizes  # reviewed in parts, each within the bound
    assert len(offered) == len(set(offered))  # no paragraph was sent with two parts
    assert job["status"] in ("COMPLETED", "FAILED")  # over its word limit it may be refused; never by an oversized request
    with pytest.raises(PermanentStageError) as refused:  # and the runner itself refuses one, for the editor as for the reviewer
        from app.runtime import get_runtime
        from app.works.ai import EDIT_SCHEMA, Edited, WorkRunner

        engine = get_runtime().store.get(job["id"]).quote.engine
        WorkRunner(get_runtime().settings, v2.models, lambda call: None, lambda: 0.0, 99.0, engine=engine)._call(
            "w_edit", {"document": "word " * (orchestration.FINAL_PART_WORDS + 10)}, EDIT_SCHEMA, Edited)
    assert refused.value.code == "REVIEW_REQUEST_TOO_LARGE"


def test_waiting_is_bounded_by_the_clock_from_the_first_pause(v2, monkeypatch):
    """Finding 9: the allowance added up planned delays, so a slow failed request or a late delivery did not count."""
    from datetime import timedelta

    from app.core.errors import RetryableStageError
    from app.jobs.models import utcnow
    from app.runtime import get_runtime

    rt = get_runtime()
    rt.settings.review_wait_sec = 1
    rt.settings.capacity_wait_limit_sec = 3600
    tries = []

    def edit(payload):
        tries.append(1)
        if len(tries) == 1:
            raise RetryableStageError("PROVIDER_UNAVAILABLE", "Unavailable.", "test: the premium model is down")
        return _verdicts(payload)

    original = rt.queue.enqueue

    def late(job_id, task_name, delay_sec=0):
        if "-w" in task_name:  # the delivery after a wait comes two hours late (the planned delay was a second)
            def back(j):
                j.waiting_since = utcnow() - timedelta(hours=2)
                return j
            rt.store.update(job_id, back)
        return original(job_id, task_name, delay_sec)

    monkeypatch.setattr(rt.queue, "enqueue", late)
    _, job, document = _draft(v2, edit=edit)
    stored = rt.store.get(job["id"])
    assert job["status"] == "FAILED" and job["failure"]["code"] == "CAPACITY_BUSY" and document is None
    assert stored.capacity_waits == 2 and stored.waited_sec > 3600  # stopped at its next delivery, before another paid request
    assert v2.models.tasks.count("w_edit") == 1


def test_the_speed_report_counts_a_work_that_stopped_before_its_draft_and_ages_jobs_to_the_report_s_time():
    """Finding 10: a failed plan with no draft vanished from the delivery figures, and an unfinished job's age was
    measured to the newest job, so the newest one was always zero minutes old."""
    import sys
    from datetime import UTC, datetime
    from pathlib import Path

    sys.path.insert(0, str(Path(__file__).parents[1]))
    from speed_report import report

    def step(work, kind, status, queued, done=None):
        return {"workId": work, "selection": {"work": kind}, "status": status, "queuedAt": queued, "completedAt": done, "createdAt": queued, "quote": {"engine": {"workflow": 2}}}

    jobs = [step("a", "PLAN", "FAILED", "2026-10-09T10:00:00+00:00"),
            step("b", "PLAN", "COMPLETED", "2026-10-09T10:00:00+00:00", "2026-10-09T10:03:00+00:00"),
            step("b", "DRAFT", "COMPLETED", "2026-10-09T10:03:00+00:00", "2026-10-09T10:10:00+00:00"),
            step("c", "PLAN", "COMPLETED", "2026-10-09T10:00:00+00:00", "2026-10-09T10:02:00+00:00"),
            step("d", "PLAN", "PROCESSING", "2026-10-09T10:30:00+00:00")]
    out = report(jobs, now=datetime(2026, 10, 9, 11, 0, tzinfo=UTC))
    assert out["documents"]["workflow 2"]["works"] == 2 and out["documents"]["workflow 2"]["delivered"] == 1  # a (stopped) and b; c has simply no draft yet
    assert out["unfinished"]["PLAN · workflow 2"] == {"jobs": 1, "oldestMinutes": 30.0}


# --- The simplified current path: workflow 2's research and plan, the earlier review, one repair (owner 2026-10-10) ----


@pytest.fixture
def simple(tmp_path, monkeypatch):
    from tests.conftest import _client

    monkeypatch.setenv("WORKFLOW", "2")
    monkeypatch.setenv("FINAL_EDITOR", "false")
    monkeypatch.setenv("WORKS_ENABLED", '["CONCEPT_NOTE","COURSEWORK","FUNDING_PROPOSAL"]')
    yield from _client(tmp_path, monkeypatch, "cost")


def _unanswered(payload, answered):
    return {"rules": [{"rule": r["rule"], "status": "PASS", "note": "Met.", "where": ""} for r in payload["rules"]],
            "coverage": [{"id": c["id"], "answered": answered, "where": ""} for c in payload["coverage"]],
            "priorities": [{"priority": p, "addressed": True, "where": ""} for p in payload["priorities"]]}


def test_the_simplified_path_researches_and_plans_the_new_way_and_keeps_the_earlier_final_review(simple):
    from app.runtime import get_runtime

    _, job, document = _draft(simple)
    assert job["status"] == "COMPLETED", job.get("failure")
    tasks = simple.models.tasks
    assert "w_final" in tasks and not {"w_edit", "w_flag", "w_resolve", "w_plan_review"} & set(tasks)  # the plan by code checks; the review as before
    engine = get_runtime().store.get(job["id"]).quote.engine
    assert (engine.workflow, engine.final_editor, engine.prompts["w_needs"]) == (2, False, "w-needs-v2") and document.approval is None
    assert get_runtime().store.get(job["id"]).topics  # the routes of workflow 2's research are recorded


def test_the_simplified_path_repairs_a_coursework_draft_once_and_no_more(simple):
    """Trial of 9 October: the drafts that went round the repair loop twice took 17 and 18 minutes."""
    repairs, reviews = [], []

    def final(payload):
        reviews.append(1)
        return _unanswered(payload, answered=False)  # every review objects: only the limit can end it

    def repair(payload):
        from tests import fake_works

        repairs.append(1)
        return fake_works.repair(payload)

    _, job, _ = _draft(simple, **{})
    assert job["status"] == "COMPLETED", job.get("failure")
    simple.models.overrides.update(w_final=final, w_repair=repair)
    from tests.test_works import _coursework, _run, _work

    work = _coursework(simple)
    _run(simple, work["id"], "PLAN")
    work = _work(simple, work["id"])
    simple.post(f"/api/works/{work['id']}/plan/approve", headers=STUDENT, json={"baseVersion": work["planVersion"]})
    _run(simple, work["id"], "DRAFT")
    assert len(repairs) == 1 and len(reviews) == 2  # reviewed, repaired once, reviewed again: then it stops


def test_the_editor_switch_is_frozen_with_the_workflow(monkeypatch):
    from app.ai.orchestration import current_engine

    def engine(**env):
        for name in ("WORKFLOW", "FINAL_EDITOR"):
            monkeypatch.delenv(name, raising=False)
        for name, value in env.items():
            monkeypatch.setenv(name, value)
        return current_engine(real_settings())

    from app.jobs.models import Engine

    assert (engine().workflow, engine().editor) == (1, False)
    assert (engine(WORKFLOW="2").workflow, engine(WORKFLOW="2").editor) == (2, True)
    assert (engine(WORKFLOW="2", FINAL_EDITOR="false").workflow, engine(WORKFLOW="2", FINAL_EDITOR="false").editor) == (2, False)
    assert engine(FINAL_EDITOR="true").editor is False  # the editor belongs to workflow 2: on workflow 1 nothing changes
    # Codex's audit of cbb99cb, finding 6: an engine frozen on workflow 2 before the switch existed carries no value
    # for it, and must keep the editor it was priced with when its job is read back and resumed.
    saved = engine(WORKFLOW="2").model_dump(by_alias=True, mode="json")
    del saved["finalEditor"]
    assert Engine.model_validate(saved).editor is True
    assert Engine.model_validate(engine(WORKFLOW="2", FINAL_EDITOR="false").model_dump(by_alias=True, mode="json")).editor is False


# --- From the trials of 10 October -------------------------------------------------------------------------------------


def test_the_editor_can_write_into_a_section_code_left_empty():
    """A chapter failed as "The section is empty": code had withheld its one paragraph, the editor wrote the section
    twice, and both corrections were refused for naming no existing paragraph."""
    from app.proposals.ai import Correction, SectionText, Table
    from app.proposals.pipeline import apply_corrections

    current = {"significance": SectionText(key="significance", paragraphs=[], table=Table(caption="", rows=[])),
               "problem": SectionText(key="problem", paragraphs=["The problem is stated here."], table=Table(caption="", rows=[]))}
    corrections = [Correction(section="significance", paragraph="p1", action="REPLACE", text="The findings will inform district planners.", kind="CONCLUSION", reason="Empty."),
                   Correction(section="significance", paragraph="", action="INSERT_AFTER", text="They will also guide later studies.", kind="CONCLUSION", reason="Empty."),
                   Correction(section="problem", paragraph="p9", action="REPLACE", text="A paragraph that does not exist.", kind="CLAIM", reason="Wrong place.")]
    applied, refused = apply_corrections(current, corrections, ["significance", "problem"])
    assert current["significance"].paragraphs == ["The findings will inform district planners.", "They will also guide later studies."]
    assert [a["paragraph"] for a in applied] == ["p1", "p2"] and list(refused) == ["problem"]  # a section with text still needs a real paragraph named


def test_the_first_draft_is_asked_for_nearly_its_whole_target_on_workflow_2(simple):
    """Told the plan's own range (70% of the target), the writer wrote body sections at 65 to 90% of it."""
    def asked_minimums(c):
        seen = []
        c.models.overrides["w_draft"] = lambda payload: (seen.extend((s["words"], s["minWords"]) for s in payload["sections"]), __import__("tests.fake_works", fromlist=["draft"]).draft(payload))[1]
        _draft(c)
        return seen

    tight = asked_minimums(simple)
    assert tight and all(minimum >= round(words * 0.92) for words, minimum in tight)


def test_the_first_draft_is_given_a_ceiling_too_so_it_does_not_overshoot_a_hard_limit(simple):
    """Round 3 of the test loop: with the floor alone the drafts went over their limit and were compressed twice."""
    seen = []

    def draft(payload):
        from tests import fake_works

        seen.extend((s["words"], s["minWords"], s["maxWords"]) for s in payload["sections"])
        return fake_works.draft(payload)

    simple.models.overrides["w_draft"] = draft
    _draft(simple)
    assert seen and all(low <= words <= high <= round(words * 1.05) for words, low, high in seen)


def test_on_the_simplified_path_a_short_draft_is_lengthened_first_and_the_one_repair_is_kept(simple):
    """Round 1 of the test loop: a 1,182-word draft for 1,500 was refused because the one repair had gone on a
    reviewer's objection and nothing was left to lengthen it."""
    from tests import fake_works

    order = []

    def short_draft(payload):
        answer = fake_works.draft(payload)
        for section in answer["sections"]:
            words = " ".join(section["paragraphs"]).split()
            section["paragraphs"] = [" ".join(words[: int(len(words) * 0.7)]).rstrip(".,;") + "."]
        return answer

    def repair(payload):
        lengthen = any("well short" in i for s in payload["sections"] for i in s["issues"])
        order.append("lengthen" if lengthen else "repair")
        out = []
        for s in payload["sections"]:
            full = fake_works._text({"key": s["key"], "heading": s["heading"], "words": s["words"], "evidence": s["evidence"]}, [])
            out.append({"key": s["key"], "paragraphs": full if lengthen else [*s["text"], "This part is now addressed."], "table": {"caption": "", "rows": []}})
        return {"sections": out}

    reviews = []

    def final(payload):
        reviews.append(1)
        return _unanswered(payload, answered=len(reviews) > 1)  # objects once: the one repair must still be there for it

    simple.models.overrides.update(w_draft=short_draft, w_repair=repair, w_final=final)
    _, job, document = _draft(simple)
    assert job["status"] == "COMPLETED", job.get("failure")
    assert order == ["lengthen", "repair"] and 1275 <= document.words <= 1500


# --- From round 2 of the test loop (10 October) -------------------------------------------------------------------------


# --- From round 3 of the test loop (10 October) -------------------------------------------------------------------------


def test_a_plan_may_state_a_descriptive_question_beside_its_hypotheses():
    """A proposal stopped at its plan: one descriptive objective stated as a question and two null hypotheses with
    their alternatives, refused three times by "Each null hypothesis needs its alternative hypothesis"."""
    from app.proposals import rulebook
    from app.proposals.pipeline import plan_statements
    from tests.test_proposals import plan

    mixed = plan(questionsKind="HYPOTHESES",
                 specificObjectives=["To describe the practices adopted.", "To assess the effect of income on adoption.", "To assess the effect of extension visits on adoption."],
                 researchQuestions=["What climate-smart practices are adopted by smallholder farmers in Mbale District?",
                                    "H01: Household income does not significantly influence adoption.", "H02: Extension visits do not significantly influence adoption."],
                 alternativeHypotheses=["Ha1: Household income significantly influences adoption.", "Ha2: Extension visits significantly influence adoption."])
    assert not [p for p in rulebook.plan_problems("ucu-2018-v2", mixed, "MASTERS") if "hypothes" in p.lower()]
    lines = plan_statements(mixed, "questions", "1.4")
    assert lines[-6:] == ["1.4.2 Research Questions and Hypotheses", "i. What climate-smart practices are adopted by smallholder farmers in Mbale District?",
                          "H01: Household income does not significantly influence adoption.", "HA1: Household income significantly influences adoption.",
                          "H02: Extension visits do not significantly influence adoption.", "HA2: Extension visits significantly influence adoption."]  # labels printed once
    unpaired = mixed.model_copy(update={"alternative_hypotheses": ["Household income significantly influences adoption."]})
    assert any("2 null hypotheses, 1 alternatives" in p for p in rulebook.plan_problems("ucu-2018-v2", unpaired, "MASTERS"))  # a real mismatch still blocks
    only_questions = mixed.model_copy(update={"research_questions": ["What is adopted?", "Why is it adopted?", "Who adopts it?"], "alternative_hypotheses": []})
    assert any("at least one null hypothesis" in p for p in rulebook.plan_problems("ucu-2018-v2", only_questions, "MASTERS"))
    plain = plan_statements(plan(questionsKind="HYPOTHESES", researchQuestions=["Income does not influence uptake."], alternativeHypotheses=["Income influences uptake."],
                                 specificObjectives=["To assess income."]), "questions", "1.4")
    assert plain[-3:] == ["1.4.2 Research Hypotheses", "H01: Income does not influence uptake.", "HA1: Income influences uptake."]  # an unmixed plan prints as before


# --- From round 4 of the test loop (10 October) -------------------------------------------------------------------------


# --- Codex's audit of cbb99cb (10 October) ------------------------------------------------------------------------------


def test_a_reviewer_is_shown_what_paperaid_places_apart_and_its_rejection_still_stands(client):
    """Finding 1. A keyword filter dropped any objection that said "remove" near "objectives", including a genuine
    one. Nothing a reviewer says is set aside now; it is told, by a new prompt version, what PaperAid places."""
    from app.runtime import get_runtime

    seen = {}

    def review(payload):
        out = []
        for s in payload["sections"]:
            if "placedByPaperAid" in s:
                seen.setdefault(s["key"], s)
            reject = s["key"] == "objectives"
            out.append({"key": s["key"], "grade": "REPAIR" if reject else "PASS", "note": "",
                        "issues": ["Remove the statement that this study will establish causation from associations; the objectives can remain."] if reject else []})
        return {"results": out}

    _, job, chapter = _chapter_one(client, p_review=review)
    assert job["status"] == "FAILED" and job["failure"]["code"] == "DOCUMENT_NOT_APPROVED" and chapter is None  # the rejection stands
    assert "p_fix" in client.models.tasks  # and was sent to the writer to put right
    objectives = seen["objectives"]
    assert objectives["placedByPaperAid"][0].endswith("General Objective") and not any("General Objective" in p for p in objectives["text"])
    assert get_runtime().store.get(job["id"]).quote.engine.prompts["p_review"] == "p-review-v5"


def test_a_job_priced_on_the_earlier_review_prompt_is_asked_as_it_was():
    from types import SimpleNamespace

    from app.proposals import pipeline
    from app.proposals.ai import SectionText, Table
    from tests.test_proposals import plan

    inp = SimpleNamespace(plan=plan(), rulebook="ucu-2018-v2", chapter=1, inputs=SimpleNamespace(level="MASTERS"))
    number = next(s.number for s in pipeline.rulebook.sections("ucu-2018-v2", 1, "MASTERS", plan()) if s.key == "objectives")
    statements = pipeline.plan_statements(plan(), "objectives", number)
    text = SectionText(key="objectives", paragraphs=["The study pursues the following objectives.", *statements], table=Table(caption="", rows=[]))
    items = {"objectives": {"key": "objectives", "number": number, "heading": "Objectives of the Study"}}
    before = pipeline.review_item(inp, items, "objectives", text, [], apart=False)
    after = pipeline.review_item(inp, items, "objectives", text, [], apart=True)
    assert before["text"] == text.paragraphs and "placedByPaperAid" not in before
    assert after["text"] == ["The study pursues the following objectives."] and after["placedByPaperAid"] == statements


def test_an_invented_figure_is_withheld_whatever_a_worked_table_shows_and_whatever_verb_the_question_uses():
    """Finding 2. A blanket exception let any number in a model-written table stand in factual prose once the question
    used a verb such as "find": "District officials recorded 900 cases." was kept. The exception is withdrawn."""
    from types import SimpleNamespace

    from app.proposals.ai import SectionText, Table
    from app.works import pipeline

    question = "Find factors associated with uptake among children aged 6-24 months."
    inp = SimpleNamespace(spec=SimpleNamespace(kind="COURSEWORK", coverage=[SimpleNamespace(text=question)]), inputs=SimpleNamespace(title="Coursework", description=question))
    table = Table(caption="Worked example", rows=[["Step", "Value"], ["Cases", "900"]], illustrative=True)
    same = SectionText(key="theme1", paragraphs=["District officials recorded 900 cases.", "Suppose a district records 900 cases in a year."], table=table)
    other = SectionText(key="conclusion", paragraphs=["District officials recorded 900 cases.", "Uptake varies between districts."], table=Table(caption="", rows=[]))
    assert pipeline._strip(inp, same, {}, "").paragraphs == ["Suppose a district records 900 cases in a year."]  # only a sentence that opens as a hypothetical
    assert pipeline._strip(inp, other, {}, "").paragraphs == ["Uptake varies between districts."]  # and never in another section
    assert not hasattr(pipeline, "_with_worked") and not hasattr(pipeline, "_calculation")


def test_missing_essential_evidence_stops_the_step_whatever_its_kind_and_on_either_path():
    """Finding 3. An exception let a missing essential theory or method through, and on the simplified path every
    missing essential topic. A compulsory method (a required instrument and its validation source) is evidence too."""
    from types import SimpleNamespace

    from app.proposals.pipeline import essential_answered

    method = _asked(1, "The validated instrument the assignment requires and its validation source", "validated instrument validation study", essential=True, category="METHOD")
    study = _asked(2, "Uptake among caregivers in Mukono", "uptake caregivers Mukono", essential=True)
    policy = _asked(3, "The national policy the question names", "national immunisation policy Uganda", essential=True, category="POLICY")
    for final_editor in (True, False, None):
        runner = SimpleNamespace(_engine=SimpleNamespace(workflow=2, final_editor=final_editor), budget_reached=False)
        for need in (method, study, policy):
            with pytest.raises(PermanentStageError) as stopped:
                essential_answered(runner, [need], [], [])
            assert stopped.value.code == "EVIDENCE_MISSING"
    assert essential_answered(_runner_v(1), [method, study, policy], [], []) is None  # workflow 1 as priced


def test_findings_on_one_part_of_the_question_never_stand_in_for_the_draft_s_own_research(simple):
    """Finding 4. A count of usable findings skipped the draft's research: six findings from one paper on one topic
    did. The draft plans its needs and only a need saved evidence can be about is left out."""
    from app.works import pipeline

    _, job, _ = _draft(simple)
    assert job["status"] == "COMPLETED", job.get("failure")
    assert simple.models.tasks.count("w_needs") == 2 and not hasattr(pipeline, "PLAN_SOURCES_ENOUGH")  # the plan's needs, then the draft's


def test_on_workflow_2_the_premium_model_approves_the_wording_that_is_delivered(simple):
    """Finding 5. After a repair the review went to the standard model, and the first review could fall back to it:
    the premium model need never have seen the final text. The recorded model is checked, not the task's name."""
    from app.runtime import get_runtime

    reviews = []

    def final(payload):
        reviews.append(1)
        return _unanswered(payload, answered=len(reviews) > 1)  # objects once, so the delivered text is the repaired one

    simple.models.overrides["w_final"] = final
    _, job, document = _draft(simple)
    assert job["status"] == "COMPLETED", job.get("failure")
    stored = get_runtime().store.get(job["id"])
    approvals = [c for c in stored.model_calls if c.task == "w_final"]
    assert len(approvals) == 2 and all(c.model == "gemini-3.1-pro-preview" for c in approvals), [(c.task, c.model) for c in approvals]
    engine = stored.quote.engine
    assert engine.vertex_signoff == ["vertex:gemini-3.1-pro-preview"] and engine.vertex_routes["w_final"] == ["vertex:gemini-3.1-pro-preview"]  # no cheaper sign-off, no stand-in
    assert engine.vertex_routes["p_review"] == ["vertex:gemini-3.1-pro-preview"] and engine.vertex_routes["w_plan_review"] == ["vertex:gemini-3.1-pro-preview"]


def test_workflow_1_keeps_the_sign_off_it_was_priced_with(client):
    from app.ai.orchestration import current_engine
    from app.runtime import get_runtime

    engine = current_engine(get_runtime().settings)
    assert engine.workflow == 1 and engine.vertex_signoff[0] != "vertex:gemini-3.1-pro-preview" and len(engine.vertex_routes["w_final"]) == 2  # as priced: unchanged


def test_on_workflow_2_an_approval_review_waits_for_the_premium_model(simple):
    from app.core.errors import RetryableStageError
    from app.runtime import get_runtime

    get_runtime().settings.review_wait_sec = 1
    calls = []

    def final(payload):
        calls.append(1)
        if len(calls) <= 2:
            raise RetryableStageError("VERTEX_MODEL_UNAVAILABLE", "Unavailable.", "test: the premium model is down")
        return _unanswered(payload, answered=True)

    simple.models.overrides["w_final"] = final
    _, job, _ = _draft(simple)
    assert job["status"] == "COMPLETED", job.get("failure")
    stored = get_runtime().store.get(job["id"])
    assert stored.capacity_waits == 2 and stored.attempts == 0 and [t.code for t in stored.timings if t.outcome == "WAITING"] == ["REVIEW", "REVIEW"]
    assert all(c.model == "gemini-3.1-pro-preview" for c in stored.model_calls if c.task == "w_final" and not c.error_code)  # nothing approved in its place


def test_a_question_among_hypotheses_is_recognised_without_its_ascii_question_mark():
    """Finding 10. Only a final ASCII "?" counted, so the same descriptive question without it, inside typographic
    quotation marks or ending in a full-width mark was read as a third null hypothesis."""
    from app.proposals import rulebook

    for question in ("What practices are adopted by smallholder farmers in Mbale District?",
                     "What practices are adopted by smallholder farmers in Mbale District",
                     "\u201cWhat practices are adopted by smallholder farmers in Mbale District?\u201d",
                     "What practices are adopted by smallholder farmers in Mbale District\uff1f",
                     "To what extent are climate-smart practices adopted in Mbale District?.",
                     "The practices adopted by smallholder farmers in Mbale District: which are they?"):
        assert rulebook.is_question(question), question
    for hypothesis in ("H01: Household income does not significantly influence adoption.", "Household income does not significantly influence adoption.",
                       "Ha2: Extension visits significantly influence adoption.", "There is no significant relationship between income and adoption.",
                       "H01: Is there no relationship between income and adoption?"):  # a labelled hypothesis is one, however it is worded
        assert not rulebook.is_question(hypothesis), hypothesis


def test_your_work_is_one_light_line_each_and_reads_no_document(v2, monkeypatch):
    """Finding 9. Every page of the app polled the three full lists, each of which builds whole views (specification,
    document and evidence files). The left panel's list comes from the stored records alone."""
    from app.runtime import get_runtime

    work, job, _ = _draft(v2)
    assert job["status"] == "COMPLETED", job.get("failure")
    files = get_runtime().files
    read = []
    monkeypatch.setattr(type(files), "get", lambda self, path, _get=type(files).get: (read.append(path), _get(self, path))[1])
    listed = v2.get("/api/me/work", headers=STUDENT)
    assert listed.status_code == 200 and read == []  # nothing read from storage
    full = v2.get(f"/api/works/{work['id']}", headers=STUDENT).json()
    assert full["status"] in ("READY", "READY_WITH_WARNINGS") and full["activeJob"] is None
    # finished: the record's claim on its step no longer counts, exactly as in the full view
    assert [(i["id"], i["section"], i["kind"], i["state"]) for i in listed.json()] == [(work["id"], "Coursework", "COURSEWORK", full["status"])]
    assert set(listed.json()[0]) == {"id", "section", "kind", "title", "updatedAt", "state", "chapters", "analyses"}  # where it is, never what it says
    assert v2.get("/api/me/work").status_code in (401, 403)  # only the signed-in person's own


# --- Codex's re-audit of d4d2bc7 (10 October) ---------------------------------------------------------------------------


def _states(client):
    return {i["id"]: i["state"] for i in client.get("/api/me/work", headers=STUDENT).json()}


def test_work_whose_latest_step_failed_or_was_stopped_is_never_listed_as_not_started(simple):
    """The light list read only the record's own failure note, so a failed plan or draft with no document showed as
    "Not started" while the workspace said why it had failed."""
    from app.core.errors import PermanentStageError as Stop
    from app.jobs.models import JobStatus
    from app.runtime import get_runtime
    from tests.test_works import _coursework, _run, _work

    def refuse(payload):
        raise Stop("PLAN_FAILED", "PaperAid could not make a plan for this. Nothing was charged.", "test")

    rt = get_runtime()
    fresh = _coursework(simple)
    assert _states(simple)[fresh["id"]] == "NOT_STARTED"  # nothing has run yet: that is what "not started" means

    simple.models.overrides["w_plan"] = refuse
    failed_plan = _coursework(simple)
    _, job = _run(simple, failed_plan["id"], "PLAN")
    assert job["status"] == "FAILED" and _states(simple)[failed_plan["id"]] == "FAILED"

    del simple.models.overrides["w_plan"]
    failed_draft = _coursework(simple)
    _run(simple, failed_draft["id"], "PLAN")
    work = _work(simple, failed_draft["id"])
    simple.post(f"/api/works/{work['id']}/plan/approve", headers=STUDENT, json={"baseVersion": work["planVersion"]})
    simple.models.overrides["w_draft"] = refuse
    _, job = _run(simple, failed_draft["id"], "DRAFT")
    assert job["status"] == "FAILED" and _states(simple)[failed_draft["id"]] == "FAILED"  # a plan exists, the draft failed: still no document

    def stopped(j):
        j.status = JobStatus.CANCELLED
        return j

    rt.store.update(job["id"], stopped)
    assert _states(simple)[failed_draft["id"]] == "FAILED"  # stopped by the person: also not "not started"
    assert _states(simple)[fresh["id"]] == "NOT_STARTED"


def test_a_proposal_or_a_data_report_that_failed_is_listed_as_failed(simple, monkeypatch):
    from datetime import timedelta
    from types import SimpleNamespace

    from app.datalab.models import DataProject
    from app.jobs.models import JobStatus, utcnow
    from app.runtime import get_runtime
    from tests.test_proposals import _create, _run

    simple.models.overrides["p_plan"] = lambda payload: (_ for _ in ()).throw(PermanentStageError("PLAN_FAILED", "No plan could be made. Nothing was charged.", "test"))
    pid = _create(simple)["id"]
    job = _run(simple, pid, "PLAN")
    assert job["status"] == "FAILED" and _states(simple)[pid] == "FAILED"

    rt = get_runtime()
    uid = rt.store.get_project(pid).owner_uid
    data = DataProject(id="dl_test", owner_uid=uid, owner_email="demo@paperaid.app", title="Survey", expires_at=utcnow() + timedelta(days=30), jobs=["job_report"])
    monkeypatch.setattr(type(rt.store), "list_datalab", lambda self, owner: [data])
    real = type(rt.store).get
    monkeypatch.setattr(type(rt.store), "get", lambda self, job_id: SimpleNamespace(status=JobStatus.FAILED) if job_id == "job_report" else real(self, job_id))
    assert _states(simple)["dl_test"] == "FAILED"  # its report could not be finished: never "No data yet"


def test_an_unsound_placed_statement_is_told_to_the_student_as_a_note_on_their_plan(client):
    """Re-audit, finding 3: p-review-v4 called the placed statements "correct as they stand". They are the student's
    approved plan and cannot be changed in the chapter, but the reviewer can still say one is unsound."""
    from pathlib import Path

    note = "Your plan: the first objective promises a causal conclusion that a cross-sectional design cannot give; word it as an association."

    def review(payload):
        return {"results": [{"key": s["key"], "grade": "PASS_WITH_WARNINGS" if s["key"] == "objectives" else "PASS", "issues": [],
                             "note": note if s["key"] == "objectives" else ""} for s in payload["sections"]]}

    _, job, chapter = _chapter_one(client, p_review=review)
    assert job["status"] == "COMPLETED", job.get("failure")
    assert any(note in w for w in chapter.warnings)  # delivered with the chapter, never dropped
    prompt = (Path(__file__).parents[1] / "app" / "ai" / "prompts" / "p-review-v5.md").read_text(encoding="utf-8")
    assert "not beyond question" in prompt and "correct as it stands" not in prompt and "Your plan:" in prompt

