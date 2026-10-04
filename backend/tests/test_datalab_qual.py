"""Qualitative Data Lab (owner decision 2026-10-04): transcripts kept only after names are replaced; codes
with quotes checked word for word; themes citing quotes by reference; the final reviewer approves the
assembled report; Word report and Excel codebook."""

import io
import json
import time

import pytest
from docx import Document

from app.datalab import qual
from tests.test_api import STUDENT

H = STUDENT
PRICES = {"DL_SMALL": 3, "DL_STANDARD": 6, "DL_LARGE": 10, "QL_SMALL": 4, "QL_STANDARD": 7, "QL_LARGE": 12}
T1 = ("Agnes Akello lives far from the health centre. The walk to the clinic takes most of the morning for women in our village. "
      "When labour starts at night there is no transport, so many mothers stay at home. You can call me on 0772123456 if you need more.")
T2 = ("The boda boda riders charge a lot when it rains and the road floods. Mothers wait for the river to go down before they travel. "
      "Agnes said the health workers are kind but the distance is the problem for everyone here.")


@pytest.fixture
def lab(tmp_path, monkeypatch):
    monkeypatch.setenv("DATALAB_ENABLED", "true")
    from app.core.config import Settings

    monkeypatch.setenv("FIXED_TOKENS", json.dumps({**Settings(_env_file=None).fixed_tokens, **PRICES}))
    from tests.conftest import _client

    yield from _client(tmp_path, monkeypatch, "fixed")


def _project(client):
    return client.post("/api/datalab", headers=H, json={"title": "Reaching care in Kamuli", "purpose": "How do mothers reach a health facility to give birth?",
                                                        "kind": "QUAL"}).json()


def _paste(client, pid, label, text, **extra):
    return client.post(f"/api/datalab/{pid}/documents/text", headers=H, json={"label": label, "text": text, "consent": True, **extra})


def test_names_and_contact_details_are_replaced_before_anything_is_stored():
    text, n = qual.pseudonymise(T1, [("Agnes Akello", "Participant A"), ("Agnes", "Participant A")])
    assert "Agnes" not in text and "Participant A lives far" in text and "0772123456" not in text and "[phone]" in text and n == 2
    assert qual.pseudonymise("Agnesia stayed home.", [("Agnes", "P")])[0] == "Agnesia stayed home."  # whole words only


def test_a_quote_is_kept_only_if_it_is_in_its_transcript_word_for_word():
    coded = qual.Coded.model_validate({"codes": [{"code": "Distance", "description": "d", "quotes": [
        {"document": "d1", "text": "The walk to the clinic takes most of the morning"},  # exact (whitespace aside)
        {"document": "d1", "text": "The walk to the clinic is very long and hard"},  # paraphrased
        {"document": "d2", "text": "The walk to the clinic takes most of the morning"},  # the wrong transcript
    ]}]})
    kept, dropped = qual.verified_quotes(coded, [{"id": "d1", "label": "A", "text": T1}, {"id": "d2", "label": "B", "text": T2}], {"d1": T1, "d2": T2})
    assert [q["text"] for q in kept[0]["quotes"]] == ["The walk to the clinic takes most of the morning"] and dropped == 2


def test_code_finds_counts_typed_quotes_and_unknown_references():
    codes = [{"code": "Distance", "description": "d", "quotes": [1, 2]}]
    quotes = {1: {"document": "d1", "text": "x y z w"}, 2: {"document": "d2", "text": "a b c d"}}
    draft = {"themes": [{"name": "Far", "definition": "d", "codes": ["Distance", "Cost"],
                         "paragraphs": ["Three participants said \"the walk to the clinic takes most of the morning for women\" and ⟦Q:9⟧ too, as ⟦Q:1⟧ shows."]}],
             "summary": ["Short."], "limitations": ["The transcripts come from a small group of participants in one district."]}
    found = " ".join(qual.problems(draft, codes, quotes))
    assert "don't exist: Cost" in found and "⟦Q:9⟧" in found and "at least two quotes" in found and "Write the summary" in found
    assert "quoted directly" in found


def _wait(client, pid, timeout=60):
    deadline = time.time() + timeout
    while time.time() < deadline:
        view = client.get(f"/api/datalab/{pid}", headers=H).json()
        if view["reports"] or view["reportFailure"]:
            return view
        time.sleep(0.2)
    raise AssertionError(view)


def test_a_qualitative_analysis_from_transcripts_to_report_and_codebook(lab):
    from openpyxl import load_workbook

    p = _project(lab)
    pid = p["id"]
    assert p["kind"] == "QUAL" and p["qualPriced"]
    refused = lab.post(f"/api/datalab/{pid}/documents/text", headers=H, json={"label": "Interview 1", "text": T1})
    assert refused.status_code == 400 and refused.json()["code"] == "CONSENT_REQUIRED"
    view = _paste(lab, pid, "Interview 1", T1, replace=[["Agnes Akello", "Participant A"], ["Agnes", "Participant A"]]).json()
    doc = Document()
    for line in T2.split(". "):
        doc.add_paragraph(line)
    buffer = io.BytesIO()
    doc.save(buffer)
    view = lab.post(f"/api/datalab/{pid}/documents", headers=H, files={"file": ("interview2.docx", buffer.getvalue(), "application/octet-stream")},
                    data={"label": "Interview 2", "consent": "true", "replace": json.dumps([["Agnes", "Participant A"]])}).json()
    assert [d["label"] for d in view["documents"]] == ["Interview 1", "Interview 2"] and view["documents"][0]["replaced"] == 2
    from app.runtime import get_runtime

    rt = get_runtime()
    stored = [rt.files.get(d.path).decode() for d in rt.store.get_datalab(pid).documents]
    assert all("Agnes" not in t and "0772123456" not in t for t in stored)  # the original never reached storage
    assert lab.post(f"/api/datalab/{pid}/dataset", headers=H, files={"file": ("d.csv", b"a,b\n1,2\n", "text/csv")}, data={"consent": "true"}).json()["code"] == "WRONG_KIND"
    job = lab.post(f"/api/datalab/{pid}/themes", headers=H)
    assert job.status_code == 200, job.json()
    view = _wait(lab, pid)
    assert view["reports"] and view["reports"][0]["kind"] == "THEMES", view["reportFailure"]
    sent = [r for t, r in zip(lab.models.tasks, lab.models.requests, strict=True) if t.startswith("q_")]
    assert sent and all("Agnes" not in r and "0772123456" not in r for r in sent)  # no model ever saw a name the researcher listed
    review = next(r for t, r in zip(lab.models.tasks, lab.models.requests, strict=True) if t == "q_review")
    assert "Appendix A. Codebook" in review and "Found in" in review  # the reviewer reads the assembled report
    report = lab.get(f"/api/datalab/{pid}/report", headers=H).json()
    headings = [s["heading"] for s in report["sections"]]
    assert headings[:5] == ["Summary", "Research question", "The data", "Method", "Themes"] and headings[-1] == "Limitations"
    theme = next(s for s in report["sections"] if s["heading"].startswith("1. "))
    assert theme["paragraphs"][1].startswith("Found in 2 of 2 transcripts") and "(Interview 1)" in " ".join(theme["paragraphs"])
    assert "⟦" not in json.dumps(report)  # every quote placed by code
    word = lab.get(f"/api/datalab/{pid}/report/export", headers=H)
    assert "Appendix A. Codebook" in "\n".join(x.text for x in Document(io.BytesIO(word.content)).paragraphs)
    book = load_workbook(io.BytesIO(lab.get(f"/api/datalab/{pid}/report/codebook", headers=H).content))
    assert book.sheetnames == ["Themes", "Codebook", "Quotes"]
    finished = rt.store.get(job.json()["id"])
    assert finished.status == "COMPLETED" and finished.billing.charged > 0 and finished.selection.datalab_band == "QL_SMALL"


def test_invented_quotes_never_reach_a_report_and_nothing_is_charged(lab):
    p = _project(lab)
    _paste(lab, p["id"], "Interview 1", T1)
    lab.models.overrides["q_code"] = lambda payload: {"codes": [{"code": "Made up", "description": "d",
                                                                 "quotes": [{"document": payload["documents"][0]["id"], "text": "Words nobody said in this interview at all"}]}]}
    job = lab.post(f"/api/datalab/{p['id']}/themes", headers=H).json()
    view = _wait(lab, p["id"])
    from app.runtime import get_runtime

    finished = get_runtime().store.get(job["id"])
    assert not view["reports"] and finished.failure.code == "NO_CODES" and finished.billing.charged == 0


def test_the_qualitative_analysis_waits_for_its_prices(tmp_path, monkeypatch):
    monkeypatch.setenv("DATALAB_ENABLED", "true")
    from tests.conftest import _client

    for client in _client(tmp_path, monkeypatch, "fixed"):
        p = _project(client)
        _paste(client, p["id"], "Interview 1", T1)
        refused = client.post(f"/api/datalab/{p['id']}/themes", headers=H)
        assert not p["qualPriced"] and refused.status_code == 400 and refused.json()["code"] == "REPORT_NOT_PRICED"
