"""Data Lab through the API (owner decision 2026-10-03): a dataset uploaded, profiled and cleaned with
confirmations, analysed by code, written up by the (stand-in) writer through number tokens and
approved by the final reviewer, then exported. Rows never reach a model or the database. Since
Codex's audit the data work runs in the worker (the page follows the project's `op`), results are
kept only on the data and settings they used, and a changed setting marks them out of date."""

import io
import json
import time

import pytest
from docx import Document

from tests.conftest import internal
from tests.test_api import OTHER, STUDENT

H = STUDENT
DL_PRICES = {"DL_SMALL": 3, "DL_STANDARD": 6, "DL_LARGE": 10}


@pytest.fixture
def lab(tmp_path, monkeypatch):
    monkeypatch.setenv("DATALAB_ENABLED", "true")
    from app.core.config import Settings

    monkeypatch.setenv("FIXED_TOKENS", json.dumps({**Settings(_env_file=None).fixed_tokens, **DL_PRICES}))
    from tests.conftest import _client

    yield from _client(tmp_path, monkeypatch, "fixed")


CSV = ("id,name,sex,age,score,district,passed\n"
       + "".join(f"{i:04d},Person {i},{'Male ' if i % 2 else 'Female'},{20 + i % 30},{50 + (i * 7) % 40},{['Gulu', 'Pader', 'Kitgum'][i % 3]},{'yes' if i % 3 else 'no'}\n"
                 for i in range(1, 121))
       + "0121,Person 121,male,N/A,70,Gulu,yes\n")


def _settle(client, pid, timeout=60):
    """The project once its data work has finished."""
    deadline = time.time() + timeout
    while time.time() < deadline:
        view = client.get(f"/api/datalab/{pid}", headers=H).json()
        if view["op"] is None or view["op"]["status"] in ("DONE", "FAILED"):
            return view
        time.sleep(0.1)
    raise AssertionError(view)


def _act(client, pid, url, **kwargs):
    """A data action: refused at once (the response), or handed to the worker (the project once done)."""
    r = client.post(url, headers=H, **kwargs)
    if r.status_code != 200:
        return r
    return _settle(client, pid)


def _project(client, data: str = CSV, name: str = "survey.csv"):
    p = client.post("/api/datalab", headers=H, json={"title": "Exam results", "purpose": "Do scores differ by sex?"}).json()
    view = _act(client, p["id"], f"/api/datalab/{p['id']}/dataset", files={"file": (name, data.encode(), "text/csv")}, data={"consent": "true"})
    assert isinstance(view, dict) and view["op"]["status"] == "DONE", view if isinstance(view, dict) else view.json()
    return view


def _analyse(client, pid, spec):
    """An analysis's result, or the refusal (a response) when it was refused at once."""
    done = _act(client, pid, f"/api/datalab/{pid}/analyses", json=spec)
    if not isinstance(done, dict):
        return done
    assert done["op"]["status"] == "DONE", done["op"]
    return client.get(f"/api/datalab/{pid}/analyses/{done['op']['result']}", headers=H).json()


def _decide(client, pid, step_id, accept):
    return _act(client, pid, f"/api/datalab/{pid}/cleaning/{step_id}", json={"accept": accept})


def _var(view, name):
    return next(v for v in view["variables"] if v["name"] == name)


def test_a_dataset_is_profiled_with_its_flags_and_the_original_kept(lab):
    view = _project(lab)
    assert view["rows"] == 121 and view["columns"] == 7 and view["source"]["name"] == "survey.csv"
    assert _var(view, "name")["excluded"] and "PERSONAL" in _var(view, "name")["flags"]  # left out by default
    assert _var(view, "id")["stored"] == "text" and "LEADING_ZEROS" in _var(view, "id")["flags"]  # 0001 stays a code
    assert _var(view, "age")["stored"] == "text"  # "N/A" keeps it text until the researcher confirms
    applied = [s for s in view["applied"] if s["automatic"]]
    assert applied and applied[0]["kind"] == "TRIM" and view["version"] == 2  # trimming is the only automatic change
    kinds = {(s["kind"], s["column"]) for s in view["pending"]}
    assert ("SET_MISSING", "age") in kinds and ("MERGE_LEVELS", "sex") in kinds
    from app.runtime import get_runtime

    rt = get_runtime()
    record = rt.store.get_datalab(view["id"])
    raw = record.model_dump_json()
    assert "Person 1" not in raw and "Gulu" not in raw  # no row or level reaches the database
    assert rt.files.get(record.source.path).decode() == CSV  # the upload exactly as received


def test_an_unreadable_file_fails_its_work_and_changes_nothing(lab):
    view = _project(lab)
    pid = view["id"]
    failed = _act(lab, pid, f"/api/datalab/{pid}/dataset", files={"file": ("bad.xlsx", b"not a workbook", "application/octet-stream")}, data={"consent": "true"})
    assert failed["op"]["status"] == "FAILED" and failed["op"]["code"] == "DATA_FORMAT"
    assert failed["version"] == view["version"] and failed["source"]["name"] == "survey.csv"  # the project as it was


def test_confirmations_make_versions_and_can_be_undone(lab):
    view = _project(lab)
    pid = view["id"]
    age = next(s for s in view["pending"] if s["kind"] == "SET_MISSING" and s["column"] == "age")
    view = _decide(lab, pid, age["id"], True)
    assert view["version"] == 3 and _var(view, "age")["stored"] == "number" and _var(view, "age")["missing"] == 1
    sex = next(s for s in view["pending"] if s["kind"] == "MERGE_LEVELS")
    view = _decide(lab, pid, sex["id"], False)
    assert not any(s["kind"] == "MERGE_LEVELS" for s in view["pending"])  # declined: not asked again
    undone = _act(lab, pid, f"/api/datalab/{pid}/undo")
    assert undone["version"] == 2 and _var(undone, "age")["stored"] == "text"
    again = lab.post(f"/api/datalab/{pid}/cleaning/{age['id']}", headers=H, json={"accept": True})
    assert again.status_code == 409  # already decided


def test_analyses_run_by_code_and_hide_small_counts(lab):
    view = _project(lab)
    pid = view["id"]
    sex = next(s for s in view["pending"] if s["kind"] == "MERGE_LEVELS")
    _decide(lab, pid, sex["id"], True)
    result = _analyse(lab, pid, {"kind": "COMPARE_TWO", "variables": ["score", "sex"], "method": "MEANS"})
    assert result["status"] in ("VALID", "VALID_WITH_WARNINGS") and result["record"]["method"].startswith("Welch") and result["chart"]
    assert lab.get(f"/api/datalab/{pid}/analyses/{result['id']}/chart.png", headers=H).content[:4] == b"\x89PNG"
    cross = _analyse(lab, pid, {"kind": "CROSSTAB", "variables": ["district", "passed"]})
    assert cross["status"] != "NOT_ESTIMABLE" and "chi2" in cross["statistics"] or "fisher_p" in cross["statistics"]
    refused = lab.post(f"/api/datalab/{pid}/analyses", headers=H, json={"kind": "DESCRIBE", "variables": ["name"]})
    assert refused.status_code == 400 and refused.json()["code"] == "VARIABLE_EXCLUDED"
    wrong = lab.post(f"/api/datalab/{pid}/analyses", headers=H, json={"kind": "CORRELATE", "variables": ["score", "district"]})
    assert wrong.status_code == 400 and wrong.json()["code"] == "VARIABLE_KIND"
    one = lab.post(f"/api/datalab/{pid}/analyses", headers=H, json={"kind": "COMPARE_TWO", "variables": ["score"]})
    assert one.status_code == 400 and one.json()["code"] == "VARIABLE_KIND"  # checked before anything reads a second variable (finding 9)


def test_a_survey_design_column_is_asked_about_then_blocks_analysis(lab):
    data = "weight,score,group\n" + "".join(f"{1 + (i % 5) / 10},{i % 17},{'a' if i % 2 else 'b'}\n" for i in range(60))
    view = _project(lab, data)
    pid = view["id"]
    assert view["survey"] == "ASK" and view["surveyColumns"] == ["weight"]
    asked = lab.post(f"/api/datalab/{pid}/analyses", headers=H, json={"kind": "DESCRIBE", "variables": ["score"]})
    assert asked.status_code == 400 and asked.json()["code"] == "SURVEY_QUESTION"
    designed = lab.post(f"/api/datalab/{pid}/variables/weight", headers=H, json={"survey": "DESIGN"}).json()
    assert designed["survey"] == "DESIGN"
    blocked = lab.post(f"/api/datalab/{pid}/analyses", headers=H, json={"kind": "DESCRIBE", "variables": ["score"]})
    assert blocked.status_code == 400 and blocked.json()["code"] == "SURVEY_UNSUPPORTED"
    assert lab.get(f"/api/datalab/{pid}/preview", headers=H).status_code == 200  # the data can still be reviewed
    other = lab.post(f"/api/datalab/{pid}/variables/weight", headers=H, json={"survey": "NOT_DESIGN"}).json()
    assert other["survey"] == "NONE"
    assert _analyse(lab, pid, {"kind": "DESCRIBE", "variables": ["score"]})["status"] == "VALID"


def test_the_preview_masks_left_out_columns(lab):
    view = _project(lab)
    page = lab.get(f"/api/datalab/{view['id']}/preview", headers=H).json()
    name_col = page["columns"].index("name")
    assert all(row[name_col] == "•••" for row in page["rows"]) and page["total"] == 121


def test_the_report_workbook_has_no_records_and_the_cleaned_data_is_its_own_file(lab):
    from openpyxl import load_workbook

    view = _project(lab)
    pid = view["id"]
    _analyse(lab, pid, {"kind": "DESCRIBE", "variables": ["district"]})
    _analyse(lab, pid, {"kind": "DESCRIBE", "variables": ["score"]})
    r = lab.get(f"/api/datalab/{pid}/workbook", headers=H)
    assert r.status_code == 200
    book = load_workbook(io.BytesIO(r.content))
    assert book.sheetnames[:3] == ["README", "Data dictionary", "Cleaning log"] and "Cleaned data" not in book.sheetnames and len(book.sheetnames) == 5
    assert book[book.sheetnames[3]]._charts  # a real Excel chart for the category percentages
    assert lab.get(f"/api/datalab/{pid}/cleaned", headers=H).status_code == 409  # not made yet
    made = _act(lab, pid, f"/api/datalab/{pid}/cleaned")
    assert made["cleanedReady"]
    cleaned = load_workbook(io.BytesIO(lab.get(f"/api/datalab/{pid}/cleaned", headers=H).content))
    assert cleaned.sheetnames == ["About this file", "Cleaned data"]
    assert "individual records" in cleaned["About this file"]["A3"].value
    header = [c.value for c in cleaned["Cleaned data"][1]]
    assert "name" not in header and "score" in header
    lab.post(f"/api/datalab/{pid}/variables/score", headers=H, json={"excluded": True})
    assert lab.get(f"/api/datalab/{pid}/cleaned", headers=H).status_code == 409  # out of date once the columns change


def test_coordinates_are_left_out_until_included_as_a_recorded_decision(lab):
    from openpyxl import load_workbook

    data = "lat,lon,score\n" + "".join(f"{0.3 + i / 1000},{32.5 + i / 1000},{i % 9}\n" for i in range(30))
    view = _project(lab, data)
    pid = view["id"]
    assert _var(view, "lat")["excluded"] and "LOCATION" in _var(view, "lat")["flags"]
    page = lab.get(f"/api/datalab/{pid}/preview", headers=H).json()
    assert page["rows"][0][0] == "•••"
    _act(lab, pid, f"/api/datalab/{pid}/cleaned")
    header = [c.value for c in load_workbook(io.BytesIO(lab.get(f"/api/datalab/{pid}/cleaned", headers=H).content))["Cleaned data"][1]]
    assert header == ["score"]
    released = lab.post(f"/api/datalab/{pid}/variables/lat", headers=H, json={"excluded": False}).json()
    assert [r["name"] for r in released["released"]] == ["lat"] and released["released"][0]["by"] == "student@example.com"


def test_a_changed_setting_makes_results_out_of_date_and_they_cant_be_reported(lab):
    view = _project(lab)
    pid = view["id"]
    result = _analyse(lab, pid, {"kind": "DESCRIBE", "variables": ["district"]})
    view = lab.get(f"/api/datalab/{pid}", headers=H).json()
    assert view["analyses"][0]["id"] == result["id"] and not view["analyses"][0]["stale"]
    view = lab.post(f"/api/datalab/{pid}/variables/district", headers=H, json={"excluded": True}).json()
    assert view["analyses"][0]["stale"]  # excluded after it ran: out of date (finding 4)
    refused = lab.post(f"/api/datalab/{pid}/report", headers=H, json={"analyses": [result["id"]]})
    assert refused.status_code == 400 and refused.json()["code"] == "STALE_ANALYSES"
    lab.post(f"/api/datalab/{pid}/variables/district", headers=H, json={"excluded": False})
    assert not lab.get(f"/api/datalab/{pid}", headers=H).json()["analyses"][0]["stale"]
    view = lab.post(f"/api/datalab/{pid}/details", headers=H, json={"threshold": 10}).json()
    assert view["analyses"][0]["stale"]  # a stricter threshold: shown again under it before reporting


def test_an_analysis_on_data_that_changed_meanwhile_is_not_kept(lab, monkeypatch):
    """Finding 1: an analysis computed on one version is never attached to another."""
    from app.datalab import service
    from app.runtime import get_runtime

    view = _project(lab)
    pid = view["id"]
    rt = get_runtime()
    real = service.HANDLERS["ANALYSE"]

    def meanwhile(rt_, p, op, written):
        commit, result = real(rt_, p, op, written)

        def bump(q):
            q.current = 1  # the data changed while the analysis ran
            return q

        rt.store.update_datalab(pid, bump)
        return commit, result

    monkeypatch.setitem(service.HANDLERS, "ANALYSE", meanwhile)
    done = _act(lab, pid, f"/api/datalab/{pid}/analyses", json={"kind": "DESCRIBE", "variables": ["score"]})
    assert done["op"]["status"] == "FAILED" and done["op"]["code"] == "DATA_CHANGED" and done["analyses"] == []
    prefix = rt.store.get_datalab(pid).storage_prefix()
    assert not [p for p, _ in rt.files.list(prefix) if "/analyses/" in p]  # its files went with it


def test_one_piece_of_data_work_at_a_time(lab):
    view = _project(lab)
    pid = view["id"]
    from app.runtime import get_runtime

    rt = get_runtime()

    def busy(q):
        q.op = q.op.model_copy(update={"status": "RUNNING", "lease_until": None})
        return q

    rt.store.update_datalab(pid, busy)
    refused = lab.post(f"/api/datalab/{pid}/analyses", headers=H, json={"kind": "DESCRIBE", "variables": ["score"]})
    assert refused.status_code == 409 and refused.json()["code"] == "OP_RUNNING"


def _wait(client, pid, done, timeout=60):
    deadline = time.time() + timeout
    while time.time() < deadline:
        view = client.get(f"/api/datalab/{pid}", headers=H).json()
        if done(view):
            return view
        time.sleep(0.2)
    raise AssertionError(view)


def test_the_report_is_written_from_tokens_approved_and_exported(lab):
    view = _project(lab)
    pid = view["id"]
    sex = next(s for s in view["pending"] if s["kind"] == "MERGE_LEVELS")
    _decide(lab, pid, sex["id"], True)
    _analyse(lab, pid, {"kind": "COMPARE_TWO", "variables": ["score", "sex"]})
    _analyse(lab, pid, {"kind": "DESCRIBE", "variables": ["district"]})
    started = lab.post(f"/api/datalab/{pid}/report", headers=H, json={})
    assert started.status_code == 200, started.json()
    view = _wait(lab, pid, lambda v: bool(v["reports"]) or bool(v["reportFailure"]))
    assert view["reports"] and not view["reportFailure"]
    assert "d_report" in lab.models.tasks and "d_report_review" in lab.models.tasks
    sent = [r for t, r in zip(lab.models.tasks, lab.models.requests, strict=True) if t.startswith("d_")]
    assert all("Person 1" not in r and "0001" not in r for r in sent)  # results only, never rows
    review = next(r for t, r in zip(lab.models.tasks, lab.models.requests, strict=True) if t == "d_report_review")
    assert "Appendix C. Data dictionary" in review and "Methods" in review  # the reviewer reads the whole document (finding 5)
    # ... knowing what code wrote and the facts it was written from, so it never asks the writer to change code's text (live check 2026-10-04)
    assert '"by": "CODE"' in review and '"by": "WRITER"' in review and '"rowsUsed"' in review and '"preparation"' in review
    doc = lab.get(f"/api/datalab/{pid}/report", headers=H).json()
    headings = [s["heading"] for s in doc["sections"]]
    assert headings[:6] == ["Executive summary", "Key findings", "The dataset", "Data preparation", "Data quality", "Methods"] and "Limitations" in headings
    assert [a["heading"] for a in doc["appendices"]] == ["Appendix A. Cleaning log", "Appendix B. Statistical output", "Appendix C. Data dictionary"]
    summary = doc["sections"][0]["paragraphs"][0]
    assert "121" in summary and "⟦" not in summary  # tokens filled by code
    dataset = " ".join(doc["sections"][2]["paragraphs"])
    assert 'use 3 of the 7 variables: "score", "sex", "district"' in dataset  # never "nothing left out" (the reviewer read it as "all analysed")
    assert "variables available for analysis" in doc["sections"][4]["paragraphs"][0]
    word = lab.get(f"/api/datalab/{pid}/report/export", headers=H)
    text = "\n".join(p.text for p in Document(io.BytesIO(word.content)).paragraphs)
    assert "Executive summary" in text and "Appendix C. Data dictionary" in text
    from app.runtime import get_runtime

    rt = get_runtime()
    job = rt.store.get(started.json()["id"])
    assert job.status == "COMPLETED" and job.billing.charged > 0
    approved = internal(rt, job, "approved.json")
    assert approved["document"]["sections"] == doc["sections"]  # exactly what the reviewer approved is what is delivered


def _review_with(rules_status):
    def review(payload):
        return {"verdict": "PASS", "rules": [{"rule": r, "status": rules_status(r), "note": "x"} for r in payload["rules"]], "issues": [], "suggestions": []}

    return review


def test_a_report_the_reviewer_does_not_approve_fails_without_charge(lab):
    view = _project(lab)
    pid = view["id"]
    _analyse(lab, pid, {"kind": "DESCRIBE", "variables": ["score"]})
    lab.models.overrides["d_report_review"] = lambda payload: {"verdict": "REPAIR", "rules": [{"rule": r, "status": "FAIL" if r == "R3" else "PASS", "note": "x"}
                                                                                              for r in payload["rules"]],
                                                               "issues": ["Remove the causal claim."], "suggestions": []}
    started = lab.post(f"/api/datalab/{pid}/report", headers=H, json={}).json()
    view = _wait(lab, pid, lambda v: bool(v["reportFailure"]) or bool(v["reports"]))
    assert not view["reports"] and "Nothing was charged" in view["reportFailure"]
    from app.runtime import get_runtime

    rt = get_runtime()
    job = rt.store.get(started["id"])
    rounds = internal(rt, job, "report_review.json")
    assert len(rounds) == 3  # reviewed, repaired and reviewed again, twice (an identical repeat is answered from the cache)
    assert job.status == "FAILED" and job.failure.code == "DOCUMENT_NOT_APPROVED" and job.billing.charged == 0


def test_not_applicable_on_every_rule_is_not_approval(lab):
    view = _project(lab)
    pid = view["id"]
    _analyse(lab, pid, {"kind": "DESCRIBE", "variables": ["score"]})
    lab.models.overrides["d_report_review"] = _review_with(lambda r: "NOT_APPLICABLE")
    started = lab.post(f"/api/datalab/{pid}/report", headers=H, json={}).json()
    view = _wait(lab, pid, lambda v: bool(v["reportFailure"]) or bool(v["reports"]))
    from app.runtime import get_runtime

    job = get_runtime().store.get(started["id"])
    assert not view["reports"] and job.failure.code == "DOCUMENT_NOT_APPROVED" and job.billing.charged == 0


def test_an_empty_narrative_is_never_delivered(lab):
    view = _project(lab)
    pid = view["id"]
    _analyse(lab, pid, {"kind": "DESCRIBE", "variables": ["score"]})
    lab.models.overrides["d_report"] = lambda payload: {"summary": [], "findings": [{"id": a["id"], "paragraphs": []} for a in payload["analyses"]],
                                                        "keyFindings": [], "limitations": [], "conclusions": []}
    started = lab.post(f"/api/datalab/{pid}/report", headers=H, json={}).json()
    view = _wait(lab, pid, lambda v: bool(v["reportFailure"]) or bool(v["reports"]))
    from app.runtime import get_runtime

    rt = get_runtime()
    job = rt.store.get(started["id"])
    assert not view["reports"] and job.failure.code == "DOCUMENT_NOT_APPROVED" and job.billing.charged == 0
    rounds = internal(rt, job, "report_rounds.json")
    assert any("Write the summary" in p for p in rounds[0]["problems"])
    assert "d_report_review" not in lab.models.tasks  # never even reviewed


def test_a_draft_that_writes_its_own_numbers_is_repaired_by_code(lab):
    view = _project(lab)
    pid = view["id"]
    _analyse(lab, pid, {"kind": "DESCRIBE", "variables": ["score"]})
    calls = []

    def writer(payload):
        calls.append(payload.get("critique"))
        from tests.fake_models import FakeModels

        good = FakeModels.default("d_report", payload)
        if len(calls) == 1:
            good["summary"] = ["The mean score was 68 and age causes better scores in this dataset."]
        return good

    lab.models.overrides["d_report"] = writer
    lab.post(f"/api/datalab/{pid}/report", headers=H, json={})
    view = _wait(lab, pid, lambda v: bool(v["reports"]) or bool(v["reportFailure"]))
    assert view["reports"] and len(calls) == 2
    assert any("digits" in c for c in calls[1]) and any("states a cause" in c for c in calls[1])


def test_settings_cant_change_while_a_report_is_written(lab):
    view = _project(lab)
    pid = view["id"]
    result = _analyse(lab, pid, {"kind": "DESCRIBE", "variables": ["district"]})
    from app.runtime import get_runtime

    rt = get_runtime()
    rt.store.set_flag("processing_enabled", False)  # the report waits in the queue
    assert lab.post(f"/api/datalab/{pid}/report", headers=H, json={}).status_code == 200
    for url, body in ((f"/api/datalab/{pid}/variables/district", {"excluded": True}), (f"/api/datalab/{pid}/details", {"threshold": 10})):
        r = lab.post(url, headers=H, json=body)
        assert r.status_code == 409 and r.json()["code"] == "STEP_RUNNING"
    assert lab.delete(f"/api/datalab/{pid}/analyses/{result['id']}", headers=H).status_code == 409  # its chart is in the report
    rt.store.set_flag("processing_enabled", True)


def test_the_report_waits_for_its_prices(tmp_path, monkeypatch):
    monkeypatch.setenv("DATALAB_ENABLED", "true")
    from tests.conftest import _client

    for client in _client(tmp_path, monkeypatch, "fixed"):
        view = _project(client)
        assert view["availability"] == "available" and not view["reportPriced"]  # analyses open, the report not priced
        _analyse(client, view["id"], {"kind": "DESCRIBE", "variables": ["score"]})
        refused = client.post(f"/api/datalab/{view['id']}/report", headers=H, json={})
        assert refused.status_code == 400 and refused.json()["code"] == "REPORT_NOT_PRICED"


def test_data_lab_is_closed_until_switched_on(client):
    r = client.post("/api/datalab", headers=H, json={"title": "x"})
    assert r.status_code == 400 and r.json()["code"] == "SERVICE_UNAVAILABLE"


def test_another_user_cannot_see_a_project_and_deletion_erases_it(lab):
    view = _project(lab)
    pid = view["id"]
    assert lab.get(f"/api/datalab/{pid}", headers=OTHER).status_code == 404
    from app.runtime import get_runtime

    rt = get_runtime()
    prefix = rt.store.get_datalab(pid).storage_prefix()
    assert rt.files.list(prefix)
    assert lab.delete(f"/api/datalab/{pid}", headers=H).status_code == 204
    assert rt.store.get_datalab(pid) is None and lab.get(f"/api/datalab/{pid}", headers=H).status_code == 404
    assert rt.files.list(prefix) == []


def test_account_deletion_erases_data_lab_projects(lab, monkeypatch):
    view = _project(lab)
    from app.runtime import get_runtime

    rt = get_runtime()
    monkeypatch.setattr(rt.settings, "credits_enabled", False)  # testing mode: a balance does not block deletion
    prefix = rt.store.get_datalab(view["id"]).storage_prefix()
    assert lab.delete("/api/me", headers=H).status_code == 204
    assert rt.store.get_datalab(view["id"]) is None and rt.files.list(prefix) == []


def test_abandoned_files_are_swept_and_named_ones_kept(lab):
    import os
    from datetime import UTC, datetime, timedelta

    from app.jobs import service
    from app.runtime import get_runtime

    view = _project(lab)
    rt = get_runtime()
    p = rt.store.get_datalab(view["id"])
    stray = f"{p.storage_prefix()}/versions/stray.parquet"
    rt.files.put(stray, b"x", "application/octet-stream")
    old = (datetime.now(UTC) - timedelta(hours=3)).timestamp()
    for path, _ in rt.files.list(p.storage_prefix()):
        os.utime(rt.files.local_path(path), (old, old))
    assert service.sweep_datalab_files(rt) == 1
    assert not rt.files.exists(stray) and all(rt.files.exists(v.path) for v in p.versions)


def test_a_left_out_column_is_shown_without_its_values(lab):
    view = _project(lab)
    name = _var(view, "name")
    assert name["excluded"] and name["levels"] == [] and name["summary"] == {}
    assert _var(view, "district")["levels"]  # included columns keep theirs


def _plan_dict():
    from tests.fake_models import PLAN

    return dict(PLAN)


def _approve_with_chapter_three(rt, pid, objectives):
    from app.proposals.models import ChapterDocument, ChapterSection, ProposalPlan, StoredChapterVersion

    def approve(p):
        p.plan = ProposalPlan.model_validate({**_plan_dict(), "specificObjectives": objectives})
        p.plan_status = "APPROVED"
        p.plan_version += 1
        doc = ChapterDocument(number=3, title="Methodology", plan_version=p.plan_version, cited=[], words=30,
                              sections=[ChapterSection(key="design", number="3.1", heading="Research design",
                                                       paragraphs=["A cross-sectional survey of students, analysed with t-tests and descriptive statistics."])])
        path = f"{p.storage_prefix()}/chapters/3-v1.json"
        rt.files.put(path, doc.model_dump_json(by_alias=True).encode(), "application/json")
        state3 = p.chapter(3)
        state3.versions = [StoredChapterVersion(version=1, job_id="job_x", words=30, plan_version=p.plan_version, path=path)]
        state3.current = 1
        return p

    rt.store.update_project(pid, approve)


def test_chapter_four_is_written_by_objective_from_a_proposal(lab):
    """A proposal with an approved plan and Chapter Three links its data; analyses answer its objectives;
    Chapter Four is written by objective, with tables numbered 4.1, 4.2 ..."""
    from app.runtime import get_runtime
    from tests.test_proposals import _create

    rt = get_runtime()
    pid = _create(lab)["id"]
    refused = lab.post(f"/api/datalab/for-proposal/{pid}", headers=H)
    assert refused.status_code == 400 and refused.json()["code"] == "PLAN_NOT_APPROVED"
    _approve_with_chapter_three(rt, pid, ["To compare scores between female and male students.", "To describe students by district."])
    view = lab.post(f"/api/datalab/for-proposal/{pid}", headers=H).json()
    assert view["proposalId"] == pid and len(view["objectives"]) == 2 and view["title"].startswith("Chapter Four")
    assert lab.post(f"/api/datalab/for-proposal/{pid}", headers=H).json()["id"] == view["id"]  # the same project next time
    did = view["id"]
    _act(lab, did, f"/api/datalab/{did}/dataset", files={"file": ("data.csv", CSV.encode(), "text/csv")}, data={"consent": "true"})
    assert lab.post(f"/api/datalab/{did}/analyses", headers=H, json={"kind": "DESCRIBE", "variables": ["district"], "objective": 3}).status_code == 400
    _analyse(lab, did, {"kind": "COMPARE_TWO", "variables": ["score", "passed"], "objective": 1})
    _analyse(lab, did, {"kind": "DESCRIBE", "variables": ["district"], "objective": 2})
    assert lab.post(f"/api/datalab/{did}/report", headers=H, json={}).status_code == 200
    done = _wait(lab, did, lambda v: bool(v["reports"]) or bool(v["reportFailure"]))
    assert done["reports"] and done["reports"][0]["kind"] == "CHAPTER_FOUR" and "d_chapter4" in lab.models.tasks
    writer = next(r for t, r in zip(lab.models.tasks, lab.models.requests, strict=True) if t == "d_chapter4")
    assert "cross-sectional survey" in writer  # Chapter Three's methods reach the writer (finding 6)
    doc = lab.get(f"/api/datalab/{did}/report", headers=H).json()
    headings = [s["heading"] for s in doc["sections"]]
    assert doc["title"] == "Chapter Four: Results" and headings[0] == "4.1 Introduction"
    assert any(h.startswith("4.2 Objective 1: To compare scores") for h in headings) and any(h.startswith("4.3 Objective 2") for h in headings)
    assert headings[-1] == "4.4 Summary of the results"
    text = "\n".join(p.text for p in Document(io.BytesIO(lab.get(f"/api/datalab/{did}/report/export", headers=H).content)).paragraphs)
    assert "Table 4.1." in text and "Table 4.2." in text and "Figure 4.1." in text


def test_chapter_four_names_missing_objectives_and_follows_changed_ones(lab):
    from app.runtime import get_runtime
    from tests.test_proposals import _create

    rt = get_runtime()
    pid = _create(lab)["id"]
    _approve_with_chapter_three(rt, pid, ["To compare scores between female and male students.", "To describe students by district.",
                                          "To relate age to scores."])
    did = lab.post(f"/api/datalab/for-proposal/{pid}", headers=H).json()["id"]
    _act(lab, did, f"/api/datalab/{did}/dataset", files={"file": ("data.csv", CSV.encode(), "text/csv")}, data={"consent": "true"})
    first = _analyse(lab, did, {"kind": "DESCRIBE", "variables": ["district"], "objective": 2})
    missing = lab.post(f"/api/datalab/{did}/report", headers=H, json={})
    assert missing.status_code == 400 and missing.json()["code"] == "MISSING_OBJECTIVES" and "1, 3" in missing.json()["message"]

    def reword(p):  # the student revises the plan: objective 2 now says something else
        p.plan = p.plan.model_copy(update={"specific_objectives": ["To compare scores between female and male students.", "To describe students by age group.",
                                                                    "To relate age to scores."]})
        return p

    rt.store.update_project(pid, reword)
    changed = lab.post(f"/api/datalab/{did}/report", headers=H, json={"missingOk": [1, 3]})
    assert changed.status_code == 400 and changed.json()["code"] == "OBJECTIVES_CHANGED"
    view = lab.get(f"/api/datalab/{did}", headers=H).json()
    assert view["objectives"][1] == "To describe students by age group."
    lab.post(f"/api/datalab/{did}/analyses/{first['id']}/objective", headers=H, json={"objective": 2})
    assert lab.post(f"/api/datalab/{did}/report", headers=H, json={"missingOk": [1, 3]}).status_code == 200
    done = _wait(lab, did, lambda v: bool(v["reports"]) or bool(v["reportFailure"]))
    assert done["reports"], done["reportFailure"]
    doc = lab.get(f"/api/datalab/{did}/report", headers=H).json()
    gap = next(s for s in doc["sections"] if s["heading"].startswith("4.2 Objective 1"))
    assert gap["paragraphs"] == ["No analysis of the data was reported for this objective."]


def test_a_district_map_runs_and_a_suggested_match_is_applied_only_when_confirmed(lab):
    data = "district,cases\n" + "".join(f"{d},{i}\n" for i, d in enumerate(["Gulu"] * 12 + ["Pader"] * 9 + ["Kampla"] * 6))
    view = _project(lab, data)
    pid = view["id"]
    first = _analyse(lab, pid, {"kind": "MAP", "variables": ["district"], "method": "COUNT"})
    assert first["unmatched"] == ["Kampla"] and first["matches"] == {"Kampla": "KAMPALA"} and first["chart"]
    assert lab.get(f"/api/datalab/{pid}/analyses/{first['id']}/chart.png", headers=H).content[:4] == b"\x89PNG"
    names = {r[0]["text"] for r in first["tables"][0]["rows"]}
    assert names == {"Gulu", "Pader"}  # never placed without confirmation
    again = _analyse(lab, pid, {"kind": "MAP", "variables": ["district"], "method": "COUNT", "aliases": first["matches"]})
    assert {r[0]["text"] for r in again["tables"][0]["rows"]} == {"Gulu", "Pader", "Kampala"} and again["unmatched"] == []
    refused = lab.post(f"/api/datalab/{pid}/analyses", headers=H, json={"kind": "MAP", "variables": ["district", "district"], "method": "MEAN"})
    assert refused.status_code == 400


def test_the_reliability_view_counts_by_service_without_any_text(lab):
    from tests.test_api import ADMIN

    view = _project(lab)
    _analyse(lab, view["id"], {"kind": "DESCRIBE", "variables": ["district"]})
    lab.post(f"/api/datalab/{view['id']}/report", headers=H, json={})
    _wait(lab, view["id"], lambda v: bool(v["reports"]))
    assert lab.get("/api/admin/reliability", headers=H).status_code == 403
    report = lab.get("/api/admin/reliability?days=7", headers=ADMIN).json()
    row = next(s for s in report["services"] if s["service"] == "Data Lab report")
    assert row["jobs"] == 1 and row["completed"] == 1 and row["completionRate"] == 1.0 and row["chargedCredits"] > 0
    raw = json.dumps(report)
    assert "Person" not in raw and "Exam results" not in raw and "survey.csv" not in raw  # no paper text, titles or file names


def test_ready_messages_are_recorded_with_the_outcome_sent_once_and_never_carry_text(lab, monkeypatch):
    from app import notify
    from app.runtime import get_runtime

    rt = get_runtime()
    sent = []
    monkeypatch.setattr(rt.settings, "sendgrid_api_key", "sg-test")
    monkeypatch.setattr(rt.settings, "notify_from", "hello@paperaid.test")
    monkeypatch.setattr(notify, "_email", lambda rt_, to, subject, text: sent.append((to, subject, text)) or True)
    assert lab.get("/api/me/notifications", headers=H).json()["available"] == {"email": True, "sms": False}
    view = _project(lab)
    _analyse(lab, view["id"], {"kind": "DESCRIBE", "variables": ["district"]})
    lab.post(f"/api/datalab/{view['id']}/report", headers=H, json={})
    _wait(lab, view["id"], lambda v: bool(v["reports"]))
    _wait(lab, view["id"], lambda v: bool(sent), timeout=10)
    assert len(sent) == 1 and sent[0][1] == "Your analysis report is ready" and f"/app/datalab/{view['id']}" in sent[0][2]
    assert "Exam results" not in sent[0][2] and "survey.csv" not in sent[0][2]  # no title, file name or text
    job = rt.store.get(rt.store.get_datalab(view["id"]).jobs[-1])
    assert job.notice.outcome == "READY" and not job.notice.pending and job.notice.channels == {"email": "SENT"}
    notify.after_job(rt, job)
    notify.sweep(rt)
    assert len(sent) == 1  # never twice for the same outcome
    off = lab.post("/api/me/notifications", headers=H, json={"notifyEmail": False}).json()
    assert off["notifyEmail"] is False
    assert lab.post("/api/me/notifications", headers=H, json={"notifySms": True}).json()["code"] == "NO_PHONE"
    assert lab.post("/api/me/notifications", headers=H, json={"phone": "0772 123"}).json()["code"] == "INVALID_PHONE"
    assert lab.post("/api/me/notifications", headers=H, json={"phone": "+256 772 123456", "notifySms": True}).json()["notifySms"] is True


def test_a_failed_send_is_tried_again_by_maintenance(lab, monkeypatch):
    """Finding 11: a provider outage doesn't lose the message."""
    from app import notify
    from app.runtime import get_runtime

    rt = get_runtime()
    attempts = []
    monkeypatch.setattr(rt.settings, "sendgrid_api_key", "sg-test")
    monkeypatch.setattr(rt.settings, "notify_from", "hello@paperaid.test")
    monkeypatch.setattr(notify, "_email", lambda rt_, to, subject, text: attempts.append(subject) and False)  # the provider is down
    view = _project(lab)
    _analyse(lab, view["id"], {"kind": "DESCRIBE", "variables": ["district"]})
    lab.post(f"/api/datalab/{view['id']}/report", headers=H, json={})
    _wait(lab, view["id"], lambda v: bool(v["reports"]))
    _wait(lab, view["id"], lambda v: bool(attempts), timeout=10)
    job_id = rt.store.get_datalab(view["id"]).jobs[-1]
    job = rt.store.get(job_id)
    assert job.notice.pending and job.notice.channels == {"email": "FAILED"} and job_id in rt.store.pending_notice_ids(10)

    def now(j):
        j.notice.next_at = None  # its wait is over
        return j

    rt.store.update(job_id, now)
    monkeypatch.setattr(notify, "_email", lambda rt_, to, subject, text: attempts.append(subject) or True)  # the provider is back
    assert notify.sweep(rt) == 1
    job = rt.store.get(job_id)
    assert not job.notice.pending and job.notice.channels == {"email": "SENT"} and len(attempts) == 2


def test_admin_retry_refuses_a_report_whose_project_is_gone(lab):
    """Finding 7: admin retry claims the Data Lab project like a student's resume, so a deleted
    project refuses before any credits are held."""
    from app.jobs.models import JobFailure, JobStatus
    from app.pricing.billing import refund_job
    from app.runtime import get_runtime
    from tests.test_api import ADMIN

    view = _project(lab)
    pid = view["id"]
    _analyse(lab, pid, {"kind": "DESCRIBE", "variables": ["district"]})
    rt = get_runtime()
    rt.store.set_flag("processing_enabled", False)
    job_id = lab.post(f"/api/datalab/{pid}/report", headers=H, json={}).json()["id"]

    def fail(j, w):
        j.failure = JobFailure(code="PROVIDER_DOWN", user_message="x", retryable=True)
        j.status = JobStatus.PROCESSING
        from app.jobs import state

        state.transition(j, JobStatus.FAILED, "Failed")
        refund_job(j, w, "test")
        return j, w

    rt.store.update_job_and_wallet(job_id, fail)
    rt.store.set_flag("processing_enabled", True)
    held = rt.store.get_wallet(rt.store.get(job_id).owner_uid).held
    rt.store.update_datalab(pid, lambda q: q.model_copy(update={"deleting": True}))
    r = lab.post(f"/api/admin/jobs/{job_id}/retry", headers=ADMIN)
    assert r.status_code == 409 and r.json()["code"] == "PROJECT_DELETED"
    job = rt.store.get(job_id)
    assert job.status == "FAILED" and rt.store.get_wallet(job.owner_uid).held == held  # nothing queued, nothing held


def test_filters_through_the_api_and_the_release_check(lab):
    view = _project(lab)
    pid = view["id"]
    sex = next(s for s in view["pending"] if s["kind"] == "MERGE_LEVELS")
    _decide(lab, pid, sex["id"], True)
    narrow = lab.post(f"/api/datalab/{pid}/analyses", headers=H, json={"kind": "DESCRIBE", "variables": ["score"],
                                                                       "filters": [{"variable": "name", "op": "IN", "values": ["Person 1"]}]})
    assert narrow.status_code == 400 and narrow.json()["code"] == "FILTER_NOT_ALLOWED"  # refused at once
    gulu = _analyse(lab, pid, {"kind": "DESCRIBE", "variables": ["score"], "filters": [{"variable": "district", "op": "IN", "values": ["Gulu"]}]})
    assert gulu["record"]["filters"] == "district is Gulu" and gulu["record"]["rowsAvailable"] == 121 and gulu["title"].endswith("(district is Gulu)")
    few = _act(lab, pid, f"/api/datalab/{pid}/analyses", json={"kind": "DESCRIBE", "variables": ["score"],
                                                               "filters": [{"variable": "score", "op": "BETWEEN", "low": "51"}]})
    assert few["op"]["status"] == "FAILED" and few["op"]["code"] == "FILTER_TOO_NARROW"  # leaves out fewer than 5: refused in the worker
    ranged = lab.post(f"/api/datalab/{pid}/analyses", headers=H, json={"kind": "DESCRIBE", "variables": ["score"],
                                                                       "filters": [{"variable": "district", "op": "BETWEEN", "low": "1"}]})
    assert ranged.status_code == 400 and ranged.json()["code"] == "FILTER_INCOMPLETE"  # a category has no range
    places = lab.get("/api/datalab/places", headers=H).json()
    assert len(places["subregions"]) == 15 and places["regions"] == ["Central", "Eastern", "Northern", "Western"]


def test_overlapping_results_cant_be_released_together(lab):
    data = "group,score\n" + "".join(f"{'a' if i < 50 else 'b' if i < 51 else 'c'},{i % 13}\n" for i in range(100))
    view = _project(lab, data)
    pid = view["id"]
    one = _analyse(lab, pid, {"kind": "DESCRIBE", "variables": ["score"], "filters": [{"variable": "group", "op": "IN", "values": ["a"]}]})
    two = _analyse(lab, pid, {"kind": "DESCRIBE", "variables": ["score"], "filters": [{"variable": "group", "op": "IN", "values": ["a", "b"]}]})
    refused = lab.post(f"/api/datalab/{pid}/report", headers=H, json={"analyses": [one["id"], two["id"]]})
    assert refused.status_code == 400 and refused.json()["code"] == "OVERLAPPING_RESULTS"
    workbook = lab.get(f"/api/datalab/{pid}/workbook", headers=H)
    assert workbook.status_code == 400 and workbook.json()["code"] == "OVERLAPPING_RESULTS"
    lab.delete(f"/api/datalab/{pid}/analyses/{two['id']}", headers=H)
    assert lab.get(f"/api/datalab/{pid}/workbook", headers=H).status_code == 200


def test_an_upload_needs_the_researchers_confirmation_and_a_cleared_country(lab):
    p = lab.post("/api/datalab", headers=H, json={"title": "x"}).json()
    files = {"file": ("d.csv", CSV.encode(), "text/csv")}
    unconfirmed = lab.post(f"/api/datalab/{p['id']}/dataset", headers=H, files=files)
    assert unconfirmed.status_code == 400 and unconfirmed.json()["code"] == "CONSENT_REQUIRED"
    kenya = lab.post(f"/api/datalab/{p['id']}/dataset", headers=H, files=files, data={"consent": "true", "country": "KEN"})
    assert kenya.status_code == 400 and kenya.json()["code"] == "COUNTRY_NOT_AVAILABLE"
    assert "Kenya's data protection law" in kenya.json()["message"]
    other = lab.post(f"/api/datalab/{p['id']}/dataset", headers=H, files=files, data={"consent": "true", "country": "OTHER"})
    assert other.json()["code"] == "COUNTRY_NOT_AVAILABLE"  # fail closed: a country not listed is frozen
    countries = lab.get("/api/datalab/countries", headers=H).json()
    assert [c["iso3"] for c in countries if c["available"]] == ["UGA"]
    view = _act(lab, p["id"], f"/api/datalab/{p['id']}/dataset", files=files, data={"consent": "true", "country": "UGA", "removed": '["2 columns: phone numbers"]'})
    from app.runtime import get_runtime

    record = get_runtime().store.get_datalab(p["id"])
    assert view["op"]["status"] == "DONE" and record.country == "UGA" and record.consent["wording"] == "datalab-upload-2026-10-04"
    assert record.source.removed == ["2 columns: phone numbers"] and record.consent["terms"]
    rules = lab.get("/api/datalab/identifier-rules", headers=H).json()
    assert "PERSONAL" in rules["headers"] and "PHONE" in rules["values"]  # the browser reads the same rules


def test_terms_must_be_accepted_before_an_upload_or_a_paid_step(lab):
    from tests.test_api import OTHER as SOMEONE

    p = lab.post("/api/datalab", headers=SOMEONE, json={"title": "x"}).json()
    from app.runtime import get_runtime

    rt = get_runtime()
    uid = rt.store.get_datalab(p["id"]).owner_uid
    rt.store.update_wallet(uid, "other@example.com", lambda w: w.model_copy(update={"terms_version": "2020-01-01"}))  # accepted older terms
    refused = lab.post(f"/api/datalab/{p['id']}/dataset", headers=SOMEONE, files={"file": ("d.csv", CSV.encode(), "text/csv")}, data={"consent": "true"})
    assert refused.status_code == 403 and refused.json()["code"] == "TERMS_REQUIRED"
    version = lab.get("/api/config", headers=SOMEONE).json()["termsVersion"]
    assert lab.post("/api/me/terms", headers=SOMEONE, json={"version": "2020-01-01"}).json()["code"] == "TERMS_CHANGED"
    assert lab.post("/api/me/terms", headers=SOMEONE, json={"version": version}).status_code == 204
    wallet = rt.store.get_wallet(uid)
    assert wallet.terms_version == version and wallet.terms_accepted_at is not None
    ok = lab.post(f"/api/datalab/{p['id']}/dataset", headers=SOMEONE, files={"file": ("d.csv", CSV.encode(), "text/csv")}, data={"consent": "true"})
    assert ok.status_code == 200
