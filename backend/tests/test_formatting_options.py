"""Formatting options (master context §16, §17, §81): the student's own settings over a preset,
an institution logo on the first page, and layouts for non-academic documents."""

import io
import struct
import zlib

from docx import Document

from app.formatting.presets import PRESETS, CustomLayout, with_custom
from tests.conftest import fixture_bytes
from tests.test_api import STUDENT, wait

FORMAT = {"writing": "NONE", "formatting": "FORMAT", "preset": "apa7", "latex": False}


def _png(width: int = 40, height: int = 20) -> bytes:
    """A small valid PNG, built without an imaging library."""
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    raw = b"".join(b"\x00" + b"\xff\x00\x00" * width for _ in range(height))
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def _draft(client, name="simple_essay.docx"):
    job_id = client.post("/api/jobs", headers=STUDENT).json()["id"]
    client.post(f"/api/jobs/{job_id}/files/source", headers=STUDENT, files={"file": (name, fixture_bytes(name), "application/octet-stream")})
    return job_id


def _run(client, job_id, selection):
    quote = client.post(f"/api/jobs/{job_id}/quote", headers=STUDENT, json={"selection": selection}).json()
    assert quote.get("quote"), quote
    client.post(f"/api/jobs/{job_id}/submit", headers=STUDENT, json={"quoteId": quote["quote"]["id"]})
    return wait(client, job_id, timeout=120)


def test_custom_settings_override_only_what_the_student_chose():
    spec = with_custom(PRESETS["apa7"], CustomLayout(font="Calibri", line_spacing=1.5, margin_cm=3.0))
    assert (spec.font, spec.line_spacing, spec.margins_cm, spec.size_pt) == ("Calibri", 1.5, (3.0, 3.0, 3.0, 3.0), 12)
    assert with_custom(PRESETS["apa7"], None) is PRESETS["apa7"]


def test_a_job_applies_the_students_settings(client):
    job = _run(client, _draft(client), {**FORMAT, "custom": {"font": "Arial", "lineSpacing": 1.5}})
    assert job["status"] == "COMPLETED"
    body = next(r["value"] for r in job["formatting"]["rules"] if r["label"] == "Body text")
    assert "Arial" in body and "1.5 line spacing" in body


def test_a_logo_is_placed_on_the_first_page_and_the_wording_is_untouched(client):
    job_id = _draft(client)
    uploaded = client.post(f"/api/jobs/{job_id}/files/logo", headers=STUDENT, files={"file": ("crest.png", _png(), "image/png")})
    assert uploaded.status_code == 200 and uploaded.json()["widthPx"] == 40
    job = _run(client, job_id, {**FORMAT, "logo": "CENTER"})
    assert job["status"] == "COMPLETED" and any(r["label"] == "Logo" for r in job["formatting"]["rules"])
    data = client.get(f"/api/jobs/{job_id}/outputs/paper", headers=STUDENT).content
    doc = Document(io.BytesIO(data))
    assert doc.inline_shapes and doc.paragraphs[0].runs and doc.paragraphs[0].text == ""  # the picture comes first


def test_a_logo_must_be_a_real_image_and_needs_formatting(client):
    job_id = _draft(client)
    bad = client.post(f"/api/jobs/{job_id}/files/logo", headers=STUDENT, files={"file": ("crest.png", b"not an image", "image/png")})
    assert bad.status_code == 422
    refused = client.post(f"/api/jobs/{job_id}/quote", headers=STUDENT, json={"selection": {**FORMAT, "logo": "CENTER"}})
    assert refused.json()["code"] == "NO_LOGO"


def test_non_academic_layouts_are_offered_and_applied(client):
    presets = {p["id"] for p in client.get("/api/config").json()["presets"] if p["available"]}
    assert {"professional", "report"} <= presets
    job = _run(client, _draft(client), {**FORMAT, "preset": "professional"})
    assert job["status"] == "COMPLETED" and "Calibri 11" in next(r["value"] for r in job["formatting"]["rules"] if r["label"] == "Body text")
