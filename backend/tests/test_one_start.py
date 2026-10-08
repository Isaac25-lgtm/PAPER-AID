"""One Start (owner decision 2026-10-01): the student gives their task and a few details and presses
Start; PaperAid reads, plans and writes by itself and the student sees the document. The plan stays
internal and continues only when PaperAid's own final review and code checks approve it (Codex
2026-10-01); the first document carries the plan's price; a figure only the student can give is
drafted as a marked gap and filled in afterwards; a minimum balance can be set per service."""

import io
import time

import pytest
from docx import Document

from app.jobs.models import ServiceSelection
from tests import fake_works
from tests.fake_models import PLAN
from tests.test_api import STUDENT
from tests.test_proposals import _create
from tests.test_works import QUESTION

H = STUDENT


@pytest.fixture
def works_client(tmp_path, monkeypatch):
    monkeypatch.setenv("WORKS_ENABLED", '["CONCEPT_NOTE","COURSEWORK","FUNDING_PROPOSAL"]')
    from tests.conftest import _client

    yield from _client(tmp_path, monkeypatch, "fixed")


def _settled(client, path, done, timeout=120):
    """Poll a work or project until `done` holds (the worker runs steps in the background)."""
    deadline = time.time() + timeout
    while time.time() < deadline:
        body = client.get(path, headers=H).json()
        if done(body):
            return body
        time.sleep(0.3)
    raise AssertionError(f"{path} did not settle: {body}")


def _coursework(client):
    work = client.post("/api/works", headers=H, json={"kind": "COURSEWORK", "variant": "ESSAY", "inputs": {"title": "Community health workers", "description": QUESTION}}).json()
    return client.post(f"/api/works/{work['id']}/answers", headers=H,
                       json={"answers": {"word_limit": "1500", "level": "LATER_UG", "ai_policy": "NOT_MENTIONED"}, "skipRest": True, "baseVersion": work["specVersion"]}).json()


def _funding(client):
    work = client.post("/api/works", headers=H, json={
        "kind": "FUNDING_PROPOSAL", "variant": "NGO_PROJECT", "mode": "COMPACT",
        "inputs": {"title": "Safer deliveries in Kamuli", "description": "Too many mothers in Kamuli deliver at home and die of preventable complications.",
                   "answers": {"problem": "Home deliveries and delayed referral.", "intervention": "Train village health teams.", "duration_months": "12", "currency": "USD"}},
    }).json()
    return client.post(f"/api/works/{work['id']}/answers", headers=H, json={"answers": {}, "skipRest": True, "baseVersion": work["specVersion"]}).json()


def _done(w):
    return bool(w["documents"]) or bool(w["autoFailure"])


# --- coursework: one Start, one charge -------------------------------------------------------------------


def test_one_start_writes_coursework_without_the_student_seeing_a_plan(works_client):
    from app.runtime import get_runtime

    client = works_client
    work = _coursework(client)
    started = client.post(f"/api/works/{work['id']}/start", headers=H)
    assert started.status_code == 200, started.json()
    assert started.json()["auto"] and started.json()["specStatus"] == "CONFIRMED"
    work = _settled(client, f"/api/works/{work['id']}", _done)
    assert work["documents"] and not work["autoFailure"] and work["planStatus"] == "APPROVED"
    store = get_runtime().store
    plan_job, draft_job = (store.get(j) for j in work["jobs"])
    assert plan_job.selection.work == "PLAN" and plan_job.selection.bundled and plan_job.billing.charged == 0
    assert draft_job.selection.work == "DRAFT" and draft_job.selection.bundled
    lines = {line.service: line.amount for line in draft_job.quote.lines}
    from app.core.config import get_settings
    from app.pricing.quote import fixed_price

    s = get_settings()
    band = draft_job.selection.work_band
    assert lines[band] == fixed_price(s, band, None) + fixed_price(s, "CW_PLAN", None)  # the plan is in the document's price
    assert draft_job.billing.charged == lines[band]


def test_a_plan_paperaid_does_not_approve_stops_without_charge_and_never_drafts(works_client):
    from app.runtime import get_runtime

    client = works_client
    client.models.overrides["w_plan_review"] = lambda payload: {"verdict": "REPAIR", "rules": [{"rule": r["rule"], "status": "PASS", "note": "Met."} for r in payload.get("rules", [])],
                                                                "issues": ["The plan does not answer the question."]}
    work = _coursework(client)
    client.post(f"/api/works/{work['id']}/start", headers=H)
    work = _settled(client, f"/api/works/{work['id']}", _done)
    assert not work["documents"] and "could not make a plan" in work["autoFailure"] and work["planStatus"] != "APPROVED"
    assert [get_runtime().store.get(j).selection.work for j in work["jobs"]] == ["PLAN"]  # no draft was started
    assert all(get_runtime().store.get(j).billing.charged == 0 for j in work["jobs"])


def test_continuing_twice_after_one_plan_starts_one_draft(works_client):
    from app.runtime import get_runtime
    from app.works import service

    client = works_client
    work = _coursework(client)
    client.post(f"/api/works/{work['id']}/start", headers=H)
    work = _settled(client, f"/api/works/{work['id']}", _done)
    plan_job = work["jobs"][0]
    service.continue_after_plan(get_runtime(), work["id"], plan_job)  # a retried export stage calls it again
    assert len(client.get(f"/api/works/{work['id']}", headers=H).json()["jobs"]) == 2


def test_a_minimum_balance_stops_start_before_any_ai_runs(works_client, monkeypatch):
    from app.runtime import get_runtime

    client = works_client
    monkeypatch.setattr(get_runtime().settings, "min_credits", {"COURSEWORK": 5000})  # more than the test account holds
    work = _coursework(client)
    refused = client.post(f"/api/works/{work['id']}/start", headers=H)
    assert refused.status_code == 402 and refused.json()["code"] == "INSUFFICIENT_CREDITS" and "credits" in refused.json()["message"]
    assert client.get(f"/api/works/{work['id']}", headers=H).json()["jobs"] == [] and "w_plan" not in client.models.tasks


def test_start_asks_for_required_answers_first(works_client):
    client = works_client
    work = client.post("/api/works", headers=H, json={"kind": "COURSEWORK", "variant": "ESSAY", "inputs": {"title": "Topic", "description": "Discuss the role of community health workers."}}).json()
    refused = client.post(f"/api/works/{work['id']}/start", headers=H)
    assert refused.status_code == 400 and refused.json()["code"] == "QUESTIONS_OPEN"


def test_bundled_pricing_frees_the_read_and_plan_and_adds_the_plan_to_the_document():
    from app.ai.orchestration import work_engine
    from app.core.config import Settings
    from app.pricing.quote import fixed_price, price

    s = Settings(pricing_mode="fixed")
    e = work_engine(s, "COURSEWORK")
    plan = price(s, ServiceSelection(work="PLAN", work_kind="COURSEWORK", work_band="CW_PLAN", bundled=True), 1500, engine=e)
    assert plan.ai_ugx == 0 and plan.budget_usd > 0  # nothing charged, the spend cap stays real
    draft = price(s, ServiceSelection(work="DRAFT", work_kind="COURSEWORK", work_band="CW_1500", bundled=True), 1500, engine=e)
    assert draft.ai_ugx == fixed_price(s, "CW_1500", None) + fixed_price(s, "CW_PLAN", None)
    unbundled = price(s, ServiceSelection(work="DRAFT", work_kind="COURSEWORK", work_band="CW_1500"), 1500, engine=e)
    assert unbundled.ai_ugx == fixed_price(s, "CW_1500", None)


# --- funding: figures only the student gives are marked gaps, filled in afterwards ----------------------------


def test_a_funding_proposal_is_drafted_with_its_missing_figures_marked_then_filled_in(works_client):
    client = works_client
    work = _funding(client)
    client.post(f"/api/works/{work['id']}/start", headers=H)
    work = _settled(client, f"/api/works/{work['id']}", _done, timeout=180)
    assert work["documents"], work["autoFailure"]
    doc = client.get(f"/api/works/{work['id']}/document", headers=H).json()
    gaps = {i["id"]: i for i in doc["readiness"] if i["status"] not in ("PASS", "NOT_APPLICABLE")}
    assert gaps["FP-027"]["basis"] == "AUTHOR" and gaps["FP-027"]["reason"] == "STUDENT_INFO_MISSING"  # the student's, not PaperAid's failure
    assert doc["status"] == "NOT_READY"
    tables = {t["caption"]: t["rows"] for t in doc["tables"]}
    assert any("[to be added]" in cell for row in tables["Monitoring and evaluation indicators"] for cell in row)
    assert any("[to be added]" in cell for row in tables["Budget summary"] for cell in row)  # never a misleading zero
    # the student adds their figures; the same text is rebuilt with them, no AI and no charge
    results = work["results"]
    for ind in results["indicators"]:
        ind["target"], ind["baseline"] = (60, 34) if ind["level"] == "outcome" else (120, 0)
    work = client.post(f"/api/works/{work['id']}/results", headers=H, json={"results": results, "baseVersion": work["resultsVersion"]}).json()
    budget = work["budget"]
    for n, line in enumerate(budget["lines"]):
        line["quantity"], line["unitCost"] = 4 + n, 250
    work = client.post(f"/api/works/{work['id']}/budget", headers=H, json={"budget": budget, "baseVersion": work["budgetVersion"]}).json()
    tasks_before = len(client.models.tasks)
    filled = client.post(f"/api/works/{work['id']}/figures", headers=H)
    assert filled.status_code == 200, filled.json()
    assert len(client.models.tasks) == tasks_before and filled.json()["current"] == 2
    doc = client.get(f"/api/works/{work['id']}/document", headers=H).json()
    assert next(i for i in doc["readiness"] if i["id"] == "FP-027")["status"] == "PASS"
    tables = {t["caption"]: t["rows"] for t in doc["tables"]}
    assert not any("[to be added]" in cell for row in tables["Budget summary"] for cell in row)
    word = client.get(f"/api/works/{work['id']}/export", headers=H)
    assert word.status_code == 200 and "[to be added]" not in "\n".join(c.text for t in Document(io.BytesIO(word.content)).tables for r in t.rows for c in r.cells)


def test_number_tokens_for_missing_figures_print_as_marked_gaps():
    from app.works import numbers
    from app.works.models import Budget, BudgetLine, Indicator, ResultsModel, WorkInputs
    from tests.test_works_golden import _spec

    spec = _spec("FUNDING_PROPOSAL", "NGO_PROJECT")
    model = ResultsModel.model_validate({**{k: v for k, v in fake_works.results({"spec": {"duration": 12}}).items() if k != "budgetLines"}})
    model.indicators = [Indicator.model_validate({**model.indicators[0].model_dump(by_alias=True), "target": None, "baseline": None})]
    tokens = numbers.values(spec, model, Budget(lines=[BudgetLine(id="B1", category="Training", description="Workshops", quantity=0, unit="workshop", unit_cost=0)]), WorkInputs(title="A project"))
    assert tokens["indicator.I1.target"][0] == "[target to be added]" and tokens["budget.total"][0] == "[total to be added]"
    assert tokens["budget.line.B1"][0] == "[amount to be added]"


# --- asking for changes with a document for context ------------------------------------------------------------


def test_a_change_request_can_bring_a_document_for_the_writer(works_client):
    client = works_client
    work = _coursework(client)
    client.post(f"/api/works/{work['id']}/start", headers=H)
    work = _settled(client, f"/api/works/{work['id']}", _done)
    note = Document()
    note.add_paragraph("Our district health officer reports that 42 village health teams were trained in Kamuli in 2025, and the supervisor wants this case used.")
    buffer = io.BytesIO()
    note.save(buffer)
    sent = client.post(f"/api/works/{work['id']}/requests/with-document", headers=H, data={"instruction": "Use the Kamuli example from my notes in the discussion", "sections": ""},
                       files={"file": ("notes.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert sent.status_code == 200, sent.json()
    request = sent.json()["requests"][0]
    assert request["contextName"] == "notes.docx" and request["context"] == ""  # the text is in file storage, never the record
    from app.runtime import get_runtime

    stored = get_runtime().store.get_work(work["id"]).requests[0]
    assert "42 village health teams" in get_runtime().files.get(stored.context_path).decode("utf-8")
    quoted = client.post(f"/api/works/{work['id']}/steps", headers=H, json={"step": "REVISE"}).json()
    assert client.post(f"/api/works/{work['id']}/steps/{quoted['job']['id']}/submit", headers=H, json={"quoteId": quoted["quote"]["id"]}).status_code == 200
    _settled(client, f"/api/works/{work['id']}", lambda w: len(w["documents"]) >= 2 or w["activeJob"] is None and "w_repair" in client.models.tasks)
    sent_to_writer = [r for t, r in zip(client.models.tasks, client.models.requests, strict=True) if t == "w_repair"]
    assert any("42 village health teams" in r and "notes.docx" in r for r in sent_to_writer)


# --- research proposals: one Start plans and writes Chapter One -------------------------------------------------


def test_one_start_plans_and_writes_chapter_one(client):
    from app.runtime import get_runtime

    project = _create(client)
    started = client.post(f"/api/projects/{project['id']}/start", headers=H, json={"acceptSampling": True})
    assert started.status_code == 200, started.json()
    project = _settled(client, f"/api/projects/{project['id']}", lambda p: any(c["number"] == 1 and c["current"] for c in p["chapters"]) or p["autoFailure"], timeout=180)
    assert not project["autoFailure"] and project["planStatus"] == "APPROVED"
    jobs = [get_runtime().store.get(j) for j in project["jobs"]]
    assert [j.selection.proposal for j in jobs] == ["PLAN", "CHAPTER_1"] and jobs[0].billing.charged == 0 and all(j.selection.bundled for j in jobs)
    if project["plan"]["samplingAssumed"]:
        assert any(a["kind"] == "SAMPLING" for a in project["acknowledgments"])  # the standard settings, as stated at Start


def test_a_proposal_plan_paperaid_does_not_approve_stops_without_a_chapter(client):
    from app.runtime import get_runtime

    client.models.overrides["p_plan_review"] = lambda payload: {"approved": False, "issues": ["Objective 2 has no matching question."]}
    project = _create(client)
    client.post(f"/api/projects/{project['id']}/start", headers=H)
    # the worker records the stop just after it publishes the plan: wait for the outcome itself
    project = _settled(client, f"/api/projects/{project['id']}", lambda p: bool(p["autoFailure"]) or any(c["current"] for c in p["chapters"]), timeout=180)
    assert "could not make a plan" in project["autoFailure"] and project["planStatus"] != "APPROVED"
    assert [get_runtime().store.get(j).selection.proposal for j in project["jobs"]] == ["PLAN"]


def test_the_conceptual_framework_is_a_figure_in_the_app_and_the_word_file(client):
    project = _create(client)
    client.post(f"/api/projects/{project['id']}/start", headers=H, json={"acceptSampling": True})
    project = _settled(client, f"/api/projects/{project['id']}", lambda p: any(c["number"] == 1 and c["current"] for c in p["chapters"]) or p["autoFailure"], timeout=180)
    if not PLAN["variables"]["independent"]:
        pytest.skip("the fake plan has no variables")
    figure = client.get(f"/api/projects/{project['id']}/framework.png", headers=H)
    assert figure.status_code == 200 and figure.content[:8] == b"\x89PNG\r\n\x1a\n"
    assert project["framework"].startswith("The study will examine the association between")
    word = client.get(f"/api/projects/{project['id']}/export", headers=H, params={"final": "false"})
    assert word.status_code == 200
    document = Document(io.BytesIO(word.content))
    pictures = document.inline_shapes
    assert len(pictures) >= 1 and pictures[0]._inline.docPr.get("descr", "").startswith("The study will examine")



# --- Codex's verification of the one-Start release (2026-10-01) ---------------------------------------------------


def _wallet(client):
    return client.get("/api/wallet", headers=H).json()


def _with_brief(client):
    work = client.post("/api/works", headers=H, json={"kind": "COURSEWORK", "variant": "ESSAY", "inputs": {"title": "CHWs", "description": QUESTION}}).json()
    client.post(f"/api/works/{work['id']}/sources/text", headers=H, json={"role": "BRIEF", "name": "Brief", "text": "Write an essay of 2,000 words that critically evaluates community health workers in Uganda."})
    return work


def test_reading_a_brief_needs_the_credits_the_document_will_take(works_client, monkeypatch):
    from app.runtime import get_runtime

    client = works_client
    monkeypatch.setattr(get_runtime().settings, "min_credits", {"COURSEWORK": 5000})
    work = _with_brief(client)
    refused = client.post(f"/api/works/{work['id']}/read", headers=H)
    assert refused.status_code == 402 and "w_read" not in client.models.tasks  # no AI before the credits are there


def test_start_reserves_the_document_price_until_the_document_holds_it(works_client):
    from app.runtime import get_runtime

    client = works_client
    before = _wallet(client)["available"]
    work = _coursework(client)
    client.post(f"/api/works/{work['id']}/start", headers=H)
    work = _settled(client, f"/api/works/{work['id']}", _done)
    store = get_runtime().store
    draft = store.get(work["jobs"][1])
    wallet = store.get_wallet(draft.owner_uid)
    assert work["documents"] and wallet.reservations == {} and wallet.held == 0  # reserved at Start, then the draft's hold, then settled
    assert before - wallet.available == draft.billing.charged > 0
    assert any(e.kind == "HOLD" and e.note == "Reserved for your document" for e in wallet.entries)


def test_a_stopped_plan_returns_what_start_reserved(works_client):
    client = works_client
    client.models.overrides["w_plan_review"] = lambda payload: {"verdict": "REPAIR", "rules": [{"rule": r["rule"], "status": "PASS", "note": "Met."} for r in payload.get("rules", [])],
                                                                "issues": ["The plan does not answer the question."]}
    before = _wallet(client)["available"]
    work = _coursework(client)
    client.post(f"/api/works/{work['id']}/start", headers=H)
    _settled(client, f"/api/works/{work['id']}", _done)
    after = _wallet(client)
    assert after["available"] == before and after["held"] == 0


def test_a_failed_read_says_so_instead_of_waiting(works_client):
    client = works_client
    client.models.refuse.add("w_read")  # the model refuses: the read fails
    work = _with_brief(client)
    client.post(f"/api/works/{work['id']}/read", headers=H)
    work = _settled(client, f"/api/works/{work['id']}", lambda w: w["activeJob"] is None and bool(w["autoFailure"]), timeout=180)
    assert work["needsRead"] and work["autoFailure"]


def test_adding_figures_refuses_anything_but_figures_and_keeps_the_approval(works_client):
    client = works_client
    work = _funding(client)
    client.post(f"/api/works/{work['id']}/start", headers=H)
    work = _settled(client, f"/api/works/{work['id']}", _done, timeout=180)
    assert work["documents"]
    results = work["results"]
    for ind in results["indicators"]:
        ind["target"] = 60
    work = client.post(f"/api/works/{work['id']}/results", headers=H, json={"results": results, "baseVersion": work["resultsVersion"]}).json()
    assert work["resultsStatus"] == "APPROVED"  # numbers only: the approval stands
    results = work["results"]
    results["goal"]["statement"] = "A different goal the student typed in"
    work = client.post(f"/api/works/{work['id']}/results", headers=H, json={"results": results, "baseVersion": work["resultsVersion"]}).json()
    refused = client.post(f"/api/works/{work['id']}/figures", headers=H)
    assert refused.status_code == 400 and refused.json()["code"] == "NOT_ONLY_FIGURES"


def test_assumed_sample_settings_no_longer_stop_the_proposal(client):
    """Owner decision 2026-10-08: no sample-size tick at Start. Assumed settings never stop Chapter One;
    nothing is recorded as the student's consent, and Chapter Three asks them to confirm (C3-ASSUMED)."""
    from tests.fake_models import PLAN as BASE

    client.models.overrides["p_finalise"] = lambda payload: {**BASE, "sampleSize": {**BASE["sampleSize"], "margin": 7, "population": 4200, "populationSource": "District records, 2025"}} if "title" in payload["draft"] else payload["draft"]
    from tests.test_proposals import DETAILS

    figures = {**DETAILS["inputs"], "populationSize": 4200, "populationSource": "District records, 2025"}  # the student's own N
    project = client.post("/api/projects", headers=H, json={**DETAILS, "inputs": figures}).json()
    client.post(f"/api/projects/{project['id']}/start", headers=H)
    project = _settled(client, f"/api/projects/{project['id']}", lambda p: bool(p["autoFailure"]) or any(c["number"] == 1 and c["current"] for c in p["chapters"]), timeout=180)
    if not project["plan"]["samplingAssumed"]:
        pytest.skip("this plan assumed no sample-size settings")
    assert not project["autoFailure"] and any(c["number"] == 1 and c["current"] for c in project["chapters"])
    assert not any(a["kind"] == "SAMPLING" for a in project["acknowledgments"])

    # Chapter Three asks for the confirmation, and the student gives it there (Codex audit 29343c2 #6)
    url = f"/api/projects/{project['id']}"
    client.post(f"{url}/chapters/1", headers=H, json={"version": next(c for c in project["chapters"] if c["number"] == 1)["current"], "approved": True})
    for step in ("CHAPTER_2", "CHAPTER_3"):
        quoted = client.post(f"{url}/steps", headers=H, json={"step": step}).json()
        assert "job" in quoted, quoted
        client.post(f"{url}/steps/{quoted['job']['id']}/submit", headers=H, json={"quoteId": quoted["quote"]["id"]})
        _settled(client, url, lambda p: not p["activeJob"], timeout=180)
    item = next(i for i in client.get(f"{url}/chapters/3", headers=H).json()["readiness"] if i["id"] == "C3-ASSUMED")
    assert item["status"] == "NEEDS_REVIEW"
    project = client.get(url, headers=H).json()
    confirmed = client.post(f"{url}/sampling", headers=H, json={"baseVersion": project["planVersion"]})
    assert confirmed.status_code == 200
    item = next(i for i in client.get(f"{url}/chapters/3", headers=H).json()["readiness"] if i["id"] == "C3-ASSUMED")
    assert item["status"] == "PASS"


def test_the_pdf_compiler_keeps_what_windows_needs():
    import os

    from app.latex.package import _tex_env

    env = _tex_env("C:/texlive/bin/pdflatex.exe" if os.name == "nt" else "/usr/bin/pdflatex", "tmp")
    assert env["openin_any"] == "p" and "PATH" in env
    if os.name == "nt":
        assert "SYSTEMROOT" in env



def test_a_short_assignment_question_does_not_block_start(works_client):
    """Live, 2026-10-01: a tester's seven-word question was refused three times by a twelve-word minimum
    the page never showed. A short question is still the question."""
    client = works_client
    work = client.post("/api/works", headers=H, json={"kind": "COURSEWORK", "variant": "ESSAY", "inputs": {"title": "Social media", "description": "Discuss the impact of social media on youth."}}).json()
    work = client.post(f"/api/works/{work['id']}/answers", headers=H, json={"answers": {"level": "LATER_UG", "citation_style": "APA7", "ai_policy": "NOT_MENTIONED"},
                                                                         "skipRest": True, "baseVersion": work["specVersion"]}).json()
    assert work["spec"]["gate"] == "PASS", work["spec"]["blockers"]  # no word limit given either: PaperAid uses the usual length
    assert client.post(f"/api/works/{work['id']}/start", headers=H).status_code == 200



# --- Codex's second verification (2026-10-01): durable continuation, reservations per attempt, reads ---------


def test_a_plan_whose_worker_stopped_before_the_draft_is_continued_by_maintenance(works_client, monkeypatch):
    from datetime import timedelta

    from app.jobs import service as jobs_service
    from app.runtime import get_runtime
    from app.works import service as works_service

    client = works_client
    real = works_service.continue_after_plan
    monkeypatch.setattr(works_service, "continue_after_plan", lambda rt, work_id, plan_job: None)  # the worker stops right after publishing
    work = _coursework(client)
    client.post(f"/api/works/{work['id']}/start", headers=H)
    work = _settled(client, f"/api/works/{work['id']}", lambda w: w["planStatus"] != "NONE" and w["activeJob"] is None)
    store = get_runtime().store
    pending = store.get_work(work["id"])
    assert pending.auto_next == work["jobs"][0] and len(work["jobs"]) == 1 and not work["documents"]
    monkeypatch.setattr(works_service, "continue_after_plan", real)
    store.update(pending.auto_next, lambda j: j.model_copy(update={"completed_at": j.completed_at - timedelta(minutes=5)}))
    assert jobs_service.resume_continuations(get_runtime()) == 1
    work = _settled(client, f"/api/works/{work['id']}", _done)
    assert work["documents"] and store.get_work(work["id"]).auto_next == ""
    assert jobs_service.resume_continuations(get_runtime()) == 0  # nothing pending: never a second draft


def test_a_start_that_loses_the_race_returns_only_its_own_reservation(works_client, monkeypatch):
    from app.core.errors import Conflict
    from app.pricing import credits
    from app.runtime import get_runtime
    from app.works import service as works_service

    client = works_client
    work = _coursework(client)
    store = get_runtime().store
    uid = store.get_work(work["id"]).owner_uid
    email = store.get_work(work["id"]).owner_email
    store.update_wallet(uid, email, lambda w: credits.reserve(w, f"work:{work['id']}:winner", 3000, "Reserved for your document"))  # the request that won

    def lost(*args, **kwargs):
        raise Conflict("PaperAid is already working on this.", code="STEP_RUNNING")

    monkeypatch.setattr(works_service, "quote_step", lost)
    refused = client.post(f"/api/works/{work['id']}/start", headers=H)
    assert refused.status_code == 409
    wallet = store.get_wallet(uid)
    assert list(wallet.reservations) == [f"work:{work['id']}:winner"] and wallet.held == 3000  # the winner's reservation stands


def test_reads_for_work_never_started_are_capped_per_day(works_client, monkeypatch):
    from app.runtime import get_runtime

    client = works_client
    monkeypatch.setattr(get_runtime().settings, "free_reads_per_day", 1)
    first = _with_brief(client)
    assert client.post(f"/api/works/{first['id']}/read", headers=H).status_code == 200
    _settled(client, f"/api/works/{first['id']}", lambda w: w["activeJob"] is None)
    second = _with_brief(client)
    refused = client.post(f"/api/works/{second['id']}/read", headers=H)
    assert refused.status_code == 400 and refused.json()["code"] == "READS_LIMIT"


def test_a_refused_read_leaves_the_work_ready_to_read_again(works_client, monkeypatch):
    """The page shows "not read yet" with Read my documents, never a spinner: no read job is left behind."""
    from app.runtime import get_runtime

    client = works_client
    work = _with_brief(client)
    monkeypatch.setattr(get_runtime().settings, "min_credits", {"COURSEWORK": 5000})
    assert client.post(f"/api/works/{work['id']}/read", headers=H).status_code == 402
    work = client.get(f"/api/works/{work['id']}", headers=H).json()
    assert work["needsRead"] and work["activeJob"] is None and work["jobs"] == []



def test_a_step_that_used_up_its_retries_never_says_it_will_keep_trying(works_client, monkeypatch):
    """Live, 2026-10-01: a provider outage outlasted the retries and the failed step still said "We'll keep trying"."""
    from app.ai.providers import UNAVAILABLE, UNAVAILABLE_FINAL
    from app.core.errors import RetryableStageError
    from app.runtime import get_runtime

    client = works_client
    monkeypatch.setattr(get_runtime().settings, "stage_max_attempts", 1)

    def down(payload):
        raise RetryableStageError("PROVIDER_UNAVAILABLE", UNAVAILABLE, "google 503")

    client.models.overrides["w_plan"] = down
    work = _coursework(client)
    client.post(f"/api/works/{work['id']}/start", headers=H)
    work = _settled(client, f"/api/works/{work['id']}", _done)
    assert work["autoFailure"] == UNAVAILABLE_FINAL and "keep trying" not in work["autoFailure"]


# --- Codex audit of 9239dd0 -----------------------------------------------------------------------------------


def test_a_start_refused_for_credits_leaves_nothing_started_and_nothing_reserved(works_client, monkeypatch):
    """The balance changed between the check and the submission: nothing is marked started (the page
    would wait for ever) and nothing is set aside, because both happen in the submission's transaction."""
    from app.pricing import credits
    from app.runtime import get_runtime
    from app.works import service as works_service

    client = works_client
    work = _coursework(client)
    monkeypatch.setattr(works_service, "check_credits", lambda *args, **kwargs: None)  # passed the check...
    store = get_runtime().store
    k = store.get_work(work["id"])
    store.update_wallet(k.owner_uid, k.owner_email, lambda w: credits.reserve(w, "elsewhere", w.available, "Spent elsewhere"))  # ...then spent
    refused = client.post(f"/api/works/{work['id']}/start", headers=H)
    assert refused.status_code == 402, refused.json()
    after = client.get(f"/api/works/{work['id']}", headers=H).json()
    assert not after["auto"] and after["activeJob"] is None and after["jobs"] == []
    assert list(store.get_wallet(k.owner_uid).reservations) == ["elsewhere"]


def test_a_start_that_fails_to_submit_reserves_nothing(works_client, monkeypatch):
    from app.core.errors import Conflict
    from app.runtime import get_runtime
    from app.works import service as works_service

    client = works_client
    work = _coursework(client)

    def lost(*args, **kwargs):
        raise Conflict("Your quote changed.", code="QUOTE_MISMATCH")

    monkeypatch.setattr(works_service, "submit", lost)
    assert client.post(f"/api/works/{work['id']}/start", headers=H).status_code == 409
    k = get_runtime().store.get_work(work["id"])
    assert not k.auto and get_runtime().store.get_wallet(k.owner_uid).reservations == {}


def test_a_proposal_start_reserves_with_its_plan_in_one_transaction(client):
    from app.runtime import get_runtime

    pid = _create(client)["id"]
    started = client.post(f"/api/projects/{pid}/start", headers=H, json={})
    assert started.status_code == 200, started.json()
    store = get_runtime().store
    p = store.get_project(pid)
    assert p.auto and p.jobs
    wallet = store.get_wallet(p.owner_uid)
    plan = store.get(p.jobs[0])
    assert plan.selection.proposal == "PLAN" and plan.selection.bundled
    assert all(key.startswith(f"project:{pid}:") for key in wallet.reservations)


def test_reads_of_new_works_have_a_firm_daily_ceiling(works_client, monkeypatch):
    """Counted atomically, whatever happens to the works (started ones included): at most twice the allowance."""
    from app.runtime import get_runtime

    client = works_client
    monkeypatch.setattr(get_runtime().settings, "free_reads_per_day", 1)
    codes = []
    for _ in range(3):
        work = _with_brief(client)
        read = client.post(f"/api/works/{work['id']}/read", headers=H)
        codes.append(read.status_code)
        _settled(client, f"/api/works/{work['id']}", lambda w: w["activeJob"] is None)
        if read.status_code == 200:
            get_runtime().store.update_work(work["id"], lambda w: w.model_copy(update={"auto": True}))  # started: not "unstarted"
    assert codes == [200, 200, 400]


def test_a_context_document_is_kept_in_file_storage_and_removed_with_its_request(works_client):
    from app.runtime import get_runtime

    client = works_client
    work = _coursework(client)
    client.post(f"/api/works/{work['id']}/start", headers=H)
    _settled(client, f"/api/works/{work['id']}", _done)
    note = Document()
    note.add_paragraph(" ".join(["evidence"] * 1500))
    buffer = io.BytesIO()
    note.save(buffer)
    sent = client.post(f"/api/works/{work['id']}/requests/with-document", headers=H, data={"instruction": "Use my notes", "sections": ""},
                       files={"file": ("notes.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}).json()
    rt = get_runtime()
    stored = rt.store.get_work(work["id"]).requests[0]
    assert stored.context == "" and len(rt.files.get(stored.context_path).decode("utf-8").split()) == 1000  # the page says: its first 1,000 words
    client.delete(f"/api/works/{work['id']}/requests/{sent['requests'][0]['id']}", headers=H)
    assert not rt.files.exists(stored.context_path)


def test_a_second_start_while_chapter_one_runs_is_refused_and_reserves_nothing(client):
    """Codex's second look at 9239dd0, kept now the sample-size stop is gone: a Start while the proposal is
    working is refused (it never marks a step pending again or stops the chapter), and nothing stays reserved."""
    from app.runtime import get_runtime

    project = _create(client)
    first = client.post(f"/api/projects/{project['id']}/start", headers=H)
    second = client.post(f"/api/projects/{project['id']}/start", headers=H)
    assert first.status_code == 200 and second.status_code == 409 and second.json()["code"] == "STEP_RUNNING"
    project = _settled(client, f"/api/projects/{project['id']}", lambda p: any(c["number"] == 1 and c["current"] for c in p["chapters"]) or bool(p["autoFailure"]), timeout=180)
    assert not project["autoFailure"]
    store = get_runtime().store
    owner = store.get_project(project["id"]).owner_uid
    assert store.get_wallet(owner) is None or store.get_wallet(owner).reservations == {}


# --- resume from the last good step (Codex roadmap item 1, 2026-10-03) ------------------------------------------


def test_start_after_an_outage_resumes_the_stopped_draft_where_it_stopped(works_client, monkeypatch):
    """A draft that stopped on a provider outage is resumed, not restarted: the same job, its finished
    stages kept (the plan and research are not redone), and the document is delivered."""
    from app.ai.providers import UNAVAILABLE
    from app.core.errors import RetryableStageError
    from app.runtime import get_runtime

    client = works_client
    monkeypatch.setattr(get_runtime().settings, "stage_max_attempts", 1)
    outage = {"on": True}
    real = client.models.overrides.get("w_draft")

    def draft(payload):
        if outage["on"]:
            raise RetryableStageError("PROVIDER_UNAVAILABLE", UNAVAILABLE, "google 503")
        from tests import fake_works

        return real(payload) if real else fake_works.answer("w_draft", payload)

    client.models.overrides["w_draft"] = draft
    work = _coursework(client)
    client.post(f"/api/works/{work['id']}/start", headers=H)
    work = _settled(client, f"/api/works/{work['id']}", _done)
    assert work["autoFailure"] and not work["documents"]
    store = get_runtime().store
    stopped = store.get(work["jobs"][-1])
    assert stopped.selection.work == "DRAFT" and stopped.status == "FAILED" and stopped.failure.retryable
    finished_before = list(stopped.completed_stages)

    def plans(w):  # this work's own planner calls: a previous test's queue thread can still reach the shared stand-in
        return sum(c.task == "w_plan" for j in w["jobs"] for c in store.get(j).model_calls)

    plans_before = plans(work)
    outage["on"] = False
    assert client.post(f"/api/works/{work['id']}/start", headers=H).status_code == 200
    work = _settled(client, f"/api/works/{work['id']}", _done)
    assert work["documents"] and not work["autoFailure"]
    assert work["jobs"][-1] == stopped.id  # the same step, resumed
    resumed = store.get(stopped.id)
    assert resumed.status == "COMPLETED" and any(a.action == "Resumed by the student" for a in resumed.admin_actions)
    assert all(s in resumed.completed_stages for s in finished_before)
    assert plans(work) == plans_before and len(work["jobs"]) == 2  # the plan was not made again


def test_a_failure_that_is_not_temporary_starts_afresh(works_client):
    client = works_client
    client.models.overrides["w_plan_review"] = lambda payload: {"verdict": "REPAIR", "rules": [{"rule": r["rule"], "status": "PASS", "note": "Met."} for r in payload.get("rules", [])],
                                                                "issues": ["The plan does not answer the question."], "suggestions": []}
    work = _coursework(client)
    client.post(f"/api/works/{work['id']}/start", headers=H)
    work = _settled(client, f"/api/works/{work['id']}", _done)
    first = work["jobs"][-1]
    client.models.overrides.pop("w_plan_review")
    client.post(f"/api/works/{work['id']}/start", headers=H)
    work = _settled(client, f"/api/works/{work['id']}", _done)
    assert work["documents"] and work["jobs"][-1] != first and len(work["jobs"]) >= 3  # a new plan and its draft, never the unapproved plan resumed
