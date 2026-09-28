"""Regressions for Codex's audit of 56c4f83 (docs/Codex_Full_Code_Audit_20260928_56c4f83.md). Each
test is one of Codex's reproductions with the assertion turned to the correct behaviour."""

import hashlib
import io
import json
from datetime import timedelta

from docx import Document

from app.core.auth import User
from app.documents.docx_io import read_docx
from app.jobs.models import Job, JobStatus, ReadinessItem, StoredFile, utcnow
from app.proposals import decisions, rulebook, service
from app.proposals.models import ChapterDocument, ChapterSection, StoredChapterVersion
from app.runtime import get_runtime
from tests import conftest as fixtures
from tests.test_api import ADMIN, STUDENT, wait
from tests.test_proposals import _create, plan

UID = "u_" + hashlib.sha256(b"student@example.com").hexdigest()[:20]
OWNER = User(uid=UID, email="student@example.com", is_admin=False)


def _paper(rt, *, words=None, scope=None, notes=None):
    data = fixtures.fixture_bytes("simple_essay.docx")
    model = read_docx(data)
    job = Job(
        id="job_auditpaper", status=JobStatus.DRAFT, owner_uid=UID, owner_email="student@example.com",
        expires_at=utcnow() + timedelta(days=30), scope_words=scope, fix_notes=notes or {},
    )
    path = f"{job.storage_prefix()}/input/source.docx"
    rt.files.put(path, data, "application/octet-stream")
    job.source = StoredFile(
        name="essay.docx", format="DOCX", size_bytes=len(data), word_count=words or model.word_count, page_estimate=1, heading_count=1,
        path=path, sha256=hashlib.sha256(data).hexdigest(),
    )
    rt.store.create(job)
    return job, model


def _project(client, number=1, *, unreviewed=False):
    """A project with an approved plan and a written, approved chapter (no AI)."""
    rt = get_runtime()
    pid = _create(client)["id"]
    p = rt.store.get_project(pid)
    approved = plan()
    sections = [
        ChapterSection(
            key=s.key, number=s.number, heading=s.heading, paragraphs=[f"The proposed study will examine caregiver access in {s.heading}."],
            depends=decisions.stamp(number, s.key, approved),
        )
        for s in rulebook.sections(p.rulebook, number, p.inputs.level, approved)
    ]
    chapter = ChapterDocument(
        number=number, title=rulebook.chapter_spec(p.rulebook, number)["title"], plan_version=1, sections=sections, cited=[], words=200,
        warnings=["Background was not reviewed."] if unreviewed else [],
        readiness=[
            ReadinessItem(
                id=f"C{number}-REVIEWED", question="Every section reviewed", basis="CODE", status="NEEDS_REVIEW" if unreviewed else "PASS",
                note="Background unreviewed" if unreviewed else "", chapter=number,
            )
        ],
    )
    path = f"{p.storage_prefix()}/chapters/{number}/seed.json"
    rt.files.put(path, chapter.model_dump_json(by_alias=True).encode(), "application/json")

    def set_up(q):
        q.plan, q.plan_status, q.plan_version = approved, "APPROVED", 1
        state = q.chapter(number)
        state.current, state.approved = 1, True
        state.versions = [StoredChapterVersion(version=1, job_id="job_seed", words=200, plan_version=1, path=path)]
        return q

    rt.store.update_project(pid, set_up)
    return pid, chapter


# --- H01 ------------------------------------------------------------------------------------------


def test_h01_admin_and_expired_views_never_carry_fix_notes(client):
    from app.jobs.service import without_paper_text

    rt = get_runtime()
    job, _ = _paper(rt, notes={"b00001": ["PRIVATE author sentence and confidential fix instruction"]})
    response = client.get(f"/api/admin/jobs/{job.id}", headers=ADMIN)
    assert response.status_code == 200 and "PRIVATE" not in json.dumps(response.json())
    assert without_paper_text(job).fix_notes == {}


def test_h01_every_job_field_is_classified_for_support_views():
    """A new Job field must be declared paper-bearing (emptied for support and expiry) or support
    metadata; this fails until it is."""
    from app.jobs.service import PAPER_FIELDS, SUPPORT_FIELDS

    assert not PAPER_FIELDS & SUPPORT_FIELDS
    assert set(Job.model_fields) == PAPER_FIELDS | SUPPORT_FIELDS


# --- H02 ------------------------------------------------------------------------------------------


def test_h02_deleting_the_account_erases_the_credit_history(client):
    rt = get_runtime()
    assert rt.store.ledger(UID, None, 10)  # the test credits
    rt.settings.credits_enabled = False
    assert client.delete("/api/me", headers=STUDENT).status_code == 204
    assert rt.store.get_wallet(UID).closing  # the tombstone still refuses late requests
    assert rt.store.ledger(UID, None, 10) == []
    rt.store.update_wallet(UID, "", lambda w: w)  # touching the tombstone adds nothing back
    assert rt.store.ledger(UID, None, 10) == []


# --- H03 ------------------------------------------------------------------------------------------


def test_h03_a_guide_arriving_after_deletion_leaves_nothing_behind(client, monkeypatch):
    rt = get_runtime()
    pid = _create(client)["id"]
    prefix = rt.store.get_project(pid).storage_prefix()
    original_put = rt.files.put
    paths = []

    def racing_put(path, data, content_type):
        if path.startswith(prefix + "/guide/"):
            service.delete(rt, OWNER, pid)
            paths.append(path)
        return original_put(path, data, content_type)

    monkeypatch.setattr(rt.files, "put", racing_put)
    doc = Document()
    doc.add_paragraph("PRIVATE institution research guide " * 120)
    out = io.BytesIO()
    doc.save(out)
    response = client.post(f"/api/projects/{pid}/guide", headers=STUDENT, files={"file": ("guide.docx", out.getvalue(), "application/octet-stream")})
    assert response.status_code == 404 and rt.store.get_project(pid) is None
    assert paths and not rt.files.exists(paths[0])


# --- H04 ------------------------------------------------------------------------------------------


def _legacy(client):
    rt = get_runtime()
    pid, chapter = _project(client)
    legacy_path = f"users/{UID}/projects/{pid}/chapters/1/legacy.json"
    rt.files.put(legacy_path, chapter.model_dump_json(by_alias=True).encode(), "application/json")

    def legacy(p):
        p.chapter(1).versions[0].path = legacy_path
        return p

    rt.store.update_project(pid, legacy)
    return rt, pid, legacy_path


def test_h04_deleting_a_project_removes_files_under_the_old_prefix(client):
    rt, pid, legacy_path = _legacy(client)
    assert client.get(f"/api/projects/{pid}/chapters/1", headers=STUDENT).status_code == 200
    assert client.delete(f"/api/projects/{pid}", headers=STUDENT).status_code == 204
    assert not rt.files.exists(legacy_path)


def test_h04_old_projects_are_migrated_out_of_the_fixed_age_prefix(client):
    from app.jobs import service as jobs

    rt, pid, legacy_path = _legacy(client)
    assert jobs.migrate_legacy_project_files(rt) == 1
    project = rt.store.get_project(pid)
    moved = project.chapter(1).versions[0].path
    assert moved.startswith(project.storage_prefix()) and rt.files.exists(moved) and not rt.files.exists(legacy_path)
    assert client.get(f"/api/projects/{pid}/chapters/1", headers=STUDENT).status_code == 200
    assert jobs.migrate_legacy_project_files(rt) == 0  # safe to repeat


def test_h04_a_step_priced_before_the_migration_still_reads_its_files(client):
    from app.jobs import service as jobs
    from app.proposals.pipeline import load_library

    rt, pid, legacy_path = _legacy(client)
    evidence_path = f"users/{UID}/projects/{pid}/evidence/job_old.json"
    rt.files.put(evidence_path, b"[]", "application/json")
    rt.store.update_project(pid, lambda p: p.model_copy(update={"evidence_files": [evidence_path]}))
    jobs.migrate_legacy_project_files(rt)
    assert not rt.files.exists(evidence_path)
    assert load_library(rt.files, [evidence_path]) == {}  # read from the moved copy, not a crash


# --- H05-H07, M09: revisions --------------------------------------------------------------------


def _comment(client, pid, key="problem", text="Clarify the problem statement."):
    response = client.post(f"/api/projects/{pid}/feedback", headers=STUDENT, json={"text": text})
    cid = response.json()["feedback"][-1]["id"]
    placed = client.post(f"/api/projects/{pid}/feedback/{cid}", headers=STUDENT, json={"chapter": 1, "sections": [key]})
    assert placed.status_code == 200, placed.text
    return cid


def _quote(client, pid):
    quoted = client.post(f"/api/projects/{pid}/steps", headers=STUDENT, json={"step": "REVISE_1"})
    assert quoted.status_code == 200, quoted.text
    return quoted.json()


def _submit(client, pid, body):
    return client.post(f"/api/projects/{pid}/steps/{body['job']['id']}/submit", headers=STUDENT, json={"quoteId": body["quote"]["id"]})


def _revise_only(keys):
    """The writer answers the comments on `keys` and returns nothing for the other sections."""

    def answer(payload):
        return {
            "sections": [
                {"key": s["key"], "paragraphs": [*s["text"], "The proposed study will address this comment."] if s["key"] in keys else [], "table": s["table"]}
                for s in payload["sections"]
            ]
        }

    return answer


def _feedback(client, pid, cid):
    return next(c for c in client.get(f"/api/projects/{pid}", headers=STUDENT).json()["feedback"] if c["id"] == cid)


def test_h05_a_revision_that_changes_nothing_fails_without_charge(fixed_client):
    pid, _ = _project(fixed_client)
    cid = _comment(fixed_client, pid)
    fixed_client.models.overrides["p_fix"] = _revise_only(set())
    body = _quote(fixed_client, pid)
    assert _submit(fixed_client, pid, body).status_code == 200
    job = wait(fixed_client, body["job"]["id"])
    assert job["status"] == "FAILED" and job["failure"]["code"] == "NOTHING_REVISED" and job["billing"]["charged"] == 0
    assert _feedback(fixed_client, pid, cid)["status"] == "OPEN"
    assert next(c for c in fixed_client.get(f"/api/projects/{pid}", headers=STUDENT).json()["chapters"] if c["number"] == 1)["current"] == 1


def test_h05_a_partial_revision_charges_its_share_and_leaves_the_rest_open(fixed_client):
    pid, _ = _project(fixed_client)
    done = _comment(fixed_client, pid, "problem")
    missed = _comment(fixed_client, pid, "scope", "Narrow the scope.")
    fixed_client.models.overrides["p_fix"] = _revise_only({"problem"})
    body = _quote(fixed_client, pid)
    assert body["quote"]["amount"] == 2000
    assert _submit(fixed_client, pid, body).status_code == 200
    job = wait(fixed_client, body["job"]["id"])
    assert job["status"] == "COMPLETED" and job["outcome"] == "PARTIAL" and job["billing"]["charged"] == 1000
    assert _feedback(fixed_client, pid, done)["status"] == "APPLIED" and _feedback(fixed_client, pid, missed)["status"] == "OPEN"


def test_h06_untouched_sections_keep_their_review_state(client):
    pid, _ = _project(client, unreviewed=True)
    _comment(client, pid)
    client.models.overrides["p_fix"] = _revise_only({"problem"})
    body = _quote(client, pid)
    _submit(client, pid, body)
    assert wait(client, body["job"]["id"])["status"] == "COMPLETED"
    chapter = client.get(f"/api/projects/{pid}/chapters/1", headers=STUDENT).json()
    assert next(r for r in chapter["readiness"] if r["id"] == "C1-REVIEWED")["status"] == "NEEDS_REVIEW"
    assert any("Background" in w for w in chapter["warnings"])


def _second_version(rt, pid, base, text="NEWER background the student chose."):
    newer = base.model_copy(deep=True)
    next(s for s in newer.sections if s.key == "background").paragraphs = [text]
    path = f"{rt.store.get_project(pid).storage_prefix()}/chapters/1/newer.json"
    rt.files.put(path, newer.model_dump_json(by_alias=True).encode(), "application/json")

    def change(p):
        p.chapter(1).versions.append(StoredChapterVersion(version=2, job_id="job_newer", words=200, plan_version=1, path=path))
        p.chapter(1).current = 2
        return p

    rt.store.update_project(pid, change)


def test_h07_a_revision_priced_on_another_version_must_be_priced_again(client):
    rt = get_runtime()
    pid, base = _project(client)
    _comment(client, pid)
    body = _quote(client, pid)
    _second_version(rt, pid, base)
    refused = _submit(client, pid, body)
    assert refused.status_code == 409 and refused.json()["code"] == "QUOTE_MISMATCH"


def test_h07_a_version_chosen_during_the_revision_stays_current(client):
    rt = get_runtime()
    pid, base = _project(client)
    _comment(client, pid)
    revise = _revise_only({"problem"})

    def choose_meanwhile(payload):
        if rt.store.get_project(pid).chapter(1).current == 1:
            _second_version(rt, pid, base)
        return revise(payload)

    client.models.overrides["p_fix"] = choose_meanwhile
    body = _quote(client, pid)
    _submit(client, pid, body)
    job = wait(client, body["job"]["id"])
    state = next(c for c in client.get(f"/api/projects/{pid}", headers=STUDENT).json()["chapters"] if c["number"] == 1)
    assert state["current"] == 2 and len(state["versions"]) == 3
    assert any("without replacing your choice" in w for w in job["warnings"])


def test_m09_moving_a_comment_after_pricing_needs_a_new_price(client):
    pid, _ = _project(client)
    cid = _comment(client, pid)
    body = _quote(client, pid)
    client.post(f"/api/projects/{pid}/feedback/{cid}", headers=STUDENT, json={"chapter": 1, "sections": ["background"]})
    refused = _submit(client, pid, body)
    assert refused.status_code == 409 and refused.json()["code"] == "QUOTE_MISMATCH"


def test_m09_a_comment_moved_during_the_revision_stays_open(client):
    rt = get_runtime()
    pid, _ = _project(client)
    cid = _comment(client, pid)
    revise = _revise_only({"problem"})

    def move_meanwhile(payload):
        service.update_feedback(rt, OWNER, pid, cid, 1, ["background"], "OPEN", "")
        return revise(payload)

    client.models.overrides["p_fix"] = move_meanwhile
    body = _quote(client, pid)
    _submit(client, pid, body)
    assert wait(client, body["job"]["id"])["status"] == "COMPLETED"
    assert _feedback(client, pid, cid)["status"] == "OPEN"


# --- H08: fix drafts -----------------------------------------------------------------------------


def _fix_draft(client):
    from app.jobs.models import AnalysisResult, Finding

    rt = get_runtime()
    job, _ = _paper(rt)
    doc = Document()
    for i in range(30):
        doc.add_paragraph(" ".join([f"passage{i}"] * 150))
    out = io.BytesIO()
    doc.save(out)
    model = read_docx(out.getvalue())
    rt.files.put(job.source.path, out.getvalue(), "application/octet-stream")
    job.source.word_count, job.source.sha256 = model.word_count, hashlib.sha256(out.getvalue()).hexdigest()
    blocks = [b.id for b in model.blocks if b.editable and b.kind == "paragraph"]
    job.status = JobStatus.COMPLETED
    finding = Finding(id="f-one", block_id=blocks[0], section="Body", reason="GENERIC_PHRASING", severity="minor", excerpt="", explanation="Clarify", suggestion="Be specific")
    job.analysis = AnalysisResult(band="LOW", confidence="HIGH", analysed_words=model.word_count, excluded_words=0, algorithm_version="audit", findings=[finding])
    rt.store.update(job.id, lambda j: job)
    rt.files.put(f"{job.storage_prefix()}/internal/document.json", model.model_dump_json(by_alias=True).encode(), "application/json")
    created = client.post(f"/api/jobs/{job.id}/fix", headers=STUDENT, json={"findingIds": ["f-one"], "safeOnly": False})
    assert created.status_code == 200, created.text
    return created.json()["id"], blocks


def test_h08_widening_a_fix_selection_is_priced_by_the_words_chosen(fixed_client):
    draft, blocks = _fix_draft(fixed_client)
    selection = {"writing": "REFINE", "academic": False, "onlyBlocks": [blocks[0]]}
    first = fixed_client.post(f"/api/jobs/{draft}/quote", headers=STUDENT, json={"selection": selection}).json()["quote"]["amount"]
    widened = fixed_client.post(f"/api/jobs/{draft}/quote", headers=STUDENT, json={"selection": {**selection, "onlyBlocks": blocks}}).json()["quote"]["amount"]
    assert first == 4000 and widened == 7000  # 150 words, then the full 4,500


def test_h08_a_selection_must_name_passages_in_the_file(fixed_client):
    draft, _ = _fix_draft(fixed_client)
    selection = {"writing": "REFINE", "academic": False, "onlyBlocks": ["b99999"]}
    refused = fixed_client.post(f"/api/jobs/{draft}/quote", headers=STUDENT, json={"selection": selection})
    assert refused.status_code == 400 and refused.json()["code"] == "INVALID_SELECTION"


def test_h08_replacing_the_file_clears_the_fix_selection(fixed_client):
    draft, _ = _fix_draft(fixed_client)
    other = {"file": ("other.docx", fixtures.fixture_bytes("simple_essay.docx"), "application/octet-stream")}
    assert fixed_client.post(f"/api/jobs/{draft}/files/source", headers=STUDENT, files=other).status_code == 200
    job = get_runtime().store.get(draft)
    assert job.fix_notes == {} and job.scope_words is None and job.selection.only_blocks == []


# --- M10-M13: credit history and cleanup ------------------------------------------------------


def test_m10_an_interrupted_wallet_write_leaves_balance_and_history_in_agreement(tmp_path, monkeypatch):
    from pathlib import Path

    import pytest

    from app.integrations.store import LocalJobStore
    from app.pricing import credits

    store = LocalJobStore(tmp_path)
    store.update_wallet("u", "u@example.com", lambda w: credits.top_up(w, 1000, "committed"))
    original = Path.replace

    def fail_wallet_replace(path, target):
        if Path(target) == store._wallet_path("u"):
            raise OSError("synthetic interrupted wallet replace")
        return original(path, target)

    monkeypatch.setattr(Path, "replace", fail_wallet_replace)
    with pytest.raises(OSError):
        store.update_wallet("u", "u@example.com", lambda w: credits.top_up(w, 5000, "interrupted"))
    monkeypatch.undo()
    wallet = store.get_wallet("u")  # the journal is replayed: the change is whole, never half
    history = store.ledger("u", None, 50)
    assert wallet.available == 6000 and [e.note for e in history] == ["interrupted", "committed"]
    assert history[0].available_after == wallet.available


def test_m10_a_torn_last_history_line_is_ignored(tmp_path):
    from app.integrations.store import LocalJobStore
    from app.pricing import credits

    store = LocalJobStore(tmp_path)
    store.update_wallet("u", "u@example.com", lambda w: credits.top_up(w, 1000, "committed"))
    with store._ledger_path("u").open("a", encoding="utf-8") as ledger:
        ledger.write('{"id": "le_torn", "at": "2026-')
    assert [e.note for e in store.ledger("u", None, 50)] == ["committed"]


def test_m11_entries_sharing_a_time_are_all_paged(client):
    from app.jobs.models import LedgerEntry

    rt = get_runtime()
    at = utcnow()

    def add(w):
        w.entries += [LedgerEntry(id=f"le_same{i:03}", kind="TOP_UP", amount=1000, at=at, note="same-time", available_after=1000, held_after=0) for i in range(60)]
        return w

    rt.store.update_wallet(UID, "student@example.com", add)
    seen, before = [], None
    while True:
        page = client.get("/api/wallet/history", headers=STUDENT, params={"before": before} if before else {}).json()
        seen += [e["id"] for e in page["entries"] if e["note"] == "same-time"]
        before = page["next"]
        if not before:
            break
    assert len(seen) == 60 == len(set(seen))


def test_m12_entries_from_before_the_complete_history_are_copied_in_once(client):
    from app.jobs import service as jobs
    from app.jobs.models import Wallet
    from app.pricing import credits

    rt = get_runtime()
    legacy = Wallet(uid="u_legacy", email="legacy@example.com")
    credits.top_up(legacy, 4000, "before-upgrade")
    rt.store._wallet_path(legacy.uid).write_text(legacy.model_dump_json(by_alias=True), encoding="utf-8")
    assert rt.store.ledger(legacy.uid, None, 50) == []
    assert jobs.backfill_ledgers(rt) >= 1
    assert [e.note for e in rt.store.ledger(legacy.uid, None, 50)] == ["before-upgrade"]
    assert jobs.backfill_ledgers(rt) == 0  # once per wallet
    assert [e.note for e in rt.store.ledger(legacy.uid, None, 50)] == ["before-upgrade"]


def test_m13_a_project_that_cannot_be_erased_does_not_hide_the_next(client):
    from app.jobs import service as jobs

    rt = get_runtime()
    first, second = _create(client)["id"], _create(client)["id"]
    for pid, days in ((first, 3), (second, 2)):
        rt.store.update_project(pid, lambda p, days=days: p.model_copy(update={"expires_at": utcnow() - timedelta(days=days)}))
    blocked = Job(id="job_blockcleanup", owner_uid=UID, owner_email="student@example.com", project_id=first, status=JobStatus.QUEUED, expires_at=utcnow() + timedelta(days=30))
    rt.store.create(blocked)
    rt.store.update_project(first, lambda p: p.model_copy(update={"active_job": blocked.id}))
    assert jobs.cleanup_expired_projects(rt) == 1
    assert rt.store.get_project(first) is not None and rt.store.get_project(second) is None
