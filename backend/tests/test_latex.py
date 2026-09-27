"""LaTeX conversion: deterministic, faithful, escaped, and honest about what it cannot convert."""

import io
import shutil
import zipfile

import pytest
from docx import Document

from app.core.config import Settings
from app.latex.convert import convert, escape
from app.latex.package import compile_pdf, package
from app.pricing.quote import latex_ugx
from tests.conftest import fixture_bytes


def _docx(*paragraphs: str) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _tex_runs() -> bool:
    """Compile tests need an engine that actually starts (a sandbox can block MiKTeX)."""
    if shutil.which("pdflatex") is None:
        return False
    pdf, _ = compile_pdf(convert(_docx("Probe.")))
    return pdf is not None


HAS_TEX = _tex_runs()


def test_the_students_text_can_never_become_latex_commands():
    tex = convert(_docx(r"Try \input{/etc/passwd} and \write18{rm -rf} with 50% of $5 & x_1 {braces} ~ ^ #")).tex
    body = tex.split(r"\begin{document}")[1]
    assert r"\input{" not in body and r"\write18" not in body
    assert r"\textbackslash{}input\{/etc/passwd\}" in body and r"50\%" in body and r"\$5" in body and r"x\_1" in body


def test_typography_and_symbols_map_to_latex():
    assert escape("“Quoted” — 10°C ≥ 3 α") == r"``Quoted'' --- 10\textdegree{}C $\geq$ 3 $\alpha$"


def test_structure_is_carried_over_faithfully():
    tex = convert(fixture_bytes("lists_and_captions.docx")).tex
    assert r"\begin{enumerate}" in tex and r"\begin{itemize}" in tex and r"\begin{tabular}" in tex
    assert r"\href{https://www.example.org/census-2024}{the census portal}" in convert(fixture_bytes("hyperlinks.docx")).tex
    assert r"\footnote{" in convert(fixture_bytes("footnotes.docx")).tex
    equation = convert(fixture_bytes("equation.docx"))
    assert (equation.equations, equation.equations_converted) == (1, 1) and r"\frac{dN}{dt}" in equation.tex


def test_headings_lose_manual_numbers_and_references_are_kept_as_written():
    doc = Document()
    doc.add_heading("2.1 Background", 2)
    doc.add_paragraph("Body text.")
    doc.add_heading("References", 1)
    doc.add_paragraph("Okello, J. (2021). Phones as tills. Journal, 5(3), 101–118.")
    out = io.BytesIO()
    doc.save(out)
    tex = convert(out.getvalue()).tex
    assert r"\subsection{Background}" in tex and r"\section*{References}" in tex
    assert r"\item[] Okello, J. (2021). Phones as tills. Journal, 5(3), 101--118." in tex


def test_an_equation_it_cannot_convert_is_reported_never_guessed():
    doc = Document()
    p = doc.add_paragraph("A matrix: ")
    from lxml import etree

    m = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    math = etree.SubElement(p._p, f"{{{m}}}oMath")
    matrix = etree.SubElement(math, f"{{{m}}}m")
    etree.SubElement(etree.SubElement(etree.SubElement(matrix, f"{{{m}}}mr"), f"{{{m}}}e"), f"{{{m}}}r").append(etree.Element(f"{{{m}}}t"))
    out = io.BytesIO()
    doc.save(out)
    result = convert(out.getvalue())
    assert (result.equations, result.equations_converted) == (1, 0)
    assert "% Equation 1 was not converted automatically" in result.tex and "Retype it" in result.warnings[0]


@pytest.mark.skipif(not HAS_TEX, reason="no LaTeX engine installed")
@pytest.mark.parametrize("name", ["simple_essay.docx", "equation.docx", "lists_and_captions.docx", "footnotes.docx", "hyperlinks.docx", "interleaved_tables.docx", "citation_fields.docx", "dissertation_long.docx"])
def test_every_conversion_compiles(name):
    pdf, problem = compile_pdf(convert(fixture_bytes(name)))
    assert pdf is not None and pdf[:4] == b"%PDF", problem


def test_the_download_holds_the_project_and_says_what_to_check():
    archive, result, compiled, _ = package(fixture_bytes("equation.docx"))
    names = set(zipfile.ZipFile(io.BytesIO(archive)).namelist())
    assert {"main.tex", "README.txt"} <= names and ("main.pdf" in names) == compiled
    readme = zipfile.ZipFile(io.BytesIO(archive)).read("README.txt").decode()
    assert "does not convert references to BibTeX" in readme


def test_latex_is_the_owners_fixed_price():
    settings = Settings(openai_api_key="sk-test", anthropic_api_key="sk-test")
    assert (latex_ugx(settings, 3000), latex_ugx(settings, 10000)) == (3000, 5100)  # UGX 150 per 300 words, minimum 3,000


def _png() -> bytes:
    import struct
    import zlib

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    raw = b"".join(b"\x00" + b"\xff\x00\x00" * 8 for _ in range(8))  # an 8x8 red square
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 8, 8, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def test_images_become_figure_files_and_compile():
    doc = Document()
    doc.add_paragraph("Figure 1 shows the sites.")
    doc.add_picture(io.BytesIO(_png()))
    doc.add_paragraph("Figure 1: Study sites", style="Caption")
    out = io.BytesIO()
    doc.save(out)
    result = convert(out.getvalue())
    assert result.figures == 1 and list(result.files) == ["figures/figure1.png"] and r"\includegraphics" in result.tex
    if HAS_TEX:
        pdf, problem = compile_pdf(result)
        assert pdf is not None, problem


def _linked_docx(target: str) -> bytes:
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    p = doc.add_paragraph("See ")
    rid = p.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "the link"
    run.append(text)
    link.append(run)
    p._p.append(link)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


HOSTILE = r"https://x.org/}{a}\typeout{PWNED}\iffalse"


def _body(target: str) -> str:
    return convert(_linked_docx(target)).tex.split(r"\begin{document}")[1]


def test_a_hostile_link_target_can_never_become_latex_commands():
    r"""Codex audit 2026-09-27 #1: a crafted target closed \href's braces and injected \typeout."""
    body = _body(HOSTILE)
    assert r"\typeout" not in body and "}{a}" not in body
    assert r"\href{https://x.org/\%7D\%7Ba\%7D\%5Ctypeout\%7BPWNED\%7D\%5Ciffalse}{the link}" in body
    for target in ("javascript:alert(1)", "file:///etc/passwd", r"\input{x}"):
        assert r"\href" not in _body(target)  # other schemes stay plain text
    assert r"\href{https://example.org/a?q=1&x=2\#top}{the link}" in _body("https://example.org/a?q=1&x=2#top")


@pytest.mark.skipif(not HAS_TEX, reason="no working LaTeX engine")
def test_hostile_and_ordinary_links_compile():
    for target in (HOSTILE, "https://example.org/a?q=1&x=2#top", "mailto:someone@example.org"):
        pdf, problem = compile_pdf(convert(_linked_docx(target)))
        assert pdf is not None, (target, problem)


def test_anything_left_out_makes_the_latex_job_partial():
    """Codex audit #14: an omitted item with a successful compile was reported as a full success."""
    from lxml import etree

    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    doc = Document()
    p = doc.add_paragraph("A chart: ")
    etree.SubElement(etree.SubElement(p._p, w + "r"), w + "object")
    out = io.BytesIO()
    doc.save(out)
    result = convert(out.getvalue())
    assert result.omitted == 1 and "left out" in result.warnings[0]
