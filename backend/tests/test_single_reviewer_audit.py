"""Codex's audit of the one-final-reviewer release (2026-10-01): no approval despite a failed rule,
proposal publication settled atomically, older works steps refused when validators changed, each
review part complete on its own, edits invalidating an earlier review, one layout for the review and
the Word file, and profile repair."""

import io

import pytest
from docx import Document

from app.jobs.models import Stage
from app.works import pipeline as works_pipeline
from tests import fake_works
from tests.test_api import STUDENT
from tests.test_single_reviewer import _drafted, _requests
from tests.test_works import _coursework, _run, _work


@pytest.fixture
def works_client(tmp_path, monkeypatch):
    monkeypatch.setenv("WORKS_ENABLED", '["CONCEPT_NOTE","COURSEWORK","FUNDING_PROPOSAL"]')
    from tests.conftest import _client

    yield from _client(tmp_path, monkeypatch, "cost")


def _funding(client):
    created = client.post("/api/works", headers=STUDENT, json={
        "kind": "FUNDING_PROPOSAL", "variant": "NGO_PROJECT", "mode": "COMPACT",
        "inputs": {"title": "Safer deliveries in Kamuli", "description": "Too many mothers in Kamuli deliver at home and die of preventable complications.",
                   "answers": {"problem": "Home deliveries and delayed referral.", "intervention": "Train village health teams.", "duration_months": "12", "currency": "USD"}},
    }).json()
    work = client.post(f"/api/works/{created['id']}/answers", headers=STUDENT, json={"answers": {}, "skipRest": True, "baseVersion": created["specVersion"]}).json()
    client.post(f"/api/works/{work['id']}/spec/confirm", headers=STUDENT, json={"baseVersion": work["specVersion"]})
    return work


# --- 1: a PASS with a failed rule, or an incomplete answer, is never an approval -------------------------


def test_a_plan_passed_with_issues_is_not_approved(works_client):
    """Codex's live-shaped case: the reviewer says PASS but still lists an issue."""
    client = works_client
    client.models.overrides["w_plan_review"] = lambda payload: {"verdict": "PASS", "rules": [{"rule": r["rule"], "status": "PASS", "note": "Met."} for r in payload.get("rules", [])],
                                                                "issues": ["The conclusion does not answer the question directly."]}
    work = _coursework(client)
    _run(client, work["id"], "PLAN")
    decision = _work(client, work["id"])["planReview"]
    assert decision["outcome"] == "OBJECTIONS" and "does not answer the question directly" in decision["objections"][0]


def test_a_plan_verdict_needs_every_rule_judged_and_none_failed():
    from app.works.ai import PlanReview, RuleVerdict

    rules = [{"rule": "CW-010", "requirement": "x"}, {"rule": "CW-011", "requirement": "y"}]
    passed = [RuleVerdict(rule="CW-010", status="PASS", note="Met."), RuleVerdict(rule="CW-011", status="PASS", note="Met.")]
    assert works_pipeline._plan_decision(PlanReview(verdict="PASS", rules=passed, issues=[]), rules, [], "REVIEW_UNAVAILABLE").outcome == "APPROVED"
    failed = [passed[0], RuleVerdict(rule="CW-011", status="FAIL", note="Not planned.")]
    decision = works_pipeline._plan_decision(PlanReview(verdict="PASS", rules=failed, issues=[]), rules, [], "REVIEW_UNAVAILABLE")
    assert decision.outcome == "OBJECTIONS" and decision.objections == ["CW-011: Not planned."]  # a PASS with a failed rule
    missing = works_pipeline._plan_decision(PlanReview(verdict="PASS", rules=passed[:1], issues=[]), rules, [], "REVIEW_UNAVAILABLE")
    assert missing.outcome == "NOT_REVIEWED" and "CW-011" in missing.objections[0]
    assert works_pipeline._plan_decision(None, rules, [], "SPEND_CAP").reason == "SPEND_CAP"


def test_a_results_model_with_a_failed_rule_or_unclassified_result_is_not_approved(works_client):
    client = works_client

    def review(payload):
        answer = fake_works.answer("w_results_review", payload)
        answer["rules"] = [{**r, "status": "FAIL", "note": "Indicators lack sources."} for r in answer["rules"]][:1] + answer["rules"][1:]
        return answer

    client.models.overrides["w_results_review"] = review
    work = _funding(client)
    _run(client, work["id"], "PLAN")
    assert _work(client, work["id"])["resultsReview"]["outcome"] == "OBJECTIONS"

    client.models.overrides["w_results_review"] = lambda payload: {**fake_works.answer("w_results_review", payload), "classified": []}  # nothing classified
    work = _funding(client)
    _run(client, work["id"], "PLAN")
    decision = _work(client, work["id"])["resultsReview"]
    assert decision["outcome"] == "NOT_REVIEWED" and "G1" in decision["objections"][0]

    client.models.overrides["w_results_review"] = lambda payload: {**fake_works.answer("w_results_review", payload), "rules": []}  # no rule judged
    work = _funding(client)
    _run(client, work["id"], "PLAN")
    assert _work(client, work["id"])["resultsReview"]["outcome"] == "NOT_REVIEWED"


# --- 2: a proposal step publishes, completes and settles in one transaction ----------------------------------


def test_an_error_after_a_proposal_is_published_never_refunds_it(client, monkeypatch):
    from app.proposals import pipeline as proposal_pipeline
    from app.runtime import get_runtime
    from tests.test_proposals import _create
    from tests.test_proposals import _run as run_step

    real = proposal_pipeline.STAGES[Stage.EXPORTING]

    def export_then_crash(ctx):
        real(ctx)
        raise RuntimeError("the worker was lost after the export committed")

    monkeypatch.setitem(proposal_pipeline.STAGES, Stage.EXPORTING, export_then_crash)
    project = _create(client)
    job = run_step(client, project["id"], "PLAN")
    rt = get_runtime()
    assert job["status"] == "COMPLETED" and rt.store.get(job["id"]).billing.state == "SETTLED"
    assert job["id"] in rt.store.get_project(project["id"]).published


# --- 3: a works step priced on older validators is refused, not run on new ones --------------------------


def test_a_step_priced_on_older_validators_is_refused_and_refunded(works_client):
    from app.runtime import get_runtime

    client = works_client
    work = _coursework(client)
    quoted = client.post(f"/api/works/{work['id']}/steps", headers=STUDENT, json={"step": "PLAN"}).json()
    rt = get_runtime()

    def older(j):
        j.quote.engine.content["validators"] = "validators-v1"  # priced before this release
        return j

    rt.store.update(quoted["job"]["id"], older)
    client.post(f"/api/works/{work['id']}/steps/{quoted['job']['id']}/submit", headers=STUDENT, json={"quoteId": quoted["quote"]["id"]})
    from tests.test_api import wait

    job = wait(client, quoted["job"]["id"], timeout=60)
    assert job["status"] == "FAILED" and job["failure"]["code"] == "ENGINE_CHANGED" and rt.store.get(job["id"]).billing.state == "RELEASED"


# --- 4: each part of a long review is complete on its own; a long section is split -------------------------


def test_a_rule_omitted_in_one_part_is_not_hidden_by_another(works_client, monkeypatch):
    from app.runtime import get_runtime

    monkeypatch.setattr(works_pipeline, "FINAL_PART_WORDS", 900)  # whole requests: about 460 words of every one are repeated context (Codex review 2026-10-07, finding 8)
    client = works_client

    def final(payload):
        answer = fake_works.answer("w_final", payload)
        if payload["part"].startswith("1 of"):
            answer["rules"] = answer["rules"][1:]  # the first part leaves a rule without a verdict
        return answer

    client.models.overrides["w_final"] = final
    work = _drafted(client)
    _, job = _run(client, work["id"], "DRAFT")
    assert job["status"] == "FAILED" and job["failure"]["code"] == "REVIEW_UNAVAILABLE"
    assert get_runtime().store.get(job["id"]).billing.state == "RELEASED"


def test_a_section_longer_than_one_part_is_split_across_parts(works_client, monkeypatch):
    from app.ai.orchestration import payload_words

    monkeypatch.setattr(works_pipeline, "FINAL_PART_WORDS", 900)  # whole requests: about 460 words of every one are repeated context (Codex review 2026-10-07, finding 8)
    client = works_client
    work = _drafted(client)
    _, job = _run(client, work["id"], "DRAFT")
    assert job["status"] == "COMPLETED", job.get("failure")
    parts = _requests(client, "w_final")
    headings = [s["heading"] for p in parts for s in p["document"]["sections"]]
    assert any("(continued, piece" in h for h in headings)
    assert all(payload_words(p) <= 900 for p in parts)


# --- 5: an edit replaces the earlier review ------------------------------------------------------------------


def test_editing_a_reviewed_plan_needs_confirmation_before_approval(works_client):
    client = works_client
    work = _coursework(client)
    _run(client, work["id"], "PLAN")
    work = _work(client, work["id"])
    assert work["planReview"]["outcome"] == "APPROVED"
    plan = work["plan"]
    plan["sections"][0]["brief"] = "A different opening, written by the student."
    work = client.post(f"/api/works/{work['id']}/plan", headers=STUDENT, json={"plan": plan, "baseVersion": work["planVersion"]}).json()
    assert work["planReview"]["outcome"] == "NOT_REVIEWED" and work["planReview"]["reason"] == "EDITED"
    refused = client.post(f"/api/works/{work['id']}/plan/approve", headers=STUDENT, json={"baseVersion": work["planVersion"]})
    assert refused.status_code == 400 and refused.json()["code"] == "ACKNOWLEDGMENT_NEEDED"
    ok = client.post(f"/api/works/{work['id']}/plan/approve", headers=STUDENT, json={"baseVersion": work["planVersion"], "acknowledge": ["PLAN_OBJECTIONS"]})
    assert ok.status_code == 200


def test_adding_the_applicants_own_targets_does_not_undo_the_results_review(works_client):
    client = works_client
    work = _funding(client)
    _run(client, work["id"], "PLAN")
    work = _work(client, work["id"])
    assert work["resultsReview"]["outcome"] == "APPROVED"
    results = work["results"]
    for ind in results["indicators"]:
        ind["target"], ind["baseline"] = 60, 34
    work = client.post(f"/api/works/{work['id']}/results", headers=STUDENT, json={"results": results, "baseVersion": work["resultsVersion"]}).json()
    assert work["resultsReview"]["outcome"] == "APPROVED"
    results["outcomes"][0]["statement"] = "A rewritten outcome."
    work = client.post(f"/api/works/{work['id']}/results", headers=STUDENT, json={"results": results, "baseVersion": work["resultsVersion"]}).json()
    assert work["resultsReview"]["reason"] == "EDITED"


def test_editing_a_reviewed_proposal_plan_needs_confirmation(client):
    from tests.test_proposals import _create
    from tests.test_proposals import _run as run_step

    project = _create(client)
    run_step(client, project["id"], "PLAN")
    view = client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()
    assert view["planReview"]["outcome"] == "APPROVED"
    saved = client.post(f"/api/projects/{project['id']}/plan", headers=STUDENT, json={"plan": {**view["plan"], "title": view["plan"]["title"] + " in Mukono"}, "baseVersion": view["planVersion"]}).json()
    assert saved["planReview"]["reason"] == "EDITED"
    refused = client.post(f"/api/projects/{project['id']}/plan/approve", headers=STUDENT, json={"baseVersion": saved["planVersion"]})
    assert refused.status_code == 400 and refused.json()["code"] == "ACKNOWLEDGMENT_NEEDED"


# --- 6: the review and the Word file come from one layout -----------------------------------------------


def test_the_review_payload_is_exactly_the_word_files_text(works_client):
    from app.works.export import DRAFT_LABEL

    client = works_client
    work = _drafted(client)
    _, job = _run(client, work["id"], "DRAFT")
    assert job["status"] == "COMPLETED"
    document = _requests(client, "w_final")[-1]["document"]
    reviewed = [*document["front"]]
    for s in document["sections"]:
        reviewed += [s["heading"], *s["text"]]
        for t in s["tables"]:
            reviewed += [t["caption"], *[c for row in t["rows"] for c in row if c]]
    for t in document["tables"]:
        reviewed += [t["caption"], *[c for row in t["rows"] for c in row if c]]
    reviewed += [*(["References"] if document["references"] else []), *document["references"], *document["notes"]]
    word = Document(io.BytesIO(client.get(f"/api/works/{work['id']}/export", headers=STUDENT).content))
    printed = [p.text for p in word.paragraphs if p.text.strip()] + [c.text for t in word.tables for r in t.rows for c in r.cells if c.text.strip()]
    printed = [t for t in printed if t != DRAFT_LABEL]  # decided after the review, from its result
    assert sorted(reviewed) == sorted(printed)


# --- 7: a rejected institution profile is repaired and reviewed again --------------------------------------


def test_a_rejected_profile_is_repaired_and_reviewed_again(client):
    from tests.test_profiles import _guide
    from tests.test_proposals import _create
    from tests.test_proposals import _run as run_step

    reviews = []

    def review(payload):
        reviews.append(1)
        return {"approved": len(reviews) > 1, "issues": [] if len(reviews) > 1 else ["Chapter Two is missing its theoretical review section."]}

    original = client.models.default

    def finalise(payload):
        answer = original("p_profile_finalise", payload)
        if payload.get("critique", {}).get("overall", "").startswith("Repair only"):
            answer = {**answer, "rules": [*answer.get("rules", []), "Chapter Two includes a theoretical review."]}
        return answer

    client.models.overrides["p_profile_review"] = review
    client.models.overrides["p_profile_finalise"] = finalise
    project = _create(client)
    client.post(f"/api/projects/{project['id']}/guide", headers=STUDENT, files={"file": ("guide.docx", _guide(), "application/octet-stream")})
    job = run_step(client, project["id"], "PROFILE")
    assert job["status"] == "COMPLETED" and len(reviews) == 2

    client.models.overrides["p_profile_review"] = lambda payload: {"approved": False, "issues": ["Not this guide's structure."]}
    project = _create(client)
    client.post(f"/api/projects/{project['id']}/guide", headers=STUDENT, files={"file": ("guide2.docx", _guide(), "application/octet-stream")})
    job = run_step(client, project["id"], "PROFILE")
    assert job["status"] == "FAILED" and job["failure"]["code"] == "PROFILE_NOT_APPROVED" and job["billing"]["charged"] == 0
