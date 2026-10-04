"""The Data Lab engine (decision 2026-10-03): reading, profiling, cleaning, statistics checked against
R's own results (tests/fixtures/datalab_reference.R), disclosure control, and the awkward cases."""

import io
import math

import numpy as np
import pandas as pd
import pytest

from app.core.errors import AppError
from app.datalab.engine import clean, disclosure, ingest, profile, stats
from app.datalab.models import AnalysisSpec

LIMITS = ingest.Limits(rows=10_000, columns=200, cells=1_000_000, expanded_bytes=50_000_000)

SLEEP_1 = [0.7, -1.6, -0.2, -1.2, -0.1, 3.4, 3.7, 0.8, 0.0, 2.0]
SLEEP_2 = [1.9, 0.8, 1.1, 0.1, -0.1, 4.4, 5.5, 1.6, 4.6, 3.4]
MPG = [21.0, 21.0, 22.8, 21.4, 18.7, 18.1, 14.3, 24.4, 22.8, 19.2, 17.8, 16.4, 17.3, 15.2, 10.4, 10.4, 14.7, 32.4, 30.4, 33.9, 21.5, 15.5, 15.2, 13.3, 19.2, 27.3,
       26.0, 30.4, 15.8, 19.7, 15.0, 21.4]
WT = [2.620, 2.875, 2.320, 3.215, 3.440, 3.460, 3.570, 3.190, 3.150, 3.440, 3.440, 4.070, 3.730, 3.780, 5.250, 5.424, 5.345, 2.200, 1.615, 1.835, 2.465, 3.520, 3.435,
      3.840, 3.845, 1.935, 2.140, 1.513, 3.170, 2.770, 3.570, 2.780]


def _ctx(columns: dict[str, list], threshold: int = 5) -> stats.Context:
    table = ingest.Table(columns=list(columns), rows=[list(r) for r in zip(*columns.values(), strict=True)])
    frame, variables = profile.infer(table)
    return stats.Context(frame, {v.name: v for v in variables}, version=1, threshold=threshold)


def _csv(text: str) -> ingest.Table:
    return ingest.read(text.encode("utf-8"), "data.csv", None, LIMITS)


# --- R reference values --------------------------------------------------------------------------------


def test_welch_t_test_matches_r():
    ctx = _ctx({"extra": SLEEP_1 + SLEEP_2, "group": ["1"] * 10 + ["2"] * 10})
    r = stats.run(ctx, AnalysisSpec(kind="COMPARE_TWO", variables=["extra", "group"], method="MEANS"))
    s = r.statistics
    assert s["t"] == pytest.approx(-1.8608, abs=1e-4) and s["df"] == pytest.approx(17.776, abs=1e-3) and s["p"] == pytest.approx(0.07939, abs=1e-5)
    assert (s["diff_low"], s["diff_high"]) == (pytest.approx(-3.3654832, abs=1e-6), pytest.approx(0.2054832, abs=1e-6))
    assert (s["mean1"], s["mean2"]) == (pytest.approx(0.75), pytest.approx(2.33))
    assert r.record.method.startswith("Welch") and "not statistically significant" in r.sentences[0]


def test_rank_sum_test_matches_r_with_ties():
    ctx = _ctx({"extra": SLEEP_1 + SLEEP_2, "group": ["1"] * 10 + ["2"] * 10})
    r = stats.run(ctx, AnalysisSpec(kind="COMPARE_TWO", variables=["extra", "group"], method="DISTRIBUTIONS"))
    assert r.statistics["w"] == pytest.approx(25.5) and r.statistics["p"] == pytest.approx(0.06933, abs=1e-5)
    assert "continuity correction" in r.record.method  # ties: R's normal approximation, not the exact test


def test_mean_and_its_interval_match_r():
    ctx = _ctx({"extra": SLEEP_1 + SLEEP_2})
    r = stats.run(ctx, AnalysisSpec(kind="DESCRIBE", variables=["extra"]))
    assert r.statistics["mean"] == pytest.approx(1.54)
    assert (r.statistics["mean_low"], r.statistics["mean_high"]) == (pytest.approx(0.5955845, abs=1e-6), pytest.approx(2.4844155, abs=1e-6))


def test_pearson_correlation_matches_r():
    ctx = _ctx({"mpg": MPG, "wt": WT})
    r = stats.run(ctx, AnalysisSpec(kind="CORRELATE", variables=["mpg", "wt"], method="PEARSON"))
    s = r.statistics
    assert s["r"] == pytest.approx(-0.8676594, abs=1e-6) and s["p"] == pytest.approx(1.294e-10, rel=1e-3)
    assert (s["r_low"], s["r_high"]) == (pytest.approx(-0.9338264, abs=1e-6), pytest.approx(-0.7440872, abs=1e-6))
    assert "strong negative" in r.sentences[0] and "does not show that one causes" in r.sentences[1]


def test_spearman_correlation_matches_r():
    ctx = _ctx({"mpg": MPG, "wt": WT})
    r = stats.run(ctx, AnalysisSpec(kind="CORRELATE", variables=["mpg", "wt"], method="SPEARMAN"))
    assert r.statistics["r"] == pytest.approx(-0.886422, abs=1e-6)
    rho, n = r.statistics["r"], 32  # R's t approximation, from its own formula
    from scipy import stats as st

    t = rho * math.sqrt((n - 2) / (1 - rho * rho))
    assert r.statistics["p"] == pytest.approx(2 * st.t.sf(abs(t), n - 2), rel=1e-9)
    assert r.statistics["p"] == pytest.approx(1.488e-11, rel=1e-2)


def test_chi_square_matches_r():
    rows = [("F", "Democrat")] * 762 + [("F", "Independent")] * 327 + [("F", "Republican")] * 468 + \
           [("M", "Democrat")] * 484 + [("M", "Independent")] * 239 + [("M", "Republican")] * 477
    ctx = _ctx({"gender": [g for g, _ in rows], "party": [p for _, p in rows]})
    r = stats.run(ctx, AnalysisSpec(kind="CROSSTAB", variables=["gender", "party"]))
    assert r.statistics["chi2"] == pytest.approx(30.07015, abs=1e-4) and r.statistics["df"] == 2 and r.statistics["p"] == pytest.approx(2.953589e-07, rel=1e-5)
    assert r.status == "VALID" and r.statistics["cramers_v"] == pytest.approx(math.sqrt(30.07015 / 2757), abs=1e-6)


def test_fisher_exact_matches_r_for_a_small_two_by_two():
    rows = [("Milk", "Milk")] * 3 + [("Milk", "Tea")] * 1 + [("Tea", "Milk")] * 1 + [("Tea", "Tea")] * 3
    ctx = _ctx({"truth": [a for a, _ in rows], "guess": [b for _, b in rows]}, threshold=1)
    r = stats.run(ctx, AnalysisSpec(kind="CROSSTAB", variables=["truth", "guess"]))
    assert r.statistics["p"] == pytest.approx(0.4857143, abs=1e-6) and r.record.method.startswith("Fisher")
    assert "answers the same question" in r.record.why  # the same question as chi-square, never a different one


def test_wilson_interval_agrees_with_an_independent_implementation():
    from statsmodels.stats.proportion import proportion_confint

    low, high = stats.wilson(15, 50)
    ref = proportion_confint(15, 50, alpha=0.05, method="wilson")
    assert (low, high) == (pytest.approx(ref[0], abs=1e-9), pytest.approx(ref[1], abs=1e-9))


# --- awkward and invalid cases -----------------------------------------------------------------------------


def test_a_variable_with_one_value_cannot_be_correlated():
    ctx = _ctx({"x": [1.0] * 10, "y": list(range(10))})
    r = stats.run(ctx, AnalysisSpec(kind="CORRELATE", variables=["x", "y"]))
    assert r.status == "NOT_ESTIMABLE" and "same value in every record" in r.warnings[0]


def test_a_group_below_the_threshold_is_not_compared():
    ctx = _ctx({"y": list(range(12)), "g": ["a"] * 9 + ["b"] * 3})
    r = stats.run(ctx, AnalysisSpec(kind="COMPARE_TWO", variables=["y", "g"]))
    assert r.status == "NOT_ESTIMABLE" and "at least 5 records" in r.warnings[0] and not r.tables


def test_more_than_two_groups_need_the_two_chosen():
    ctx = _ctx({"y": list(range(30)), "g": ["a", "b", "c"] * 10})
    assert stats.run(ctx, AnalysisSpec(kind="COMPARE_TWO", variables=["y", "g"])).status == "NOT_ESTIMABLE"
    chosen = stats.run(ctx, AnalysisSpec(kind="COMPARE_TWO", variables=["y", "g"], groups=["a", "c"]))
    assert chosen.status != "NOT_ESTIMABLE" and chosen.record.rows_used == 20 and any("other groups" in x for x in chosen.record.left_out)


def test_a_failed_normality_check_warns_but_never_switches_the_method():
    skewed = [1, 1, 1, 1, 2, 2, 3, 40, 90, 200] + [1, 2, 2, 3, 3, 4, 5, 60, 120, 300]
    ctx = _ctx({"y": skewed, "g": ["a"] * 10 + ["b"] * 10})
    r = stats.run(ctx, AnalysisSpec(kind="COMPARE_TWO", variables=["y", "g"], method="MEANS"))
    assert r.record.method.startswith("Welch") and r.status == "VALID_WITH_WARNINGS" and "comparing the distributions" in r.warnings[0]


def test_large_tables_with_small_expected_counts_warn():
    ctx = _ctx({"a": ["x", "y", "z"] * 4 + ["x"] * 3, "b": ["p", "q", "r"] * 5}, threshold=1)
    r = stats.run(ctx, AnalysisSpec(kind="CROSSTAB", variables=["a", "b"]))
    assert r.status == "VALID_WITH_WARNINGS" and "Grouping small categories" in r.warnings[0]


def test_missing_values_are_left_out_and_counted():
    ctx = _ctx({"y": [1.0, 2.0, None, 4.0, 5.0, 6.0, None, 8.0], "x": [2.0, 4.0, 6.0, 8.0, None, 12.0, 14.0, 16.0]})
    r = stats.run(ctx, AnalysisSpec(kind="CORRELATE", variables=["y", "x"]))
    assert r.record.rows_used == 5 and r.record.rows_available == 8
    assert sum("Fewer than 5 records have no value" in x for x in r.record.left_out) == 2  # small counts are never printed (finding 2)


def test_significance_wording_follows_the_level_set():
    ctx = _ctx({"extra": SLEEP_1 + SLEEP_2, "group": ["1"] * 10 + ["2"] * 10})
    ctx.alpha = 0.1  # p = 0.079 is below 0.1
    r = stats.run(ctx, AnalysisSpec(kind="COMPARE_TWO", variables=["extra", "group"]))
    assert "not statistically significant" not in r.sentences[0] and r.record.alpha == 0.1


# --- disclosure control ------------------------------------------------------------------------------------


def test_small_counts_are_hidden_with_the_count_that_would_reveal_them():
    one_way = disclosure.protect(np.array([50, 3, 20]), 5, row_totals=True, column_totals=False).cells[0]
    assert list(one_way) == [False, True, True]  # 3 would follow from the total, so 20 is hidden too
    counts = np.array([[30, 2, 40], [25, 20, 35]])
    hidden = disclosure.protect(counts, 5).cells
    assert hidden[0, 1] and hidden.sum(axis=1)[0] >= 2 and hidden.sum(axis=0)[1] >= 2  # never alone in its row or column


def test_a_total_can_never_give_a_hidden_count_away():
    """Codex's case (finding 2): a 1 beside a 0 had a row total of 1 on show."""
    guard = disclosure.protect(np.array([[1, 0], [20, 30]]), 5)
    assert guard.cells[0, 0] and guard.rows[0]  # the total of 1 is a small count itself
    assert disclosure._width(*_grid(np.array([[1, 0], [20, 30]]), guard), 0, 0) >= 5


def _grid(counts, guard):
    r, c = counts.shape
    grid = np.zeros((r + 1, c + 1))
    grid[:r, :c], grid[:r, c], grid[r, :c], grid[r, c] = counts, counts.sum(axis=1), counts.sum(axis=0), counts.sum()
    hidden = np.zeros_like(grid, dtype=bool)
    hidden[:r, :c], hidden[:r, c], hidden[r, :c], hidden[r, c] = guard.cells, guard.rows, guard.columns, guard.total
    return grid, hidden, np.ones_like(grid, dtype=bool), r, c


def _feasible(counts, guard, cell, columns_shown=True):
    """Every value the hidden cell could take, by brute force over whole numbers: an independent check
    of the protection (the engine itself reasons with linear programming)."""
    import itertools

    r, c = counts.shape
    unknown = [(i, j) for i in range(r) for j in range(c) if guard.cells[i, j]]
    n = int(counts.sum())
    seen = set()
    for values in itertools.product(range(n + 1), repeat=len(unknown)):
        trial = counts.copy()
        for (i, j), v in zip(unknown, values, strict=True):
            trial[i, j] = v
        if any(not guard.rows[i] and trial[i].sum() != counts[i].sum() for i in range(r)):
            continue
        if columns_shown and any(not guard.columns[j] and trial[:, j].sum() != counts[:, j].sum() for j in range(c)):
            continue
        if not guard.total and trial.sum() != n:
            continue
        seen.add(int(trial[cell]))
    return seen


@pytest.mark.parametrize("seed", range(40))
def test_no_hidden_count_can_be_narrowed_down(seed):
    rng = np.random.default_rng(seed)
    shape = (2, 2) if seed % 2 else (1, 4)
    counts = rng.integers(0, 9, size=shape)
    if counts.sum() < 5:
        counts[0, 0] += 5
    guard = disclosure.protect(counts, 5, row_totals=True, column_totals=shape[0] > 1)
    for i, j in zip(*np.nonzero((counts > 0) & (counts < 5)), strict=True):
        assert guard.cells[i, j]
        values = _feasible(counts, guard, (i, j), columns_shown=shape[0] > 1)
        assert max(values) - min(values) >= 5, (counts.tolist(), (i, j), sorted(values))


def test_a_hidden_count_is_hidden_everywhere_in_the_result():
    ctx = _ctx({"answer": ["yes"] * 40 + ["no"] * 30 + ["maybe"] * 3})
    r = stats.run(ctx, AnalysisSpec(kind="DESCRIBE", variables=["answer"]))
    hidden_rows = [row for row in r.tables[0].rows if row[1].suppressed]
    assert {row[0].text for row in hidden_rows} == {"maybe", "no"}  # "no" too: it would reveal "maybe" from the total
    assert all(row[1].text == disclosure.HIDDEN and row[2].text == disclosure.HIDDEN for row in hidden_rows)  # one neutral mark for every hidden cell
    assert "count:maybe" not in r.statistics and "pct:maybe" not in r.statistics  # the text can't use it either
    from app.datalab.engine import charts

    png = charts.for_result(r)
    assert png is None or png[:4] == b"\x89PNG"


# --- reading, profiling and cleaning --------------------------------------------------------------------------


def test_identifiers_with_leading_zeros_stay_text():
    frame, variables = profile.infer(_csv("code,score\n00123,4\n00456,5\n00789,6\n"))
    by = {v.name: v for v in variables}
    assert by["code"].stored == "text" and "LEADING_ZEROS" in by["code"].flags and list(frame["code"]) == ["00123", "00456", "00789"]
    assert by["score"].stored == "number"


def test_semicolon_files_and_repeated_names_are_read():
    t = _csv("Age;Age;\n30;31;x\n40;41;y\n")
    assert t.columns == ["Age", "Age (2)", "Column 3"] and len(t.rows) == 2


def test_rows_wider_than_the_header_are_refused_with_the_row_number():
    with pytest.raises(AppError) as exc:
        _csv("a,b\n1,2\n3,4,5\n")
    assert "Row 3" in exc.value.message


def test_size_limits_are_checked_while_reading():
    small = ingest.Limits(rows=5, columns=10, cells=100, expanded_bytes=10_000)
    with pytest.raises(AppError) as exc:
        ingest.read(("a\n" + "1\n" * 10).encode(), "x.csv", None, small)
    assert exc.value.code == "DATA_TOO_LONG"


def test_macro_and_old_workbooks_are_refused():
    for name, code in (("book.xlsm", "DATA_MACROS"), ("book.xls", "DATA_OLD_EXCEL")):
        with pytest.raises(AppError) as exc:
            ingest.read(b"x", name, None, LIMITS)
        assert exc.value.code == code


def test_an_excel_sheet_is_read_as_values():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["district", "cases", "population"])
    ws.append(["Pader", 120, 230000])
    ws.append(["Gulu", 300, 325000])
    other = wb.create_sheet("Notes")
    other.append(["note"])
    other.append(["x"])
    buffer = io.BytesIO()
    wb.save(buffer)
    t = ingest.read(buffer.getvalue(), "data.xlsx", None, LIMITS)
    assert t.columns == ["district", "cases", "population"] and t.sheet == "Data" and t.sheets == ["Data", "Notes"]
    assert t.rows[0] == ["Pader", 120, 230000]
    assert ingest.read(buffer.getvalue(), "data.xlsx", "Notes", LIMITS).columns == ["note"]
    ws["D1"] = "rate"
    ws["D2"] = "=B2/C2*1000"  # a formula: only its saved result can be read, and openpyxl saves none
    buffer = io.BytesIO()
    wb.save(buffer)
    with pytest.raises(AppError) as exc:  # never a silent blank (finding 10)
        ingest.read(buffer.getvalue(), "data.xlsx", None, LIMITS)
    assert exc.value.code == "DATA_FORMULAS" and "1 cell" in exc.value.message


def test_values_are_never_changed_silently():
    frame, variables = profile.infer(_csv("when,big\n2026-10-01,9007199254740993\n2026-10-02 12:00,12\n"))
    by = {v.name: v for v in variables}
    assert by["when"].stored == "date" and frame["when"].iloc[1] == pd.Timestamp("2026-10-02 12:00")  # a date-time beside a date is kept
    assert by["big"].stored == "text" and "LONG_NUMBER" in by["big"].flags and frame["big"].iloc[0] == "9007199254740993"  # never rounded


def test_constant_groups_are_not_estimable_never_infinite():
    ctx = _ctx({"y": [5.0] * 10 + [7.0] * 10, "g": ["a"] * 10 + ["b"] * 10})
    r = stats.run(ctx, AnalysisSpec(kind="COMPARE_TWO", variables=["y", "g"], method="MEANS"))
    assert r.status == "NOT_ESTIMABLE" and "same value" in r.warnings[0]
    import json

    json.dumps(r.model_dump(mode="json"), allow_nan=False)  # always serialisable


def test_personal_identifiers_and_survey_columns_are_flagged():
    frame, variables = profile.infer(_csv("Respondent name,phone,v005,cluster,age\n"
                                          + "".join(f"Person {i},07{i:08d},{1000000 + i},{i % 3},{20 + i}\n" for i in range(25))))
    by = {v.name: v for v in variables}
    assert "PERSONAL" in by["Respondent name"].flags and by["Respondent name"].excluded
    assert "PERSONAL" in by["phone"].flags and by["phone"].excluded
    assert "SURVEY_DESIGN" in by["v005"].flags and "SURVEY_DESIGN" in by["cluster"].flags and not by["age"].flags


def test_cleaning_is_proposed_not_applied_except_trimming():
    text = "sex,age,score\n" + "Male ,25,3\nmale,30,4\nFemale,N/A,5\nF,40,-9\nFemale,200,5\nMale ,25,3\n"
    frame, variables = profile.infer(_csv(text))
    steps = clean.proposals(frame, variables)
    kinds = {(s.kind, s.column) for s in steps}
    assert ("TRIM", "sex") in kinds and ("MERGE_LEVELS", "sex") in kinds and ("SET_MISSING", "age") in kinds and ("DUPLICATES", "") in kinds
    assert all(s.automatic == (s.kind == "TRIM") for s in steps) and all(s.question for s in steps if not s.automatic)
    assert ("SET_MISSING", "score") in kinds  # -9 in a column otherwise never negative
    trimmed, n = clean.apply(frame, next(s for s in steps if s.kind == "TRIM"))
    assert n == 2 and set(trimmed["sex"].dropna()) == {"Male", "male", "Female", "F"} and set(frame["sex"].dropna()) >= {"Male "}  # the original untouched
    merged, _ = clean.apply(trimmed, next(s for s in clean.proposals(trimmed, variables) if s.kind == "MERGE_LEVELS"))
    assert set(merged["sex"].dropna()) == {"Male", "Female"}
    ages, k = clean.apply(frame, next(s for s in steps if s.kind == "SET_MISSING" and s.column == "age"))
    assert k == 1 and ages["age"].dtype == "float64"  # with "N/A" gone the column is numbers
    deduped, d = clean.apply(frame, next(s for s in steps if s.kind == "DUPLICATES"))
    assert d == 1 and len(deduped) == len(frame) - 1


def test_out_of_range_ages_are_proposed():
    frame, variables = profile.infer(_csv("age\n" + "\n".join(str(a) for a in [25, 30, 130, -2, 45]) + "\n"))
    step = next(s for s in clean.proposals(frame, variables) if s.kind == "OUT_OF_RANGE")
    out, n = clean.apply(frame, step)
    assert n == 2 and out["age"].isna().sum() == 2 and pd.isna(frame["age"]).sum() == 0


def test_a_perfect_correlation_reports_itself_as_its_interval():
    ctx = _ctx({"x": [1.0, 2.0, 3.0, 4.0, 5.0, 6.0], "y": [2.0, 4.0, 6.0, 8.0, 10.0, 12.0]})
    for method in ("PEARSON", "SPEARMAN"):
        r = stats.run(ctx, AnalysisSpec(kind="CORRELATE", variables=["x", "y"], method=method))
        assert r.statistics["r"] == pytest.approx(1.0) and r.statistics["r_low"] == r.statistics["r_high"] == pytest.approx(1.0)


# --- district maps ------------------------------------------------------------------------------------------


def test_district_names_match_exactly_by_known_spelling_or_are_suggested_never_guessed():
    from app.datalab.engine import maps

    matched, suggested, unmatched = maps.match(["Gulu", "gulu district", "Madi Okollo", "Luwero", "Kampla", "Atlantis"], {})
    assert matched == {"Gulu": "GULU", "gulu district": "GULU", "Madi Okollo": "MADI-OKOLLO", "Luwero": "LUWEERO"}
    assert suggested == {"Kampla": "KAMPALA"} and unmatched == ["Kampla", "Atlantis"]  # a near miss is suggested, not applied
    confirmed, _, still = maps.match(["Kampla"], {"Kampla": "KAMPALA"})
    assert confirmed == {"Kampla": "KAMPALA"} and still == []
    assert len(maps.available_districts()) == 146 and "GULU" in maps.available_districts("Northern") and "GULU" not in maps.available_districts("Central")


def test_a_district_map_hides_small_districts_and_reports_unmatched_names():
    rows = {"district": ["Gulu"] * 30 + ["Pader"] * 20 + ["Kitgum"] * 3 + ["Kampala"] * 25 + ["Kampla"] * 2, "score": list(range(80))}
    ctx = _ctx(rows)
    from app.datalab.engine import maps

    result, png = maps.run(ctx, AnalysisSpec(kind="MAP", variables=["district"]))
    assert png is not None and png[:4] == b"\x89PNG"
    hidden = {r[0].text for r in result.tables[0].rows if r[1].suppressed}
    assert "Kitgum" in hidden and len(hidden) >= 2  # 3 records hidden, with the next smallest so the total can't reveal it
    assert result.unmatched == ["Kampla"] and result.matches == {"Kampla": "KAMPALA"} and result.status == "VALID_WITH_WARNINGS"
    assert any("Fewer than 5 records have district names" in x for x in result.record.left_out)
    northern, _ = maps.run(ctx, AnalysisSpec(kind="MAP", variables=["district", "score"], method="MEAN", region="Northern"))
    names = {r[0].text for r in northern.tables[0].rows}
    assert "Kampala" not in names and any("outside the Northern region" in x for x in northern.record.left_out)


def test_released_map_layers_never_change():
    """A map's fingerprint names its layer, so a released layer file is frozen like a released prompt:
    a new build writes a new file and is added here."""
    import hashlib
    import json
    from pathlib import Path

    geo = Path(__file__).parents[1] / "app" / "datalab" / "geo"
    released = json.loads((geo / "released.json").read_text(encoding="utf-8"))["files"]
    assert {name: hashlib.sha256((geo / name).read_bytes().replace(b"\r\n", b"\n")).hexdigest() for name in released} == released
    assert {p.name for p in geo.glob("*.geojson")} <= set(released)


def test_the_shared_identifier_rules_flag_the_expected_columns():
    """The browser (web/src/features/datalab/upload.test.ts) checks the same fixture against the same list."""
    import json
    from pathlib import Path

    folder = Path(__file__).parent / "fixtures"
    table = ingest.read((folder / "identifiers_fixture.csv").read_bytes(), "f.csv", None, LIMITS)
    _, variables = profile.infer(table)
    flagged = [v.name for v in variables if "PERSONAL" in v.flags or "LOCATION" in v.flags]
    assert flagged == json.loads((folder / "identifiers_expected.json").read_text(encoding="utf-8"))["flagged"]
