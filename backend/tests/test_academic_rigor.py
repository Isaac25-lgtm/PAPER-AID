"""Claims, statistics and the proposal's front pages (the academic audit of 10 October 2026, as agreed with the owner
on the 11th). Each test is a defect the audit found in a delivered trial document."""

import io
import zipfile

import pytest
from docx import Document

from app.proposals import evidence
from app.proposals import pipeline as proposals
from app.proposals.models import EvidenceItem, EvidenceSource
from app.works import pipeline as works
from tests.test_api import STUDENT
from tests.test_proposal_v2 import _with_chapter_one


def test_a_claim_that_nothing_has_been_studied_is_sent_back():
    """A proposal said factors in Mukono "remain uninvestigated" and baseline evidence "remains completely lacking"
    while a registered study and a dissertation on the topic existed."""
    for claim in ("The factors driving completion in Mukono District remain uninvestigated.",
                  "Baseline evidence on dose completion remains completely lacking.",
                  "No study has examined caregivers in Mukono since the vaccine was introduced.",
                  "This relationship has never been studied in Uganda.",
                  "There is no empirical evidence on this outcome in Uganda.",
                  "There is no published research available on this population."):
        problems = evidence.gap_problems(claim)
        assert len(problems) == 1 and "never that no study exists" in problems[0], claim
    for sound in ("Published evidence on dose completion in Mukono District is limited.",
                  "The studies found did not examine caregivers in Mukono District.",
                  "No significant difference was reported between the two groups.",
                  "The study will examine factors that have not yet received much attention.",
                  "Data on stock-outs were completely recorded at each facility.",
                  # Not claims about research: the check names research or evidence itself.
                  "In many rural schools electricity is completely lacking.",
                  "Undocumented migrants were excluded from the sampling frame.",
                  "There is no research department in the organisation.",
                  "The role of fathers remains largely unstudied in the work found.",  # hedged: a limit, not an absolute
                  # Codex's audit of 1a3b1e7, finding 5: saying a fact is not known is honest, and is never a claim about research
                  "The eligible population remains unknown.",
                  "The district disease burden remains undocumented.",
                  "Data on the eligible population are entirely missing from the district register.",
                  "No data exist on how many facilities offer the service.",
                  "Whether the service is available in every sub-county is unknown.",
                  "There is no evidence that the vaccine causes harm.",
                  # the claim is reported or denied, not made
                  "Earlier reviews claimed that no study has examined caregivers, which later work showed to be wrong.",
                  "One author wrote that “no study has examined this question”, a claim this review does not accept.",
                  # Codex's re-audit of 458e65b, finding 3: a sentence that rejects the claim
                  "The topic is not unstudied.",
                  "It would be misleading to say that no study has examined this question."):
        assert evidence.gap_problems(sound) == [], sound
        assert evidence.strip_overclaims(sound) is sound
    kept = "The eligible population remains unknown. The district disease burden remains undocumented."
    assert evidence.strip_overclaims(kept) is kept


def test_a_sentence_still_overclaiming_after_repair_is_withheld_before_the_final_review():
    """The repair rounds are given both checks' findings. What they leave is removed as an untraceable figure is, so
    a pattern never fails a whole chapter and never reaches the student."""
    paragraph = ("Uptake of the vaccine varies between districts. No study has examined caregivers in Mukono District. "
                 "Education showed a significant association with uptake (OR = 1.27; 95% CI: 0.56–2.89; p = 0.043). The study will address these questions.")
    assert evidence.strip_overclaims(paragraph) == "Uptake of the vaccine varies between districts. The study will address these questions."
    assert evidence.strip_overclaims(paragraph, gaps=False) == ("Uptake of the vaccine varies between districts. No study has examined caregivers in Mukono District. "
                                                               "The study will address these questions.")  # coursework: the statistic only
    sound = "Published evidence is limited.  The interval reported (0.56–2.89) includes 1."
    assert evidence.strip_overclaims(sound) is sound  # nothing removed: the approved spacing is untouched
    for denied in ("The result was not statistically significant (OR = 1.27; 95% CI: 0.56–2.89; p = 0.043).",
                   "The result did not reach statistical significance (OR = 1.27; 95% CI: 0.56–2.89; p = 0.043).",
                   "There was no significant association (OR = 1.27; 95% CI: 0.56–2.89; p = 0.043)."):
        assert evidence.statistic_problems(denied) == [], denied


def test_a_ratio_whose_interval_includes_one_is_not_called_significant():
    """ "OR = 1.27; 95% CI: 0.56–2.89; p = 0.043", called a significant association. The figures were the
    publisher's own, so code corrects none of them: it stops the word "significant". """
    copied = "Caregivers with secondary education showed a significant association with uptake (OR = 1.27; 95% CI: 0.56–2.89; p = 0.043)."
    problems = evidence.statistic_problems(copied)
    assert len(problems) == 1 and "0.56–2.89" in problems[0] and "Do not call this result significant" in problems[0]
    for sound in ("The study reported OR = 1.27 (95% CI: 0.56–2.89; p = 0.043), an interval that includes 1.",  # told as reported
                  "The association was not significant (OR = 1.27; 95% CI: 0.56–2.89; p = 0.043).",
                  "Uptake was significantly higher among urban caregivers (OR = 2.10; 95% CI: 1.30–3.40; p = 0.002).",  # the two agree
                  "The odds ratio was 1.27 with a 95% confidence interval of 0.56 to 2.89 and p = 0.43, which is not significant.",
                  "The result was significant at the five per cent level."):
        assert evidence.statistic_problems(sound) == [], sound
    for worded in ("A significant effect was found, with an odds ratio of 1.27, 95% confidence interval 0.56 to 2.89, p < 0.05.",
                   "Education was a significant predictor: the odds ratio was 1.27 with a 95% confidence interval of 0.56 to 2.89 (p = 0.043)."):
        assert len(evidence.statistic_problems(worded)) == 1, worded


def test_the_writers_are_told_beforehand_what_code_will_check():
    """A rule the writer first hears of in a repair round costs a round. New work is written and repaired on prompts
    that carry them; work priced before keeps the prompts it was priced with."""
    from pathlib import Path

    from app.ai.orchestration import STEPS

    folder = Path(__file__).parents[1] / "app" / "ai" / "prompts"
    for step in ("p_draft", "p_fix", "w_draft", "w_repair"):
        prompt = (folder / f"{STEPS[step].prompt}.md").read_text(encoding="utf-8")
        assert '"caution"' in prompt and "that no study exists" in prompt and "value of no difference (1 for" in prompt, step
        assert "interval includes 1 is never" not in prompt and "when its reported interval includes 1" not in prompt, step  # a difference's value of no difference is 0
    plan = (folder / f"{STEPS['p_plan'].prompt}.md").read_text(encoding="utf-8")
    assert "relationships and associations" in plan and "never list one variable in two roles" in plan
    # Codex's audit of 1a3b1e7, finding 8: what a study can show is not decided by what its design is called
    assert "a cohort included, examine relationships and associations" in plan and "stays descriptive" in plan and "tell the student which word changed" in plan
    for step in ("p_draft", "p_fix"):
        prompt = (folder / f"{STEPS[step].prompt}.md").read_text(encoding="utf-8")
        assert "longitudinal findings" not in prompt and "identify a cause" in prompt, step
    assert "command word asks for and no more" in (folder / f"{STEPS['w_draft'].prompt}.md").read_text(encoding="utf-8")


def test_a_question_that_asks_to_describe_is_planned_and_researched_as_one():
    """ "What are input devices? Describe five examples and explain what each is used for" was answered with a
    microphone's readout circuit and touchscreens in aircraft cockpits (round of 11 October): the research topics and
    the plan set an answer's depth, so both are told, on every path."""
    from pathlib import Path

    from app.ai.orchestration import STEPS, WORKFLOW_PROMPTS
    from app.proposals.ai import NEEDS_SCHEMA, NEEDS_SCHEMA_V2, needs_schema

    folder = Path(__file__).parents[1] / "app" / "ai" / "prompts"
    for prompt in (STEPS["w_needs"].prompt, WORKFLOW_PROMPTS[2]["w_needs"]):
        assert "never for specialist research on one narrow device" in (folder / f"{prompt}.md").read_text(encoding="utf-8"), prompt
    assert needs_schema(STEPS["w_needs"].prompt) == NEEDS_SCHEMA and needs_schema(WORKFLOW_PROMPTS[2]["w_needs"]) == NEEDS_SCHEMA_V2  # each keeps its path's answer format
    plan = (folder / f"{STEPS['w_plan'].prompt}.md").read_text(encoding="utf-8")
    assert "at the depth its command words and the student's level ask for, and no deeper" in plan and "plan a direct answer" in plan


def _item(**details) -> EvidenceItem:
    return EvidenceItem(**{"id": "E0a1b2c", "chapter": 1, "need": "n", "statement": "Moodle is used by 98.7% of institutions.",
                           "passage": "98.7% of the universities sampled in Kampala used Moodle", "scope": "Kampala institutions sampled, 2021", "access": "FULL_TEXT",
                           "verified": True, "support": "PARTLY_SUPPORTED", "source": EvidenceSource(url="https://example.org/a", title="E-learning in Kampala", year="2021"),
                           "retrieved_on": "2026-10-11", **details})


def test_a_partly_supported_finding_reaches_the_writer_with_what_it_must_not_claim():
    """98.7% of a Kampala sample was written as "98.7% of accredited institutions" nationally. A partly supported
    finding stays usable, so its writer and reviewer are now told where the statement goes beyond its source."""
    caution = "The passage is about universities sampled in Kampala, not all accredited institutions in Uganda."
    for given in (proposals._for_model([_item(caution=caution)], passages=True)[0], works._for_model([_item(caution=caution)], passages=True)[0]):
        assert given["caution"].startswith("Only partly supported: The passage is about universities sampled in Kampala") and "people, place and period" in given["caution"]
        assert given["scope"] == "Kampala institutions sampled, 2021" and given["sourceWords"].startswith("98.7% of the universities sampled in Kampala")
    assert "caution" not in proposals._for_model([_item(support="SUPPORTED")], passages=True)[0]  # a supported finding, and evidence saved before, are given as before


def test_the_support_check_s_reservation_is_kept_with_the_finding(client):
    seen = {}
    client.models.overrides["verify"] = lambda payload: {"results": [{"id": c["id"], "support": "PARTLY_SUPPORTED", "note": "The source covers  rural caregivers in one district only."}
                                                                     for c in payload["claims"]]}

    def draft(payload):
        from tests.fake_models import FakeModels

        seen.setdefault("evidence", [e for s in payload["sections"] for e in s["evidence"]])
        return FakeModels.proposal("p_draft", payload)

    client.models.overrides["p_draft"] = draft
    project_id = _with_chapter_one(client)
    library = client.get(f"/api/projects/{project_id}/evidence", headers=STUDENT).json()
    assert library and all(i["caution"] == "The source covers rural caregivers in one district only." for i in library if i["support"] == "PARTLY_SUPPORTED")
    assert seen["evidence"] and all("one district only" in e["caution"] for e in seen["evidence"])


def test_the_contents_page_lists_the_document_s_headings_and_never_an_instruction(client):
    """Every proposal printed "Word fills in this table when the document is opened (or right-click it and choose
    Update Field)" where its contents should be."""
    project_id = _with_chapter_one(client)
    exported = client.get(f"/api/projects/{project_id}/export", headers=STUDENT)
    assert exported.status_code == 200
    xml = zipfile.ZipFile(io.BytesIO(exported.content)).read("word/document.xml").decode("utf-8")
    assert "Word fills in this table" not in xml and "Update Field" not in xml
    assert ' TOC \\o "1-3"' in xml and 'w:dirty="true"' in xml  # still a real contents field, which Word fills with page numbers
    paragraphs = [p.text for p in Document(io.BytesIO(exported.content)).paragraphs]
    contents = paragraphs[paragraphs.index("Table of Contents") + 1]
    # its own headings until then, read from the finished document: every numbered section, the statements placed
    # under Objectives, and the references (Codex's audit of 1a3b1e7: the list given to it had left those out)
    assert "CHAPTER ONE\nGENERAL INTRODUCTION\n    1.0 Introduction\n    1.1 Background to the Study\n    1.2 Statement of the Problem\n" in contents
    assert "\n        1.3.1 General Objective\n        1.3.2 Specific Objectives\n" in contents and contents.endswith("\nREFERENCES")
    headings = [p.text.strip() for p in Document(io.BytesIO(exported.content)).paragraphs if p.style.name in ("Heading 1", "Heading 2", "Heading 3") and p.text.strip() != "Table of Contents"]
    assert [line.strip() for line in contents.split("\n")] == headings


def test_the_declaration_and_approval_pages_say_they_are_forms_still_to_sign(client):
    project_id = _with_chapter_one(client)
    paragraphs = [p.text for p in Document(io.BytesIO(client.get(f"/api/projects/{project_id}/export", headers=STUDENT).content)).paragraphs]
    declaration, approval = paragraphs.index("Declaration"), paragraphs.index("Approval")
    assert paragraphs[declaration + 1] == "(A form for you to read, complete and sign before you submit. It is not yet signed.)"
    assert paragraphs[approval + 1] == "(A form for your supervisor to sign once they have approved the proposal. It is not yet signed.)"
    assert any(p.startswith("Signature: ____") for p in paragraphs[declaration:approval])


# --- Codex's audit of 1a3b1e7 ------------------------------------------------------------------------------------

RATIO = "OR = 1.27; {level} CI: 0.56–2.89; p = 0.043"


def test_the_statistics_check_judges_only_what_cannot_be_true():
    """Finding 2: a 99% interval that includes 1 at p = 0.02 can be right; the level was never read."""
    said = "Education showed a significant association with uptake ({})."
    assert len(evidence.statistic_problems(said.format(RATIO.format(level="95%")))) == 1
    assert len(evidence.statistic_problems(said.format(RATIO.format(level="90%")))) == 1  # narrower still: it cannot include 1 at p below 0.05
    for level in ("99%", "99.9%", ""):  # can agree, or the level is not stated: left to the reviewer
        assert evidence.statistic_problems(said.format(RATIO.format(level=level).replace("  ", " "))) == [], level
    assert evidence.statistic_problems("Education showed a significant association (OR = 1.27; 99% CI: 0.56–2.89; p = 0.02).") == []
    # a difference or a coefficient: the value of no difference is 0, so an interval that includes 1 says nothing
    for other in ("The mean difference was significant (difference = 1.4; 95% CI: 0.6–2.2; p = 0.001).",
                  "Income was a significant predictor (β = 0.9; 95% CI: 0.2–1.6; p = 0.012)."):
        assert evidence.statistic_problems(other) == [], other


def test_the_word_significant_must_belong_to_the_result_it_is_judged_against():
    """Finding 2: a sentence with a significant age result and "education was not significant" was deleted."""
    mixed = ("Caregiver age was significantly associated with uptake (OR = 2.10; 95% CI: 1.30–3.40; p = 0.002), whereas education was not significant "
             "(OR = 1.27; 95% CI: 0.56–2.89; p = 0.043).")
    assert evidence.statistic_problems(mixed) == [] and evidence.strip_overclaims(mixed) is mixed
    reversed_order = "Education was not associated with uptake (OR = 1.27; 95% CI: 0.56–2.89; p = 0.043), while age was significant (OR = 2.10; 95% CI: 1.30–3.40; p = 0.002)."
    assert evidence.statistic_problems(reversed_order) == []
    wrong = "Age was significant (OR = 2.10; 95% CI: 1.30–3.40; p = 0.002) and education was significant too (OR = 1.27; 95% CI: 0.56–2.89; p = 0.043)."
    assert len(evidence.statistic_problems(wrong)) == 1
    after = "For education the estimate was OR = 1.27 (95% CI: 0.56–2.89; p = 0.043), which the authors describe as statistically significant."
    assert len(evidence.statistic_problems(after)) == 1
    # Codex's re-audit of 458e65b, finding 3: the sentence itself says the result should not be called significant
    caution = "The result is inconsistent and should not be described as statistically significant (OR = 1.27; 95% CI: 0.56–2.89; p = 0.043)."
    assert evidence.statistic_problems(caution) == [] and evidence.strip_overclaims(caution) is caution
    for denied in ("This cannot be called a significant association (OR = 1.27; 95% CI: 0.56–2.89; p = 0.043).",
                   "It would be wrong to describe the association as significant (OR = 1.27; 95% CI: 0.56–2.89; p = 0.043).",
                   "The association (OR = 1.27; 95% CI: 0.56–2.89; p = 0.043) isn’t significant once the interval is read."):
        assert evidence.statistic_problems(denied) == [], denied


def test_a_table_is_read_as_prose_is_and_a_row_with_its_headings():
    """Finding 6: the same sentence passed unchanged in a caption or a cell, and a statistic spread across columns
    was never read as one."""
    sentence = "Education showed a significant association with uptake (OR = 1.27; 95% CI: 0.56–2.89; p = 0.043)."
    header = ["Factor", "OR", "95% CI", "p-value", "Interpretation"]
    wrong = ["Education", "1.27", "0.56–2.89", "0.043", "Significant"]
    sound = ["Age", "2.10", "1.30–3.40", "0.002", "Significant"]
    denied = ["Income", "1.27", "0.56–2.89", "0.043", "Not significant"]
    assert evidence.row_text(header, wrong) == "Factor = Education; OR = 1.27; 95% CI = 0.56–2.89; p-value = 0.043; Interpretation = Significant"
    assert evidence.row_overclaims(header, wrong) == [4] and evidence.row_overclaims(header, sound) == [] and evidence.row_overclaims(header, denied) == []
    assert len(evidence.table_problems([header, wrong, sound, denied])) == 1 and "Education" in evidence.table_problems([header, wrong, sound, denied])[0]
    assert evidence.strip_table_overclaims([header, wrong, sound, denied]) == [header, ["Education", "1.27", "0.56–2.89", "0.043", ""], sound, denied]
    untouched = [header, sound, denied]
    assert evidence.strip_table_overclaims(untouched) is untouched
    # Codex's re-audit of 458e65b, finding 4: the claim made by an answer under a heading that asks
    asks = ["Factor", "OR", "95% CI", "p-value", "Statistically significant?"]
    yes = ["Education", "1.27", "0.56–2.89", "0.043", "Yes"]
    assert evidence.row_overclaims(asks, yes) == [4] and len(evidence.table_problems([asks, yes])) == 1
    assert evidence.strip_table_overclaims([asks, yes]) == [asks, ["Education", "1.27", "0.56–2.89", "0.043", ""]]
    for answer in ("No", "Not significant", "NS", "n.s."):
        assert evidence.row_overclaims(asks, ["Education", "1.27", "0.56–2.89", "0.043", answer]) == [], answer
    assert evidence.row_overclaims(asks, ["Age", "2.10", "1.30–3.40", "0.002", "Yes"]) == []  # a result that is significant
    assert evidence.row_overclaims(["Factor", "OR", "95% CI", "p-value", "Not significant?"], ["Education", "1.27", "0.56–2.89", "0.043", "Yes"]) == []

    # the proposal's checks and its last line of defence
    from app.proposals.ai import SectionText, Table
    from app.proposals.models import StepInput
    from app.proposals.pipeline import _checks, _strip_table

    inp = StepInput.model_construct(chapter=2, plan=None)
    figures = "1.27 95% 0.56 2.89 0.043 2.10 1.30 3.40 0.002"  # the table's own numbers are traced: only the claim is at issue
    for table in (Table(caption=sentence, rows=[]), Table(caption="Factors", rows=[["Finding"], [sentence]]), Table(caption="Factors", rows=[header, wrong, sound])):
        text = SectionText(key="empirical", paragraphs=["The studies found are summarised below."], table=table)
        found = _checks(inp, "empirical", text, {}, figures)
        assert len(found) == 1 and found[0].startswith("In the table: ") and "do not agree" in found[0], (table, found)
        cleaned = _strip_table(inp, text, {}, set(), figures)
        assert cleaned != table and "ignificant association" not in cleaned.caption and all("ignificant" not in c or r == sound for r in cleaned.rows[1:] for c in r)
    gap = SectionText(key="empirical", paragraphs=["A summary follows."], table=Table(caption="No study has examined caregivers in Mukono District.", rows=[]))
    assert len(_checks(inp, "empirical", gap, {}, "")) == 1 and _strip_table(inp, gap, {}, set(), "").caption == ""


def test_coursework_tables_get_the_same_reading(client):
    from app.proposals.ai import SectionText as WorkText
    from app.proposals.ai import Table as WorkTable
    from app.works import pipeline as works_pipeline
    from app.works.models import PlanSection

    header = ["Factor", "OR", "95% CI", "p-value", "Interpretation"]
    wrong = ["Education", "1.27", "0.56–2.89", "0.043", "Significant"]
    sentence = "Education showed a significant association with uptake (OR = 1.27; 95% CI: 0.56–2.89; p = 0.043)."
    assert evidence.table_problems([header, wrong]) and evidence.statistic_problems(sentence)
    section = PlanSection.model_construct(key="theme1", heading="Evidence", brief="", words=300, min_words=0, max_words=0, field_id="", coverage=[], criteria=[])
    for table in (WorkTable(caption=sentence, rows=[header, ["Age", "2.10", "1.30–3.40", "0.002", "Significant"]]), WorkTable(caption="Factors", rows=[header, wrong])):
        text = WorkText(key="theme1", paragraphs=["The studies are summarised in the table."], table=table)
        assert any("do not agree" in p for p in works_pipeline._checks(_work_input(), section, text, {}, "", {})), table


def _work_input():
    from app.works.models import ResolvedSpec, WorkStepInput

    spec = ResolvedSpec.model_construct(kind="COURSEWORK", variant="ESSAY", mode="", level="LATER_UG", source_policy="OPEN", fields=[], citation_style="APA7", directives=[], coverage=[],
                                        active_rules=[], flags={}, limits=[], requirements=[], target_words=1500)
    return WorkStepInput.model_construct(spec=spec, plan=None, inputs=None, results=None, budget=None, kind="COURSEWORK")


@pytest.fixture
def works_client(tmp_path, monkeypatch):
    monkeypatch.setenv("WORKS_ENABLED", '["CONCEPT_NOTE","COURSEWORK","FUNDING_PROPOSAL"]')
    from tests.conftest import _client

    yield from _client(tmp_path, monkeypatch, "cost")


def test_the_coursework_evaluator_reads_each_section_beside_the_evidence_it_cites(works_client):
    """Finding 3: the caution reached the writer and the repair, and no reviewer. The request the evaluator is
    actually sent is read here, not the helper that builds an evidence entry."""
    from tests.test_works import _coursework, _run, _work

    client = works_client
    reservation = "The passage is about universities sampled in Kampala, not Uganda nationally."
    client.models.overrides["w_verify"] = lambda payload: {"results": [{"id": c["id"], "support": "PARTLY_SUPPORTED", "note": reservation} for c in payload["claims"]]}
    asked: list[dict] = []

    def evaluate(payload):
        asked.append(payload)
        return client.models.default("w_evaluate", payload)

    client.models.overrides["w_evaluate"] = evaluate
    work = _coursework(client)
    _run(client, work["id"], "PLAN")
    work = _work(client, work["id"])
    client.post(f"/api/works/{work['id']}/plan/approve", headers=STUDENT, json={"baseVersion": work["planVersion"]})
    _, job = _run(client, work["id"], "DRAFT")
    assert job["status"] == "COMPLETED", job.get("failure")
    sections = [s for payload in asked for s in payload["sections"]]
    citing = [s for s in sections if "⟦E" in " ".join(s["text"])]
    assert citing and all(s["evidence"] for s in citing)
    for section in citing:
        cited = set(evidence.cited_ids(" ".join(section["text"])))
        assert {e["id"] for e in section["evidence"]} == cited  # what it cites, and nothing else
        assert all(reservation in e["caution"] and e["sourceWords"] and "scope" in e for e in section["evidence"])
    assert all("evidence" in s for s in sections)  # a section that cites nothing is told so with an empty list


def test_page_numbers_carry_on_into_the_references(client):
    """Found when Word filled the contents page of a trial proposal (2026-10-11): "REFERENCES ... 1". A new section
    starts as a copy of the one before, its restart at page 1 included."""
    import re

    project_id = _with_chapter_one(client)
    exported = client.get(f"/api/projects/{project_id}/export", headers=STUDENT)
    xml = zipfile.ZipFile(io.BytesIO(exported.content)).read("word/document.xml").decode("utf-8")
    assert "REFERENCES" in xml
    numbering = re.findall(r"<w:pgNumType[^>]*>", xml)
    assert len(numbering) == 3 and ['w:start="1"' in n for n in numbering] == [True, True, False]  # the preliminary pages, the first chapter, then on
    assert 'w:fmt="lowerRoman"' in numbering[0] and all('w:fmt="decimal"' in n for n in numbering[1:])


# --- the owner's go-ahead of 11 October: what still needed fixing on the live path -----------------------------------

def test_a_finding_from_an_old_source_is_given_to_the_writer_and_reviewers_as_dated():
    """A 2003 paper was cited for what malaria "accounts for" today (round of 11 October)."""
    old = _item(support="SUPPORTED", source=EvidenceSource(url="https://example.org/k", title="Malaria morbidity and mortality in Uganda", year="2003"))
    for given in (proposals._for_model([old], passages=True, dated=True)[0], works._for_model([old], passages=True, dated=True)[0]):
        assert given["dated"].startswith("Published in 2003.") and "never as the present position" in given["dated"] and "caution" not in given
    for year in ("2016", "2021", "", "n.d."):  # ten years or fewer before it was gathered, or no year to go by
        recent = _item(support="SUPPORTED", source=EvidenceSource(url="https://example.org/r", title="A study", year=year))
        assert "dated" not in proposals._for_model([recent], dated=True)[0] and evidence.dated(recent) == "", year
    # Codex's audit of 6e513af: only a step priced with it is told, so one priced before is asked exactly as it was
    from app.proposals.models import StepInput
    from app.works.models import WorkStepInput

    assert "dated" not in proposals._for_model([old], passages=True)[0] and "dated" not in works._for_model([old], passages=True)[0]
    assert StepInput.model_fields["tell_dated"].default is False and WorkStepInput.model_fields["tell_dated"].default is False
    assert evidence.dated(_item(source=EvidenceSource(url="https://example.org/o", title="A study", year="2015"))) != ""  # eleven years


def test_the_plan_asks_for_the_diagram_a_subject_explains_a_point_with_and_basic_answers_are_plain():
    from pathlib import Path

    from app.ai.orchestration import STEPS

    folder = Path(__file__).parents[1] / "app" / "ai" / "prompts"
    plan = (folder / f"{STEPS['w_plan'].prompt}.md").read_text(encoding="utf-8")
    assert "asks for the graph by name and says what it must show, whether or not the question uses the word" in plan and "none where a diagram would only decorate" in plan
    draft = (folder / f"{STEPS['w_draft'].prompt}.md").read_text(encoding="utf-8")
    assert "Say a thing the plain way first" in draft and "Plain is not casual" in draft and "give each one as \"figure\" data" in draft
    assert STEPS["w_search"].prompt == "w-search-v2" and STEPS["p_search"].prompt == "p-search-v1"  # proposals search as before
    search = (folder / "w-search-v2.md").read_text(encoding="utf-8")
    assert "open textbooks" in search and "use a company's page only for a fact about that company's own product" in search


# --- Codex's audit of 6e513af, finding 2: the claim's own clause decides, and what code cannot judge is read ----------

DISPUTED = "(OR = 1.27; 95% CI: 0.56–2.89; p = 0.043)"


def test_a_word_in_another_clause_never_hides_a_claim():
    """ "There is no electricity in the district, and no study has examined vaccine uptake": the first "no" let the
    claim through."""
    for claim in ("There is no electricity in the district, and no study has examined vaccine uptake.",
                  "Clinics are not staffed at night; no study has examined what that costs caregivers.",
                  "Although the district is not remote, this relationship has never been studied there.",
                  "Because there is no register, no study has examined the eligible population."):
        assert len(evidence.gap_problems(claim)) == 1, claim
        assert evidence.strip_overclaims(claim) == ""


def test_reporting_a_significant_result_still_calls_it_significant():
    """ "The study suggests a significant association" beside the disputed figures passed: "suggests" was taken for
    a doubt. It is how the first such sentence found was written, faithfully, from a publisher's abstract."""
    for verb in ("suggests", "states", "claims", "reports", "found"):
        sentence = f"The study {verb} a significant association between education and uptake {DISPUTED}."
        assert len(evidence.statistic_problems(sentence)) == 1, verb
    said = f"The authors state that the association is significant {DISPUTED}."
    assert len(evidence.statistic_problems(said)) == 1
    listed = f"Education showed a significant association with both knowledge and uptake {DISPUTED}."
    assert evidence.statistic_findings(listed)[0] == [] and len(evidence.statistic_findings(listed)[1]) == 1  # a list's "and" may hide a new clause: read, not judged


def test_what_code_cannot_judge_is_given_to_the_reviewer_and_never_removed():
    denials = ("The topic is not unstudied.",
               "It would be misleading to say that no study has examined this question.",
               "Earlier reviews claimed that no study has examined caregivers, which later work showed to be wrong.",
               "One author wrote that “no study has examined this question”, a claim this review does not accept.",
               "It is not, however, unstudied.",
               "It is not true, as some claim, that no study has examined this question.")
    for sentence in denials:
        problems, notes = evidence.gap_findings(sentence)
        assert problems == [] and len(notes) == 1 and sentence in notes[0] and "could not tell" in notes[0], sentence
        assert evidence.strip_overclaims(sentence) is sentence
    cautions = (f"The result is inconsistent and should not be described as statistically significant {DISPUTED}.",
                f"The association was not significant {DISPUTED}.",
                f"Among women with no schooling the association was significant {DISPUTED}.",  # a "no" that is not about the result: the reviewer reads it
                f"Age was significant and education showed a weaker pattern {DISPUTED}.")  # the word is another clause's: code cannot tie it to this result
    for sentence in cautions:
        problems, notes = evidence.statistic_findings(sentence)
        assert problems == [] and len(notes) == 1 and "0.56–2.89" in notes[0], sentence
        assert evidence.strip_overclaims(sentence) is sentence
    assert evidence.for_the_reviewer([denials[0], cautions[0], "Published evidence is limited."]) == [evidence.gap_findings(denials[0])[1][0], evidence.statistic_findings(cautions[0])[1][0]]
    assert evidence.for_the_reviewer([denials[0], cautions[0]], gaps=False) == [evidence.statistic_findings(cautions[0])[1][0]]  # coursework: the statistic only
    assert evidence.for_the_reviewer(["The role of fathers remains largely unstudied.", "The eligible population remains unknown."]) == []


def test_the_chapter_reviewer_is_sent_the_sentences_to_read_and_the_chapter_keeps_them(client):
    from tests.fake_models import FakeModels

    denial = "It would be misleading to say that no study has examined this question."
    asked: list[dict] = []

    def draft(payload):
        drafted = FakeModels.proposal("p_draft", payload)
        drafted["sections"][0]["paragraphs"][0] += " " + denial
        return drafted

    def review(payload):
        asked.append(payload)
        return FakeModels.proposal("p_review", payload)

    client.models.overrides["p_draft"], client.models.overrides["p_review"] = draft, review
    project_id = _with_chapter_one(client)
    sections = [s for payload in asked for s in payload["sections"]]
    told = [s for s in sections if s.get("forYourJudgement")]
    assert told and all(len(s["forYourJudgement"]) == 1 and denial in s["forYourJudgement"][0] for s in told)
    assert all(not any("claims that nothing has been studied" in c for c in s["paperaidChecks"]) for s in told)  # not a check: it decides nothing
    assert all("forYourJudgement" not in s for s in sections if s not in told)
    chapter = client.get(f"/api/projects/{project_id}/chapters/1", headers=STUDENT).json()
    assert any(denial in p for s in chapter["sections"] for p in s["paragraphs"])  # delivered as written
    from pathlib import Path

    from app.ai.orchestration import STEPS

    folder = Path(__file__).parents[1] / "app" / "ai" / "prompts"
    for step in ("p_review", "w_evaluate"):
        assert '"forYourJudgement"' in (folder / f"{STEPS[step].prompt}.md").read_text(encoding="utf-8"), step


# --- the owner's quality round of 11 October: what recurred in the delivered documents -------------------------------

PASSAGE = ("In a cross-sectional study, the investigator measures the outcome and the exposures in the study participants at the same time. "
           "Unlike in case-control studies or cohort studies, the participants are selected based on the inclusion and exclusion criteria set for the study.")


def _cited(passage: str = PASSAGE) -> dict[str, EvidenceItem]:
    return {"E0a1b2c": _item(support="SUPPORTED", passage=passage, statement="Exposure and outcome are measured at one time.")}


def test_a_source_s_own_wording_is_paraphrased_or_quoted_never_carried_over():
    """In one coursework 20 of 57 sentences carried eight or more of a source's words in a row, one of them 22, each
    with its citation and none in quotation marks."""
    library = _cited()
    copied = "In a cross-sectional study, the investigator measures the outcome and the exposures in the study participants at the same time ⟦E0a1b2c⟧."
    problems = evidence.copied_problems(copied, library)
    assert len(problems) == 1 and "21 words in a row" in problems[0] and "“In a cross-sectional study, the investigator ...”" in problems[0]
    for sound in ("Because exposure and outcome are recorded together, a cross-sectional study cannot show which came first ⟦E0a1b2c⟧.",  # in the writer's own words
                  "Setia describes a design in which “the investigator measures the outcome and the exposures in the study participants at the same time” ⟦E0a1b2c⟧.",
                  "The investigator measures the outcome and the exposures together ⟦E0a1b2c⟧.",  # a short run is how a subject is spoken of
                  "In a cross-sectional study, the investigator measures the outcome and the exposures in the study participants at the same time."):  # cites nothing: not this check's
        assert evidence.copied_problems(sound, library) == [], sound
    # a sentence without a citation of its own is read against what its paragraph cites
    paragraph = "The design is simple ⟦E0a1b2c⟧. Unlike in case-control studies or cohort studies, the participants are selected based on the inclusion and exclusion criteria."
    assert len(evidence.copied_problems(paragraph, library)) == 1
    # a citation token breaks a run: words on either side of one are not one borrowed phrase
    broken = "The investigator measures the outcome and the exposures ⟦E0a1b2c⟧ in the study participants at the same time."
    assert evidence.copied_problems(broken, library) == []


def test_copied_wording_the_repairs_left_is_marked_as_a_quotation():
    library = _cited()
    left = "Its weakness is plain. In a cross-sectional study, the investigator measures the outcome and the exposures in the study participants at the same time ⟦E0a1b2c⟧. This limits it."
    marked = evidence.quote_copied(left, library)
    assert marked == ("Its weakness is plain. “In a cross-sectional study, the investigator measures the outcome and the exposures in the study participants at the same time” ⟦E0a1b2c⟧. "
                      "This limits it.")
    assert evidence.copied_problems(marked, library) == [] and evidence.quote_copied(marked, library) is marked
    own = "Because exposure and outcome are recorded together, the order of events is unknown ⟦E0a1b2c⟧."
    assert evidence.quote_copied(own, library) is own


def test_a_sentence_is_said_once_in_a_document():
    """ "In a cross-sectional study, the investigator measures ... (Setia, 2016)." stood word for word in the
    introduction and in two later sections."""
    said = "Paradigms of inquiry are distinguished through their ontology, epistemology, and methodology ⟦E0a1b2c⟧."
    found = evidence.repeated_sentences([("introduction", [f"Research has two traditions. {said}"]), ("theme1", [f"{said} Each has its own procedures."]),
                                         ("conclusion", ["The two approaches differ in their assumptions and in their use of context, as shown."])])
    assert list(found) == ["theme1"] and "already said, word for word" in found["theme1"][0] and "Paradigms of inquiry" in found["theme1"][0]
    assert evidence.repeated_sentences([("introduction", [said, said])]) == {}  # within one section it is the section's own reviewer's matter
    assert evidence.repeated_sentences([("a", ["It is limited."]), ("b", ["It is limited."])]) == {}  # a short sentence may recur


def test_only_a_figure_goes_out_of_date():
    """ "Methodological literature from 2015 noted that quantitative research requires large samples": the dated
    note had been applied to a statement with nothing in it that dates."""
    old = EvidenceSource(url="https://example.org/m", title="Sampling in qualitative research", year="2011")
    method = _item(support="SUPPORTED", source=old, statement="Quantitative research requires samples large enough to give precise estimates.")
    figure = _item(support="SUPPORTED", source=old, statement="The youth labour force was growing at 4.7 percent a year.")
    dated_year_only = _item(support="SUPPORTED", source=old, statement="Sackett described the bias in 1979.")
    assert evidence.dated(method) == "" and evidence.dated(dated_year_only) == "" and evidence.dated(figure).startswith("Published in 2011.")


def _funding_input():
    from app.works.models import ResolvedSpec, WorkStepInput

    spec = ResolvedSpec.model_construct(kind="FUNDING_PROPOSAL", variant="NGO_PROJECT", mode="COMPACT", level="", source_policy="OPEN", fields=[], citation_style="APA7",
                                        directives=[], coverage=[], active_rules=[], flags={}, limits=[], requirements=[], target_words=3000)
    return WorkStepInput.model_construct(spec=spec, plan=None, inputs=None, results=None, budget=None, kind="FUNDING_PROPOSAL")


def test_a_funding_document_never_speaks_to_the_student():
    """ "However, key operational details remain gaps for the student to confirm. The student must establish
    baseline data ..." was printed in a proposal an organisation would send to a funder."""
    from app.proposals.ai import SectionText, Table
    from app.works import pipeline as works_pipeline
    from app.works.models import PlanSection

    notes = ("However, key operational details remain gaps for the student to confirm.", "The student must establish baseline data from facility records.",
             "Staff salary allocations must be supplied by the student.", "[The student should provide the organisation’s registration number.]",
             "PaperAid has left the targets blank.")
    sound = ("The project will train 40 village health teams.", "Students in upper primary classes will form health clubs.", "Each student receives a reusable pad kit.",
             "The baseline will be established in the first month from facility records.")
    funding, coursework = _funding_input(), _work_input()
    for sentence in notes:
        assert works_pipeline._notes_to_student(funding, sentence) == [sentence], sentence
        assert works_pipeline._notes_to_student(coursework, sentence) == []  # an essay may speak of students
    for sentence in sound:
        assert works_pipeline._notes_to_student(funding, sentence) == [], sentence
    section = PlanSection.model_construct(key="implementation", heading="Implementation", brief="", words=300, min_words=0, max_words=0, field_id="", coverage=[], criteria=[])
    text = SectionText(key="implementation", paragraphs=[f"{sound[0]} {notes[1]} {sound[3]}"], table=Table(caption="", rows=[]))
    found = works_pipeline._checks(funding, section, text, {}, "40", {})
    assert any("is a note to the student" in p and "establish baseline data" in p for p in found)
    assert works_pipeline._strip(funding, text, {}, "40").paragraphs == [f"{sound[0]} {sound[3]}"]  # what a repair left is never delivered
    assert works_pipeline._strip(coursework, text, {}, "40").paragraphs == text.paragraphs


def test_a_results_model_s_text_is_never_cut_in_the_middle_of_a_word():
    """One funding proposal in three stopped at its plan: "Semi-annually upon placement completion for each cohort
    (Mon" was what code had kept of a frequency, the reviewer asked for the truncated text to be completed, and no
    repair could."""
    from app.works.models import Indicator
    from app.works.pipeline import _clip_strings

    frequency = "Semi-annually upon placement completion for each cohort (Months 6 and 12)"
    assert _clip_strings({"indicators": [{"frequency": frequency}]}) == {"indicators": [{"frequency": frequency}]}
    assert Indicator.model_fields["frequency"].metadata[0].max_length == 160
    long = "Quarterly, " + "at the close of each training cycle and again " * 6
    cut = _clip_strings({"frequency": long})["frequency"]
    assert len(cut) <= 160 and long.startswith(cut) and long[len(cut)] == " " and not cut.endswith((",", " "))
    assert _clip_strings({"unit": "x" * 80})["unit"] == "x" * 60  # one unbroken word has nowhere else to end


def test_the_prompts_carry_what_the_round_showed():
    from pathlib import Path

    from app.ai.orchestration import STEPS

    folder = Path(__file__).parents[1] / "app" / "ai" / "prompts"

    def prompt(step: str) -> str:
        return (folder / f"{STEPS[step].prompt}.md").read_text(encoding="utf-8")

    for step in ("p_draft", "w_draft"):
        assert "Put what a source says in your own words" in prompt(step) and "not as wording to reuse" in prompt(step), step
    for step in ("p_fix", "w_repair"):
        assert "words carried over from a source" in prompt(step), step
    for step in ("w_draft", "w_repair"):
        assert "It never mentions \"the student\", PaperAid, or what still has to be supplied" in prompt(step), step
    assert "what the things have in common as well as how they differ" in prompt("w_plan")
    # a Quality Control section left out the planned pilot through two repairs, and its chapter was refused
    assert "Saying what will be done is not inventing a result" in prompt("p_draft") and "describing a planned step is not inventing a result" in prompt("p_fix")
    review = prompt("p_plan_review")
    assert "is not invented merely because neither the student nor the evidence stated it" in review and "Never ask for the plan's proposals to be emptied" in review
    assert STEPS["p_plan_review"].prompt == STEPS["p_plan_review_peer"].prompt == "p-plan-review-v3"


def test_a_sign_off_objection_is_something_a_marker_would_count():
    """The premium model, re-checking a repair, refused a whole chapter because "These enterprises" had its noun a
    sentence away (quality round, 11 October). New engines carry signoff-v2; one priced before keeps its own."""
    from pathlib import Path

    from app.ai.orchestration import VERTEX_ADDENDA

    assert VERTEX_ADDENDA["signoff"] == "signoff-v2"
    folder = Path(__file__).parents[1] / "app" / "ai" / "prompts"
    old, new = (folder / "signoff-v1.md").read_text(encoding="utf-8"), (folder / "signoff-v2.md").read_text(encoding="utf-8")
    assert "A small matter of wording that a repair left behind" in new and "is not an objection" in new
    assert "Every blocking finding in `previousAudit` is resolved" in new and "Approve only when nothing blocking remains" in new  # what v1 required still stands
    assert new.replace(new[new.index("An objection at sign-off"):new.index("Report every unresolved")], "") == old
