"""The independent academic audit of 12 October 2026 (21 delivered documents), first patch: rendering and source
identity. Each test is a defect the audit named, with the cases that must stay as they are beside it."""

import io

import pytest
from docx import Document

from app.analysis import fetch
from app.proposals import evidence, export, pipeline
from app.proposals.models import EvidenceItem, EvidenceSource
from tests.test_api import STUDENT
from tests.test_one_start import works_client  # noqa: F401  (a fixture)
from tests.test_proposal_v2 import _with_chapter_one


def _source(**details) -> EvidenceSource:
    return EvidenceSource(**{"url": "https://example.org/a", "title": "Observational research methods", "year": "2003", **details})


def _library(*sources: EvidenceSource) -> dict[str, EvidenceItem]:
    return {f"E{n:06x}": EvidenceItem(id=f"E{n:06x}", chapter=1, need="n", statement="s", passage="p", scope="", access="ABSTRACT", verified=True, support="SUPPORTED",
                                      source=source, retrieved_on="2026-10-12") for n, source in enumerate(sources, start=1)}


class _Finding:
    def __init__(self, url, title, publisher, published="2025"):
        self.url, self.title, self.publisher, self.published = url, title, publisher, published


@pytest.fixture
def offline(monkeypatch):
    monkeypatch.setattr(fetch, "resolve_doi", lambda url: "")
    monkeypatch.setattr(fetch, "crossref_work", lambda doi: None)
    monkeypatch.setattr(fetch, "crossref_search", lambda text, rows=3: [])
    monkeypatch.setattr(fetch, "page_meta", lambda url: {})
    monkeypatch.setattr(fetch, "page_title", lambda url: "")
    return monkeypatch


def test_a_name_typed_before_a_narrative_citation_is_printed_once():
    """ "As Mann Mann (2003) observes", "As Kim Kim (2023) points out", "The Centers for Disease Control and
    Prevention Centers for Disease Control and Prevention (2012) notes". """
    mann = _source(url="https://example.org/mann", authors=["Mann, C. J."])
    kim = _source(url="https://example.org/kim", authors=["Kim, S.", "Lee, J."], year="2023")
    cdc = _source(url="https://example.org/cdc", organisation="Centers for Disease Control and Prevention", year="2012")
    many = _source(url="https://example.org/many", authors=["Okello, J.", "Namara, P.", "Achen, R."], year="2021")
    citer = evidence.Citer(_library(mann, kim, cdc, many), "APA7")
    assert citer.render("As Mann ⟦E000001|n⟧ observes, these studies are quick.") == "As Mann (2003) observes, these studies are quick."
    assert citer.render("As Kim and Lee ⟦E000002|n⟧ point out, order is unknown.") == "As Kim and Lee (2023) point out, order is unknown."
    assert citer.render("The Centers for Disease Control and Prevention ⟦E000003|n⟧ notes that prevalence differs.") == (
        "The Centers for Disease Control and Prevention (2012) notes that prevalence differs.")
    assert citer.render("Okello et al. ⟦E000004|n⟧ disagree, and Mann ⟦E000001|n⟧ highlights the limit.") == "Okello et al. (2021) disagree, and Mann (2003) highlights the limit."
    # nothing else is touched: a different name, the same name further away, a parenthetical citation, a token alone
    assert citer.render("Smith and ⟦E000001|n⟧ agree.") == "Smith and Mann (2003) agree."
    assert citer.render("Mann wrote the first account, and ⟦E000001|n⟧ later revised it.") == "Mann wrote the first account, and Mann (2003) later revised it."
    assert citer.render("Hermann ⟦E000001|n⟧ is another author's name.") == "Hermann Mann (2003) is another author's name."
    assert citer.render("Mann is cited here ⟦E000001⟧.") == "Mann is cited here (Mann, 2003)."
    assert citer.render("⟦E000001|n⟧ observes this.") == "Mann (2003) observes this."
    # APA 6 names three to five authors in full at a first citation: the name typed in either form is not doubled
    six = evidence.Citer(_library(mann, kim, cdc, many), "APA6")
    assert six.render("Okello, Namara, and Achen ⟦E000004|n⟧ disagree.") == "Okello, Namara, and Achen (2021) disagree."
    assert six.render("Okello et al. ⟦E000004|n⟧ disagree again.") == "Okello et al. (2021) disagree again."


def test_code_numbers_a_table_once():
    """ "Table 3.1: Table 3.1: Work Plan and Activity Schedule", in the list of tables and above the table."""
    assert evidence.caption_text("Table 3.1: Work Plan and Activity Schedule") == "Work Plan and Activity Schedule"
    assert evidence.caption_text("Table 2 - Studies reviewed") == "Studies reviewed" and evidence.caption_text("Figure 1.1. Conceptual framework") == "Conceptual framework"
    for own in ("Work Plan and Activity Schedule", "Tables of contents in ten proposals", "Table salt consumption by district", "", "Timetable 3: fieldwork"):
        assert evidence.caption_text(own) == own, own


def test_a_table_s_label_is_printed_once_in_the_export(client):
    from tests.fake_models import FakeModels

    def draft(payload):
        drafted = FakeModels.proposal("p_draft", payload)
        for section in drafted["sections"]:
            if section.get("table") and section["table"].get("rows"):
                section["table"]["caption"] = "Table 3.1: " + (section["table"].get("caption") or "Work plan")
        return drafted

    client.models.overrides["p_draft"] = draft
    project_id = _with_chapter_one(client)
    exported = client.get(f"/api/projects/{project_id}/export", headers=STUDENT)
    text = "\n".join(p.text for p in Document(io.BytesIO(exported.content)).paragraphs)
    assert "Table 3.1: Table" not in text and "Table 1.1: Table" not in text


def test_a_draft_s_cover_says_what_it_holds_and_what_is_still_to_write(client):
    """A draft of two chapters was read as a proposal that claimed to be whole and had lost its methodology."""
    project_id = _with_chapter_one(client)
    paragraphs = [p.text for p in Document(io.BytesIO(client.get(f"/api/projects/{project_id}/export", headers=STUDENT).content)).paragraphs]
    cover = paragraphs[:paragraphs.index("Declaration")]
    assert any("not yet complete for submission" in p for p in cover)
    assert "Written so far: Chapter One of three." in cover and "Still to write: Chapters Two and Three." in cover

    class _Doc:
        def __init__(self, missing=()):
            self.missing = list(missing)

    assert export.draft_scope({1: _Doc(), 2: _Doc()}) == ["Written so far: Chapters One and Two of three.", "Still to write: Chapter Three."]
    assert export.draft_scope({1: _Doc(), 2: _Doc(), 3: _Doc()}) == ["Written so far: Chapters One, Two and Three of three."]
    assert export.draft_scope({1: _Doc(["1.7", "1.8"])}) == ["Written so far: Chapter One of three.", "Still to write: Chapters Two and Three.",
                                                             "Not yet finished: Chapter One (2 of its sections)."]
    assert export.draft_scope({}) == ["No chapter is written yet.", "Still to write: Chapters One, Two and Three."]


def test_a_site_s_address_is_never_an_author(offline):
    """ "gu.ac.ug. (2025). Cogent Social Sciences." and "(gu.ac.ug, 2025; icpac.net, 2019)": a study by four named
    researchers, cited by the address of the university page that listed it."""
    for host in ("gu.ac.ug", "icpac.net", "https://www.gu.ac.ug/"):
        assert pipeline._web_source(_Finding("https://gu.ac.ug/publication/determinants-of-adoption", title="Cogent Social Sciences", publisher=host)) is None, host
    # the work itself, when its registration can be found by its title, is cited by its authors
    offline.setattr(fetch, "crossref_search", lambda text, rows=3: [{
        "doi": "10.1080/23311886.2025.2588856", "title": "Determinants of adoption of climate-smart agricultural technologies among smallholder coffee farmers in Western Uganda",
        "subtitle": "", "authors": "Kule, E. K.; Kyohangirwe, A.; Midamba, D. C.; Byakatonda, J.", "year": "2025", "container": "Cogent Social Sciences", "volume": "11",
        "issue": "1", "pages": ""}])
    found = pipeline._web_source(_Finding("https://gu.ac.ug/publication/determinants-of-adoption", publisher="gu.ac.ug",
                                          title="Determinants of adoption of climate-smart agricultural technologies among smallholder coffee farmers in Western Uganda"))
    assert found.authors[0] == "Kule, E. K." and len(found.authors) == 4 and found.doi == "10.1080/23311886.2025.2588856" and found.organisation == ""
    # a body that wrote its own page stays its author: this is about an address, not about universities
    own = pipeline._web_source(_Finding("https://gu.ac.ug/research-policy", title="Research and innovations policy", publisher="Gulu University"))
    assert own is not None and own.organisation == "Gulu University"


def test_a_book_is_cited_by_its_authors_when_its_whole_title_is_registered(offline):
    """ "Princeton University Press. (2012). A Tale of Two Cultures: Qualitative and Quantitative Research in the
    Social Sciences." A book's subtitle is registered apart from its title."""
    url, title = "https://assets.press.princeton.edu/chapters/s9898.pdf", "A Tale of Two Cultures: Qualitative and Quantitative Research in the Social Sciences"
    book = {"doi": "10.23943/princeton/9780691149707.001.0001", "title": "A Tale of Two Cultures", "subtitle": "Qualitative and Quantitative Research in the Social Sciences",
            "authors": "Goertz, G.; Mahoney, J.", "year": "2012", "container": "", "volume": "", "issue": "", "pages": ""}
    offline.setattr(fetch, "crossref_search", lambda text, rows=3: [book])
    found = pipeline._web_source(_Finding(url, title=title, publisher="Princeton University Press", published="2012"))
    assert found.authors == ["Goertz, G.", "Mahoney, J."] and found.title == title and found.organisation == ""
    # the main title alone is not the whole title: no looser match is taken for the same work
    other = pipeline._web_source(_Finding(url, title="A Tale of Two Cultures in Rural Schools: A Report", publisher="Princeton University Press", published="2012"))
    assert other.authors == [] and other.organisation == "Princeton University Press"
    wrong_year = pipeline._web_source(_Finding(url, title=title, publisher="Princeton University Press", published="2019"))
    assert wrong_year.authors == []
    record = fetch._crossref_record({"title": ["A Tale of Two Cultures"], "subtitle": ["Qualitative and Quantitative Research in the Social Sciences"],
                                     "author": [{"family": "Goertz", "given": "Gary"}]})
    assert record["subtitle"] == "Qualitative and Quantitative Research in the Social Sciences" and fetch._crossref_record({"title": ["A study"]})["subtitle"] == ""


# --- second patch: an exact provision of a law on the word of a source that is not the law ---------------------------

def _cited(title: str, url: str = "https://educationoutloud.org/brief.pdf", organisation: str = "Education Out Loud") -> dict[str, EvidenceItem]:
    return _library(_source(url=url, title=title, year="2026", organisation=organisation))


BRIEF = "Policy Brief: Enhancing School Feeding in Uganda"
ACT = "Education (Pre-Primary, Primary and Post-Primary) Act, 2008"


def test_a_provision_s_number_is_printed_only_on_the_word_of_the_law_itself():
    """ "Section 13(5)(2c) of the Education Act 2008" (the duty is in section 6(2)(c)): copied faithfully from a
    policy brief that numbered it wrongly. PaperAid's quote check had confirmed that the brief says so."""
    brief, law = _cited(BRIEF), _cited(ACT, "https://example.go.ug/education-act-2008.pdf", "Parliament of Uganda")
    cases = {
        # the law is still named, code supplies no number of its own, and the account is the cited source's
        "School feeding operates under Section 13(5)(2c) of the Education Act 2008, which places the duty on parents ⟦E000001⟧.":
            "The cited source reports: School feeding operates under the Education Act 2008, which places the duty on parents ⟦E000001⟧.",
        "Section 13(5)(2c) of the Education Act 2008 places the duty on parents ⟦E000001⟧.": "The cited source reports that the Education Act 2008 places the duty on parents ⟦E000001⟧.",
        "The guidelines conform to the Education Act (Article 13, sub-section 5[2c]) ⟦E000001⟧.": "The cited source reports that the guidelines conform to the Education Act ⟦E000001⟧.",
        # a name keeps its capital letters: only a word that is never a name is lowered
        "Uganda's Education Act (Section 13) places the duty on parents ⟦E000001⟧.": "The cited source reports: Uganda's Education Act places the duty on parents ⟦E000001⟧.",
        "UNESCO guidance follows Article 28 of the Convention ⟦E000001⟧.": "The cited source reports: UNESCO guidance follows the Convention ⟦E000001⟧.",
    }
    for sentence, plain in cases.items():
        problems, notes = evidence.provision_findings(sentence, brief)
        assert len(problems) == 1 and notes == [] and "on the word of a source that is not the law itself" in problems[0], sentence
        assert evidence.drop_provisions(sentence, brief) == plain
        assert evidence.provision_findings(sentence, law)[0]  # a title alone does not verify a provision
    # the law's own publisher is the law, whatever the page is called
    on_ulii = _cited("Act 13 of 2008", "https://ulii.org/akn/ug/act/2008/13/eng@2023-12-31", "Uganda Legal Information Institute")
    first = next(iter(cases))
    assert evidence.provision_findings(first, on_ulii)[0]  # the host cannot vouch for an unidentified instrument
    # one law never vouches for another: the Education Act says nothing of an article of the Constitution
    article = "Article 34 of the Constitution guarantees the right ⟦E000001⟧."
    constitution = _cited("Constitution of the Republic of Uganda, 1995", "https://example.go.ug/constitution.pdf", "Parliament of Uganda")
    for other in (brief, law):
        assert len(evidence.provision_findings(article, other)[0]) == 1
        assert evidence.drop_provisions(article, other) == "The cited source reports that the Constitution guarantees the right ⟦E000001⟧."
    constitution["E000001"].passage = "Article 34 of the Constitution addresses children's rights."
    assert evidence.provision_findings(article, constitution) == ([], []) and evidence.drop_provisions(article, constitution) is article


def test_two_cited_sources_are_named_as_two():
    library = _library(_source(url="https://educationoutloud.org/brief.pdf", title=BRIEF, organisation="Education Out Loud"),
                       _source(url="https://example.org/circular", title="Ministry circular on school feeding", organisation="Ministry of Education and Sports"))
    sentence = "In practice, schools rely on Section 13(5)(2c) of the Education Act 2008 and on circulars ⟦E000001⟧⟦E000002⟧."
    assert evidence.drop_provisions(sentence, library) == "The cited sources report that in practice, schools rely on the Education Act 2008 and on circulars ⟦E000001⟧⟦E000002⟧."


def test_only_an_exact_provision_of_a_named_law_is_this_check_s_matter():
    brief = _cited(BRIEF)
    for sound in ("Parents must provide food under the Education Act 2008 ⟦E000001⟧.",  # the law, no number
                  "Section 3 of the report describes the sample ⟦E000001⟧.",  # no law is named
                  "The Act was amended in 2023 ⟦E000001⟧."):
        assert evidence.provision_findings(sound, brief) == ([], []), sound
        assert evidence.drop_provisions(sound, brief) is sound
    # written a way code cannot safely shorten: left exactly as it is, and given to the reviewer to read
    other = "The Education Act 2008, at s. 6(2)(c), places the duty on parents ⟦E000001⟧."
    problems, notes = evidence.provision_findings(other, brief)
    assert problems == [] and len(notes) == 1 and other in notes[0] and evidence.drop_provisions(other, brief) is other
    assert evidence.for_the_reviewer([other], library=brief) == notes and evidence.for_the_reviewer([other]) == []
    # one paragraph: only the sentence at fault changes, and its neighbours keep their wording
    paragraph = "Meals matter for attendance ⟦E000001⟧. Section 13(5)(2c) of the Education Act 2008 places the duty on parents ⟦E000001⟧. Uptake varies."
    assert evidence.drop_provisions(paragraph, brief) == "Meals matter for attendance ⟦E000001⟧. The cited source reports that the Education Act 2008 places the duty on parents ⟦E000001⟧. Uptake varies."


def test_a_chapter_is_checked_and_its_reviewer_told_before_a_number_is_taken_out():
    from app.proposals.ai import SectionText, Table
    from app.proposals.models import StepInput

    brief = _cited(BRIEF)
    inp = StepInput.model_construct(chapter=1, plan=None, academic_policy=1)
    text = SectionText(key="background", paragraphs=["Section 13(5)(2c) of the Education Act 2008 places the duty on parents ⟦E000001⟧."], table=Table(caption="", rows=[]))
    found = pipeline._checks(inp, "background", text, brief, "13 5 2c 2008")
    assert any("exact provision of the Education Act" in p for p in found)
    unreadable = SectionText(key="background", paragraphs=["The Education Act 2008, at s. 6(2)(c), places the duty on parents ⟦E000001⟧."], table=Table(caption="", rows=[]))
    assert not any("exact provision" in p for p in pipeline._checks(inp, "background", unreadable, brief, "6 2 2008"))
    assert len(pipeline._to_read(unreadable, brief)) == 1 and pipeline._to_read(unreadable) == []


def test_a_work_s_provision_number_is_taken_out_before_its_last_review():
    from app.proposals.ai import SectionText, Table
    from app.works import pipeline as works_pipeline

    brief = _cited(BRIEF)
    current = {"context": SectionText(key="context", paragraphs=["Section 13(5)(2c) of the Education Act 2008 places the duty on parents ⟦E000001⟧.", "A second paragraph."],
                                      table=Table(caption="", rows=[])),
               "approach": SectionText(key="approach", paragraphs=["Parents must provide food under the Education Act 2008 ⟦E000001⟧."], table=Table(caption="", rows=[]))}
    untouched = current["approach"]
    assert works_pipeline._unnumber(current, ["context", "approach"], brief) == ["context"]
    assert current["context"].paragraphs == ["The cited source reports that the Education Act 2008 places the duty on parents ⟦E000001⟧.", "A second paragraph."] and current["approach"] is untouched
    assert works_pipeline._unnumber(current, ["context", "approach"], brief) == []


# --- third patch: funding documents ------------------------------------------------------------------------


def _line(n: int, **figures):
    from app.works.models import BudgetLine

    return BudgetLine(**{"id": f"B{n}", "category": "Training", "description": f"Line {n}", "quantity": 0, "unit": "session", "unit_cost": 0, **figures})


def test_an_empty_budget_is_not_a_complete_one_and_a_line_given_in_kind_is_not_a_gap():
    from app.works import budget as budget_engine
    from app.works import numbers
    from app.works.export import IN_KIND, _budget_tables
    from app.works.models import Budget, ResolvedSpec, WorkInputs
    from app.works.service import BUDGET_FIGURES

    assert not budget_engine.complete(Budget(lines=[]))  # nothing costed: `all` of no lines used to say complete
    assert not budget_engine.complete(Budget(lines=[_line(1, quantity=4, unit_cost=250), _line(2)]))  # a cost of nothing is "not entered yet"
    assert not budget_engine.complete(Budget(lines=[_line(1, quantity=4)]))
    assert budget_engine.complete(Budget(lines=[_line(1, quantity=4, unit_cost=250)]))
    # the applicant says a line costs the budget nothing: known, at no cost, and it adds nothing to the total
    hall = _line(2, description="Training hall", quantity=12, unit="month", in_kind=True)
    budget = Budget(lines=[_line(1, quantity=4, unit_cost=250), hall])
    assert budget_engine.complete(budget) and budget_engine.known(hall) and budget_engine.line_total(hall) == 0
    assert budget_engine.totals(budget).total == 1000
    summary, detail = _budget_tables(budget)
    assert summary[-1] == ["Total", "USD 1,000"] and detail[2][2:6] == ["12", "month", IN_KIND, "USD 0"]
    assert "to be added" not in str(summary) + str(detail)
    spec = ResolvedSpec(version=1, kind="FUNDING_PROPOSAL", variant="NGO_PROJECT", mode="STANDARD", rules_version="rules-v1", target_words=5000)
    assert numbers.values(spec, None, budget, WorkInputs(title="Safer deliveries"))["budget.total"][0] == "USD 1,000"
    # an unmarked line beside it is still a gap, and so is every total it affects
    summary, detail = _budget_tables(Budget(lines=[hall, _line(3)]))
    assert summary[-1] == ["Total", "[to be added]"] and detail[2][5] == "[to be added]"
    assert numbers.values(spec, None, Budget(lines=[hall, _line(3)]), WorkInputs(title="Safer deliveries"))["budget.total"][0] == "[total to be added]"
    assert "in_kind" in BUDGET_FIGURES  # marking a line is adding a figure: no AI, no charge, the text unchanged


def _funding_document(status: str, budget, exploratory: bool = False):
    from app.works.models import ResolvedSpec, WorkDocument, WorkSection

    spec = ResolvedSpec(version=1, kind="FUNDING_PROPOSAL", variant="NGO_PROJECT", mode="STANDARD", rules_version="rules-v1", target_words=5000, duration_months=12)
    document = WorkDocument(kind="FUNDING_PROPOSAL", variant="NGO_PROJECT", title="Safer deliveries", spec_version=1, plan_version=1, status=status, exploratory=exploratory,
                            sections=[WorkSection(key="budget_narrative", heading="Budget", paragraphs=["Costs follow the activities."])])
    return document, spec


def test_a_file_that_prints_a_marked_gap_says_on_its_front_that_it_is_a_draft():
    from app.works import export as works_export
    from app.works.export import DRAFT_LABEL, EXPLORATORY, GAPS_LABEL
    from app.works.models import Budget

    gaps, whole = Budget(lines=[_line(1)]), Budget(lines=[_line(1, quantity=4, unit_cost=250)])

    def labels(status, budget, draft=True, exploratory=False):
        document, spec = _funding_document(status, budget, exploratory)
        return [b.text for b in works_export.layout(document, spec, None, budget, {}, {}, draft) if b.kind == "label" and b.text in (DRAFT_LABEL, GAPS_LABEL, EXPLORATORY)]

    assert labels("READY_WITH_WARNINGS", gaps) == [GAPS_LABEL] == labels("READY", gaps)  # its checks pass, and it is still not ready to send
    assert labels("READY_WITH_WARNINGS", whole) == [] == labels("READY", whole)  # nothing is missing: no label
    assert labels("NOT_READY", gaps) == [DRAFT_LABEL] and labels("READY", gaps, exploratory=True) == [EXPLORATORY]  # one label, never two
    assert labels("READY_WITH_WARNINGS", gaps, draft=False) == []  # the final review reads the text; the label is decided after it
    document, spec = _funding_document("READY_WITH_WARNINGS", gaps)
    word = Document(io.BytesIO(works_export.build(document, spec, None, gaps, {}, {}, draft=True)))
    front = [p.text for p in word.paragraphs][:4]
    assert GAPS_LABEL in front and front.index(GAPS_LABEL) < front.index("Budget")  # under the title, before the text


def _schedule(*months: tuple[int, int]):
    from app.works.models import Activity, ResultsModel

    names = ("Recruit and train mentors", "First cohort of placements", "Second cohort of placements", "Follow up every graduate")
    return ResultsModel(activities=[Activity(id=f"A{n}", statement=names[n - 1], start_month=a or None, end_month=b or None) for n, (a, b) in enumerate(months, start=1)])


def test_the_reviewer_is_given_the_printed_schedule_and_code_decides_nothing_about_it():
    """The audit's case: set-up, two six-month cohorts one after the other and a follow-up six months later were
    promised inside twelve months, while each activity's own months fitted. Code cannot read a promise; it gives the
    months to the reviewer, alike for a schedule that cannot hold and for one that overlaps soundly."""
    from app.works.results import schedule_note

    promise = "After two months of set-up, two six-month cohorts follow one another, and each graduate is followed up six months later."
    impossible, overlapping = _schedule((1, 2), (3, 8), (6, 11), (12, 12)), _schedule((1, 2), (3, 8), (5, 10), (9, 12))
    note = schedule_note(impossible, 12, promise, False)
    assert "for a project of 12 months" in note and "A2 (First cohort of placements): months 3 to 8" in note and "A3 (Second cohort of placements): months 6 to 11" in note
    assert "A4 (Follow up every graduate): month 12" in note
    sound = schedule_note(overlapping, 12, "Two cohorts run in parallel for part of the year.", True)
    assert "A3 (Second cohort of placements): months 5 to 10" in sound and "Activities may run at the same time" in sound
    for text in (note, sound):  # the same neutral statement for both: no verdict is code's to give
        assert not any(word in text.lower() for word in ("impossible", "cannot fit", "does not fit", "infeasible", "wrong"))
    # only a section that prints the workplan or speaks of weeks or months is read beside it
    assert schedule_note(impossible, 12, "The project strengthens referral between health centres.", False) == ""
    assert schedule_note(impossible, 12, "The project strengthens referral between health centres.", True) != ""
    assert schedule_note(impossible, 12, "Mentors meet their group in months 3 to 8.", False) != ""
    assert schedule_note(None, 12, promise, True) == "" and schedule_note(_schedule(), 12, promise, True) == ""  # no schedule: nothing to read
    assert "A1 (Recruit and train mentors): no months given" in schedule_note(_schedule((0, 0)), None, promise, True)


def test_the_reviewer_is_told_when_a_headline_amount_stands_beside_a_budget_with_no_total():
    from app.works.models import Budget
    from app.works.numbers import uncosted_note

    ceiling, asked = "The call funds up to ⟦N:ceiling⟧ for each project.", "We propose a request of ⟦N:answer.amount_requested⟧."
    for budget in (None, Budget(lines=[]), Budget(lines=[_line(1)])):  # no budget, no lines, or lines without costs
        assert "the call's ceiling" in uncosted_note(ceiling, budget) and "the amount the applicant said they request" in uncosted_note(asked, budget)
        assert "the call's ceiling and the amount the applicant said they request" in uncosted_note(ceiling + " " + asked, budget)
        assert uncosted_note("The budget pays for training and transport.", budget) == ""  # no headline amount: nothing to read
        assert uncosted_note("The total is ⟦N:budget.total⟧.", budget) == ""  # prints as a marked gap, which says so itself
    costed = Budget(lines=[_line(1, quantity=4, unit_cost=250)])
    assert uncosted_note(ceiling + " " + asked, costed) == ""  # a total exists: code's own checks compare it with the ceiling and the request


def test_a_funding_draft_s_reviewer_reads_the_schedule_and_a_coursework_s_reads_none(works_client):  # noqa: F811
    from app.ai.orchestration import STEPS
    from tests.test_one_start import H, _coursework, _done, _funding, _settled
    from tests.test_single_reviewer import _requests

    assert (STEPS["w_draft"].prompt, STEPS["w_repair"].prompt, STEPS["w_evaluate"].prompt) == ("w-draft-v10", "w-repair-v10", "w-evaluate-v6")
    client = works_client
    work = _funding(client)
    client.post(f"/api/works/{work['id']}/start", headers=H)
    work = _settled(client, f"/api/works/{work['id']}", _done, timeout=180)
    assert work["documents"], work["autoFailure"]
    read = [note for request in _requests(client, "w_evaluate") for section in request["sections"] for note in section.get("forYourJudgement", [])]
    schedules = [note for note in read if note.startswith("The schedule this proposal commits to")]
    assert schedules and all("for a project of 12 months" in note for note in schedules)
    first = work["results"]["activities"][0]
    assert all(f"{first['id']} (" in note for note in schedules)  # every activity, as the workplan prints it
    seen = len(client.models.tasks)
    work = _coursework(client)
    client.post(f"/api/works/{work['id']}/start", headers=H)
    work = _settled(client, f"/api/works/{work['id']}", _done, timeout=180)
    assert work["documents"], work["autoFailure"]
    later = [r for t, r in list(zip(client.models.tasks, client.models.requests, strict=True))[seen:] if t == "w_evaluate"]
    assert later and not any("The schedule this proposal commits to" in r or "No budget total exists yet" in r for r in later)


# --- the audited records themselves ------------------------------------------------------------------------
# tests/fixtures/independent_audit.json: the passages the audit named and the evidence they cite, as the live
# algorithm stored them in the quality round of 11 October (public bibliographic records and page passages only).


def _audited() -> dict:
    import json
    from pathlib import Path

    return json.loads((Path(__file__).parent / "fixtures" / "independent_audit.json").read_text(encoding="utf-8"))


def _stored(record: dict) -> dict[str, EvidenceItem]:
    return {item["id"]: EvidenceItem.model_validate(item) for item in record["evidence"]}


def test_the_audited_sentences_print_each_author_s_name_once():
    import re

    records = _audited()["doubled"]
    assert len(records) >= 4
    for record in records:
        library = _stored(record)
        for style in ("APA7", "APA6", "HARVARD"):
            citer = evidence.Citer(library, style)
            printed = citer.render(record["paragraph"])
            for token in re.findall(r"⟦(E[0-9a-f]{6})\|n⟧", record["paragraph"]):
                for first in (True, False):
                    name = evidence.author_label(library[token].source, style, first, True)
                    bare = re.sub(r"\s*\([^()]*\)$", "", name)  # a body's name without the abbreviation it prints with
                    assert f"{name} {name}" not in printed and f"{bare} {bare}" not in printed, (style, name)
            assert "⟦" not in printed
    # the sentences that had the fault are among them, and the wording around the name is untouched
    centres = next(r for r in records if "The Centers for Disease Control and Prevention ⟦" in r["paragraph"])
    assert "The Centers for Disease Control and Prevention (" in evidence.Citer(_stored(centres), "APA7").render(centres["paragraph"])
    kim = next(r for r in records if "As Kim ⟦" in r["paragraph"])
    assert re.search(r"As Kim \(\d{4}\) points out", evidence.Citer(_stored(kim), "APA7").render(kim["paragraph"]))


def test_the_audited_passage_names_the_education_act_without_the_number_its_brief_gave():
    records = _audited()["provision"]
    assert records
    for record in records:
        library, paragraph = _stored(record), record["paragraph"]
        problems = evidence.provision_problems(paragraph, library)
        assert problems and all("on the word of a source that is not the law itself" in p for p in problems)
        told = evidence.drop_provisions(paragraph, library)
        assert "13(5)" not in told and "5[2c]" not in told and "Education" in told and "Act" in told  # the law is still named
        assert evidence.provision_findings(told, library) == ([], [])
        before, after = evidence.sentences(paragraph), evidence.sentences(told)
        assert len(before) == len(after) and sum(a != b for a, b in zip(before, after, strict=True)) == len(problems)  # only the sentences at fault
        assert evidence.cited_ids(told) == evidence.cited_ids(paragraph)  # every citation stays where it was


def test_the_audited_caption_and_site_addresses(offline):
    audited = _audited()
    assert [evidence.caption_text(r["caption"]) for r in audited["caption"]] == ["Work Plan and Activity Schedule"]
    hosts = set()
    for record in audited["site_author"]:
        source = record["source"]
        hosts.add(source["organisation"])
        assert pipeline._web_source(_Finding(source["url"], source["title"], source["organisation"], source["year"])) is None  # an address wrote nothing
    assert hosts == {"gu.ac.ug", "icpac.net"}


# --- verification of the takeover of 13 October -------------------------------------------------------------


def test_the_methods_contract_is_off_until_a_trial_proves_it():
    """A new plan with the contract depends on the model returning a methods object in a strict shape, and a plan
    without it fails. No real model has been asked for one, so an engine priced now asks for a plan exactly as the
    live path did (the prompts it was proven with) unless METHODS_CONTRACT is on; nothing else of the engine differs."""
    from app.ai.orchestration import METHODS_OFF, STEPS, current_engine
    from app.core.config import Settings
    from app.proposals.ai import plan_schema

    off, on = current_engine(Settings()), current_engine(Settings(methods_contract=True))
    assert Settings().methods_contract is False
    assert {k: off.prompts[k] for k in METHODS_OFF} == {"p_plan": "p-plan-v5", "p_finalise": "p-finalise-v4", "p_plan_review": "p-plan-review-v3", "p_plan_review_peer": "p-plan-review-v3"}
    assert {k: on.prompts[k] for k in METHODS_OFF} == {k: STEPS[k].prompt for k in METHODS_OFF} == {"p_plan": "p-plan-v6", "p_finalise": "p-finalise-v5", "p_plan_review": "p-plan-review-v4",
                                                                                                 "p_plan_review_peer": "p-plan-review-v4"}
    assert {k: v for k, v in off.prompts.items() if k not in METHODS_OFF} == {k: v for k, v in on.prompts.items() if k not in METHODS_OFF}
    assert "methods" not in plan_schema(off.prompts["p_plan"])["properties"] and "methods" in plan_schema(on.prompts["p_plan"])["properties"]


def test_a_plan_without_the_contract_is_given_to_a_model_as_it_was_before_the_new_fields():
    from app.proposals.models import ProposalPlan
    from tests.test_framework import PLAN

    plan = ProposalPlan.model_validate(PLAN)
    data = pipeline._plan_data(plan)
    assert "methods" not in data and not {"designEffect", "nonresponseRate", "adjustmentSource"} & set(data["sampleSize"])
    assert data == {k: v for k, v in plan.model_dump(by_alias=True).items() if k != "methods"} | {"sampleSize": data["sampleSize"]}
    adjusted = plan.model_copy(update={"sample_size": plan.sample_size.model_copy(update={"design_effect": 1.5, "adjustment_source": "The district survey's report"})})
    assert pipeline._plan_data(adjusted)["sampleSize"]["designEffect"] == 1.5  # what the student supplied is given

