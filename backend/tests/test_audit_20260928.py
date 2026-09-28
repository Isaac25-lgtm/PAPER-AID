"""Regressions for Codex's full audit of 655cda5 (docs/Codex_Full_Code_Audit_20260928_655cda5.md).
Each test asserts the correct behaviour where Codex's diagnostic reproduced the defect; the finding
number is in each name. The browser findings (#9, #15) are covered by the web journeys."""

import io
from datetime import timedelta

import httpx
import pytest
from docx import Document

from app.jobs.models import ReadinessItem, utcnow
from app.proposals import decisions, evidence, export, pipeline
from app.proposals.ai import ProposalRunner
from app.proposals.models import ChapterDocument, ChapterSection, ProposalInputs, StepInput
from app.runtime import get_runtime
from tests.fake_models import PLAN
from tests.test_api import ADMIN, STUDENT
from tests.test_proposals import DETAILS, LIB, _approved, _create, _run, plan


def _chapter_ready(client):
    project = _create(client)
    _run(client, project["id"], "PLAN")
    size = {**PLAN["sampleSize"], "populationSource": "DHO records"}
    _approved(client, project["id"], sampleSize=size)
    return project["id"]


# --- 1: source fetching connects only to the address it validated -----------------------------


def test_01_a_host_is_resolved_once_and_the_request_goes_to_that_address(monkeypatch):
    import socket

    from app.analysis import fetch

    calls, seen = [], []
    original = socket.getaddrinfo

    def resolver(host, port, *args, **kwargs):
        if host == "rebind.example":
            calls.append(host)
            return original("93.184.216.34" if len(calls) == 1 else "127.0.0.1", port, *args, **kwargs)
        return original(host, port, *args, **kwargs)

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append((request.url.host, request.headers["host"], request.extensions.get("sni_hostname")))
        return httpx.Response(200, text="page", headers={"content-type": "text/plain"})

    real_client = httpx.Client
    monkeypatch.setattr(socket, "getaddrinfo", resolver)
    monkeypatch.setattr(fetch.httpx, "Client", lambda **kw: real_client(transport=httpx.MockTransport(handler), **kw))
    got = fetch._get("https://rebind.example/report")
    assert got is not None and got[0] == b"page"
    assert calls == ["rebind.example"]  # one resolution: nothing can rebind between check and connect
    assert seen == [("93.184.216.34", "rebind.example", "rebind.example")]  # pinned IP, real Host and TLS name


def test_01_a_host_resolving_to_a_private_address_is_never_requested(monkeypatch):
    import socket

    from app.analysis import fetch

    seen = []
    original = socket.getaddrinfo
    monkeypatch.setattr(socket, "getaddrinfo", lambda host, port, *a, **k: original("127.0.0.1" if host == "inside.example" else host, port, *a, **k))
    real_client = httpx.Client
    monkeypatch.setattr(fetch.httpx, "Client", lambda **kw: real_client(transport=httpx.MockTransport(lambda r: seen.append(r) or httpx.Response(200)), **kw))
    assert fetch._get("http://inside.example/secret") is None and seen == []


# --- 2: formatting keeps a later section's own header and footer ---------------------------------


def test_02_formatting_keeps_a_later_sections_distinct_header_and_footer():
    from docx.enum.section import WD_SECTION

    from app.documents.docx_io import read_docx
    from app.formatting.apply import apply_formatting
    from app.formatting.presets import PRESETS

    doc = Document()
    doc.add_paragraph("The original body wording is unchanged.")
    doc.add_section(WD_SECTION.NEW_PAGE)
    later = doc.sections[1]
    later.header.is_linked_to_previous = False
    later.header.paragraphs[0].text = "Distinct chapter header"
    later.footer.is_linked_to_previous = False
    later.footer.paragraphs[0].text = "Distinct chapter footer"
    source = io.BytesIO()
    doc.save(source)
    output, result = apply_formatting(source.getvalue(), PRESETS["apa7"], read_docx(source.getvalue()))
    produced = Document(io.BytesIO(output))
    assert "Distinct chapter header" in " ".join(p.text for p in produced.sections[1].header.paragraphs)
    assert "Distinct chapter footer" in " ".join(p.text for p in produced.sections[1].footer.paragraphs)
    assert any("left it unchanged" in w for w in result.warnings)  # told, not silently numbered


# --- 3: project files are outside the bucket's fixed-age deletion ----------------------------


def test_03_project_files_live_outside_the_job_file_backstop(client):
    from pathlib import Path

    project = get_runtime().store.get_project(_create(client)["id"])
    assert project.storage_prefix().startswith("projects/")
    guide = (Path(__file__).parents[2] / "docs" / "deployment.md").read_text(encoding="utf-8")
    assert '"matchesPrefix": ["users/"]' in guide


# --- 4: a project cannot be deleted between its claim and the credits being held -------------


def test_04_a_deletion_that_wins_the_race_refuses_the_step_and_holds_nothing(client, monkeypatch):
    rt = get_runtime()
    pid = _create(client)["id"]
    quoted = client.post(f"/api/projects/{pid}/steps", headers=STUDENT, json={"step": "PLAN"}).json()
    original = rt.store.update_job_wallet_and_project

    def delete_first(job_id, project_id, mutate):
        assert client.delete(f"/api/projects/{pid}", headers=STUDENT).status_code == 204
        return original(job_id, project_id, mutate)

    monkeypatch.setattr(rt.store, "update_job_wallet_and_project", delete_first)
    owner = rt.store.get_project(pid).owner_uid
    held_before = rt.store.get_wallet(owner).held
    response = client.post(f"/api/projects/{pid}/steps/{quoted['job']['id']}/submit", headers=STUDENT, json={"quoteId": quoted["quote"]["id"]})
    assert response.status_code in (404, 409)
    assert rt.store.get(quoted["job"]["id"]) is None  # the deletion took the priced step with it
    assert rt.store.get_wallet(owner).held == held_before


def test_04_a_step_that_wins_the_race_blocks_the_deletion(client, monkeypatch):
    rt = get_runtime()
    pid = _create(client)["id"]
    monkeypatch.setattr(rt.queue, "enqueue", lambda *a, **k: None)  # keep the step queued
    quoted = client.post(f"/api/projects/{pid}/steps", headers=STUDENT, json={"step": "PLAN"}).json()
    assert client.post(f"/api/projects/{pid}/steps/{quoted['job']['id']}/submit", headers=STUDENT, json={"quoteId": quoted["quote"]["id"]}).status_code == 200
    deletion = client.delete(f"/api/projects/{pid}", headers=STUDENT)
    assert deletion.status_code == 409 and rt.store.get_project(pid) is not None
    assert rt.store.get_project(pid).active_job == quoted["job"]["id"]


def test_04_an_admin_cannot_retry_a_step_onto_a_deleted_project(client, monkeypatch):
    rt = get_runtime()
    pid = _create(client)["id"]
    client.models.refuse.add("p_needs")  # the step fails permanently
    job = _run(client, pid, "PLAN")
    assert job["status"] == "FAILED"

    def retryable(j):
        j.failure = j.failure.model_copy(update={"retryable": True})
        return j

    rt.store.update(job["id"], retryable)
    project = rt.store.get_project(pid)
    rt.store.update_project(pid, lambda p: p.model_copy(update={"deleting": True}))
    response = client.post(f"/api/admin/jobs/{job['id']}/retry", headers=ADMIN)
    assert response.status_code == 409 and response.json()["code"] == "PROJECT_DELETED"
    assert project is not None and rt.store.get(job["id"]).status.value == "FAILED"


# --- 5: expiry never deletes a project renewed after it was listed; every page is swept ----------


def test_05_a_project_renewed_after_listing_survives_cleanup(client, monkeypatch):
    from app.jobs.service import cleanup_expired_projects

    rt = get_runtime()
    pid = _create(client)["id"]
    rt.store.update_project(pid, lambda p: p.model_copy(update={"expires_at": utcnow() - timedelta(seconds=1)}))
    original = rt.store.expired_projects

    def renew_meanwhile(cutoff, limit):
        listed = original(cutoff, limit)
        assert client.post(f"/api/projects/{pid}/details", headers=STUDENT, json=DETAILS).status_code == 200
        return listed

    monkeypatch.setattr(rt.store, "expired_projects", renew_meanwhile)
    assert cleanup_expired_projects(rt) == 0 and rt.store.get_project(pid) is not None


def test_05_cleanup_sweeps_past_the_first_page(client, monkeypatch):
    from app.jobs import service

    rt = get_runtime()
    ids = [_create(client)["id"] for _ in range(3)]
    for pid in ids:
        rt.store.update_project(pid, lambda p: p.model_copy(update={"expires_at": utcnow() - timedelta(days=1)}))
    monkeypatch.setattr(service, "PAGE_SIZE", 1)
    assert service.cleanup_expired_projects(rt) == 3 and all(rt.store.get_project(i) is None for i in ids)


# --- 6: deleting a project erases its steps' private inputs and saved answers ------------------


def test_06_project_deletion_erases_every_step_job(client):
    rt = get_runtime()
    pid = _create(client)["id"]
    job = _run(client, pid, "PLAN")
    prefix = rt.store.get(job["id"]).storage_prefix()
    assert b"Namukasa" in rt.files.get(f"{prefix}/internal/proposal_input.json")
    assert client.delete(f"/api/projects/{pid}", headers=STUDENT).status_code == 204
    assert rt.store.get(job["id"]) is None and not rt.files.exists(f"{prefix}/internal/proposal_input.json")


# --- 7: tables and captions are checked, repaired and rendered like prose ---------------------


def test_07_tables_and_captions_are_checked_and_rendered(client):
    pid = _chapter_ready(client)
    original = client.models.default

    def bad_table(payload):
        answer = original("p_draft", payload)
        for section in answer["sections"]:
            if section["key"] == "workplan":
                section["table"] = {"caption": "Smith (2019) found 73%", "rows": [["Activity", "Finding"], ["Interview", "73% ⟦E999999⟧"], ["Pilot", "Month 5"]]}
        return answer

    client.models.overrides["p_draft"] = bad_table
    job = _run(client, pid, "CHAPTER_3")
    assert job["outcome"] == "PARTIAL"
    doc = Document(io.BytesIO(client.get(f"/api/projects/{pid}/export", headers=STUDENT).content))
    text = " ".join(p.text for p in doc.paragraphs) + " " + " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "73%" not in text and "⟦" not in text and "Smith" not in text
    assert "Month 5" in text  # a month of the student's own timeline is not an invented figure


# --- 8: a complete export needs every required section and no integrity failures --------------


def test_08_complete_export_refuses_empty_chapters_and_integrity_failures(client):
    pid = _chapter_ready(client)
    project = get_runtime().store.get_project(pid)
    chapters = {}
    for n in (1, 2, 3):
        project.chapter(n).approved = True
        chapters[n] = ChapterDocument(
            number=n, title="Chapter", plan_version=project.plan_version, sections=[], cited=[], words=0,
            readiness=[
                ReadinessItem(id="X-MISSING", question="Required element", basis="CODE", status="MISSING", note="Absent.", chapter=n),
                ReadinessItem(id="C2-REFS30", question="30 references", basis="CODE", status="NEEDS_REVIEW", note="12.", chapter=n),
            ],
        )
    blockers = export.final_blockers(project, chapters)
    assert any("missing required sections" in b for b in blockers)
    assert any("Required element" in b for b in blockers)
    assert not any("30 references" in b for b in blockers)  # a recommendation stays advisory


# --- 10, 11, 16: review coverage, repair bounds and reviewer notes -----------------------------


def test_10_an_unreviewed_chapter_is_partial_and_says_so(client, monkeypatch):
    pid = _chapter_ready(client)

    def no_review(self, *args, **kwargs):
        self.budget_reached = True
        return {}

    monkeypatch.setattr(ProposalRunner, "grade", no_review)
    job = _run(client, pid, "CHAPTER_1")
    assert job["outcome"] == "PARTIAL" and any("could not be reviewed" in w for w in job["warnings"])
    chapter = client.get(f"/api/projects/{pid}/chapters/1", headers=STUDENT).json()
    assert next(r for r in chapter["readiness"] if r["id"] == "C1-REVIEWED")["status"] == "NEEDS_REVIEW"


def test_11_there_are_at_most_two_fixes_and_the_last_text_is_reviewed(client, monkeypatch):
    pid = _chapter_ready(client)
    client.models.overrides["p_review"] = lambda p: {"results": [{"key": s["key"], "grade": "REPAIR", "issues": ["Correct the wording."], "note": ""} for s in p["sections"]]}
    sequence = []
    grade, fix = ProposalRunner.grade, ProposalRunner.fix
    monkeypatch.setattr(ProposalRunner, "grade", lambda self, *a, **k: (sequence.append("review"), grade(self, *a, **k))[1])
    monkeypatch.setattr(ProposalRunner, "fix", lambda self, *a, **k: (sequence.append("fix"), fix(self, *a, **k))[1])
    job = _run(client, pid, "CHAPTER_1")
    assert sequence == ["review", "fix", "review", "fix", "review"]
    assert job["outcome"] == "PARTIAL" and any("still had concerns" in w for w in job["warnings"])


def test_16_reviewer_notes_reach_the_student(client):
    pid = _chapter_ready(client)
    client.models.overrides["p_review"] = lambda p: {"results": [{"key": s["key"], "grade": "PASS_WITH_WARNINGS", "issues": [], "note": "Confirm the sub-counties."} for s in p["sections"]]}
    job = _run(client, pid, "CHAPTER_1")
    chapter = client.get(f"/api/projects/{pid}/chapters/1", headers=STUDENT).json()
    assert any("Confirm the sub-counties." in w for w in chapter["warnings"])
    assert any("Confirm the sub-counties." in w for w in job["warnings"]) and job["outcome"] == "FULL"


# --- 12: adding an objective marks the sections that cover every objective -----------------------


def test_12_adding_an_objective_marks_whole_set_sections_and_missing_reviews(client):
    before = plan()
    section = ChapterSection(key="objectives", number="1.4", heading="Objectives", paragraphs=["x"], depends=decisions.stamp(1, "objectives", before))
    theory = ChapterSection(key="theory", number="2.1", heading="Theory", paragraphs=["x"], depends=decisions.stamp(2, "theory", before))
    review = ChapterSection(key="empirical1", number="2.2", heading="Distance", paragraphs=["x"], depends=decisions.stamp(2, "empirical1", before))
    after = plan(
        specificObjectives=[*PLAN["specificObjectives"], "To examine affordability"],
        researchQuestions=[*PLAN["researchQuestions"], "How does cost affect uptake?"],
        alignment=[*PLAN["alignment"], {"objective": 4, "data": "Costs", "collection": "Questionnaire", "analysis": "Regression"}],
    )
    one = ChapterDocument(number=1, title="t", plan_version=1, sections=[section], cited=[], words=1)
    two = ChapterDocument(number=2, title="t", plan_version=1, sections=[theory, review], cited=[], words=1)
    assert decisions.stale(one, after) == ["1.4 Objectives"]
    assert decisions.stale(two, after) == []  # theory and objective 1's review are untouched


# --- 13: only the student's own figure is a population size --------------------------------------


def _step(notes: str, **figures) -> StepInput:
    return StepInput(project_id="prj_a", step="PLAN", chapter=0, rulebook="ucu-2018-v1", plan_version=0, inputs=ProposalInputs.model_validate({**DETAILS["inputs"], "notes": notes, **figures}))


def test_13_a_year_is_never_a_population_but_the_students_figure_is():
    generated = plan(sampleSize={**PLAN["sampleSize"], "population": 2026})
    assert pipeline._student_figures_only(generated, _step("The study will be undertaken in 2026.")).sample_size.population is None
    entered = _step("", populationSize=2026, populationSource="Parish register, 2025")
    kept = pipeline._student_figures_only(generated, entered).sample_size
    assert kept.population == 2026 and kept.population_source == "Parish register, 2025"


# --- 14: APA 6 first citations follow reading order ---------------------------------------------


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("⟦E00000b|n⟧ found low uptake. It remains low ⟦E00000b⟧.", "Okello, Namara, and Kato (2022) found low uptake. It remains low (Okello et al., 2022)."),
        ("It remains low ⟦E00000b⟧. ⟦E00000b|n⟧ found this.", "It remains low (Okello, Namara, & Kato, 2022). Okello et al. (2022) found this."),
    ],
)
def test_14_apa6_first_citation_is_the_first_in_the_text(text, expected):
    assert evidence.Citer(LIB, "APA6").render(text) == expected
