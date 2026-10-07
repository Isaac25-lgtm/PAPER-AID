"""An unanswered passage is not a LOW judgment, including in legacy percentage recovery."""

import io

import pytest
from docx import Document

from app.ai import costs, orchestration
from app.ai.orchestration import AIRunner
from app.core.config import Settings
from app.jobs.models import Engine
from app.pricing.quote import bound_quote, price
from tests.test_ai import ScriptedProvider, real_settings, runner_with
from tests.test_api import STUDENT, _submit


@pytest.fixture(autouse=True)
def earlier_workflow(monkeypatch):
    """These tests cover the multi-provider algorithm that engines priced before the Gemini workflow
    (owner decision 2026-10-07) keep running on, and that GEMINI_WORKFLOW=false brings back."""
    monkeypatch.setenv("GEMINI_WORKFLOW", "false")


def _passage(bid):
    return {"id": bid, "text": "A passage to assess.", "signals": []}


def _judgment(bid):
    return {"id": bid, "riskBand": "low", "reasons": [], "explanation": "", "suggestion": "", "excerpt": "", "confirmed": [], "rejected": [], "preserve": False, "risk": ""}


def test_omitted_and_unknown_ids_are_not_reviewed(monkeypatch):
    retry = {"blocks": [_judgment("unknown")]}
    runner, _ = runner_with(monkeypatch, ScriptedProvider({"analyse": [{"blocks": [_judgment("b1"), _judgment("unknown")]}, retry]}))
    answers, seen = runner.analyse([_passage("b1"), _passage("b2")], [], {})
    assert set(answers) == seen == {"b1"}


def test_an_empty_answer_does_not_imply_low_risk(monkeypatch):
    runner, _ = runner_with(monkeypatch, ScriptedProvider({"analyse": [{"blocks": []}, {"blocks": []}]}))
    answers, seen = runner.analyse([_passage("b1")], [], {})
    assert answers == {} and seen == set()


def test_an_answer_for_a_different_batch_does_not_count(monkeypatch):
    monkeypatch.setattr(orchestration, "BATCH_WORDS", 4)
    swapped = [{"blocks": [_judgment("b2")]}, {"blocks": [_judgment("b1")]}]
    runner, _ = runner_with(monkeypatch, ScriptedProvider({"analyse": swapped * 2}))  # and again in the retry
    answers, seen = runner.analyse([_passage("b1"), _passage("b2")], [], {})
    assert answers == {} and seen == set()


def test_a_frozen_old_engine_retains_its_model_and_prompt(monkeypatch):
    scripted = ScriptedProvider({"analyse": [{"blocks": []}]})
    monkeypatch.setattr(orchestration, "provider_for", lambda ref, settings: (scripted, ref.split(":")[-1]))
    old = Engine(lead_model="openai:gpt-6-sol", writer_model="anthropic:claude-opus-5-5", prompts={"analyse": "analyse-v2"})
    runner = AIRunner(real_settings(), lambda c: None, lambda: 0, 5, engine=old)
    assert runner.model_for("analyse") == "openai:gpt-6-sol"
    answers, seen = runner.analyse([_passage("b1")], [], {})
    assert answers == {} and seen == {"b1"}  # v2 explicitly permitted omitted LOW passages


def test_luna_uses_its_verified_price():
    assert costs.price_for("openai", "gpt-6-luna") == (0.10, 0.50, 0.01)
    assert costs.cost_usd("openai", "gpt-6-luna", 1_000_000, 100_000, 0) == pytest.approx(0.15)


def test_bound_quote_prices_the_frozen_checker():
    from app.jobs.models import ServiceSelection

    settings = Settings(_env_file=None, pricing_mode="fixed")
    selection = ServiceSelection(writing="AI_CHECK", academic=False)
    old = Engine(lead_model="openai:gpt-6-sol", writer_model="anthropic:claude-opus-5-5", prompts={"analyse": "analyse-v2"})
    quote = bound_quote(settings, selection, "source", 1000, old)
    assert quote.budget_usd == pytest.approx(price(settings, selection, 1000, engine=old).budget_usd)
    assert quote.budget_usd != pytest.approx(price(settings, selection, 1000).budget_usd)  # one checker, the older prompt
    assert quote.engine.ai_check_model is None


@pytest.mark.parametrize("returned", [0, 1])
def test_incomplete_analysis_has_no_overall_percentage_or_low_headline(client, returned):
    from app.runtime import get_runtime

    client.models.overrides["analyse"] = lambda p: {"blocks": [_judgment(b["id"]) for b in p["blocks"][:returned]]}
    job_id, job = _submit(client, {"writing": "AI_CHECK", "academic": False})
    assert job["status"] == "COMPLETED" and job["outcome"] == "PARTIAL"
    assert job["analysis"]["coverageComplete"] is False and "percent" not in job["analysis"]
    assert get_runtime().store.get(job_id).analysis.percent is None  # computed internally, never exposed to students
    document = client.get(f"/api/jobs/{job_id}/document", headers=STUDENT).json()
    assert "percent" not in document  # saved rule scores must not recreate or expose an overall percentage
    report = Document(io.BytesIO(client.get(f"/api/jobs/{job_id}/outputs/writing-report", headers=STUDENT).content))
    text = "\n".join(p.text for p in report.paragraphs)
    assert "Not every passage could be assessed" in text and "Low (confidence:" not in text


def test_a_paper_without_eligible_passages_is_not_zero_percent_low(client):
    from app.runtime import get_runtime
    from tests.test_api import get_quote, wait

    doc = Document()
    for i in range(6):
        doc.add_paragraph(f"Short paragraph {i} contains only ten words for this check.")
    source = io.BytesIO()
    doc.save(source)
    job_id = client.post("/api/jobs", headers=STUDENT).json()["id"]
    uploaded = client.post(f"/api/jobs/{job_id}/files/source", headers=STUDENT, files={"file": ("short.docx", source.getvalue(), "application/octet-stream")})
    assert uploaded.status_code == 200
    rt = get_runtime()  # uploaded before scorable words were counted: priced and run as before
    rt.store.update(job_id, lambda j: j.model_copy(update={"source": j.source.model_copy(update={"scorable_words": None})}))
    quote = get_quote(client, job_id, {"writing": "AI_CHECK", "academic": False})
    client.post(f"/api/jobs/{job_id}/submit", headers=STUDENT, json={"quoteId": quote["id"]})
    job = wait(client, job_id)
    assert job["outcome"] == "PARTIAL"
    assert job["analysis"]["coverageComplete"] is False and "percent" not in job["analysis"]
    assert get_runtime().store.get(job_id).analysis.percent is None
    assert any("no body passages" in w for w in job["warnings"])
