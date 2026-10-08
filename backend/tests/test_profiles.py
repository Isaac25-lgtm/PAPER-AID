"""Institution profiles from an uploaded research guide: read by the lead, critiqued by the writer,
finalised by the lead, then validated by code into a rulebook the whole proposal follows."""

import io

import pytest
from docx import Document

from app.proposals import profile, rulebook
from tests.fake_models import profile_answer
from tests.test_api import STUDENT
from tests.test_proposals import _approved, _create, _run


def _guide(text: str = "research proposal") -> bytes:
    doc = Document()
    doc.add_heading(f"Kyambogo University Graduate School: guidelines for the {text}", level=1)
    for i in range(40):
        doc.add_paragraph(f"Section {i}: the {text} shall follow the structure set out here, with chapters, headings and references as described in this guide.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _answer():
    return profile_answer({"guide": "research proposal", "reference": profile.reference()})


def test_a_profile_is_validated_into_a_rulebook_with_gaps_filled_from_the_default():
    answer = _answer()
    answer["chapters"][0]["sections"].append(dict(answer["chapters"][0]["sections"][0]))  # a repeated key
    answer["citation"] = "OTHER"
    book = profile.build(answer, "guide.docx")
    assert book["institution"] == "Kyambogo University" and book["id"].startswith("custom-")
    keys = [s["key"] for s in book["chapters"][0]["sections"]]
    assert len(keys) == len(set(keys)) and "definitions" in keys
    assert abs(sum(c["share"] for c in book["chapters"][:3]) - 1) < 0.01
    assert abs(sum(s["share"] for s in book["chapters"][1]["sections"]) - 1) < 0.01
    assert book["levels"]["MASTERS"]["pages"] == [20, 40] and book["levels"]["PHD"] == rulebook.load(rulebook.DEFAULT)["levels"]["PHD"]
    assert book["formatting"]["font"] == "Times New Roman" and book["objectives"]["min"] == 3
    assert book["chapters"][3]["kind"] == "CONCEPT"  # the concept paper keeps the default layout
    assert book["vetting"]["questions"]["3"][0]["question"] == "Is the sampling procedure justified?"
    assert book["vetting"]["questions"]["1"] == rulebook.load(rulebook.DEFAULT)["vetting"]["questions"]["1"]  # none in the guide
    assert any("citation style" in u for u in book["unclear"]) and book["default_citation"] == "APA7"


def test_a_maximum_alone_is_kept_and_what_code_fills_in_is_named_for_the_reviewer(client):
    """Live run 2026-10-07: a guide saying objectives "number no more than four" lost its four (min 0 was
    refused), and the final reviewer rejected the standard values code fills in three times over."""
    answer = _answer()
    answer["objectives"] = {"min": 0, "max": 4}
    answer["formatting"] = {"font": "Times New Roman", "sizePt": 12, "lineSpacing": 1.5, "marginsIn": 0}  # 2.5 cm and 3.5 cm: no single margin
    answer["levels"] = [{"level": "MASTERS", "pagesMin": 0, "pagesMax": 25}]  # "must not exceed 25 pages"
    book = profile.build(answer, "guide.docx")
    assert (book["objectives"]["min"], book["objectives"]["max"]) == (2, 4) and book["objectives"]["source"] != rulebook.load(rulebook.DEFAULT)["objectives"]["source"]
    assert book["levels"]["MASTERS"]["pages"] == [15, 25]
    standard = " | ".join(book["from_standard"])
    assert "levels BACHELORS, PGD, PHD" in standard and "Chapter 4" in standard and "preliminary_pages" not in standard
    assert "formatting margins_in: one margin" in standard and "formatting page_numbers" in standard and "size_pt" not in standard  # 12 pt was given
    assert any("set them in Word" in u for u in book["unclear"])  # the student is told about the margins
    assert "Trebuchet" not in book["words_per_page_source"]

    project = _create(client)
    client.post(f"/api/projects/{project['id']}/guide", headers=STUDENT, files={"file": ("KyU guide.docx", _guide(), "application/octet-stream")})
    assert _run(client, project["id"], "PROFILE")["status"] == "COMPLETED"
    review = next(r for t, r in zip(client.models.tasks, client.models.requests, strict=True) if t == "p_profile_review")
    assert '"fromStandard"' in review and "concept note" in review and "preliminary_pages" not in review  # never used, never judged


def test_a_text_that_is_not_a_guide_is_refused():
    with pytest.raises(profile.NotAGuide):
        profile.build({**_answer(), "chapters": []}, "notes.docx")


def test_the_proposal_follows_the_students_institution(client):
    project = _create(client)
    uploaded = client.post(f"/api/projects/{project['id']}/guide", headers=STUDENT, files={"file": ("KyU guide.docx", _guide(), "application/octet-stream")})
    assert uploaded.status_code == 200 and uploaded.json()["guideName"] == "KyU guide.docx"
    quoted = client.post(f"/api/projects/{project['id']}/steps", headers=STUDENT, json={"step": "PROFILE"}).json()
    assert quoted["quote"]["lines"][0]["label"].startswith("Your institution's guide")
    job = _run(client, project["id"], "PROFILE")
    assert job["status"] == "COMPLETED", job
    # One final reviewer; Opus's guidance is optional and off (owner decision 2026-09-30)
    profile_calls = [t for t in client.models.tasks if t.startswith("p_profile")]  # another step's calls may interleave under load
    assert profile_calls[-4:] == ["p_profile", "p_profile_critique", "p_profile_finalise", "p_profile_review"]
    project = client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()
    assert project["rulebook"].startswith("custom-") and project["institution"] == "Kyambogo University" and project["citation"] == "APA7"
    assert "The guide does not say how long the literature review should be." in project["institutionNotes"] and project["guideRead"]

    _run(client, project["id"], "PLAN")
    _approved(client, project["id"])  # Chapter One follows the new structure
    chapter = client.get(f"/api/projects/{project['id']}/chapters/1", headers=STUDENT).json()
    assert chapter["title"] == "Introduction" and any(s["heading"] == "Definition of Key Terms" for s in chapter["sections"])
    doc = Document(io.BytesIO(client.get(f"/api/projects/{project['id']}/export", headers=STUDENT).content))
    assert any("OF KYAMBOGO UNIVERSITY" in p.text for p in doc.paragraphs) and doc.styles["Normal"].font.name == "Times New Roman"

    from app.runtime import get_runtime

    stored = rulebook.stored_path(project["rulebook"])
    assert get_runtime().files.exists(stored)
    assert client.delete(f"/api/projects/{project['id']}", headers=STUDENT).status_code == 204
    assert not get_runtime().files.exists(stored)  # the profile goes with the project


def test_a_guide_without_a_proposal_structure_fails_without_charge(client):
    project = _create(client)
    client.post(f"/api/projects/{project['id']}/guide", headers=STUDENT, files={"file": ("notes.docx", _guide("lecture notes"), "application/octet-stream")})
    job = _run(client, project["id"], "PROFILE")
    assert job["status"] == "FAILED" and job["failure"]["code"] == "NOT_A_GUIDE" and job["paymentStatus"] != "PAID"
    assert client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()["rulebook"] == rulebook.DEFAULT


def test_a_profile_needs_a_guide_first(client):
    project = _create(client)
    refused = client.post(f"/api/projects/{project['id']}/steps", headers=STUDENT, json={"step": "PROFILE"})
    assert refused.status_code == 400 and refused.json()["code"] == "NO_GUIDE"


def test_a_guide_added_after_chapter_one_restructures_it_and_keeps_the_earlier_version(client):
    """Owner decision 2026-10-08: the institution's guide is added after Chapter One to align the work. Code
    restructures the written chapter to the guide (order, numbering, headings; the approved statements under
    their new numbers), keeps the earlier version, leaves the guide's new sections to write and lists those
    whose requirement differs; writing and revising them follow as the student's own steps."""
    from tests.fake_models import PLAN

    project = _create(client)
    _run(client, project["id"], "PLAN")
    _approved(client, project["id"])
    url = f"/api/projects/{project['id']}"
    before = client.get(f"{url}/chapters/1", headers=STUDENT).json()
    uploaded = client.post(f"{url}/guide", headers=STUDENT, files={"file": ("KyU guide.docx", _guide(), "application/octet-stream")})
    assert uploaded.status_code == 200
    assert _run(client, project["id"], "PROFILE")["status"] == "COMPLETED"
    project = client.get(url, headers=STUDENT).json()
    assert project["rulebook"].startswith("custom-") and project["citation"] == "APA6"  # written chapters keep their style
    state = next(c for c in project["chapters"] if c["number"] == 1)
    assert len(state["versions"]) == 2 and state["current"] == 2 and state["versions"][-1]["note"] == "Restructured to your institution's guide"
    chapter = client.get(f"{url}/chapters/1", headers=STUDENT).json()
    planned = rulebook.sections(project["rulebook"], 1, "MASTERS", __import__("app.proposals.models", fromlist=["ProposalPlan"]).ProposalPlan.model_validate(project["plan"]))
    order = [s.key for s in planned]
    keys = [s["key"] for s in chapter["sections"]]
    assert keys == [k for k in order if k in keys]  # the guide's order
    assert "1.4 Definition of Key Terms" in chapter["missing"] or any("Definition of Key Terms" in m for m in chapter["missing"])
    objectives = next(s for s in chapter["sections"] if s["key"] == "objectives")
    assert f"{objectives['number']}.1 General Objective" in objectives["paragraphs"] and PLAN["purpose"] in objectives["paragraphs"]
    earlier = client.get(f"{url}/chapters/1?version=1", headers=STUDENT).json()
    assert [s["key"] for s in earlier["sections"]] == [s["key"] for s in before["sections"]]  # the earlier version is kept
    aligned = next(i for i in chapter["readiness"] if i["id"] == "C1-ALIGNED")
    assert aligned["status"] == "NEEDS_REVIEW"

    # the guide's new sections, written by finishing the chapter
    assert _run(client, project["id"], "COMPLETE_1")["status"] == "COMPLETED"
    finished = client.get(f"{url}/chapters/1", headers=STUDENT).json()
    assert not finished["missing"] and any(s["heading"] == "Definition of Key Terms" for s in finished["sections"])
    assert finished["toAlign"] == chapter["toAlign"]  # still to revise to the guide



# --- Codex audit of 29343c2 (2026-10-08) -------------------------------------------------------------------------


def _aligned(client):
    project = _create(client)
    _run(client, project["id"], "PLAN")
    _approved(client, project["id"])
    url = f"/api/projects/{project['id']}"
    client.post(f"{url}/guide", headers=STUDENT, files={"file": ("KyU guide.docx", _guide(), "application/octet-stream")})
    assert _run(client, project["id"], "PROFILE")["status"] == "COMPLETED"
    return project["id"], url


def test_a_chapter_still_to_align_with_the_guide_is_neither_approved_nor_exported_complete(client):
    """#1: an ethics section written to the old requirement could be approved and exported as complete."""
    project_id, url = _aligned(client)
    _run(client, project_id, "COMPLETE_1")  # the guide's new sections
    chapter = client.get(f"{url}/chapters/1", headers=STUDENT).json()
    assert chapter["toAlign"] and not chapter["missing"]
    state = next(c for c in client.get(url, headers=STUDENT).json()["chapters"] if c["number"] == 1)
    refused = client.post(f"{url}/chapters/1", headers=STUDENT, json={"version": state["current"], "approved": True})
    assert refused.status_code == 400 and refused.json()["code"] == "CHAPTER_NOT_ALIGNED"
    assert any("revise to your institution's guide" in b for b in client.get(url, headers=STUDENT).json()["blockers"])

    # revised to the guide on the student's request: what changed and passed review is aligned
    project = client.get(url, headers=STUDENT).json()
    before = {c["id"] for c in project["feedback"]}
    asked = client.post(f"{url}/chapters/1/request", headers=STUDENT, json={"instruction": "Revise this section to my institution's guide.",
                                                                            "sections": chapter["toAlign"]})
    assert asked.status_code == 200, asked.json()
    mine = next(c["id"] for c in asked.json()["feedback"] if c["id"] not in before)
    quoted = client.post(f"{url}/steps", headers=STUDENT, json={"step": "REVISE_1", "comments": [mine]}).json()
    client.post(f"{url}/steps/{quoted['job']['id']}/submit", headers=STUDENT, json={"quoteId": quoted["quote"]["id"]})
    from tests.test_api import wait

    assert wait(client, quoted["job"]["id"], STUDENT)["status"] == "COMPLETED"
    revised = client.get(f"{url}/chapters/1", headers=STUDENT).json()
    assert len(revised["toAlign"]) < len(chapter["toAlign"])


def _with_chapter_one(answer, sections):
    answer["chapters"][0]["sections"] = sections
    return answer


def test_guide_sections_are_matched_by_their_role_and_the_statements_placed_by_code():
    """#2: "Study Aims" under another key lost the objectives and was left for a model to write."""
    base = [s for s in _answer()["chapters"][0]["sections"] if s["key"] not in ("objectives", "questions")]
    answer = _with_chapter_one(_answer(), [*base,
        {"key": "aim", "heading": "Aim of the Study", "brief": "The aim.", "share": 0.05, "perObjective": False, "table": False},
        {"key": "studyaims", "heading": "Study Aims", "brief": "The aims.", "share": 0.05, "perObjective": False, "table": False},
        {"key": "rq", "heading": "Research Hypotheses", "brief": "The hypotheses.", "share": 0.05, "perObjective": False, "table": False}])
    book = profile.build(answer, "guide.docx")
    placed = {s["key"]: s.get("from_plan") for s in book["chapters"][0]["sections"] if s.get("from_plan")}
    assert placed == {"purpose": "purpose", "objectives": "specific_objectives", "questions": "questions"}

    twice = _with_chapter_one(_answer(), [*base,
        {"key": "objectives", "heading": "Objectives of the Study", "brief": "x", "share": 0.05, "perObjective": False, "table": False},
        {"key": "specific", "heading": "Specific Objectives", "brief": "y", "share": 0.05, "perObjective": False, "table": False}])
    with pytest.raises(profile.AmbiguousGuide):
        profile.build(twice, "guide.docx")


def test_a_purpose_section_the_earlier_structure_lacked_is_placed_by_code_when_aligning(client, monkeypatch):
    from app.proposals import pipeline
    from app.proposals.models import ChapterDocument, ProposalPlan

    project_id, url = _aligned(client)
    from app.runtime import get_runtime

    stored = get_runtime().store.get_project(project_id)
    plan = stored.plan
    assert isinstance(plan, ProposalPlan)
    real = rulebook.sections

    def with_purpose(rulebook_id, chapter, level, p):  # the guide's Chapter One gives the general objective a section of its own
        out = real(rulebook_id, chapter, level, p)
        if rulebook_id == stored.rulebook and chapter == 1:
            purpose = rulebook.SectionPlan("purpose", "1.2", "Aim of the Study", "The aim.", 100, from_plan="purpose")
            out = [*out[:2], purpose, *out[2:]]
        return out

    monkeypatch.setattr(rulebook, "sections", with_purpose)
    doc = ChapterDocument(number=1, title="Introduction", plan_version=1, sections=[], cited=[], words=0)
    aligned = pipeline.align_document(doc, "ucu-2018-v2", stored.rulebook, "MASTERS", plan)
    purpose = next(s for s in aligned.sections if s.key == "purpose")
    assert purpose.paragraphs == [plan.purpose] and "purpose" not in aligned.missing


def test_a_guide_silent_on_objective_counts_keeps_the_owners_cap_and_the_handbook_safeguards():
    """#4: a custom profile dropped the cap, the primary question and the hypothesis pairs."""
    silent = {**_answer(), "objectives": {"min": 0, "max": 0}}
    book = profile.build(silent, "guide.docx")
    assert book["objectives_by_level"]["MASTERS"]["max"] == 3 and book["primary_question"] and book["hypothesis_pairs"]
    assert book["concept_objectives"] == {"min": 3, "max": 5, "source": "§1.4, p. 4"}
    stated = profile.build(_answer(), "guide.docx")  # the guide sets its own number: its rule replaces the cap
    assert "objectives_by_level" not in stated and stated["primary_question"]


def test_a_finish_follows_the_current_plan_after_a_confirmed_setting(client):
    """#3: correcting a proposed setting after the guide made "write the sections your guide adds" refuse with PLAN_CHANGED."""
    project_id, url = _aligned(client)
    project = client.get(url, headers=STUDENT).json()
    confirmed = client.post(f"{url}/setting", headers=STUDENT, json={"studyArea": "Wakiso District", "population": "Caregivers",
                                                                    "baseVersion": project["planVersion"]})
    assert confirmed.status_code == 200
    assert _run(client, project_id, "COMPLETE_1")["status"] == "COMPLETED"
    assert not client.get(f"{url}/chapters/1", headers=STUDENT).json()["missing"]


def test_a_plan_changed_while_the_guide_is_read_aligns_nothing(client, monkeypatch):
    """#5: the chapters would have been restructured with the earlier plan."""
    from app.proposals import pipeline
    from app.runtime import get_runtime

    project = _create(client)
    _run(client, project["id"], "PLAN")
    _approved(client, project["id"])
    url = f"/api/projects/{project['id']}"
    client.post(f"{url}/guide", headers=STUDENT, files={"file": ("KyU guide.docx", _guide(), "application/octet-stream")})
    store, real = get_runtime().store, pipeline.align_document

    def edited_meanwhile(*args):  # the student saves a plan edit while the chapters are being restructured
        def bump(p):
            p.plan_version += 1
            return p

        store.update_project(project["id"], bump)
        return real(*args)

    monkeypatch.setattr(pipeline, "align_document", edited_meanwhile)
    job = _run(client, project["id"], "PROFILE")
    assert job["status"] == "FAILED" and job["failure"]["code"] == "INPUTS_CHANGED"
    state = next(c for c in client.get(url, headers=STUDENT).json()["chapters"] if c["number"] == 1)
    assert len(state["versions"]) == 1  # nothing aligned
