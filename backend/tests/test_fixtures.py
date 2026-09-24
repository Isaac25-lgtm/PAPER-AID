"""Guards the synthetic fixture set itself: if the generator drifts, engine tests would test nothing."""

import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

GENERATED = Path(__file__).parent / "fixtures" / "generated"


@pytest.fixture(scope="session", autouse=True)
def generated() -> None:
    subprocess.run([sys.executable, str(GENERATED.parent / "generate.py")], check=True, capture_output=True)


def manifest() -> list[dict]:
    return json.loads((GENERATED / "manifest.json").read_text())


def test_every_accepted_docx_opens():
    for entry in manifest():
        if entry["expect"] == "accept" and entry["file"].endswith(".docx"):
            assert Document(GENERATED / entry["file"]).paragraphs, entry["file"]


def test_every_rejected_fixture_names_a_reason():
    assert all(e["rejectReason"] for e in manifest() if e["expect"] == "reject")


def test_citation_fields_are_real_fields():
    xml = zipfile.ZipFile(GENERATED / "citation_fields.docx").read("word/document.xml").decode()
    assert xml.count('w:fldCharType="begin"') == 4
    assert "ZOTERO_ITEM" in xml and "CSL_CITATION" in xml


def test_footnotes_part_and_reference_exist():
    z = zipfile.ZipFile(GENERATED / "footnotes.docx")
    assert "word/footnotes.xml" in z.namelist()
    assert "footnoteReference" in z.read("word/document.xml").decode()


def test_landscape_section_sits_between_portrait_sections():
    orientations = [s.orientation for s in Document(GENERATED / "landscape_section.docx").sections]
    assert [int(o) for o in orientations] == [0, 1, 0]


def test_dissertation_is_long_enough_to_need_chunking():
    words = sum(len(p.text.split()) for p in Document(GENERATED / "dissertation_long.docx").paragraphs)
    assert words > 13_000


def test_hostile_archives_have_their_hostile_property():
    bomb = zipfile.ZipFile(GENERATED / "zip_bomb.docx").infolist()
    assert sum(i.file_size for i in bomb) / sum(i.compress_size for i in bomb) > 100
    assert len(zipfile.ZipFile(GENERATED / "many_entries.docx").namelist()) > 5000
    assert any(".." in n for n in zipfile.ZipFile(GENERATED / "path_traversal.docx").namelist())
    assert "word/vbaProject.bin" in zipfile.ZipFile(GENERATED / "macro_renamed.docx").namelist()
    assert (GENERATED / "encrypted_office.docx").read_bytes()[:8] == bytes.fromhex("D0CF11E0A1B11AE1")
    assert not zipfile.is_zipfile(GENERATED / "renamed_exe.docx")


def test_pdfs_behave_as_labelled():
    text = PdfReader(GENERATED / "text_based.pdf")
    assert len(text.pages) == 2 and len(text.pages[0].extract_text()) > 500
    assert PdfReader(GENERATED / "scanned.pdf").pages[0].extract_text().strip() == ""
    assert PdfReader(GENERATED / "encrypted.pdf").is_encrypted
