"""One accountable final reviewer (owner decision 2026-09-30, superseding the dual veto), and the live
cases that led to it: proposal plans (prj_c704b0c370f7), works documents, plans and Results Models,
Paper Check rewrites, the coursework questions screen (wrk_64b3b916e144), page limits, delivery and
number formats. Older engines keep both reviewers (tests/test_four_models.py)."""

import io
import json

import pytest
from docx import Document

from app.ai import orchestration
from app.rules.compliance import overall, report
from app.rules.extract import value_in_quote
from app.rules.validators import VALIDATORS, Context
from app.works import pipeline as works_pipeline
from app.works.models import WorkDocument, WorkInputs, WorkSection
from tests import fake_works
from tests.fake_models import PLAN
from tests.test_api import STUDENT, wait
from tests.test_works import QUESTION, _coursework, _run, _work
from tests.test_works_golden import _spec


@pytest.fixture(autouse=True)
def earlier_workflow(monkeypatch):
    """These tests cover the multi-provider algorithm that engines priced before the Gemini workflow
    (owner decision 2026-10-07) keep running on, and that GEMINI_WORKFLOW=false brings back."""
    monkeypatch.setenv("GEMINI_WORKFLOW", "false")

# What the live final reviewer said about the owner's plan (job_76db5523ed82, 2026-09-30).
LIVE_OBJECTIONS = [
    "Eligibility requires use of only one specified feature, but the analyses include separate quality measures for recommendations and conversational tools. "
    "Specify that respondents will rate only features they have used and how feature-specific analyses will handle those who have not used each category.",
    "Recent feature-use frequency is listed as an intervening variable even though the plan explicitly treats it as a descriptive measure or covariate, not a mediator. "
    "Reclassify it to keep the variable definitions aligned with the analyses.",
]


@pytest.fixture
def works_client(tmp_path, monkeypatch):
    monkeypatch.setenv("WORKS_ENABLED", '["CONCEPT_NOTE","COURSEWORK","FUNDING_PROPOSAL"]')
    from tests.conftest import _client

    yield from _client(tmp_path, monkeypatch, "cost")


def _project(client):
    from tests.test_proposals import _create

    return _create(client)


def _plan_step(client, project_id):
    from tests.test_proposals import _run as run_step

    return run_step(client, project_id, "PLAN")


def _requests(client, task):
    return [json.loads(r[len(task):]) for t, r in zip(client.models.tasks, client.models.requests, strict=True) if t == task]


def _repairs_change_the_plan(client):
    """A targeted repair really changes what it was asked to fix (identical text would replay the
    earlier verdict from the response cache, correctly)."""
    rounds = []

    def finalise(payload):
        plan = dict(payload["draft"])
        if any("issue" in i for i in payload["critique"].get("items", [])):
            rounds.append(1)
            plan["inclusion"] = (plan.get("inclusion") or "") + f" Respondents rate only the features they have used (repair {len(rounds)})."
        return plan

    client.models.overrides["p_finalise"] = finalise


def _work_repairs_change_the_plan(client):
    rounds = []

    def plan(payload):
        answer = fake_works.plan(payload)
        if payload.get("critique"):
            rounds.append(1)
            answer["sections"][-1]["brief"] = f"Answer the question directly (repair {len(rounds)})."
        return answer

    client.models.overrides["w_plan"] = plan


# --- routing ---------------------------------------------------------------------------------------


def test_sol_is_the_one_final_reviewer_and_never_approves_what_it_finalised():
    from app.core.config import Settings

    engine = orchestration.current_engine(Settings(_env_file=None))
    assert engine.single_reviewer and not engine.frontier_guidance  # Opus's guidance is optional, off by default
    assert orchestration.model_for_engine(engine, "p_plan_review") == "openai:gpt-6-sol"
    assert orchestration.model_for_engine(engine, "p_finalise") == "anthropic:claude-sonnet-5-5"  # moved from Sol
    assert orchestration.model_for_engine(engine, "p_review") == "openai:gpt-6-sol"
    work = orchestration.work_engine(Settings(_env_file=None, openai_api_key="k", anthropic_api_key="k", gemini_api_key="k"), "FUNDING_PROPOSAL")
    for task in ("w_final", "w_plan_review", "w_results_review"):
        assert orchestration.model_for_engine(work, task) == "openai:gpt-6-sol"
    assert orchestration.model_for_engine(work, "w_evaluate") == "anthropic:claude-sonnet-5-5"  # Opus no longer reviews sections
    older = engine.model_copy(update={"single_reviewer": False})  # a job priced before the decision keeps its routing
    assert orchestration.model_for_engine(older, "p_finalise") == "openai:gpt-6-sol"


# --- proposal plans (live case prj_c704b0c370f7) ------------------------------------------------------


def test_the_live_plan_objections_are_repaired_and_the_repaired_plan_reviewed_again(client):
    reviews = []

    def review(payload):
        reviews.append(payload["finalPlan"])
        return {"approved": False, "issues": LIVE_OBJECTIONS} if len(reviews) == 1 else {"approved": True, "issues": []}

    client.models.overrides["p_plan_review"] = review
    _repairs_change_the_plan(client)
    project = _project(client)
    job = _plan_step(client, project["id"])
    assert job["status"] == "COMPLETED" and job["outcome"] == "FULL", job
    assert len(reviews) == 2 and "p_plan_review_peer" not in client.models.tasks  # one reviewer, no second veto
    repair = _requests(client, "p_finalise")[-1]
    assert [i["issue"] for i in repair["critique"]["items"]] == LIVE_OBJECTIONS  # repaired for exactly what was raised
    view = client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()
    assert view["planReview"]["outcome"] == "APPROVED" and view["planStatus"] == "DRAFT"


def test_a_typed_citation_is_converted_only_when_it_matches_confirmed_evidence(client):
    """The live plan cited "(Uganda Communications Commission, 2025)" beside an evidence id. Code turns
    a typed citation of a confirmed source into its id; any other goes to the repair by name."""
    finals = []

    def finalise(payload):
        finals.append(1)
        plan = dict(payload["draft"])
        if len(finals) == 1:
            plan["problem"] = plan["problem"] + " Caregivers cite distance (Okello et al., 2022) and cost (Uganda Communications Commission, 2025)."
        else:  # the targeted repair removes what it was told to
            plan["problem"] = plan["problem"].replace(" and cost (Uganda Communications Commission, 2025)", "")
        return plan

    client.models.overrides["p_finalise"] = finalise
    project = _project(client)
    job = _plan_step(client, project["id"])
    assert job["status"] == "COMPLETED", job
    critique = _requests(client, "p_finalise")[-1]["critique"]["items"]
    assert any("Uganda Communications Commission, 2025" in i["issue"] for i in critique)  # unmatched: named for repair, never invented
    problem = client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()["plan"]["problem"]
    assert "Okello et al., 2022" not in problem and "Uganda Communications Commission" not in problem
    assert "(E" in problem  # the matched citation became its evidence id


def test_objections_left_after_two_repairs_keep_an_unapproved_plan_that_costs_nothing(fixed_client):
    from app.runtime import get_runtime

    client = fixed_client
    client.models.overrides["p_plan_review"] = lambda payload: {"approved": False, "issues": LIVE_OBJECTIONS[:1]}
    _repairs_change_the_plan(client)
    project = _project(client)
    job = _plan_step(client, project["id"])
    assert job["status"] == "COMPLETED" and job["outcome"] == "PARTIAL"
    assert client.models.tasks.count("p_plan_review") == 3  # the review, then two repairs each reviewed again
    stored = get_runtime().store.get(job["id"])
    assert stored.billing.charged == 0 and stored.billing.state == "SETTLED"  # an unusable plan is not charged
    view = client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()
    assert view["planReview"]["outcome"] == "OBJECTIONS" and view["planReview"]["objections"] == LIVE_OBJECTIONS[:1]
    assert view["autoChapterOne"] is False  # never starts Chapter One by itself
    refused = client.post(f"/api/projects/{project['id']}/plan/approve", headers=STUDENT, json={"baseVersion": view["planVersion"]})
    assert refused.status_code == 400 and refused.json()["code"] == "ACKNOWLEDGMENT_NEEDED"
    approved = client.post(f"/api/projects/{project['id']}/plan/approve", headers=STUDENT, json={"baseVersion": view["planVersion"], "acknowledge": ["OBJECTIONS"]})
    assert approved.status_code == 200 and approved.json()["planStatus"] == "APPROVED"
    ack = get_runtime().store.get_project(project["id"]).acknowledgments[-1]
    assert ack.kind == "OBJECTIONS" and ack.plan_version == view["planVersion"] and len(ack.text_sha256) == 64
    assert approved.json()["chapters"][0]["versions"] == [] and not approved.json()["activeJob"]


def test_a_plan_review_that_cannot_complete_is_not_approval(fixed_client):
    from app.runtime import get_runtime

    client = fixed_client
    client.models.truncate.add("p_plan_review")  # cut off by the output limit, every time
    project = _project(client)
    job = _plan_step(client, project["id"])
    assert job["status"] == "COMPLETED" and job["outcome"] == "PARTIAL"
    assert get_runtime().store.get(job["id"]).billing.charged == 0
    review = client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()["planReview"]
    assert review["outcome"] == "NOT_REVIEWED" and review["reason"] == "REVIEW_UNAVAILABLE"


def test_standard_sampling_settings_need_an_explicit_versioned_acknowledgment(client):
    from app.runtime import get_runtime

    sampled = {**PLAN, "sampleSize": {**PLAN.get("sampleSize", {}), "method": "COCHRAN", "margin": 0, "proportion": 0}}
    client.models.overrides["p_finalise"] = lambda payload: sampled if payload.get("kind") == "plan" else payload["draft"]  # a chapter's briefs are finalised too
    project = _project(client)
    assert _plan_step(client, project["id"])["status"] == "COMPLETED"
    view = client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()
    assert view["plan"]["samplingAssumed"]
    refused = client.post(f"/api/projects/{project['id']}/plan/approve", headers=STUDENT, json={"baseVersion": view["planVersion"]})
    assert refused.status_code == 400 and "sample size settings" in refused.json()["message"]
    ok = client.post(f"/api/projects/{project['id']}/plan/approve", headers=STUDENT, json={"baseVersion": view["planVersion"], "acknowledge": ["SAMPLING"]})
    assert ok.status_code == 200
    if ok.json().get("activeJob"):  # Chapter One follows an approved plan: let it finish before the stand-ins go
        wait(client, ok.json()["activeJob"], timeout=120)
    ack = get_runtime().store.get_project(project["id"]).acknowledgments[-1]
    assert ack.kind == "SAMPLING" and ack.plan_version == view["planVersion"]


# --- Paper Check ---------------------------------------------------------------------------------------


def test_paper_check_rewrites_have_one_final_reviewer(client):
    from tests.test_api import _submit

    client.models.overrides["review"] = lambda payload: {"results": []}  # Sol gives no approval
    _, job = _submit(client, {"writing": "REFINE", "academic": False})
    assert "review_peer" not in client.models.tasks
    assert job["status"] == "COMPLETED" and job["outcome"] == "PARTIAL" and job["refinement"]["refinedBlocks"] == 0  # the student's wording is kept


# --- works: the final review of the exact deliverable ---------------------------------------------------


def _drafted(client, ai="BANNED"):
    work = _coursework(client, ai_answer=ai)
    _run(client, work["id"], "PLAN")
    work = _work(client, work["id"])
    client.post(f"/api/works/{work['id']}/plan/approve", headers=STUDENT, json={"baseVersion": work["planVersion"]})
    return work


def _with_table(payload):
    drafted = fake_works.draft(payload)
    drafted["sections"][1]["table"] = {"caption": "Coverage by district", "rows": [["District", "Visits"], ["Kamuli", "Weekly"]]}
    return drafted


def test_the_final_reviewer_sees_exactly_what_will_be_delivered(works_client):
    client = works_client
    client.models.overrides["w_draft"] = _with_table
    work = _drafted(client)
    _, job = _run(client, work["id"], "DRAFT")
    assert job["status"] == "COMPLETED", job.get("failure")
    final = _requests(client, "w_final")[-1]
    document = final["document"]
    assert any(t["caption"] == "Coverage by district" for s in document["sections"] for t in s["tables"])
    assert document["references"] and document["notes"] == ["This document was drafted by an AI-assisted third party."]
    assert final["part"] == "1 of 1" and {m["key"] for m in final["manifest"]} == {s["key"] for s in document["sections"]}
    assert not any("⟦" in p for s in document["sections"] for p in s["text"])  # citations and figures as they print


def test_a_missing_final_verdict_is_not_reviewed_and_nothing_is_delivered_or_charged(works_client):
    from app.runtime import get_runtime

    client = works_client

    def final(payload):
        answer = fake_works.answer("w_final", payload)
        answer["rules"] = answer["rules"][1:]  # one rule left without a verdict
        return answer

    client.models.overrides["w_final"] = final
    work = _drafted(client)
    _, job = _run(client, work["id"], "DRAFT")
    assert job["status"] == "FAILED" and job["failure"]["code"] == "REVIEW_UNAVAILABLE"
    assert get_runtime().store.get(job["id"]).billing.state == "RELEASED" and not _work(client, work["id"])["documents"]


def test_a_long_deliverable_is_reviewed_in_parts_with_a_manifest_of_the_whole(works_client, monkeypatch):
    monkeypatch.setattr(works_pipeline, "FINAL_PART_WORDS", 900)  # whole requests: about 460 words of every one are repeated context (Codex review 2026-10-07, finding 8)
    client = works_client
    work = _drafted(client)
    _, job = _run(client, work["id"], "DRAFT")
    assert job["status"] == "COMPLETED", job.get("failure")
    parts = [r for r in _requests(client, "w_final")]
    count = int(parts[0]["part"].split(" of ")[1])
    assert count > 1 and [p["part"] for p in parts[:count]] == [f"{n} of {count}" for n in range(1, count + 1)]
    assert all(len(p["manifest"]) == len(parts[0]["manifest"]) for p in parts)
    assert parts[count - 1]["document"]["references"] and not parts[0]["document"]["references"]  # the whole, once


# --- works: plans and Results Models are reviewed again after repair ---------------------------------------


def test_a_repaired_works_plan_is_reviewed_again(works_client):
    client = works_client
    reviewed = []

    def review(payload):
        reviewed.append(payload["plan"])
        verdict = "REPAIR" if len(reviewed) == 1 else "PASS"
        return {"verdict": verdict, "rules": [{"rule": r["rule"], "status": "PASS", "note": "Met."} for r in payload.get("rules", [])],
                "issues": ["The conclusion must answer the question directly."] if verdict == "REPAIR" else []}

    client.models.overrides["w_plan_review"] = review
    _work_repairs_change_the_plan(client)
    work = _coursework(client)
    _, job = _run(client, work["id"], "PLAN")
    assert job["status"] == "COMPLETED" and job["outcome"] == "FULL"
    assert len(reviewed) == 2 and _work(client, work["id"])["planReview"]["outcome"] == "APPROVED"


def test_a_works_plan_still_objected_to_is_kept_unapproved_and_costs_nothing(works_client):
    from app.runtime import get_runtime

    client = works_client
    client.models.overrides["w_plan_review"] = lambda payload: {"verdict": "REPAIR", "rules": [], "issues": ["Theme 2 repeats Theme 1."]}
    _work_repairs_change_the_plan(client)
    work = _coursework(client)
    _, job = _run(client, work["id"], "PLAN")
    assert job["status"] == "COMPLETED" and job["outcome"] == "PARTIAL" and client.models.tasks.count("w_plan_review") == 3
    assert get_runtime().store.get(job["id"]).delivery == {"CW_PLAN": 0.0}
    work = _work(client, work["id"])
    assert work["planReview"]["objections"] == ["Theme 2 repeats Theme 1."]
    refused = client.post(f"/api/works/{work['id']}/plan/approve", headers=STUDENT, json={"baseVersion": work["planVersion"]})
    assert refused.status_code == 400 and refused.json()["code"] == "ACKNOWLEDGMENT_NEEDED"
    ok = client.post(f"/api/works/{work['id']}/plan/approve", headers=STUDENT, json={"baseVersion": work["planVersion"], "acknowledge": ["PLAN_OBJECTIONS"]})
    assert ok.status_code == 200 and ok.json()["planStatus"] == "APPROVED"


def test_a_repaired_results_model_is_reviewed_again(works_client):
    client = works_client
    reviews = []

    def review(payload):
        reviews.append(1)
        answer = fake_works.answer("w_results_review", payload)
        if len(reviews) == 1:
            answer["issues"] = ["Outcome 1 is an activity, not a change."]
        return answer

    client.models.overrides["w_results_review"] = review

    def results(payload):
        answer = fake_works.results(payload)
        if payload.get("critique"):
            answer["outcomes"][0]["statement"] = "Mothers in the four sub-counties deliver with a skilled attendant."
        return answer

    client.models.overrides["w_results"] = results
    created = client.post("/api/works", headers=STUDENT, json={
        "kind": "FUNDING_PROPOSAL", "variant": "NGO_PROJECT", "mode": "COMPACT",
        "inputs": {"title": "Safer deliveries in Kamuli", "description": "Too many mothers in Kamuli deliver at home and die of preventable complications.",
                   "answers": {"problem": "Home deliveries and delayed referral.", "intervention": "Train village health teams.", "duration_months": "2 years", "currency": "USD"}},
    }).json()
    work = client.post(f"/api/works/{created['id']}/answers", headers=STUDENT, json={"answers": {}, "skipRest": True, "baseVersion": created["specVersion"]}).json()
    assert work["inputs"]["answers"]["duration_months"] == "24"  # "2 years" saved as months
    client.post(f"/api/works/{work['id']}/spec/confirm", headers=STUDENT, json={"baseVersion": work["specVersion"]})
    _, job = _run(client, work["id"], "PLAN")
    assert job["status"] == "COMPLETED" and len(reviews) == 2
    assert _work(client, work["id"])["resultsReview"]["outcome"] == "APPROVED"


# --- the coursework questions screen (live case wrk_64b3b916e144) ---------------------------------------


def test_a_word_limit_typed_with_words_is_saved_and_applied(works_client):
    client = works_client
    work = client.post("/api/works", headers=STUDENT, json={"kind": "COURSEWORK", "variant": "ESSAY", "inputs": {"title": "CHWs", "description": QUESTION}}).json()
    saved = client.post(f"/api/works/{work['id']}/answers", headers=STUDENT, json={"answers": {"word_limit": "3,000 words"}, "baseVersion": work["specVersion"]}).json()
    assert saved["inputs"]["answers"]["word_limit"] == "3000"
    assert next(q for q in saved["spec"]["questions"] if q["id"] == "word_limit")["answered"]
    assert saved["spec"]["limits"][0]["max"] == 3000 and saved["spec"]["targetWords"] == 2910
    bad = client.post(f"/api/works/{work['id']}/answers", headers=STUDENT, json={"answers": {"word_limit": "a lot"}, "baseVersion": saved["specVersion"]})
    assert bad.status_code == 400 and bad.json()["code"] == "INVALID_ANSWER" and "for example 3,000" in bad.json()["message"]


def test_no_word_limit_given_is_an_answer_and_never_an_invented_limit(works_client):
    client = works_client
    work = client.post("/api/works", headers=STUDENT, json={"kind": "COURSEWORK", "variant": "ESSAY", "inputs": {"title": "CHWs", "description": QUESTION}}).json()
    saved = client.post(f"/api/works/{work['id']}/answers", headers=STUDENT,
                        json={"answers": {"word_limit": "NO_LIMIT", "level": "LATER_UG", "ai_policy": "NOT_MENTIONED"}, "baseVersion": work["specVersion"]}).json()
    assert saved["spec"]["gate"] == "PASS" and not saved["spec"]["limits"]
    assert any(a.startswith("Your brief gives no word limit") for a in saved["spec"]["assumptions"])
    wrong = client.post(f"/api/works/{work['id']}/answers", headers=STUDENT, json={"answers": {"level": "NO_LIMIT"}, "baseVersion": saved["specVersion"]})
    assert wrong.status_code == 400


# --- page limits, delivery, numbers ---------------------------------------------------------------------


def test_a_page_limited_work_is_never_ready_from_an_estimate():
    from tests.test_works_golden import _req

    spec = _spec("CONCEPT_NOTE", "FUNDING_CONCEPT", reqs=[_req("limit.pages", "no more than 2 pages", 2, "pages")])
    doc = WorkDocument(kind="CONCEPT_NOTE", variant="FUNDING_CONCEPT", title="t", spec_version=1, plan_version=1,
                       sections=[WorkSection(key="summary", heading="Summary", paragraphs=["word " * 300])])
    estimated = report(Context(spec=spec, stage="FINAL", inputs=WorkInputs(title="A work"), doc=doc), ("DRAFT", "FINAL", "RENDER"))
    unmeasured = [i for i in estimated if i.reason == "PAGE_COUNT_UNMEASURED"]
    assert unmeasured and overall(estimated) == "NOT_READY" and "page count" in unmeasured[0].action
    measured = report(Context(spec=spec, stage="FINAL", inputs=WorkInputs(title="A work"), doc=doc, pages=1.0), ("DRAFT", "FINAL", "RENDER"))
    assert not [i for i in measured if i.reason == "PAGE_COUNT_UNMEASURED"]
    over = VALIDATORS["limits.rendered_pages"](Context(spec=spec, stage="FINAL", inputs=WorkInputs(title="A work"), doc=doc, pages=3), {"id": "SH-005"})
    assert over[0] == "FAIL"


def test_every_download_is_exactly_the_word_file_checked_before_completion(works_client):
    from app.runtime import get_runtime

    client = works_client
    work = _drafted(client)
    _, job = _run(client, work["id"], "DRAFT")
    assert job["status"] == "COMPLETED"
    rt = get_runtime()
    entry = rt.store.get_work(work["id"]).documents[-1]
    assert entry.docx_path and rt.files.exists(entry.docx_path)
    first = client.get(f"/api/works/{work['id']}/export", headers=STUDENT).content
    second = client.get(f"/api/works/{work['id']}/export", headers=STUDENT).content
    assert first == second == rt.files.get(entry.docx_path)
    paragraphs = [p.text for p in Document(io.BytesIO(first)).paragraphs if p.text.strip()]
    assert paragraphs[-1] == "This document was drafted by an AI-assisted third party."


def test_a_completed_work_converts_to_pdf_or_is_refused_as_incomplete(works_client):
    from app.latex.package import compile_pdf

    client = works_client
    work = _drafted(client)
    _run(client, work["id"], "DRAFT")
    response = client.get(f"/api/works/{work['id']}/export.pdf", headers=STUDENT)
    if response.status_code == 200:
        assert response.content[:4] == b"%PDF"
    else:  # no LaTeX on this machine: refused plainly, never an empty or partial PDF
        assert response.json()["code"] in ("PDF_FAILED", "PDF_INCOMPLETE"), response.json()
        assert compile_pdf is not None


@pytest.mark.parametrize(("key", "number", "quote", "verified"), [
    ("limit.words", 1500, "no more than 1 500 words", True),
    ("limit.words", 1500, "no more than 1 500 words", True),
    ("limit.words", 1500, "no more than 1 500 words", True),
    ("limit.words", 500, "no more than 1 500 words", False),
    ("limit.pages", 50, "a 5-page limit applies to all 50 applicants", False),
    ("ceiling", 50, "grants of up to USD 50k", False),
])
def test_numbers_in_quotes_are_read_whole_and_fail_closed(key, number, quote, verified):
    assert value_in_quote(key, str(number), number, quote) is verified



def test_the_reviewer_never_sees_sampling_settings_the_method_does_not_use():
    """Real-model pilot 2026-09-30: a power-analysis plan gave a margin of 0 (unused), PaperAid held the
    standard 0.05 in its place, and the reviewer read the two as a contradiction on every round."""
    from app.proposals.models import ProposalPlan
    from app.proposals.pipeline import _as_reviewed

    stated = ProposalPlan.model_validate({**PLAN, "sampleSize": {**PLAN.get("sampleSize", {}), "method": "AUTHOR_STATED", "stated": 384}})
    assert not {"margin", "proportion", "confidence"} & set(_as_reviewed(stated)["sampleSize"])
    cochran = ProposalPlan.model_validate({**PLAN, "sampleSize": {**PLAN.get("sampleSize", {}), "method": "COCHRAN"}})
    assert {"margin", "proportion", "confidence"} <= set(_as_reviewed(cochran)["sampleSize"])


def test_a_budget_that_does_not_match_the_amount_requested_is_caught_before_paying(works_client):
    """Real-model pilot 2026-09-30: the amount the student answered (USD 9,000) was never compared with
    the budget's total, so a draft could be priced on an inconsistent budget."""
    client = works_client
    created = client.post("/api/works", headers=STUDENT, json={
        "kind": "FUNDING_PROPOSAL", "variant": "NGO_PROJECT", "mode": "COMPACT",
        "inputs": {"title": "Safer deliveries in Kamuli", "description": "Too many mothers in Kamuli deliver at home and die of preventable complications.",
                   "answers": {"problem": "Home deliveries and delayed referral.", "intervention": "Train village health teams.", "duration_months": "12",
                               "currency": "USD", "amount_requested": "USD 9,000"}},
    }).json()
    work = client.post(f"/api/works/{created['id']}/answers", headers=STUDENT, json={"answers": {}, "skipRest": True, "baseVersion": created["specVersion"]}).json()
    client.post(f"/api/works/{work['id']}/spec/confirm", headers=STUDENT, json={"baseVersion": work["specVersion"]})
    _run(client, work["id"], "PLAN")
    work = _work(client, work["id"])
    results = work["results"]
    for ind in results["indicators"]:
        ind["target"], ind["baseline"] = (60, 34) if ind["level"] == "outcome" else (120, None)
    work = client.post(f"/api/works/{work['id']}/results", headers=STUDENT, json={"results": results, "baseVersion": work["resultsVersion"]}).json()
    work = client.post(f"/api/works/{work['id']}/results/approve", headers=STUDENT, json={"baseVersion": work["resultsVersion"]}).json()
    budget = work["budget"]
    for li, (qty, cost) in zip(budget["lines"], [(4, 1000), (12, 250), (4, 300)], strict=True):  # USD 8,200, not the 9,000 requested
        li["quantity"], li["unitCost"] = qty, cost
    work = client.post(f"/api/works/{work['id']}/budget", headers=STUDENT, json={"budget": budget, "baseVersion": work["budgetVersion"]}).json()
    client.post(f"/api/works/{work['id']}/plan/approve", headers=STUDENT, json={"baseVersion": work["planVersion"]})
    refused = client.post(f"/api/works/{work['id']}/steps", headers=STUDENT, json={"step": "DRAFT"})
    assert refused.status_code == 400 and refused.json()["code"] == "NOT_READY_TO_DRAFT" and "differs from the budget's total" in refused.json()["message"]


def test_a_funding_workplan_converts_to_pdf():
    """Real-model pilot 2026-09-30: a funding draft's PDF failed on the workplan's "■" (U+25A0)."""
    from docx import Document as NewDocument

    from app.latex.convert import convert
    from app.latex.package import compile_pdf

    doc = NewDocument()
    table = doc.add_table(rows=2, cols=3)
    for c, text in enumerate(["Activity", "Month 1", "Month 2"]):
        table.cell(0, c).text = text
    for c, text in enumerate(["Train village health teams", "■", ""]):
        table.cell(1, c).text = text
    buffer = io.BytesIO()
    doc.save(buffer)
    converted = convert(buffer.getvalue())
    assert r"\rule{1.2ex}{1.2ex}" in converted.tex and not converted.omitted
    pdf, problem = compile_pdf(converted)
    if pdf is None:  # a machine without LaTeX cannot compile; the character must not be the reason
        assert "U+25A0" not in (problem or ""), problem


def test_every_final_review_request_fits_the_bound_whole(works_client, monkeypatch):
    """Codex review 2026-10-07, finding 8: the bound covered the document part, not the specification, rules,
    context and manifest every part repeats; the whole request is now measured."""
    from app.ai.orchestration import payload_words

    monkeypatch.setattr(works_pipeline, "FINAL_PART_WORDS", 900)  # about 460 words of every request are repeated context
    client = works_client
    work = _drafted(client)
    _, job = _run(client, work["id"], "DRAFT")
    assert job["status"] == "COMPLETED", job.get("failure")
    sent = _requests(client, "w_final")
    assert len(sent) > 2 and all(payload_words(p) <= 900 for p in sent), [payload_words(p) for p in sent]
