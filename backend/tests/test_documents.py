import io
import re

import pytest
from docx import Document
from lxml import etree

from app.core.errors import InvalidDocument
from app.documents import protect
from app.documents.docx_io import W, apply_revisions, body_text_fingerprint, read_docx
from app.documents.intake import inspect_upload
from tests.conftest import fixture_bytes, manifest


@pytest.mark.parametrize("entry", manifest(), ids=lambda e: e["file"])
def test_intake_matches_manifest(entry):
    data = fixture_bytes(entry["file"])
    if entry["expect"] == "accept":
        assert inspect_upload(data, entry["file"], 20 * 1024 * 1024, 25_000, 150).word_count > 0
    else:
        with pytest.raises(InvalidDocument) as err:
            inspect_upload(data, entry["file"], 20 * 1024 * 1024, 25_000, 150)
        assert err.value.code == entry["rejectReason"]


def _xml(data: bytes, xpath_tag: str) -> list[bytes]:
    doc = Document(io.BytesIO(data))
    return [etree.tostring(el) for el in doc.element.body.iter(W + xpath_tag)]


def _edit_every_editable_block(name: str) -> tuple[bytes, bytes, dict[str, str]]:
    source = fixture_bytes(name)
    model = read_docx(source)
    revisions = {b.id: (b.masked or "").replace(" in ", " within ", 1) for b in model.blocks if b.editable}
    return source, apply_revisions(source, revisions), revisions


def test_citation_fields_survive_rewriting_byte_for_byte():
    source, patched, revisions = _edit_every_editable_block("citation_fields.docx")
    assert revisions, "fixture must have editable paragraphs"
    for tag in ("instrText", "fldChar"):
        assert _xml(source, tag) == _xml(patched, tag)
    assert "(Author 1, 2010)" in " ".join(body_text_fingerprint(patched))


@pytest.mark.parametrize(
    ("name", "tag"),
    [("footnotes.docx", "footnoteReference"), ("hyperlinks.docx", "hyperlink"), ("equation.docx", "oMath")],
)
def test_opaque_elements_survive_rewriting(name, tag):
    source, patched, _ = _edit_every_editable_block(name)
    if tag == "oMath":
        math = "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"
        find = lambda d: [etree.tostring(e) for e in Document(io.BytesIO(d)).element.body.iter(math)]  # noqa: E731
        assert find(source) == find(patched) and find(source)
    else:
        assert _xml(source, tag) == _xml(patched, tag) and _xml(source, tag)


def test_differently_formatted_runs_are_locked():
    source, patched, _ = _edit_every_editable_block("mixed_formatting.docx")
    italic = [r.text for p in Document(io.BytesIO(patched)).paragraphs for r in p.runs if r.italic]
    assert "Plasmodium falciparum" in italic


def test_unchanged_blocks_are_untouched_and_changed_blocks_change():
    source = fixture_bytes("simple_essay.docx")
    model = read_docx(source)
    target = next(b for b in model.blocks if b.editable)
    patched = apply_revisions(source, {target.id: "A completely new sentence for this paragraph."})
    before, after = body_text_fingerprint(source), body_text_fingerprint(patched)
    diffs = [i for i, (x, y) in enumerate(zip(before, after, strict=True)) if x != y]
    assert len(before) == len(after) and len(diffs) == 1


def test_paragraphs_with_existing_tracked_changes_are_not_editable():
    model = read_docx(fixture_bytes("existing_tracked_changes.docx"))
    tracked = [b for b in model.blocks if "undergraduate" in b.text]
    assert tracked and not tracked[0].editable
    assert any("tracked changes" in w for w in model.warnings)


def test_fake_headings_and_references_are_classified():
    headings = [b for b in read_docx(fixture_bytes("fake_headings.docx")).blocks if b.kind == "heading"]
    assert [b.level for b in headings] == [1, 1, 2, 1] and all(b.detected_heading for b in headings)
    kinds = {b.kind for b in read_docx(fixture_bytes("references_heavy.docx")).blocks}
    assert "reference" in kinds


def test_table_order_is_preserved():
    texts = [b.text for b in read_docx(fixture_bytes("interleaved_tables.docx")).blocks]
    assert texts.index("Variable") < texts.index("Group 01")


def test_protect_round_trip_and_checks():
    text = 'As Ssempala (2019) found, 62% of students agreed (Okello, 2021). See https://doi.org/10.1/x and "quoted words here".'
    masked, originals = protect.mask(text)
    assert "Ssempala (2019)" not in masked and "(Okello, 2021)" not in masked and "https://" not in masked
    assert protect.unmask(masked, originals) == text
    assert protect.check_rewrite(masked, masked) == []
    assert protect.check_rewrite(masked, masked.replace("62", "65"))
    assert protect.check_rewrite(masked, re.sub(r"⟦P1⟧", "", masked))
    assert protect.check_rewrite(masked, masked + " (Kato, 2020)")
