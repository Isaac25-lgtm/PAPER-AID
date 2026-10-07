"""Regressions for Codex's re-check of the 56c4f83 fixes (seven remaining issues) and its review of the
one-screen redesign (five issues). Each test is the reported failure with the correct outcome."""

import base64
import io
import json

import pytest
from docx import Document

from app.analysis import signals
from app.documents.docx_io import read_docx
from app.jobs.models import JobStatus, Stage
from app.jobs.pipeline import StageContext, _select_targets
from app.proposals import profile, rulebook, service
from app.runtime import get_runtime
from tests.fake_models import profile_answer
from tests.test_api import STUDENT, enable_score, wait
from tests.test_audit_56c4f83 import OWNER, UID, _comment, _paper, _project, _quote, _submit
from tests.test_proposals import _create
from tests.test_studio import _run


def _guide_bytes(text="research proposal"):
    doc = Document()
    for i in range(40):
        doc.add_paragraph(f"Section {i}: the {text} shall follow the structure set out here, with chapters and headings.")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# --- H03: a guide copy never outlives a failed attachment --------------------------------------------


def test_h03_any_failure_while_attaching_a_guide_removes_its_copy(client, monkeypatch):
    rt = get_runtime()
    pid = _create(client)["id"]
    written = []
    original_put = rt.files.put
    monkeypatch.setattr(rt.files, "put", lambda path, data, ct: (written.append(path), original_put(path, data, ct))[1])

    def broken(project_id, mutate):
        raise RuntimeError("database unavailable")

    monkeypatch.setattr(rt.store, "update_project", broken)
    with pytest.raises(RuntimeError):
        service.upload_guide(rt, OWNER, pid, "guide.docx", _guide_bytes())
    guide = next(p for p in written if "/guides/" in p)
    assert guide.startswith(f"users/{UID}/guides/")  # under the bucket's fixed-age backstop
    assert not rt.files.exists(guide)


# --- H05: a revision that did not resolve its comment neither marks it nor charges for it ------------


def test_h05_an_unresolved_revision_leaves_the_comment_open_and_costs_nothing(fixed_client):
    pid, _ = _project(fixed_client)
    cid = _comment(fixed_client, pid)
    fixed_client.models.overrides["p_fix"] = lambda p: {"sections": [{"key": s["key"], "paragraphs": [*s["text"], "Changed but not good enough."], "table": s["table"]} for s in p["sections"]]}
    fixed_client.models.overrides["p_review"] = lambda p: {"results": [{"key": s["key"], "grade": "REPAIR", "issues": ["The comment is still not answered."], "note": ""} for s in p["sections"]]}
    body = _quote(fixed_client, pid)
    _submit(fixed_client, pid, body)
    job = wait(fixed_client, body["job"]["id"])
    assert job["status"] == "FAILED" and job["failure"]["code"] == "NOTHING_REVISED" and job["billing"]["charged"] == 0
    feedback = fixed_client.get(f"/api/projects/{pid}", headers=STUDENT).json()["feedback"]
    assert next(c for c in feedback if c["id"] == cid)["status"] == "OPEN"


# --- M10: a torn history line never swallows the next entry --------------------------------------------


def test_m10_a_credit_after_a_torn_line_is_still_in_the_history(tmp_path):
    from app.integrations.store import LocalJobStore
    from app.pricing import credits

    store = LocalJobStore(tmp_path)
    store.update_wallet("u", "u@example.com", lambda w: credits.top_up(w, 1000, "first"))
    with store._ledger_path("u").open("a", encoding="utf-8") as ledger:
        ledger.write('{"id": "le_torn", "at": "2026-')  # a crash mid-append
    store.update_wallet("u", "u@example.com", lambda w: credits.top_up(w, 5000, "second"))
    history = store.ledger("u", None, 50)
    assert store.get_wallet("u").available == 6000 and [e.note for e in history] == ["second", "first"]
    store._wallet_path("u").write_text(store.get_wallet("u").model_copy(update={"ledger_backfilled": False}).model_dump_json(by_alias=True), encoding="utf-8")
    with store._ledger_path("u").open("a", encoding="utf-8") as ledger:
        ledger.write("{broken")
    assert store.backfill_ledger("u") is True  # a damaged line does not stop the backfill
    assert {e.note for e in store.ledger("u", None, 50)} == {"first", "second"}


# --- H04: a guide step priced before the migration reads the moved guide -------------------------------


def test_h04_a_profile_step_priced_before_migration_still_reads_its_guide(client):
    from app.jobs import service as jobs
    from app.proposals.models import GuideFile

    rt = get_runtime()
    pid = _create(client)["id"]
    legacy = f"users/{UID}/projects/{pid}/guide/old.txt"
    text = " ".join(p.text for p in Document(io.BytesIO(_guide_bytes())).paragraphs)
    rt.files.put(legacy, text.encode(), "text/plain")
    rt.store.update_project(pid, lambda p: p.model_copy(update={"guide": GuideFile(name="guide.docx", words=600, sha256="x", path=legacy)}))
    quoted = client.post(f"/api/projects/{pid}/steps", headers=STUDENT, json={"step": "PROFILE"}).json()
    assert jobs.migrate_legacy_project_files(rt) == 1 and not rt.files.exists(legacy)
    client.post(f"/api/projects/{pid}/steps/{quoted['job']['id']}/submit", headers=STUDENT, json={"quoteId": quoted["quote"]["id"]})
    assert wait(client, quoted["job"]["id"])["status"] == "COMPLETED"


# --- M27: the PDF and its cache key come from one snapshot ---------------------------------------------


def test_m27_an_edit_during_a_pdf_never_caches_the_old_document_under_the_new_details(client, monkeypatch):
    from app.proposals import service as proposals

    pid, _ = _project(client)
    rt = get_runtime()
    compiled = []
    monkeypatch.setattr(proposals, "compile_pdf", lambda result: (compiled.append(1), (b"%PDF-1.4 " + str(len(compiled)).encode(), ""))[1])
    original = proposals._export
    edited = []

    def export_then_edit(rt_, user, p, final):
        out = original(rt_, user, p, final)
        if not edited:  # the student renames themselves while the first PDF is being made
            edited.append(1)
            rt.store.update_project(pid, lambda q: q.model_copy(update={"title_page": q.title_page.model_copy(update={"student_name": "New Name"})}))
        return out

    monkeypatch.setattr(proposals, "_export", export_then_edit)
    first = client.get(f"/api/projects/{pid}/export.pdf", headers=STUDENT).content
    second = client.get(f"/api/projects/{pid}/export.pdf", headers=STUDENT).content
    assert first != second and len(compiled) == 2  # the new details get a new PDF


# --- M16: no malformed profile escapes as an ordinary error ---------------------------------------------


@pytest.mark.parametrize(
    "damage",
    [
        lambda a: a["chapters"][0].update(number=[1]),
        lambda a: a.update(institution=None),
        lambda a: a["formatting"].update(sizePt=10**400),
        lambda a: a.update(chapters=[None, "x", 3]),
        lambda a: a["chapters"][0]["sections"][0].update(key=None, heading=["x"], share={"a": 1}),
        lambda a: a.update(levels=[{"level": ["MASTERS"], "pagesMin": 10**500, "pagesMax": "30"}]),
        lambda a: a.update(vetting=[{"chapter": [1], "question": 5}], rules=[None, 3, "Write in the future tense."]),
    ],
)
def test_m16_malformed_profiles_are_built_or_refused_as_not_a_guide(damage):
    answer = profile_answer({"guide": "research proposal", "reference": profile.reference()})
    damage(answer)
    try:
        book = profile.build(answer, "guide.docx")
    except profile.NotAGuide:
        return
    assert book["chapters"][0]["sections"] and isinstance(book["institution"], str)


# --- M17: the page knows the profile is missing, so it can offer recovery -------------------------------


def test_m17_a_missing_profile_is_flagged_for_recovery(client):
    rt = get_runtime()
    pid, _ = _project(client)
    rt.store.update_project(pid, lambda p: p.model_copy(update={"rulebook": "custom-gone0000"}))
    rulebook.load.cache_clear()
    view = client.get(f"/api/projects/{pid}", headers=STUDENT).json()
    assert view["profileMissing"] is True
    assert view["chapters"][0]["current"] == 1
    restored = client.post(f"/api/projects/{pid}/rulebook/default", headers=STUDENT)
    assert restored.status_code == 200 and restored.json()["profileMissing"] is False
    assert restored.json()["chapters"][0]["current"] == 1


# --- Review #1: continuing uses the student's current keep/undo choices ---------------------------------


def test_review1_continuing_never_brings_back_rejected_wording(fixed_client):
    refined_id, refined = _run(fixed_client, {"writing": "REFINE", "academic": False})
    change = next(c for c in refined["refinement"]["changes"] if not c["kept"])
    fixed_client.post(f"/api/jobs/{refined_id}/rebuild", headers=STUDENT)
    rejected = fixed_client.post(f"/api/jobs/{refined_id}/changes/{change['blockId']}", headers=STUDENT, json={"accepted": False}).json()
    assert not any(o["id"] == "paper-reviewed" for o in rejected["outputs"])  # the old "your choices" file is withdrawn
    draft = fixed_client.post(f"/api/jobs/{refined_id}/continue", headers=STUDENT, json={"origin": "result"}).json()
    text = " ".join(b["text"] for b in fixed_client.get(f"/api/jobs/{draft['id']}/document", headers=STUDENT).json()["blocks"])
    full = next(c for c in fixed_client.get(f"/api/jobs/{refined_id}/document", headers=STUDENT).json()["changes"] if c["blockId"] == change["blockId"])
    assert full["before"] != full["after"] and full["before"] in text and full["after"] not in text


def test_reviewed_download_is_not_published_if_choices_change_during_its_save(fixed_client, monkeypatch):
    rt = get_runtime()
    job_id, refined = _run(fixed_client, {"writing": "REFINE", "academic": False})
    change = next(c for c in refined["refinement"]["changes"] if not c["kept"])
    assert fixed_client.post(f"/api/jobs/{job_id}/rebuild", headers=STUDENT).status_code == 200
    previous = next(o.path for o in rt.store.get(job_id).outputs if o.id == "paper-reviewed")
    original_update, original_put = rt.store.update, rt.files.put
    written = []

    def put(path, data, content_type):
        if "/output/paper-reviewed-" in path:
            written.append(path)
        return original_put(path, data, content_type)

    def competing_update(target, mutate):
        def reject(j):
            j.rejected_changes = [change["blockId"]]
            j.outputs = [o for o in j.outputs if o.id != "paper-reviewed"]
            return j
        original_update(target, reject)
        return original_update(target, mutate)

    monkeypatch.setattr(rt.files, "put", put)
    monkeypatch.setattr(rt.store, "update", competing_update)
    refused = fixed_client.post(f"/api/jobs/{job_id}/rebuild", headers=STUDENT)
    assert refused.status_code == 409 and refused.json()["code"] == "CHOICES_CHANGED"
    assert written and previous not in written
    assert all(not rt.files.exists(path) for path in written)
    assert rt.files.exists(previous)
    assert not any(o.id == "paper-reviewed" for o in rt.store.get(job_id).outputs)


# --- Review #2: selections are exact, and "whole paper" is the whole paper -----------------------------


def test_review2_an_invalid_passage_is_refused_not_widened(fixed_client):
    checked_id, _ = _run(fixed_client, {"writing": "AI_CHECK", "academic": False})
    refused = fixed_client.post(f"/api/jobs/{checked_id}/continue", headers=STUDENT, json={"origin": "original", "instruction": "Shorter.", "blocks": ["b99999"]})
    assert refused.status_code == 400 and refused.json()["code"] == "INVALID_SELECTION"


def test_review2_the_whole_paper_means_every_passage(fixed_client):
    rt = get_runtime()
    job, _ = _paper(rt)
    doc = Document()
    for i in range(305):
        doc.add_paragraph(f"Paragraph {i} of a long paper with enough words to count as a passage here.")
    out = io.BytesIO()
    doc.save(out)
    rt.files.put(job.source.path, out.getvalue(), "application/octet-stream")
    rt.store.update(job.id, lambda j: j.model_copy(update={"status": JobStatus.COMPLETED}))
    draft = fixed_client.post(f"/api/jobs/{job.id}/continue", headers=STUDENT, json={"origin": "original", "instruction": "Use plainer words."}).json()
    assert len(draft["selection"]["onlyBlocks"]) == 305 and list(draft["fixNotes"]) == ["*"]
    model = read_docx(out.getvalue())
    analysis, measured = signals.analyse(model)
    # the draft as a running stage owns it: artifacts are written only by the attempt holding the job
    running = rt.store.update(draft["id"], lambda j: j.model_copy(update={"status": JobStatus.PROCESSING, "stage": Stage.PLANNING, "lease_owner": "attempt"}))
    ctx = StageContext(rt, running)
    ctx.put_json("analysis.json", {"result": analysis.model_dump(by_alias=True), "scores": {s.block.id: s.score for s in measured}})
    targets = _select_targets(ctx, model)
    assert {t.id for t in targets} == {b.id for b in model.blocks if b.editable}


def test_project_deletion_erases_replaced_and_legacy_guides_but_keeps_other_projects(client):
    rt = get_runtime()
    pid, other = _create(client)["id"], _create(client)["id"]
    copies = []
    for project_id in (pid, pid, other):
        response = client.post(f"/api/projects/{project_id}/guide", headers=STUDENT, files={"file": ("guide.docx", _guide_bytes(), "application/octet-stream")})
        assert response.status_code == 200
        copies.append(rt.store.get_project(project_id).guide.path)
    legacy = f"users/{UID}/guides/{pid}-old-attempt.txt"
    rt.files.put(legacy, b"private old guide", "text/plain")
    assert client.delete(f"/api/projects/{pid}", headers=STUDENT).status_code == 204
    assert all(not rt.files.exists(path) for path in [*copies[:2], legacy])
    assert rt.files.exists(copies[2])


@pytest.mark.parametrize("writing", ["REFINE", "NONE"])
@pytest.mark.parametrize("fixture", ["simple_essay.docx", "dissertation_long.docx"])
def test_selected_passages_keep_their_identity_after_formatting_and_a_logo(fixed_client, writing, fixture):
    from tests.conftest import fixture_bytes

    rt = get_runtime()
    job_id = fixed_client.post("/api/jobs", headers=STUDENT).json()["id"]
    # Identical paragraphs deliberately rule out matching by text alone.
    doc = Document(io.BytesIO(fixture_bytes(fixture)))
    duplicate = "The proposed study will examine access to services in the district using household responses and local records to explain differences between groups."
    doc.add_heading("Additional findings", level=1)  # editable body text after the fixture's references
    for _ in range(3):
        doc.add_paragraph(duplicate)
    source = io.BytesIO()
    doc.save(source)
    assert fixed_client.post(f"/api/jobs/{job_id}/files/source", headers=STUDENT, files={"file": ("essay.docx", source.getvalue(), "application/octet-stream")}).status_code == 200
    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
    assert fixed_client.post(f"/api/jobs/{job_id}/files/logo", headers=STUDENT, files={"file": ("crest.png", png, "image/png")}).status_code == 200
    selection = {"writing": writing, "academic": False, "formatting": "FORMAT", "preset": "harvard", "logo": "CENTER"}
    quoted = fixed_client.post(f"/api/jobs/{job_id}/quote", headers=STUDENT, json={"selection": selection}).json()["quote"]
    assert fixed_client.post(f"/api/jobs/{job_id}/submit", headers=STUDENT, json={"quoteId": quoted["id"]}).status_code == 200
    assert wait(fixed_client, job_id)["status"] == "COMPLETED"
    view = fixed_client.get(f"/api/jobs/{job_id}/document", headers=STUDENT).json()
    rewritten = {c["blockId"]: c["after"] for c in view["changes"] if not c["kept"]}
    duplicates = [b for b in view["blocks"] if b["text"] == duplicate]
    assert len(duplicates) == 3
    picked = [next(b for b in view["blocks"] if b["kind"] == "paragraph"), duplicates[1]]
    request = fixed_client.post(f"/api/jobs/{job_id}/continue", headers=STUDENT, json={"origin": "result", "instruction": "Shorten only these passages.", "blocks": [b["id"] for b in picked]})
    assert request.status_code == 200, request.json()
    draft = request.json()
    current = read_docx(rt.files.get(rt.store.get(draft["id"]).source.path))
    selected = [current.by_id()[bid].text for bid in draft["selection"]["onlyBlocks"]]
    assert selected == [rewritten.get(b["id"], b["text"]) for b in picked]
    new_duplicates = [b.id for b in current.blocks if b.text == rewritten.get(duplicates[1]["id"], duplicate)]
    assert draft["selection"]["onlyBlocks"][-1] == new_duplicates[1]


def test_old_refined_results_recover_both_percentages_without_a_model_call(fixed_client):
    enable_score()
    rt = get_runtime()
    job_id, original = _run(fixed_client, {"writing": "REFINE", "academic": False})
    calls = len(fixed_client.models.requests)

    def old(j):
        j.analysis = j.analysis.model_copy(update={"percent": None})
        j.analysis_after = j.analysis_after.model_copy(update={"percent": None})
        return j

    rt.store.update(job_id, old)
    view = fixed_client.get(f"/api/jobs/{job_id}/document", headers=STUDENT).json()
    assert view["percent"] == original["analysis"]["percent"]
    assert view["percentAfter"] == original["analysisAfter"]["percent"]
    assert len(fixed_client.models.requests) == calls


# --- Review #3 lives in the web wording; review #4: a priced request is priced alone --------------------


def test_review4_an_abandoned_request_never_joins_a_later_price(client):
    rt = get_runtime()
    pid, _ = _project(client)
    client.post(f"/api/projects/{pid}/chapters/1/request", headers=STUDENT, json={"instruction": "Rewrite the background.", "sections": ["background"]})
    later = client.post(f"/api/projects/{pid}/chapters/1/request", headers=STUDENT, json={"instruction": "Sharpen the problem.", "sections": ["problem"]}).json()
    request_id = later["feedback"][-1]["id"]
    quoted = client.post(f"/api/projects/{pid}/steps", headers=STUDENT, json={"step": "REVISE_1", "comments": [request_id]}).json()
    job = rt.store.get(quoted["job"]["id"])
    frozen = json.loads(rt.files.get(f"{job.storage_prefix()}/internal/proposal_input.json"))
    assert list(frozen["revise"]) == ["problem"] and frozen["revise"]["problem"] == ["Sharpen the problem."]
    # without naming requests, only supervisor comments are revised: a student's abandoned request never is
    none = client.post(f"/api/projects/{pid}/steps", headers=STUDENT, json={"step": "REVISE_1"})
    assert none.status_code == 400 and none.json()["code"] == "NO_FEEDBACK"


# --- Old results get their percentage -----------------------------------------------------------------


def test_results_from_before_the_percentage_show_one(fixed_client):
    enable_score()
    rt = get_runtime()
    checked_id, checked = _run(fixed_client, {"writing": "AI_CHECK", "academic": False})
    rt.store.update(checked_id, lambda j: j.model_copy(update={"analysis": j.analysis.model_copy(update={"percent": None})}))
    doc = fixed_client.get(f"/api/jobs/{checked_id}/document", headers=STUDENT).json()
    assert doc["percent"] == checked["analysis"]["percent"]
