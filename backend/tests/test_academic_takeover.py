"""Takeover regressions: real failure boundaries, frozen contracts and unchanged approval/accounting paths.
Provider calls are disabled by conftest; no model usage, paid call or production data is involved.
"""

import hashlib
import io
import json
from copy import deepcopy

import pytest
from docx import Document
from pydantic import ValidationError

from app.proposals import decisions, evidence, export, framework, methods, pipeline, sampling
from app.proposals.ai import PLAN_SCHEMA_V3, Table, plan_schema
from app.proposals.models import (
    ChapterDocument,
    ChapterSection,
    MethodsContract,
    MethodVariable,
    ObjectiveMethod,
    ProposalInputs,
    ProposalPlan,
    SampleSize,
    StepInput,
)
from tests.fake_models import PLAN
from tests.test_independent_audit import ACT, BRIEF, _cited


def study() -> ProposalPlan:
    plan = ProposalPlan.model_validate(deepcopy(PLAN))
    plan.sample_size = SampleSize(method="YAMANE", population=5200, population_source="Student's eligible frame, 2026")
    plan.methods = MethodsContract(
        variables=[MethodVariable(id="V1", label="Uptake", role="OUTCOME", scale="BINARY", definition="Receipt recorded on the card", unit="eligible child", instrument="Child card", period="Survey date"),
                   MethodVariable(id="V2", label="Distance", role="PREDICTOR", scale="CONTINUOUS", definition="Travel distance", unit="km", instrument="Questionnaire", period="Survey date")],
        objectives=[ObjectiveMethod(objective=i, outcome_id="V1", predictor_ids=["V2"], estimator="DESCRIPTIVE" if i == 1 else "BINARY_LOGISTIC") for i in (1, 2, 3)],
        population_frame="Eligible children in the named study district, registry dated 2026", population_origin="USER_SUPPLIED")
    plan.variables = methods.framework_variables(plan)
    return plan


def test_legacy_schema_and_decision_hashes_do_not_acquire_new_assumptions():
    plan = ProposalPlan.model_validate(deepcopy(PLAN))
    assert plan.methods is None and plan_schema("p-plan-v5") == PLAN_SCHEMA_V3
    raw = PLAN["sampleSize"]
    sample = SampleSize.model_validate(raw)
    old = {k: v for k, v in sample.model_dump().items() if k not in ("design_effect", "nonresponse_rate", "adjustment_source")}
    assert decisions.hashes(plan)["sample_size"] == hashlib.sha256(json.dumps(old, sort_keys=True).encode()).hexdigest()[:12]
    inp = StepInput.model_construct(plan=plan, inputs=ProposalInputs(topic="A study of uptake", level="MASTERS"))
    assert "methods" not in pipeline._plan(inp) and "designEffect" not in pipeline._plan(inp)["sampleSize"]
    assert "samplingAdjustmentSource" not in pipeline._study(inp)
    assert "methods" in plan_schema("p-plan-v6")["required"]


@pytest.mark.parametrize("defect", ["duplicate_variable", "duplicate_objective", "missing_objective", "undefined_outcome", "undefined_predictor", "same_predictor", "wrong_scale"])
def test_structural_methods_contradictions_cannot_be_approved(defect):
    plan = study()
    contract = plan.methods
    if defect == "duplicate_variable":
        contract.variables.append(contract.variables[0].model_copy())
    elif defect == "duplicate_objective":
        contract.objectives.append(contract.objectives[0].model_copy())
    elif defect == "missing_objective":
        contract.objectives.pop()
    elif defect == "undefined_outcome":
        contract.objectives[1].outcome_id = "V9"
    elif defect == "undefined_predictor":
        contract.objectives[1].predictor_ids = ["V9"]
    elif defect == "same_predictor":
        contract.objectives[1].predictor_ids = ["V1"]
    else:
        contract.variables[0].scale = "COUNT"
    assert methods.problems(plan)


def test_descriptive_intensity_can_be_distinct_from_binary_primary_inference():
    plan = study()
    plan.methods.variables.append(MethodVariable(id="V3", label="Practice intensity", role="EXPLORATORY", scale="COUNT", definition="Number of listed practices used"))
    plan.methods.objectives[0].outcome_id = "V3"
    assert methods.problems(plan) == []
    plan.methods.objectives[1].outcome_id = "V3"
    assert any("requires a binary outcome" in p for p in methods.problems(plan))
    plan.methods.objectives[1].transformation = "Binary: one or more of the specified practices; rationale to be reviewed."
    assert methods.problems(plan) == []  # a proposal for semantic review, not a certification by code
    assert "clinical cutoff" in " ".join(methods.context(plan)["reviewQuestions"])


def test_model_cannot_invent_a_registry_or_sampling_adjustments_for_a_real_user_number():
    plan = study()
    plan.sample_size.population_source = "An invented government registry"
    plan.sample_size.design_effect, plan.sample_size.nonresponse_rate = 2, .2
    inp = StepInput.model_construct(inputs=ProposalInputs(topic="A study of uptake", level="MASTERS", population_size=5200), academic_policy=1)
    settled = pipeline._student_figures_only(plan, inp)
    assert settled.sample_size.population == 5200 and settled.sample_size.population_source == ""
    assert settled.sample_size.design_effect is None and settled.sample_size.nonresponse_rate is None
    assert settled.methods.population_origin == "USER_SUPPLIED"
    assert any("eligible study population" in q for q in settled.questions_for_student)
    assert "not independently verified" in " ".join(methods.context(settled)["reviewQuestions"])


def test_sampling_adjustments_are_calculated_only_from_explicit_supplied_assumptions():
    sample = SampleSize(method="YAMANE", population=5200)
    base = sampling.calculate(sample)
    assert base.size == 372 and "design effect" not in base.steps
    adjusted = sample.model_copy(update={"design_effect": 1.5, "nonresponse_rate": .1, "adjustment_source": "Supervisor-approved design rationale"})
    result = sampling.calculate(adjusted)
    assert result.size == 620 and "372" in result.steps and "558" in result.steps and "620" in result.steps
    assert "not evidence of statistical power" in result.steps
    assert sampling.calculate(adjusted.model_copy(update={"adjustment_source": ""})).missing
    assert sampling.calculate(adjusted.model_copy(update={"population": 50})).size is None
    assert sampling.calculate(adjusted.model_copy(update={"method": "CENSUS"})).missing


@pytest.mark.parametrize("value", [float("nan"), float("inf"), -1, 0])
def test_invalid_design_effect_cannot_enter_an_input_or_plan(value):
    with pytest.raises(ValidationError):
        ProposalInputs(topic="A study of uptake", level="MASTERS", design_effect=value)
    with pytest.raises(ValidationError):
        SampleSize(design_effect=value)


def test_changed_methods_invalidate_affected_text_but_legacy_plan_hashes_remain_valid():
    plan = study()
    doc = ChapterDocument(number=3, title="Methodology", plan_version=1, cited=[], words=10,
                          sections=[ChapterSection(key="analysis", number="3.9", heading="Analysis", paragraphs=["The study will analyse uptake."], depends=decisions.stamp(3, "analysis", plan))])
    assert decisions.stale(doc, plan) == []
    updated = plan.model_copy(deep=True)
    updated.methods.variables[0].period = "A different observation window"
    assert decisions.stale(doc, updated) == ["3.9 Analysis"]
    legacy = ProposalPlan.model_validate(deepcopy(PLAN))
    doc.sections[0].depends = decisions.stamp(3, "analysis", legacy)
    assert decisions.stale(doc, legacy) == []
    assert decisions.stale(doc, plan) == ["3.9 Analysis"]


def test_framework_additional_factors_do_not_draw_an_unjustified_moderation_arrow(monkeypatch):
    plan = study()
    plan.methods.variables.append(MethodVariable(id="V3", label="Household resources", role="ADJUSTMENT", scale="CONTINUOUS", rationale="Proposed confounding adjustment"))
    plan.variables = methods.framework_variables(plan)
    arrows = []
    old = framework._arrow
    def draw(*args, **kwargs):
        arrows.append(kwargs)
        return old(*args, **kwargs)
    monkeypatch.setattr(framework, "_arrow", draw)
    assert framework.figure(plan).startswith(b"\x89PNG")
    assert not any(a.get("dashed") for a in arrows)
    assert "adjustment" in framework.describe_plan(plan) and "No moderation or mediation arrow is implied" in framework.describe_plan(plan)
    arrows.clear()
    assert framework.draw(plan.variables)
    assert any(a.get("dashed") for a in arrows)  # legacy drawings are not silently reinterpreted


def test_a_revision_of_another_section_preserves_its_framework_snapshot():
    plan = study()
    base = ChapterDocument(number=1, title="Introduction", plan_version=1, cited=[], words=10, framework_plan=plan,
                           sections=[ChapterSection(key="framework", number="1.8", heading="Framework", paragraphs=["Original framework."]),
                                     ChapterSection(key="background", number="1.1", heading="Background", paragraphs=["Original background."])])
    later = plan.model_copy(deep=True)
    later.variables.dependent = ["A different outcome"]
    revised = ChapterDocument(number=1, title="Introduction", plan_version=2, cited=[], words=10, framework_plan=later,
                              sections=[ChapterSection(key="background", number="1.1", heading="Background", paragraphs=["New background."])])
    combined, _ = pipeline._merge(base, revised, {"background"})
    assert combined.framework_plan == plan
    revised.sections = [base.sections[0].model_copy(update={"paragraphs": ["New framework."]})]
    combined, _ = pipeline._merge(base, revised, {"framework"})
    assert combined.framework_plan == later


def test_ragged_proposal_tables_keep_every_cell():
    result = pipeline._table(Table(caption="Activities", rows=[["Activity", "Month"], ["Pilot", "One", "Permission pending"]]))
    assert result == [["Activity", "Month", ""], ["Pilot", "One", "Permission pending"]]


def test_framework_renaming_keeps_ids_and_removal_cannot_leave_an_undefined_outcome():
    plan = study()
    edited = plan.variables.model_copy(update={"independent": ["Travel distance"]})
    contract = methods.edit_framework(plan, edited)
    assert next(v for v in contract.variables if v.role == "PREDICTOR").id == "V2"
    assert contract.objectives[1].predictor_ids == ["V2"]
    with pytest.raises(ValueError, match="undefined outcome"):
        methods.edit_framework(plan, plan.variables.model_copy(update={"dependent": ["A different outcome", "A second outcome"]}))


def test_framework_links_follow_the_analysis_ids_rather_than_box_order():
    plan = study()
    plan.methods.variables.append(MethodVariable(id="V3", label="Another outcome", role="OUTCOME", scale="BINARY"))
    for row in plan.methods.objectives:
        row.outcome_id = "V3"
    plan.variables = methods.framework_variables(plan)
    assert framework.method_links(plan) == [(0, 1)]
    assert "Distance and Another outcome" in framework.describe_plan(plan)
    assert "Distance and Uptake" not in framework.describe_plan(plan)


def test_legal_host_or_title_alone_never_certifies_a_provision():
    text = "Section 6(2)(c) of the Education Act 2008 places a duty on parents ⟦E000001⟧."
    for title, host, passage in [(ACT, "https://ulii.org/a", "An unrelated provision"),
                                 ("Policy Brief about the Education Act 2008", "https://parliament.go.ug/a", "Section 6(2)(c) is discussed"),
                                 ("Civil Procedure Act 2008", "https://ulii.org/a", "Section 6(2)(c) discusses hearings"),
                                 (ACT, "https://ulii.org.evil.example/a", "Section 6(2)(c) discusses food")]:
        library = _cited(title, host)
        library["E000001"].passage = passage
        assert evidence.provision_problems(text, library)
    library = _cited(ACT, "https://ulii.org/a")
    library["E000001"].passage = "Section 6(2)(c) states the duty in the named Act."
    assert evidence.provision_findings(text, library) == ([], [])
    library["E000001"].verified = False
    assert evidence.provision_problems(text, library)


def test_legal_attribution_does_not_delete_conditions_quotes_or_criticism():
    library = _cited(BRIEF)
    for text in ['If Section 13 of the Education Act 2008 applied here, the claim would require checking ⟦E000001⟧.',
                 'The brief claims that Section 13 of the Education Act 2008 places a duty on parents ⟦E000001⟧.',
                 '“Section 13 of the Education Act 2008” is incorrectly numbered in the brief ⟦E000001⟧.']:
        problems, notes = evidence.provision_findings(text, library)
        assert not problems and notes and evidence.drop_provisions(text, library) == text
    text = "Section 13 of the Education Act 2008 imposes a duty ⟦E000001⟧."
    assert evidence.drop_provisions(text, library).startswith("The cited source reports that ")
    assert evidence.cited_ids(evidence.drop_provisions(text, library)) == ["E000001"]
    uncited = "Section 13 of the Education Act 2008 imposes a duty."
    assert evidence.provision_problems(uncited, library)
    assert evidence.drop_provisions(uncited, library) == uncited  # never invent attribution for an uncited sentence


def test_a_supported_provision_does_not_vouch_for_a_second_number():
    library = _cited(ACT, "https://ulii.org/a")
    library["E000001"].passage = "Section 6 of the Education Act describes a duty."
    text = "Section 6 of the Education Act 2008 and Section 99 of the Education Act 2008 impose duties ⟦E000001⟧."
    assert evidence._unproved_provision(text, text, library) == ("Section 99", "Education Act")
    assert evidence.provision_problems(text, library)


def _final_case():
    from app.works.models import PlanSection, ResolvedSpec, WorkDocument, WorkInputs, WorkPlan, WorkSection, WorkStepInput
    spec = ResolvedSpec(version=1, kind="COURSEWORK", variant="ESSAY", rules_version="rules-v1", target_words=1500)
    plan = WorkPlan(title="Methods", sections=[PlanSection(key="methods", heading="Methods", words=1000)])
    inp = WorkStepInput(work_id="wrk_test", step="DRAFT", kind="COURSEWORK", variant="ESSAY", inputs=WorkInputs(title="Methods"), spec=spec, plan=plan, academic_policy=1)
    document = WorkDocument(kind="COURSEWORK", variant="ESSAY", title="Methods", spec_version=1, plan_version=1, status="READY",
                            sections=[WorkSection(key="methods", heading="Methods", paragraphs=["Travel distance was associated with uptake ⟦E000001⟧."])], cited=["E000001"])
    library = _cited("Observational study", "https://example.org/article")
    library["E000001"].passage = "Distance was associated with uptake among surveyed caregivers."
    library["E000001"].scope = "Caregivers in a single district"
    return inp, document, library


class _Final:
    def __init__(self):
        self.calls = []
    def final(self, request):
        from app.works.ai import Final
        self.calls.append(request)
        return Final(rules=[], coverage=[], priorities=[])


def test_final_reviewer_receives_actual_cited_passages_and_their_limits():
    from app.works.pipeline import _final_review
    inp, document, library = _final_case()
    runner = _Final()
    assert _final_review(runner, inp, document, library, {}, []) is not None
    assert len(runner.calls) == 1
    request = runner.calls[0]
    assert request["citedEvidence"][0]["sourceWords"] == library["E000001"].passage
    assert request["citedEvidence"][0]["scope"] == library["E000001"].scope
    assert request["citedEvidence"][0]["reference"] in request["document"]["references"]
    assert request["sectionEvidence"] == {"methods": ["E000001"]}
    inp.academic_policy = 0
    runner = _Final()
    assert _final_review(runner, inp, document, library, {}, []) is not None
    assert "citedEvidence" not in runner.calls[0]  # frozen earlier request


def test_unavailable_or_oversized_final_evidence_never_sends_a_paid_request():
    from app.works.pipeline import _final_review
    inp, document, library = _final_case()
    library["E000001"].verified = False
    runner = _Final()
    assert _final_review(runner, inp, document, library, {}, []) is None and not runner.calls
    library["E000001"].verified = True
    library["E000001"].passage = "evidence " * 8000
    assert _final_review(runner, inp, document, library, {}, []) is None and not runner.calls


def test_population_frame_confirmation_is_owned_versioned_idempotent_and_not_billable(client):
    from app.runtime import get_runtime
    from tests.test_api import OTHER, STUDENT
    from tests.test_proposals import _create
    project = _create(client)
    rt = get_runtime()
    def prepared(p):
        p.plan, p.plan_version, p.plan_status = study(), 1, "APPROVED"
        return p
    rt.store.update_project(project["id"], prepared)
    before = len(client.models.tasks)
    path = f"/api/projects/{project['id']}/population-frame"
    assert client.post(path, headers=OTHER, json={"baseVersion": 1}).status_code == 404
    assert client.post(path, headers=STUDENT, json={"baseVersion": 0}).status_code == 409
    assert client.post(path, headers=STUDENT, json={"baseVersion": 1}).status_code == 200
    assert methods.frame_confirmed(rt.store.get_project(project["id"]))
    assert client.post(path, headers=STUDENT, json={"baseVersion": 1}).status_code == 200
    saved = rt.store.get_project(project["id"])
    assert len([a for a in saved.acknowledgments if a.kind == "FRAME"]) == 1 and len(client.models.tasks) == before
    saved.plan.sample_size.population = 5100
    assert not methods.frame_confirmed(saved)
    saved = rt.store.get_project(project["id"])
    saved.plan_version = 2
    assert not methods.frame_confirmed(saved)


def test_submission_frame_blocker_does_not_block_a_chapter_one_draft(client):
    from app.runtime import get_runtime
    from tests.test_api import STUDENT
    from tests.test_proposal_v2 import _with_chapter_one
    project_id = _with_chapter_one(client)
    rt = get_runtime()
    def prepared(p):
        p.plan = study()
        return p
    rt.store.update_project(project_id, prepared)
    project = rt.store.get_project(project_id)
    assert any("eligible sampling frame" in p for p in export.final_blockers(project, {}))
    response = client.get(f"/api/projects/{project_id}/export", headers=STUDENT)
    assert response.status_code == 200
    text = "\n".join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)
    assert "user supplied" in text and "Written so far: Chapter One of three." in text
