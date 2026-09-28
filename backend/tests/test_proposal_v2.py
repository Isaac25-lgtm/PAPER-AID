"""Proposal V2: supervisor feedback placed by code and applied by revising only those sections, the
response report, version comparison and the conceptual framework figure."""

import io

from docx import Document

from app.proposals import feedback
from tests.test_api import STUDENT
from tests.test_proposals import _approved, _create, _run

WRITTEN = [
    feedback.Written(1, "background", "1.1", "Background to the Study"),
    feedback.Written(1, "problem", "1.2", "Statement of the Problem"),
    feedback.Written(2, "empirical1", "2.2", "Distance to health facilities"),
    feedback.Written(2, "empirical2", "2.3", "Caregivers' attitudes"),
    feedback.Written(3, "sampling", "3.5", "Sample Size and Sampling Techniques"),
]


def test_pasted_feedback_is_split_into_comments():
    text = "Comments on your proposal:\n1. The problem statement lacks local figures\n   from Mukono.\n2) Check section 3.5.\n\n- Use the future tense."
    comments = [c.text for c in feedback.split(text)]
    assert comments == ["Comments on your proposal:", "The problem statement lacks local figures from Mukono.", "Check section 3.5.", "Use the future tense."]
    assert [c.text for c in feedback.split("First point.\n\nSecond point\ncontinued.")] == ["First point.", "Second point continued."]


def test_comments_are_placed_by_number_heading_or_the_words_they_use():
    def place(text, anchor=""):
        return feedback.suggest(feedback.Read(text, anchor), WRITTEN)

    assert place("Check section 3.5 again.") == (3, ["sampling"])
    assert place("Too long.", anchor="1.2 Statement of the Problem") == (1, ["problem"])
    assert place("The statement of the problem needs local data.") == (1, ["problem"])
    assert place("Synthesise the empirical literature instead of listing studies.") == (2, ["empirical1", "empirical2"])
    assert place("Good work overall.") == (None, [])
    assert place("Explain the ethics approval process.") == (None, [])  # not written yet: never suggested


def test_word_comments_are_read_with_the_heading_they_sit_under():
    doc = Document()
    doc.add_heading("1.2 Statement of the Problem", level=2)
    paragraph = doc.add_paragraph("Malaria remains a major problem in the district.")
    doc.add_comment(paragraph.runs, text="Give the district's figures here.", author="Supervisor")
    buffer = io.BytesIO()
    doc.save(buffer)
    read = feedback.read_file("marked.docx", buffer.getvalue())
    assert read == [feedback.Read("Give the district's figures here.", "1.2 Statement of the Problem")]
    assert feedback.suggest(read[0], WRITTEN) == (1, ["problem"])


def _with_chapter_one(client):
    project = _create(client)
    _run(client, project["id"], "PLAN")
    _approved(client, project["id"])  # Chapter One starts on approval
    return project["id"]


def test_feedback_revises_only_its_sections_and_the_report_says_what_was_done(client):
    project_id = _with_chapter_one(client)
    added = client.post(f"/api/projects/{project_id}/feedback", headers=STUDENT, json={"text": "1. In 1.2, show the size of the problem in Mukono.\n2. Good overall; no other changes."})
    assert added.status_code == 200, added.json()
    first, second = added.json()["feedback"]
    assert (first["chapter"], first["sections"], first["status"], first["round"]) == (1, ["problem"], "OPEN", 1)
    assert second["chapter"] is None

    def revise(payload):  # the writer answers the supervisor's comment
        return {"sections": [{"key": s["key"], "paragraphs": [*s["text"], "In Mukono, the problem affects many households."], "table": s["table"]} for s in payload["sections"]]}

    client.models.overrides["p_fix"] = revise
    quoted = client.post(f"/api/projects/{project_id}/steps", headers=STUDENT, json={"step": "REVISE_1"}).json()
    assert "revised from your supervisor" in quoted["quote"]["lines"][0]["label"]
    job = _run(client, project_id, "REVISE_1")
    assert job["status"] == "COMPLETED", job
    fixed = [t for t in client.models.tasks if t == "p_fix"]
    assert fixed and "p_draft" not in client.models.tasks[client.models.tasks.index("p_fix") :]  # nothing redrafted

    project = client.get(f"/api/projects/{project_id}", headers=STUDENT).json()
    chapter = next(c for c in project["chapters"] if c["number"] == 1)
    assert chapter["current"] == 2 and chapter["versions"][-1]["note"] == "Revised from your supervisor's comments"
    assert project["feedback"][0]["status"] == "APPLIED" and project["feedback"][0]["appliedIn"] == 2

    diff = client.get(f"/api/projects/{project_id}/chapters/1/compare?older=1&newer=2", headers=STUDENT).json()
    assert diff["changed"] == 1, [(s["heading"], s["status"], s["pieces"][:3]) for s in diff["sections"] if s["status"] != "SAME"]
    changed = next(s for s in diff["sections"] if s["status"] == "CHANGED")
    assert changed["heading"] == "Statement of the Problem"
    assert any(p["op"] == "added" and "Mukono" in p["text"] for p in changed["pieces"])

    declined = client.post(f"/api/projects/{project_id}/feedback/{second['id']}", headers=STUDENT, json={"status": "DECLINED", "response": "Thank you."})
    assert declined.status_code == 200
    report = client.get(f"/api/projects/{project_id}/feedback/report", headers=STUDENT)
    assert report.status_code == 200
    rows = [[c.text for c in r.cells] for r in Document(io.BytesIO(report.content)).tables[0].rows]
    assert rows[1][2] == "Chapter 1: 1.2 Statement of the Problem" and rows[1][3] == "Revised in Chapter 1 (version 2)."
    assert rows[2][2] == "The whole proposal" and rows[2][3] == "Thank you."

    nothing = client.post(f"/api/projects/{project_id}/steps", headers=STUDENT, json={"step": "REVISE_1"})
    assert nothing.status_code == 400 and nothing.json()["code"] == "NO_FEEDBACK"


def test_feedback_can_only_be_placed_on_written_sections(client):
    project_id = _with_chapter_one(client)
    comment = client.post(f"/api/projects/{project_id}/feedback", headers=STUDENT, json={"text": "Rework this."}).json()["feedback"][0]
    wrong = client.post(f"/api/projects/{project_id}/feedback/{comment['id']}", headers=STUDENT, json={"chapter": 3, "sections": ["sampling"]})
    assert wrong.status_code == 400 and wrong.json()["code"] == "UNKNOWN_SECTION"
    applied = client.post(f"/api/projects/{project_id}/feedback/{comment['id']}", headers=STUDENT, json={"chapter": 1, "sections": ["problem"], "status": "APPLIED"})
    assert applied.status_code == 400 and applied.json()["code"] == "NOT_APPLIED"
    unwritten = client.post(f"/api/projects/{project_id}/steps", headers=STUDENT, json={"step": "REVISE_2"})
    assert unwritten.json()["code"] == "NOTHING_TO_REVISE"


def test_the_conceptual_framework_is_drawn_from_the_plan(client):
    project_id = _with_chapter_one(client)
    chapter = client.get(f"/api/projects/{project_id}/chapters/1", headers=STUDENT).json()
    assert [c["label"] for c in chapter["framework"]] == ["Independent variables", "Dependent variable"]
    assert chapter["framework"][0]["items"] == ["distance", "attitudes", "health worker advice"]
    doc = Document(io.BytesIO(client.get(f"/api/projects/{project_id}/export", headers=STUDENT).content))
    texts = [p.text for p in doc.paragraphs]
    assert "Figure 1.1: Conceptual framework" in texts and "List of Figures" in texts
    figure = next(t for t in doc.tables if "Independent variables" in t.cell(0, 0).text)
    assert "vaccine uptake" in figure.cell(0, 2).text and figure.cell(0, 1).text == "→"
