"""Proposal V1: rulebook, decisions, evidence rules, the project lifecycle and the review service.
Invariants first (Codex review, §8.6): stale edits, dependency invalidation, invented figures and
citations, author-owned facts, privacy of searches, deletion and expiry."""

import io
from datetime import timedelta

from docx import Document

from app.proposals import decisions, evidence, rulebook, sampling
from app.proposals.models import EvidenceItem, EvidenceSource, ProposalPlan, SampleSize
from tests.fake_models import PLAN
from tests.test_api import OTHER, STUDENT, wait

DETAILS = {
    "inputs": {"topic": "Malaria vaccine uptake among caregivers in Mukono", "level": "MASTERS", "programme": "Master of Public Health", "faculty": "Faculty of Health Sciences", "studyArea": "Mukono District"},
    "titlePage": {"studentName": "Grace Namukasa", "regNumber": "M24/U001", "supervisor": "Dr. John Okot", "submissionDate": "October 2026"},
    "citation": "APA6",
}


def plan(**changes) -> ProposalPlan:
    return ProposalPlan.model_validate({**PLAN, **changes})


# --- rulebook, decisions and sampling ---------------------------------------------------------


def test_sections_follow_the_manual_and_the_study_type():
    quantitative = rulebook.sections("ucu-2018-v1", 3, "MASTERS", plan())
    qualitative = rulebook.sections("ucu-2018-v1", 3, "MASTERS", plan(studyType="QUALITATIVE"))
    assert "variables" in {s.key for s in quantitative} and "variables" not in {s.key for s in qualitative}
    review = rulebook.sections("ucu-2018-v1", 2, "MASTERS", plan())
    assert [s.key for s in review if s.key.startswith("empirical")] == ["empirical1", "empirical2", "empirical3"]
    assert review[2].heading == "The effect of distance on uptake"
    assert [s.number for s in review][:2] == ["2.0", "2.1"]


def test_plan_rules_block_structure_but_only_advise_on_counts():
    assert rulebook.plan_problems("ucu-2018-v1", plan()) == []
    uneven = plan(researchQuestions=["Only one?"])
    assert any("own research question" in p for p in rulebook.plan_problems("ucu-2018-v1", uneven))
    no_row = plan(alignment=PLAN["alignment"][:2])
    assert any("Objective 3" in p for p in rulebook.plan_problems("ucu-2018-v1", no_row))
    six = plan(specificObjectives=[f"To study {i}" for i in range(6)], researchQuestions=[f"Q{i}?" for i in range(6)], alignment=[{"objective": i, "data": "d", "collection": "c", "analysis": "a"} for i in range(1, 7)])
    assert rulebook.plan_problems("ucu-2018-v1", six) == [] and rulebook.plan_advice("ucu-2018-v1", six)


def test_sample_sizes_are_calculated_by_code_from_the_students_figures():
    assert sampling.calculate(SampleSize(method="YAMANE", population=2400)).size == 343
    assert sampling.calculate(SampleSize(method="KREJCIE_MORGAN", population=2400)).size == 332
    assert sampling.calculate(SampleSize(method="COCHRAN")).size == 385
    assert sampling.calculate(SampleSize(method="COCHRAN", population=2400)).size == 332
    missing = sampling.calculate(SampleSize(method="YAMANE"))
    assert missing.size is None and "accessible population" in missing.missing
    assert rulebook.chapter_blockers(plan(sampleSize={"method": "YAMANE"}), 3)
    assert rulebook.chapter_blockers(plan(sampleSize={"method": "YAMANE"}), 1) == []
    assert rulebook.chapter_blockers(plan(sampleSize={"method": "SATURATION", "stated": 20}), 3)  # quantitative work needs a calculation
    assert rulebook.chapter_blockers(plan(studyType="QUALITATIVE", questionsKind="QUESTIONS", sampleSize={"method": "SATURATION", "stated": 20}), 3) == []


def test_changing_one_decision_flags_only_the_sections_built_on_it():
    from app.proposals.models import ChapterDocument, ChapterSection

    before = plan()
    doc = ChapterDocument(
        number=2, title="Literature Review", plan_version=1, cited=[], words=10,
        sections=[ChapterSection(key=k, number=n, heading=h, paragraphs=["x"], depends=decisions.stamp(2, k, before)) for k, n, h in (("theory", "2.1", "Theoretical Review"), ("empirical2", "2.3", "Attitudes"), ("empirical1", "2.2", "Distance"))],
    )
    after = plan(specificObjectives=[PLAN["specificObjectives"][0], "To examine caregivers' trust in the vaccine", PLAN["specificObjectives"][2]])
    assert decisions.stale(doc, before) == []
    assert decisions.stale(doc, after) == ["2.3 Attitudes"]
    assert "O2" in decisions.changed(before, after)


# --- evidence and citation rules -------------------------------------------------------------


def _item(id_, passage, **source) -> EvidenceItem:
    return EvidenceItem(id=id_, chapter=1, need="n", statement="s", passage=passage, verified=True, support="SUPPORTED", source=EvidenceSource(**source), retrieved_on="2026-09-28")


LIB = {
    "E00000a": _item("E00000a", "malaria accounts for 12% of all deaths in children", url="https://who.int/r", title="World malaria report 2024", organisation="World Health Organization", year="2024"),
    "E00000b": _item("E00000b", "uptake was 45.5% among caregivers in rural districts", url="https://doi.org/10.1/x", title="Uptake", authors=["Okello, J.", "Namara, A.", "Kato, P."], year="2022", container="Malaria Journal", volume="21", issue="3", pages="1–10", doi="10.1/x", kind="ARTICLE"),
}


def test_citations_render_in_the_projects_apa_edition():
    text = "Uptake was 45.5% ⟦E00000b⟧. ⟦E00000b|n⟧ found it low."
    apa6 = evidence.Citer(LIB, "APA6").render(text)
    assert apa6 == "Uptake was 45.5% (Okello, Namara, & Kato, 2022). Okello et al. (2022) found it low."
    assert evidence.Citer(LIB, "APA7").render(text) == "Uptake was 45.5% (Okello et al., 2022). Okello et al. (2022) found it low."
    references = evidence.reference_list([s.source for s in LIB.values()], "APA6")
    assert references[0] == "Okello, J., Namara, A., & Kato, P. (2022). Uptake. Malaria Journal, 21(3), 1–10. doi:10.1/x"
    assert references[1] == "World Health Organization. (2024). World malaria report 2024. Retrieved from https://who.int/r"
    assert evidence.reference_list([LIB["E00000b"].source], "APA7")[0].endswith("https://doi.org/10.1/x")


def test_only_confirmed_evidence_and_its_figures_may_appear():
    usable = set(LIB)
    assert evidence.citation_problems("Fine ⟦E00000a⟧.", usable) == []
    assert evidence.citation_problems("Invented ⟦E999999⟧.", usable)
    assert evidence.citation_problems("As Smith (2019) showed.", usable)
    assert evidence.citation_problems("As shown (Smith, 2019).", usable)
    assert evidence.figure_problems("Deaths were 12% ⟦E00000a⟧.", LIB, "") == []
    assert evidence.figure_problems("Deaths were 14% ⟦E00000a⟧.", LIB, "")
    assert evidence.figure_problems("A sample of 343 caregivers.", LIB, '{"population": 2400} 343') == []
    kept = evidence.strip_unsupported("Deaths were 12% ⟦E00000a⟧. In towns it was 60%. Uptake will be measured.", LIB, usable, "")
    assert kept == "Deaths were 12% ⟦E00000a⟧. Uptake will be measured."


def test_the_planned_study_must_be_in_the_future_tense():
    assert evidence.tense_problems("Data were collected using a questionnaire.")
    assert evidence.tense_problems("The study will use a questionnaire.") == []


# --- the project lifecycle over the API ------------------------------------------------------


def _create(client, headers=STUDENT):
    response = client.post("/api/projects", headers=headers, json=DETAILS)
    assert response.status_code == 200, response.json()
    return response.json()


def _run(client, project_id, step, note="", headers=STUDENT):
    quoted = client.post(f"/api/projects/{project_id}/steps", headers=headers, json={"step": step, "note": note})
    assert quoted.status_code == 200, quoted.json()
    body = quoted.json()
    submitted = client.post(f"/api/projects/{project_id}/steps/{body['job']['id']}/submit", headers=headers, json={"quoteId": body["quote"]["id"]})
    assert submitted.status_code == 200, submitted.json()
    return wait(client, body["job"]["id"], headers)


def _approved(client, project_id, **plan_changes):
    project = client.get(f"/api/projects/{project_id}", headers=STUDENT).json()
    edited = {**project["plan"], **plan_changes}
    saved = client.post(f"/api/projects/{project_id}/plan", headers=STUDENT, json={"plan": edited, "baseVersion": project["planVersion"]})
    assert saved.status_code == 200, saved.json()
    approved = client.post(f"/api/projects/{project_id}/plan/approve", headers=STUDENT, json={"baseVersion": saved.json()["planVersion"]})
    assert approved.status_code == 200, approved.json()
    body = approved.json()
    if body["activeJob"]:  # Chapter One starts on the first approval (owner request 2026-09-28)
        wait(client, body["activeJob"])
    return body


def test_a_plan_is_researched_negotiated_and_never_invents_the_students_figures(client):
    project = _create(client)
    job = _run(client, project["id"], "PLAN")
    assert job["status"] == "COMPLETED" and job["projectId"] == project["id"]
    project = client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()
    assert project["planStatus"] == "DRAFT" and project["planVersion"] == 1
    assert project["plan"]["sampleSize"]["population"] is None  # the model's 2,400 was never the student's figure
    assert any("accessible population" in q for q in project["plan"]["questionsForStudent"])
    assert project["evidenceCount"] == 2
    tasks = client.models.tasks
    assert tasks.index("p_needs") < tasks.index("p_extract") < tasks.index("verify") < tasks.index("p_plan") < tasks.index("p_critique") < tasks.index("p_finalise")
    library = client.get(f"/api/projects/{project['id']}/evidence", headers=STUDENT).json()
    article = next(i for i in library if i["source"]["doi"])
    assert article["verified"] and article["support"] == "SUPPORTED" and article["source"]["metadata"] == "CROSSREF"
    assert article["source"]["title"] == "Vaccine Uptake Among Caregivers in Central Uganda"  # registered details, not the model's


def test_an_over_long_field_from_the_model_is_shortened_to_a_whole_sentence(client):
    """Found in a real-model run: a long analysis plan must not fail the step. It ends at a complete
    sentence (a sentence cut mid-way made both reviewers refuse a plan: Phase 4 pilot)."""
    long_plan = {**PLAN, "alignment": [{**PLAN["alignment"][0], "analysis": "Binary logistic regression will be used. " * 60}, *PLAN["alignment"][1:]]}
    client.models.overrides["p_finalise"] = lambda payload: long_plan
    project = _create(client)
    job = _run(client, project["id"], "PLAN")
    assert job["status"] == "COMPLETED", job
    analysis = client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()["plan"]["alignment"][0]["analysis"]
    assert len(analysis) <= 1000 and analysis.endswith("used.")


def test_an_over_long_run_on_field_fails_the_plan_without_charge(fixed_client):
    """No sentence end to cut at: never delivered cut mid-way (Codex review 2026-09-30 #7)."""
    long_plan = {**PLAN, "alignment": [{**PLAN["alignment"][0], "analysis": "Binary logistic regression " * 80}, *PLAN["alignment"][1:]]}
    fixed_client.models.overrides["p_finalise"] = lambda payload: long_plan
    project = _create(fixed_client)
    job = _run(fixed_client, project["id"], "PLAN")
    assert job["status"] == "FAILED" and job["failure"]["code"] == "PLAN_INVALID" and job["billing"]["charged"] == 0


def test_the_plans_price_includes_chapter_one_which_starts_on_approval_once(client):
    project = _create(client)
    quoted = client.post(f"/api/projects/{project['id']}/steps", headers=STUDENT, json={"step": "PLAN"}).json()
    assert [line["label"] for line in quoted["then"]] == ["Chapter 1, started when you approve the plan (up to)"] and quoted["then"][0]["amount"] > 0
    assert client.post(f"/api/projects/{project['id']}/steps/{quoted['job']['id']}/submit", headers=STUDENT, json={"quoteId": quoted["quote"]["id"]}).status_code == 200
    wait(client, quoted["job"]["id"])
    approved = _approved(client, project["id"], sampleSize={**PLAN["sampleSize"], "populationSource": "DHO records"})
    assert approved["activeJob"] and approved["notice"] is None
    chapters = {c["number"]: c for c in client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()["chapters"]}
    assert chapters[1]["current"] == 1
    again = _approved(client, project["id"], title="A refined title")  # a later approval starts nothing
    assert again["activeJob"] is None


def test_chapter_one_that_cannot_start_says_why(client):
    from app.runtime import get_runtime

    rt = get_runtime()
    project = _create(client)
    _run(client, project["id"], "PLAN")
    owner = rt.store.get_project(project["id"]).owner_uid
    rt.store.update_wallet(owner, "student@example.com", lambda w: w.model_copy(update={"available": 1}))  # almost nothing left
    approved = _approved(client, project["id"])
    assert approved["activeJob"] is None and "could not start automatically" in approved["notice"] and "tokens" in approved["notice"]


def test_a_stale_plan_edit_is_refused_and_approval_needs_a_complete_plan(client):
    project = _create(client)
    _run(client, project["id"], "PLAN")
    current = client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()
    first = client.post(f"/api/projects/{project['id']}/plan", headers=STUDENT, json={"plan": {**current["plan"], "title": "Tab one"}, "baseVersion": current["planVersion"]})
    assert first.status_code == 200
    second = client.post(f"/api/projects/{project['id']}/plan", headers=STUDENT, json={"plan": {**current["plan"], "title": "Tab two"}, "baseVersion": current["planVersion"]})
    assert second.status_code == 409 and second.json()["code"] == "PLAN_CHANGED"
    broken = {**first.json()["plan"], "researchQuestions": ["Only one?"]}
    saved = client.post(f"/api/projects/{project['id']}/plan", headers=STUDENT, json={"plan": broken, "baseVersion": first.json()["planVersion"]}).json()
    assert saved["planProblems"]
    refused = client.post(f"/api/projects/{project['id']}/plan/approve", headers=STUDENT, json={"baseVersion": saved["planVersion"]})
    assert refused.status_code == 400 and refused.json()["code"] == "PLAN_INCOMPLETE"


def test_chapters_need_an_approved_plan_and_chapter_three_needs_the_students_figures(client):
    project = _create(client)
    _run(client, project["id"], "PLAN")
    early = client.post(f"/api/projects/{project['id']}/steps", headers=STUDENT, json={"step": "CHAPTER_1"})
    assert early.status_code == 400 and early.json()["code"] == "PLAN_NOT_APPROVED"
    _approved(client, project["id"])
    blocked = client.post(f"/api/projects/{project['id']}/steps", headers=STUDENT, json={"step": "CHAPTER_3"})
    assert blocked.status_code == 400 and blocked.json()["code"] == "AUTHOR_INPUT_NEEDED"
    size = {**client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()["plan"]["sampleSize"], "population": 2400, "populationSource": "Mukono DHO records, 2025"}
    _approved(client, project["id"], sampleSize=size)
    job = _run(client, project["id"], "CHAPTER_3")
    assert job["status"] == "COMPLETED", job
    chapter = client.get(f"/api/projects/{project['id']}/chapters/3", headers=STUDENT).json()
    assert [s["number"] for s in chapter["sections"]][:2] == ["3.0", "3.1"]
    figures = next(r for r in chapter["readiness"] if r["id"] == "C3-FIGURES")
    assert figures["basis"] == "AUTHOR" and "343" in figures["note"]
    workplan = next(s for s in chapter["sections"] if s["heading"] == "Work Plan")
    assert workplan["table"] == [["Activity", "Months"], ["Data collection", "Month 3"]]


def test_a_chapter_cites_only_confirmed_evidence_and_shows_what_needs_review_after_a_change(client):
    project = _create(client)
    _run(client, project["id"], "PLAN")
    _approved(client, project["id"])
    job = _run(client, project["id"], "CHAPTER_2")
    assert job["status"] == "COMPLETED" and job["outcome"] == "FULL", job
    chapter = client.get(f"/api/projects/{project['id']}/chapters/2", headers=STUDENT).json()
    text = " ".join(p for s in chapter["sections"] for p in s["paragraphs"])
    # APA 6 from Crossref's details. Chapter One (started on approval) cited the work first, so here it is "et al."
    assert "(Okello et al., 2022)" in text and "⟦" not in text
    assert chapter["references"] == ["Okello, J., Namara, A., & Kato, P. (2022). Vaccine Uptake Among Caregivers in Central Uganda. Malaria Journal, 21(3), 1–10. doi:10.1186/s12936-022-0001"]
    refs = next(r for r in chapter["readiness"] if r["id"] == "C2-REFS30")
    assert refs["basis"] == "CODE" and refs["status"] == "NEEDS_REVIEW"
    order = client.models.tasks
    assert order.index("p_brief") < order.index("p_draft") < order.index("p_review") < order.index("p_readiness")
    # the student changes objective 2 after the chapter was written
    changed = client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()["plan"]
    changed["specificObjectives"][1] = "To examine caregivers' trust in the vaccine"
    _approved(client, project["id"], specificObjectives=changed["specificObjectives"])
    project = client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()
    review = next(c for c in project["chapters"] if c["number"] == 2)["needsReview"]
    # the introduction, the review of objective 2 and the gap depend on objective 2; theory and objectives 1 and 3 do not
    assert review == ["2.0 Introduction", "2.3 Caregivers' attitudes towards the vaccine", "2.5 Summary and Research Gap"]


def test_invented_figures_and_typed_citations_never_reach_the_chapter(client):
    project = _create(client)
    _run(client, project["id"], "PLAN")
    _approved(client, project["id"])

    def draft(payload):
        out = []
        for s in payload["sections"]:
            out.append({"key": s["key"], "paragraphs": ["Malaria kills 73% of children in the district. As Smith (2019) showed, it matters.", "The study will follow the plan."], "table": {"caption": "", "rows": []}})
        return {"sections": out}

    client.models.overrides["p_draft"] = draft  # the fixes return the same text, so the rule must remove it
    job = _run(client, project["id"], "CHAPTER_1")
    assert job["status"] == "COMPLETED" and job["outcome"] == "PARTIAL"
    assert any("removed sentences" in w for w in job["warnings"])
    chapter = client.get(f"/api/projects/{project['id']}/chapters/1", headers=STUDENT).json()
    text = " ".join(p for s in chapter["sections"] for p in s["paragraphs"])
    assert "73%" not in text and "Smith" not in text and "The study will follow the plan." in text
    trace = next(r for r in chapter["readiness"] if r["id"] == "C1-TRACE")
    assert trace["status"] == "NEEDS_REVIEW"
    # the removal happened before the final approval: both reviewers approved the delivered wording
    for task in ("p_review", "p_review_peer"):
        last = [r for t, r in zip(client.models.tasks, client.models.requests, strict=True) if t == task][-1]
        assert "73%" not in last and "Smith" not in last and "The study will follow the plan." in last


def test_web_sources_get_their_year_and_pmc_articles_their_registered_details(client):
    """Found in a real-model run: "8 December 2025" gave n.d., and a PubMed Central article was
    cited as a web page with its journal as the author."""
    from tests.fake_models import SOURCE_PAGE, WORK

    pmc = "https://pmc.ncbi.nlm.nih.gov/articles/PMC1234567/"
    client.models.pages[pmc] = SOURCE_PAGE
    client.models.dois[pmc] = WORK["doi"]
    client.models.works = []  # no scholarly results: both needs go to the web search
    found = [
        {"url": "https://stats.example.org/report-2022", "title": "Annual report", "publisher": "Statistics office", "published": "8 December 2025", "access": "FULL_TEXT", "statement": "s", "passage": "The report gives the figure for 2022.", "scope": ""},
        {"url": pmc, "title": "Uptake", "publisher": "BMC Public Health", "published": "", "access": "FULL_TEXT", "statement": "t", "passage": "It covers every district in the country.", "scope": ""},
    ]
    client.models.opened = [f["url"] for f in found]
    client.models.pages[pmc] = "<p>It covers every district in the country.</p>"
    client.models.overrides["p_search"] = lambda payload: {"findings": found}
    project = _create(client)
    _run(client, project["id"], "PLAN")
    library = {i["source"]["url"]: i["source"] for i in client.get(f"/api/projects/{project['id']}/evidence", headers=STUDENT).json()}
    assert library["https://stats.example.org/report-2022"]["year"] == "2025"
    article = library[pmc]
    assert article["metadata"] == "CROSSREF" and article["authors"][0] == "Okello, J." and article["organisation"] == ""


def test_searches_never_carry_the_students_name(client):
    project = _create(client)
    client.models.overrides["p_needs"] = lambda payload: {"needs": [{"id": "n1", "need": "Uptake", "kind": "LITERATURE", "query": "Namukasa malaria vaccine"}]}
    _run(client, project["id"], "PLAN")
    assert client.models.searched == []  # the query named the student, so it was never sent
    requests = [r for r in client.models.requests if r.startswith("p_")]
    assert not any("Namukasa" in r or "M24/U001" in r for r in requests if not r.startswith("p_needs"))
    assert not any("M24/U001" in r for r in requests)


def test_a_step_priced_on_an_old_plan_must_be_priced_again(client):
    project = _create(client)
    _run(client, project["id"], "PLAN")
    before = client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()
    quoted = client.post(f"/api/projects/{project['id']}/steps", headers=STUDENT, json={"step": "PLAN"}).json()
    client.post(f"/api/projects/{project['id']}/plan", headers=STUDENT, json={"plan": {**before["plan"], "title": "My own title"}, "baseVersion": before["planVersion"]})
    stale = client.post(f"/api/projects/{project['id']}/steps/{quoted['job']['id']}/submit", headers=STUDENT, json={"quoteId": quoted["quote"]["id"]})
    assert stale.status_code == 409 and stale.json()["code"] == "QUOTE_MISMATCH"


def test_a_plan_made_while_the_student_edited_theirs_is_kept_as_a_candidate(client):
    project = _create(client)
    _run(client, project["id"], "PLAN")

    def student_edits_meanwhile(payload):
        current = client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()
        edit = client.post(f"/api/projects/{project['id']}/plan", headers=STUDENT, json={"plan": {**current["plan"], "title": "My own title"}, "baseVersion": current["planVersion"]})
        assert edit.status_code == 200
        return payload["draft"]

    client.models.overrides["p_finalise"] = student_edits_meanwhile
    job = _run(client, project["id"], "PLAN")
    assert job["status"] == "COMPLETED" and any("kept alongside yours" in w for w in job["warnings"])
    current = client.get(f"/api/projects/{project['id']}", headers=STUDENT).json()
    assert current["plan"]["title"] == "My own title" and current["candidatePlan"]["title"] == PLAN["title"]
    taken = client.post(f"/api/projects/{project['id']}/plan/candidate", headers=STUDENT, json={"accept": True}).json()
    assert taken["plan"]["title"] == PLAN["title"] and taken["candidatePlan"] is None and taken["planStatus"] == "DRAFT"


def test_the_sample_size_preview_is_calculated_by_the_server(client):
    project = _create(client)
    preview = client.post(f"/api/projects/{project['id']}/sample-size", headers=STUDENT, json={"method": "YAMANE", "population": 2400}).json()
    assert preview["size"] == 343 and "Yamane" in preview["steps"] and preview["missing"] == ""
    missing = client.post(f"/api/projects/{project['id']}/sample-size", headers=STUDENT, json={"method": "YAMANE"}).json()
    assert missing["size"] is None and "accessible population" in missing["missing"]


def test_a_step_is_submitted_only_through_its_project(client):
    project = _create(client)
    quoted = client.post(f"/api/projects/{project['id']}/steps", headers=STUDENT, json={"step": "PLAN"}).json()
    direct = client.post(f"/api/jobs/{quoted['job']['id']}/submit", headers=STUDENT, json={"quoteId": quoted["quote"]["id"]})
    assert direct.status_code == 400 and direct.json()["code"] == "PROPOSAL_STEP"
    assert client.get("/api/jobs?status=DRAFT", headers=STUDENT).json()["items"] == []


def test_other_users_cannot_reach_a_project(client):
    project = _create(client)
    for response in (
        client.get(f"/api/projects/{project['id']}", headers=OTHER),
        client.post(f"/api/projects/{project['id']}/steps", headers=OTHER, json={"step": "PLAN"}),
        client.get(f"/api/projects/{project['id']}/evidence", headers=OTHER),
        client.get(f"/api/projects/{project['id']}/export", headers=OTHER),
    ):
        assert response.status_code == 404
    assert client.get("/api/projects", headers=OTHER).json() == []


def test_the_export_is_the_ucu_layout_and_the_final_version_needs_everything(client):
    project = _create(client)
    _run(client, project["id"], "PLAN")
    _approved(client, project["id"])
    _run(client, project["id"], "CHAPTER_1")
    final = client.get(f"/api/projects/{project['id']}/export?final=true", headers=STUDENT)
    assert final.status_code == 400 and "Write Chapter 2" in final.json()["message"]
    draft = client.get(f"/api/projects/{project['id']}/export", headers=STUDENT)
    assert draft.status_code == 200 and draft.content[:2] == b"PK"
    doc = Document(io.BytesIO(draft.content))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == PLAN["title"].upper() and "Grace Namukasa" in texts
    assert "CHAPTER ONE" in texts and "REFERENCES" in texts
    assert doc.styles["Normal"].font.name == "Trebuchet MS" and doc.styles["Normal"].paragraph_format.line_spacing == 2.0
    assert any("not yet complete for submission" in t for t in texts)
    assert not any("⟦" in t for t in texts)


def test_deleting_a_project_or_the_account_removes_its_chapters_and_evidence(client):
    from app.runtime import get_runtime

    rt = get_runtime()
    project = _create(client)
    _run(client, project["id"], "PLAN")
    stored = rt.store.get_project(project["id"])
    assert stored and stored.evidence_files and rt.files.exists(stored.evidence_files[0])
    assert client.delete(f"/api/projects/{project['id']}", headers=STUDENT).status_code == 204
    assert rt.store.get_project(project["id"]) is None and not rt.files.exists(stored.evidence_files[0])
    fresh = {"Authorization": "Dev leaving@example.com"}  # no credit, so the account may be deleted
    second = _create(client, fresh)
    assert client.delete("/api/me", headers=fresh).status_code == 204
    assert rt.store.get_project(second["id"]) is None
    assert client.post("/api/projects", headers=fresh, json=DETAILS).status_code == 409  # the closed account stays closed


def test_projects_expire_thirty_days_after_the_students_last_action(client):
    from app.jobs import service
    from app.runtime import get_runtime

    rt = get_runtime()
    project = _create(client)
    assert project["expiresAt"]

    def age(p):
        p.expires_at = p.expires_at - timedelta(days=31)
        return p

    rt.store.update_project(project["id"], age)
    client.get(f"/api/projects/{project['id']}", headers=STUDENT)  # viewing is not an action: no renewal
    assert service.cleanup_expired_projects(rt) == 1
    assert rt.store.get_project(project["id"]) is None


# --- review of an uploaded proposal ----------------------------------------------------------


def _proposal_docx(methods_past: bool = True) -> bytes:
    doc = Document()
    doc.add_heading("Determinants of vaccine uptake in Mukono", 0)
    doc.add_heading("CHAPTER ONE: INTRODUCTION", 1)
    for heading, text in (
        ("Background to the Study", "Malaria remains a major cause of illness among children under five in Uganda and the region."),
        ("Statement of the Problem", "Uptake of the vaccine remains below the national target in the district."),
        ("Specific Objectives", ""),
    ):
        doc.add_heading(heading, 2)
        if text:
            doc.add_paragraph(text * 3)
    for objective in ("To assess the effect of distance on uptake", "To examine caregivers' attitudes"):
        doc.add_paragraph(objective, style="List Bullet")
    doc.add_heading("Research Questions", 2)
    doc.add_paragraph("How does distance affect uptake?", style="List Bullet")
    doc.add_heading("CHAPTER TWO: LITERATURE REVIEW", 1)
    doc.add_paragraph("Studies in East Africa report that distance limits uptake (Okello, 2022). " * 4)
    doc.add_heading("CHAPTER THREE: METHODOLOGY", 1)
    doc.add_heading("Research Design", 2)
    doc.add_paragraph(("Data were collected using a structured questionnaire. " if methods_past else "Data will be collected using a questionnaire. ") * 3)
    doc.add_heading("References", 1)
    doc.add_paragraph("Okello, J. (2022). Vaccine uptake in Uganda. Malaria Journal, 21, 1-10.")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_an_uploaded_proposal_is_reviewed_against_the_rulebook_without_changing_it(client):
    job_id = client.post("/api/jobs", headers=STUDENT).json()["id"]
    upload = client.post(f"/api/jobs/{job_id}/files/source", headers=STUDENT, files={"file": ("proposal.docx", _proposal_docx(), "application/octet-stream")})
    assert upload.status_code == 200, upload.json()
    quote = client.post(f"/api/jobs/{job_id}/quote", headers=STUDENT, json={"selection": {"proposal": "REVIEW", "level": "MASTERS"}}).json()["quote"]
    assert client.post(f"/api/jobs/{job_id}/submit", headers=STUDENT, json={"quoteId": quote["id"]}).status_code == 200
    job = wait(client, job_id)
    assert job["status"] == "COMPLETED", job
    items = {i["id"]: i for i in job["proposalReview"]["items"]}
    assert items["S-background"]["status"] == "PASS" and items["S-ethics"]["status"] == "MISSING"
    assert items["N-alignment"]["status"] == "NEEDS_REVIEW"  # two objectives, one question
    assert items["N-tense"]["status"] == "NEEDS_REVIEW" and "Data were collected" in items["N-tense"]["note"]
    assert items["C1-BROAD"]["basis"] == "AI" and items["N-references"]["basis"] == "CODE"
    assert job["proposalReview"]["findings"][0]["kind"] == "ALIGNMENT"
    assert [o["id"] for o in job["outputs"]] == ["proposal-review"]
    report = Document(io.BytesIO(client.get(f"/api/jobs/{job_id}/outputs/proposal-review", headers=STUDENT).content))
    assert any("Proposal review" in p.text for p in report.paragraphs)


def test_a_review_runs_alone(client):
    job_id = client.post("/api/jobs", headers=STUDENT).json()["id"]
    client.post(f"/api/jobs/{job_id}/files/source", headers=STUDENT, files={"file": ("proposal.docx", _proposal_docx(False), "application/octet-stream")})
    mixed = client.post(f"/api/jobs/{job_id}/quote", headers=STUDENT, json={"selection": {"proposal": "REVIEW", "writing": "AI_CHECK"}})
    assert mixed.status_code == 400 and mixed.json()["code"] == "REVIEW_ALONE"
    step = client.post(f"/api/jobs/{job_id}/quote", headers=STUDENT, json={"selection": {"proposal": "PLAN"}})
    assert step.status_code == 400 and step.json()["code"] == "PROPOSAL_STEP"
