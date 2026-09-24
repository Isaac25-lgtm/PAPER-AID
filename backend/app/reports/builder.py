"""Downloadable Word reports: the writing report (AI Check) and the change report (refinement)."""

import io
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor

from app.jobs.models import AnalysisResult, RefinementResult

GREEN = RGBColor(0x0F, 0x63, 0x3E)
MUTED = RGBColor(0x46, 0x55, 0x4D)
DISCLAIMER = (
    "This is an estimate of formulaic writing patterns produced by PaperAid. It is not a detector verdict and not proof of "
    "authorship: AI detectors often disagree with each other and can flag human writing. PaperAid is not affiliated with "
    "Turnitin or any other detection service and cannot predict their results."
)
REASON_LABELS = {
    "GENERIC_PHRASING": "Generic phrasing",
    "UNIFORM_STRUCTURE": "Repetitive structure",
    "LOW_SPECIFICITY": "Vague claim",
    "FORMULAIC_TRANSITIONS": "Formulaic transition",
    "OVER_HEDGING": "Over-hedging",
    "UNSUPPORTED_SUMMARY": "Unsupported summary",
}


def _document(title: str, paper_name: str, when: datetime):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    brand = doc.add_paragraph()
    run = brand.add_run("PaperAid")
    run.bold, run.font.color.rgb, run.font.size = True, GREEN, Pt(14)
    doc.add_heading(title, level=1)
    meta = doc.add_paragraph()
    meta_run = meta.add_run(f"{paper_name} · {when:%d %B %Y, %H:%M} UTC")
    meta_run.font.color.rgb = MUTED
    return doc


def _note(doc, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic, r.font.size, r.font.color.rgb = True, Pt(9.5), MUTED


def _save(doc) -> bytes:
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def writing_report(paper_name: str, when: datetime, analysis: AnalysisResult, after: AnalysisResult | None) -> bytes:
    doc = _document("Writing report", paper_name, when)
    doc.add_heading("Estimated AI-likeness", level=2)
    summary = doc.add_paragraph()
    summary.add_run(f"{analysis.band.capitalize()}").bold = True
    summary.add_run(f" (confidence: {analysis.confidence.lower()})")
    if after:
        summary.add_run(" before refinement; ")
        summary.add_run(f"{after.band.capitalize()}").bold = True
        summary.add_run(" after refinement.")
    doc.add_paragraph(
        f"Based on {analysis.analysed_words:,} analysed words. References, quotations, tables and headings "
        f"({analysis.excluded_words:,} words) were excluded. Method: {analysis.method}; algorithm {analysis.algorithm_version}."
    )
    _note(doc, DISCLAIMER)

    doc.add_heading(f"Findings ({len(analysis.findings)})", level=2)
    if not analysis.findings:
        doc.add_paragraph("No formulaic writing patterns were flagged.")
    for finding in analysis.findings:
        heading = doc.add_paragraph()
        label = heading.add_run(REASON_LABELS.get(finding.reason, finding.reason))
        label.bold = True
        heading.add_run(f"  ·  {finding.section} · {finding.severity}").font.color.rgb = MUTED
        quote = doc.add_paragraph(style="Intense Quote" if "Intense Quote" in [s.name for s in doc.styles] else None)
        quote.add_run(finding.excerpt).italic = True
        doc.add_paragraph(finding.explanation)
        tip = doc.add_paragraph()
        tip.add_run("Suggestion: ").bold = True
        tip.add_run(finding.suggestion)
    return _save(doc)


def change_report(paper_name: str, when: datetime, refinement: RefinementResult) -> bytes:
    doc = _document("Change report", paper_name, when)
    doc.add_paragraph(
        f"{refinement.refined_blocks} passages refined · {refinement.kept_original} kept in your original wording · "
        f"{refinement.untouched_blocks} paragraphs untouched."
    )
    doc.add_paragraph(
        "Citations, quotations, numbers, links and footnotes were locked during editing and checked afterwards. "
        "Every change passed an accuracy check before it was applied to your document."
    )
    _note(doc, f"Refinement method: {refinement.method}.")
    for change in refinement.changes:
        doc.add_heading(f"{change.section or 'Body'} — {'kept original' if change.kept else 'refined'}", level=3)
        label = doc.add_paragraph()
        label.add_run("Your original").bold = True
        doc.add_paragraph(change.before)
        label = doc.add_paragraph()
        label.add_run("Proposed (not applied)" if change.kept else "Refined").bold = True
        revised = doc.add_paragraph(change.after)
        if change.kept:
            for run in revised.runs:
                run.font.strike = True
        if change.reason:
            _note(doc, f"Why: {change.reason}")
        if change.note:
            _note(doc, change.note)
    _note(doc, "Review every change before you submit. The final paper is your responsibility.")
    return _save(doc)
