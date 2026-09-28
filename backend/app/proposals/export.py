"""The proposal as a Word document in the UCU layout (manual §2.2, pp. 9-10): Trebuchet MS 12,
double spacing, one-inch margins, page numbers at the bottom centre; preliminary pages numbered in
Roman numerals, the chapters in Arabic; a reference list built from the evidence actually cited.

Nothing is invented to fill a gap. A draft export leaves out what is missing (a title-page detail,
an unwritten chapter) and says so on its cover line; a final export requires everything
(`final_blockers`). The approval page is a form for the supervisor to sign; PaperAid never signs or
claims approval. Word updates the contents table's page numbers when the file is opened."""

import io
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt

from app.formatting.apply import _page_field_paragraph, _set_font, _set_page_numbering
from app.proposals import decisions, evidence, rulebook
from app.proposals.models import ChapterDocument, EvidenceItem, Project

ORDINALS = {1: "ONE", 2: "TWO", 3: "THREE"}


def final_blockers(project: Project, chapters: dict[int, ChapterDocument]) -> list[str]:
    """What must be resolved before a complete (submission) export."""
    problems = []
    page = project.title_page
    for label, value in (("your name", page.student_name), ("your registration number", page.reg_number), ("your supervisor's name", page.supervisor), ("the submission date", page.submission_date)):
        if not value.strip():
            problems.append(f"Add {label} to the title page details.")
    if project.plan is None or project.plan_status != "APPROVED":
        problems.append("Approve your plan.")
    for n in (1, 2, 3):
        doc = chapters.get(n)
        if doc is None:
            problems.append(f"Write Chapter {n}.")
            continue
        if not project.chapter(n).approved:
            problems.append(f"Approve Chapter {n}.")
        if project.plan is None:
            continue
        if decisions.stale(doc, project.plan):
            problems.append(f"Chapter {n} has sections written from decisions you have since changed; revise them.")
        # Structure and integrity block a complete export (Codex audit 2026-09-28 #8); the AI's
        # judgements and recommendations (such as 30 references) stay advisory.
        written = {s.key for s in doc.sections if any(p.strip() for p in s.paragraphs)}
        expected = rulebook.sections(project.rulebook, n, project.inputs.level, project.plan)
        missing = [f"{s.number} {s.heading}" for s in expected if s.key not in written]
        if missing:
            problems.append(f"Chapter {n} is missing required sections: {', '.join(missing)}.")
        for item in doc.readiness:
            blocking = item.basis == "CODE" and (item.status in ("MISSING", "BLOCKED") or item.id.endswith("-REVIEWED") and item.status != "PASS")
            if blocking:
                problems.append(f"Chapter {n}: {item.question} ({item.note})")
    return problems


def _setup(doc) -> None:
    book = rulebook.load(rulebook.DEFAULT)["formatting"]
    normal = doc.styles["Normal"]
    _set_font(normal, book["font"], book["size_pt"])
    normal.paragraph_format.line_spacing = book["line_spacing"]
    normal.paragraph_format.space_after = Pt(0)
    for level, size in ((1, 14), (2, 12), (3, 12)):
        style = doc.styles[f"Heading {level}"]
        _set_font(style, book["font"], size)
        style.font.bold, style.font.italic = True, False
        style.font.color.rgb = None
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(12), Pt(6)
        style.paragraph_format.line_spacing = book["line_spacing"]
        if level == 1:
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    section = doc.sections[0]
    margin = Inches(book["margins_in"])
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = margin
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        settings.append(parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>'))


def _centered(doc, text: str, bold: bool = False, caps: bool = False, space_after: int = 0) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text.upper() if caps else text)
    run.bold = bold


def _title_page(doc, project: Project, draft: bool) -> None:
    plan, page, inputs = project.plan, project.title_page, project.inputs
    book = rulebook.load(project.rulebook)
    level = book["levels"][inputs.level]["label"]
    _centered(doc, (plan.title if plan else inputs.topic), bold=True, caps=True, space_after=24)
    for line in (page.student_name, page.reg_number):
        if line.strip():
            _centered(doc, line)
    faculty = f" TO THE {inputs.faculty.upper()}" if inputs.faculty.strip() else ""
    programme = inputs.programme.upper() if inputs.programme.strip() else level.upper()
    _centered(doc, "")
    _centered(doc, f"A RESEARCH PROPOSAL SUBMITTED{faculty} IN PARTIAL FULFILMENT OF THE REQUIREMENTS FOR THE AWARD OF THE {programme} OF UGANDA CHRISTIAN UNIVERSITY")
    _centered(doc, "")
    if page.supervisor.strip():
        _centered(doc, f"Supervisor: {page.supervisor}")
    if page.submission_date.strip():
        _centered(doc, page.submission_date)
    if draft:
        _centered(doc, "")
        _centered(doc, f"Draft prepared with PaperAid, {date.today():%d %B %Y}: not yet complete for submission.")


def _form_page(doc, heading: str, lines: list[str]) -> None:
    doc.add_heading(heading, level=1)
    for line in lines:
        doc.add_paragraph(line)


TOC_RUNS = (
    '<w:fldChar w:fldCharType="begin" w:dirty="true"/>',
    '<w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText>',
    '<w:fldChar w:fldCharType="separate"/>',
    "<w:t>Word fills in this table when the document is opened (or right-click it and choose Update Field).</w:t>",
    '<w:fldChar w:fldCharType="end"/>',
)


def _contents(doc) -> None:
    doc.add_heading("Table of Contents", level=1)
    paragraph = doc.add_paragraph()._p
    for inner in TOC_RUNS:
        paragraph.append(parse_xml(f'<w:r {nsdecls("w")}>{inner}</w:r>'))


def _footer_numbers(section, fmt: str) -> None:
    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    for run in list(footer.runs):
        run._r.getparent().remove(run._r)
    _page_field_paragraph(footer, WD_ALIGN_PARAGRAPH.CENTER)
    _set_page_numbering(section._sectPr, fmt)


def build(project: Project, chapters: dict[int, ChapterDocument], library: dict[str, EvidenceItem], draft: bool) -> bytes:
    doc = Document()
    _setup(doc)
    _title_page(doc, project, draft)

    prelim = doc.add_section(WD_SECTION.NEW_PAGE)
    name = project.title_page.student_name.strip()
    _form_page(
        doc,
        "Declaration",
        [
            (f"I, {name}, " if name else "I ") + "declare that this research proposal is my original work and has not been submitted to any other institution for any award.",
            "",
            "Signature: ______________________________        Date: ____________________",
        ],
    )
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    supervisor = project.title_page.supervisor.strip()
    _form_page(
        doc,
        "Approval",
        [
            "This research proposal has been submitted with my approval as the University supervisor.",
            "",
            f"Name: {supervisor}" if supervisor else "Name: ______________________________",
            "Signature: ______________________________        Date: ____________________",
        ],
    )
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _contents(doc)
    tables = [(n, s) for n, c in sorted(chapters.items()) for s in c.sections if s.table]
    if tables:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading("List of Tables", level=1)
        for i, (n, s) in enumerate(tables, start=1):
            doc.add_paragraph(f"Table {n}.{i}: {s.table_caption or s.heading}")
    _footer_numbers(prelim, "lowerRoman")

    citer = evidence.Citer(library, project.citation)
    cited: list[str] = []
    table_count = 0
    for n in (1, 2, 3):
        chapter = chapters.get(n)
        if chapter is None:
            continue
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        if not any(c < n for c in chapters):
            _footer_numbers(section, "decimal")
        doc.add_heading(f"CHAPTER {ORDINALS[n]}", level=1)
        doc.add_heading(chapter.title.upper(), level=1)
        for s in chapter.sections:
            doc.add_heading(f"{s.number} {s.heading}", level=2)
            for paragraph in s.paragraphs:
                cited += [i for i in evidence.cited_ids(paragraph) if i not in cited]
                p = doc.add_paragraph(citer.render(paragraph))
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if s.table:
                table_count += 1
                for field in [s.table_caption, *[c for row in s.table for c in row]]:
                    cited += [i for i in evidence.cited_ids(field) if i not in cited]
                caption = doc.add_paragraph()
                caption.add_run(f"Table {n}.{table_count}: {citer.render(s.table_caption) or s.heading}").bold = True
                grid = doc.add_table(rows=len(s.table), cols=len(s.table[0]))
                grid.style = "Table Grid"
                for r, row in enumerate(s.table):
                    for c, value in enumerate(row):
                        cell = grid.cell(r, c)
                        cell.text = citer.render(value)
                        if r == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
    if cited:
        doc.add_section(WD_SECTION.NEW_PAGE)
        doc.add_heading("REFERENCES", level=1)
        for entry in evidence.reference_list([library[i].source for i in cited if i in library], project.citation):
            p = doc.add_paragraph(entry)
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
