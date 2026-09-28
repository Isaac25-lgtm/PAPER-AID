"""Supervisor feedback (Proposal V2, master context §47). Code reads the comments and suggests
where each applies; the student confirms; a revision step then fixes exactly those sections. The
response report lists every comment with what was done about it. No AI runs here."""

import io
import re
from dataclasses import dataclass
from pathlib import PurePosixPath

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.errors import InvalidDocument
from app.documents.intake import OLE_MAGIC, ZIP_MAGIC, _check_docx_archive

MAX_COMMENTS = 100  # per project
MAX_TEXT = 1500
MAX_FILE_BYTES = 15 * 1024 * 1024
BULLET = re.compile(r"^\s*(?:[-•*▪◦]|\(?\d{1,3}[.)]|\(?[a-zA-Z][.)]|[ivxIVX]{1,5}[.)])\s+")
SECTION_NUMBER = re.compile(r"(?<![\d.])([123])\.(\d{1,2})(?![\d])")
HEADING_STYLE = re.compile(r"^(heading|title)", re.I)
NUMBERED_HEADING = re.compile(r"^(?:chapter\s+\w+|\d+(?:\.\d+){0,2})\s+\S")

# Words a comment uses for a section, when it names neither a number nor a heading.
KEYWORDS: list[tuple[re.Pattern[str], int, str]] = [
    (re.compile(r"statement of (?:the )?problem|problem statement", re.I), 1, "problem"),
    (re.compile(r"\bbackground\b", re.I), 1, "background"),
    (re.compile(r"\bpurpose\b|general objective", re.I), 1, "purpose"),
    (re.compile(r"specific objectives?|\bobjectives?\b", re.I), 1, "objectives"),
    (re.compile(r"research questions?|hypothes[ie]s|propositions?", re.I), 1, "questions"),
    (re.compile(r"\bscope\b", re.I), 1, "scope"),
    (re.compile(r"\bjustification\b", re.I), 1, "justification"),
    (re.compile(r"\bsignificance\b", re.I), 1, "significance"),
    (re.compile(r"conceptual framework", re.I), 1, "framework"),
    (re.compile(r"theor(?:y|etical)", re.I), 2, "theory"),
    (re.compile(r"empirical|literature review|\bliterature\b", re.I), 2, "empirical"),
    (re.compile(r"research gap|\bgap\b", re.I), 2, "gap"),
    (re.compile(r"research design|\bdesign\b", re.I), 3, "design"),
    (re.compile(r"study area|area of study", re.I), 3, "area"),
    (re.compile(r"sample size|sampling", re.I), 3, "sampling"),
    (re.compile(r"study population|target population|\bpopulation\b", re.I), 3, "population"),
    (re.compile(r"questionnaire|interview guide|\binstruments?\b", re.I), 3, "instruments"),
    (re.compile(r"validity|reliability|quality control|pilot", re.I), 3, "quality"),
    (re.compile(r"data analysis|\banalysis\b|\bSPSS\b|\bSTATA\b|regression", re.I), 3, "analysis"),
    (re.compile(r"\bethic", re.I), 3, "ethics"),
    (re.compile(r"work ?plan|time ?line|\bbudget\b", re.I), 3, "workplan"),
]


@dataclass(frozen=True)
class Written:
    """A section the student has: where a comment can be applied."""

    chapter: int
    key: str
    number: str
    heading: str


@dataclass(frozen=True)
class Read:
    text: str
    anchor: str = ""


def _clean(text: str) -> str:
    return " ".join(text.split())


def split(text: str) -> list[Read]:
    """Pasted feedback as comments: one per numbered or bulleted item, otherwise one per paragraph.
    A line that continues an item (no bullet, no blank line before it) stays with it."""
    comments: list[str] = []
    current: list[str] = []
    listed = any(BULLET.match(line) for line in text.splitlines())

    def close() -> None:
        if current:
            comments.append(_clean(" ".join(current)))
            current.clear()

    for line in text.splitlines():
        if not line.strip():
            close()
        elif BULLET.match(line):
            close()
            current.append(BULLET.sub("", line, count=1))
        elif listed and not current and comments:
            current.append(line)  # a heading or lead-in line between items
        else:
            current.append(line)
    close()
    return [Read(c[:MAX_TEXT]) for c in comments if len(c) >= 3]


def _style_name(paragraph, styles: dict[str, str]) -> str:
    ppr = paragraph.find(qn("w:pPr"))
    style = ppr.find(qn("w:pStyle")) if ppr is not None else None
    return styles.get(style.get(qn("w:val")), "") if style is not None else ""


def _text(element) -> str:
    return _clean("".join(t.text or "" for t in element.iter(qn("w:t"))))


def from_docx(data: bytes) -> list[Read]:
    """A marked-up Word file: each Word comment with the heading above the passage it is attached
    to. A file without comments is read as pasted feedback (its paragraphs, in order)."""
    if not data.startswith(ZIP_MAGIC):
        raise InvalidDocument("This isn't a valid Word document. Open it in Word and save it as .docx again.", code="INVALID_CONTAINER")
    _check_docx_archive(data)
    doc = Document(io.BytesIO(data))
    styles = {s.style_id: s.name for s in doc.styles}
    notes: dict[str, str] = {}
    for rel in doc.part.rels.values():
        if rel.reltype.endswith("/comments") and not rel.is_external:
            root = etree.fromstring(rel.target_part.blob, etree.XMLParser(resolve_entities=False, no_network=True))
            for comment in root.iter(qn("w:comment")):
                notes[comment.get(qn("w:id"))] = _text(comment)
    if not notes:
        return split("\n\n".join(p.text for p in doc.paragraphs))
    found: list[Read] = []
    heading = ""
    for paragraph in doc.element.body.iter(qn("w:p")):
        text = _text(paragraph)
        if text and (HEADING_STYLE.match(_style_name(paragraph, styles)) or (len(text) < 90 and NUMBERED_HEADING.match(text))):
            heading = text
        ids = [e.get(qn("w:id")) for e in paragraph.iter(qn("w:commentRangeStart"), qn("w:commentReference"))]
        for cid in dict.fromkeys(ids):
            if cid in notes and notes[cid]:
                found.append(Read(notes.pop(cid)[:MAX_TEXT], anchor=(heading or text)[:300]))
    found += [Read(t[:MAX_TEXT]) for t in notes.values() if t]  # comments whose anchor was not found
    return found


def from_pdf(data: bytes) -> list[Read]:
    """A PDF: its comment annotations (sticky notes, highlights with notes); without any, its text."""
    if b"%PDF-" not in data[:1024]:
        raise InvalidDocument("This isn't a valid PDF file.", code="INVALID_CONTAINER")
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise InvalidDocument("This PDF is password-protected. Remove the password and upload it again.", code="ENCRYPTED")
        found: list[Read] = []
        text: list[str] = []
        for number, page in enumerate(reader.pages[:200], start=1):
            for annot in page.get("/Annots") or []:
                note = annot.get_object().get("/Contents")
                if isinstance(note, str) and note.strip():
                    found.append(Read(_clean(note)[:MAX_TEXT], anchor=f"Page {number}"))
            if len(text) < 60:
                text.append(page.extract_text() or "")
    except PdfReadError as exc:
        raise InvalidDocument("This PDF could not be read. Save it again and upload it.", code="INVALID_CONTAINER") from exc
    return found or split("\n".join(text))


def read_file(filename: str, data: bytes) -> list[Read]:
    ext = PurePosixPath(filename.lower()).suffix
    if not data:
        raise InvalidDocument("This file is empty.", code="EMPTY_FILE")
    if len(data) > MAX_FILE_BYTES:
        raise InvalidDocument("This file is larger than 15 MB.", code="FILE_TOO_LARGE")
    if data.startswith(OLE_MAGIC):
        raise InvalidDocument("This file is password-protected or in the old .doc format. Save it as .docx.", code="ENCRYPTED")
    if ext == ".docx":
        return from_docx(data)
    if ext == ".pdf":
        return from_pdf(data)
    raise InvalidDocument("Upload the feedback as a Word document (.docx) or a PDF, or paste it.", code="UNSUPPORTED_TYPE")


def _norm(text: str) -> str:
    return re.sub(r"[^a-z ]", "", text.lower()).strip()


def suggest(comment: Read, written: list[Written]) -> tuple[int | None, list[str]]:
    """Where a comment probably applies: a section number it names, a heading it names or sits
    under, or a word it uses for one. Only sections the student has written are suggested."""
    if not written:
        return None, []
    by_number = {w.number: w for w in written}
    for source in (comment.anchor, comment.text):
        numbers = [f"{a}.{b}" for a, b in SECTION_NUMBER.findall(source)]
        hits = [by_number[n] for n in numbers if n in by_number]
        if hits:
            chapter = hits[0].chapter
            return chapter, list(dict.fromkeys(h.key for h in hits if h.chapter == chapter))[:3]
    for source in (comment.anchor, comment.text):
        text = _norm(source)
        hits = [w for w in written if len(_norm(w.heading)) > 4 and _norm(w.heading) in text]
        if hits:
            best = max(hits, key=lambda w: len(w.heading))
            return best.chapter, [best.key]
    for pattern, chapter, key in KEYWORDS:
        # a per-objective section ("empirical1", "empirical2"…) answers to its family's word
        hits = [w.key for w in written if w.chapter == chapter and (w.key == key or (w.key.startswith(key) and w.key[len(key):].isdigit()))]
        if pattern.search(comment.text) and hits:
            return chapter, hits
    return None, []
