"""Review of an uploaded proposal against the institution's rulebook (Proposal V1, milestone 1).

Code checks what code can know: which required elements have a heading, how many objectives and
questions there are, planned work described in the past tense, citations against the reference
list, and the length for the level. The lead then answers the rulebook's vetting questions and
lists what a supervisor would raise, pointing at paragraphs. Nothing in the proposal is changed.

  EXTRACTING (shared) → ANALYSING (this module) → EXPORTING (report)"""

import io
import re
from datetime import datetime
from pathlib import PurePosixPath
from typing import TYPE_CHECKING

from docx import Document
from docx.shared import Pt, RGBColor

from app.analysis import paper_checks
from app.jobs.models import Job, ProposalReview, ReadinessItem, ReviewFinding, Stage, StoredOutput, utcnow
from app.proposals import evidence, rulebook
from app.proposals.ai import ProposalRunner

if TYPE_CHECKING:
    from app.documents.model import Block, DocumentModel
    from app.jobs.pipeline import StageContext

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
METHOD = "Structure, counts, tense and citations checked by PaperAid; vetting questions and findings by PaperAid's AI"

CHAPTER_STARTS = {2: re.compile(r"literature|related (?:studies|works)"), 3: re.compile(r"methodolog|research methods|materials and methods")}
REFERENCES = re.compile(r"^\s*(?:list of )?(?:references|bibliography|works cited)\b")

# (chapter, id, what the manual requires, heading pattern, study types it may not apply to)
ELEMENTS: list[tuple[int, str, str, str, tuple[str, ...]]] = [
    (1, "background", "Background to the study", r"background", ()),
    (1, "problem", "Statement of the problem", r"problem", ()),
    (1, "purpose", "Purpose or general objective", r"purpose|general objective|main objective|aim of the study", ()),
    (1, "objectives", "Specific objectives", r"objective", ()),
    (1, "questions", "Research questions, hypotheses or propositions", r"question|hypothes|proposition", ()),
    (1, "scope", "Scope of the study", r"scope", ()),
    (1, "justification", "Justification (rationale)", r"justification|rationale", ()),
    (1, "significance", "Significance of the study", r"significance|importance of the study", ()),
    (1, "framework", "Conceptual or theoretical framework", r"framework", ()),
    (2, "review", "Literature review", r"literature|review", ()),
    (3, "design", "Research design", r"design", ()),
    (3, "area", "Area of study", r"area|site|setting|location", ("non-empirical",)),
    (3, "population", "Study population", r"population", ("non-empirical",)),
    (3, "sampling", "Sample size and sampling techniques", r"sampl", ("non-empirical",)),
    (3, "instruments", "Data collection instruments", r"instrument|tool", ("non-empirical",)),
    (3, "quality", "Quality control (validity and reliability)", r"validity|reliab|quality|trustworth", ("non-empirical",)),
    (3, "analysis", "Data processing and analysis", r"analys|processing", ()),
    (3, "ethics", "Ethical considerations", r"ethic", ()),
    (3, "constraints", "Anticipated methodological constraints", r"limitation|constraint", ()),
    (3, "workplan", "Work plan or timeline", r"work ?plan|time ?line|time ?frame|schedule|gantt", ()),
]


def _headings(model: "DocumentModel") -> list[tuple[int, "Block"]]:
    """Each heading with the chapter it falls in (1 until the literature review, then 2, then 3;
    0 once the references start)."""
    chapter, out = 1, []
    for block in model.blocks:
        if block.kind not in ("heading", "title") and not block.detected_heading:
            continue
        text = block.text.lower()
        if REFERENCES.search(text):
            chapter = 0
        elif chapter and chapter < 3 and CHAPTER_STARTS[3].search(text):
            chapter = 3
        elif chapter == 1 and CHAPTER_STARTS[2].search(text):
            chapter = 2
        out.append((chapter, block))
    return out


def _chapter_of(model: "DocumentModel") -> dict[str, int]:
    """Block id → chapter, by position between chapter headings."""
    starts = {b.id: c for c, b in _headings(model)}
    chapter, out = 1, {}
    for block in model.blocks:
        chapter = starts.get(block.id, chapter)
        out[block.id] = chapter
    return out


def _listed_under(model: "DocumentModel", pattern: str) -> list[str]:
    """List items, or sentences starting "To …"/ending "?", under the first heading matching `pattern`."""
    items: list[str] = []
    inside = False
    for block in model.blocks:
        if block.kind == "heading" or block.detected_heading:
            if inside and items:
                break
            inside = bool(re.search(pattern, block.text.lower()))
            continue
        if not inside:
            continue
        text = block.text.strip()
        if block.kind == "list_item" or re.match(r"^(?:\(?[ivx\d]+[.)]\s*)?to\s", text, re.I) or text.endswith("?"):
            items.append(text)
    return items


def code_checks(model: "DocumentModel", level: str, rulebook_id: str) -> tuple[list[ReadinessItem], dict[str, int]]:
    book = rulebook.load(rulebook_id)
    headings = _headings(model)
    chapters = _chapter_of(model)
    items: list[ReadinessItem] = []
    for chapter, key, label, pattern, may_skip in ELEMENTS:
        found = next((b for c, b in headings if c == chapter and re.search(pattern, b.text.lower())), None)
        if found is not None:
            items.append(ReadinessItem(id=f"S-{key}", question=label, status="PASS", basis="CODE", note=f"Found under “{found.text[:80]}”.", where=found.id, chapter=chapter))
        else:
            skip = f" Not needed for {may_skip[0]} studies." if may_skip else ""
            items.append(
                ReadinessItem(
                    id=f"S-{key}", question=label, status="NEEDS_REVIEW" if may_skip else "MISSING", basis="CODE", chapter=chapter,
                    note=f"No heading for this was found in chapter {chapter}. If the text covers it, a sub-heading helps examiners find it.{skip}",
                )
            )
    objectives = _listed_under(model, r"specific objective|objectives of the study|^objectives")
    questions = _listed_under(model, r"question|hypothes|proposition")
    low, high = book["objectives"]["min"], book["objectives"]["max"]
    if objectives:
        ok = low <= len(objectives) <= high
        note = f"{len(objectives)} specific objectives found." + ("" if ok else f" {low} to {high} are generally expected.")
        items.append(ReadinessItem(id="N-objectives", question="Number of specific objectives", status="PASS" if ok else "NEEDS_REVIEW", basis="CODE", note=note, chapter=1))
    if objectives and questions:
        same = len(objectives) == len(questions)
        items.append(
            ReadinessItem(
                id="N-alignment", question="One research question per objective", status="PASS" if same else "NEEDS_REVIEW", basis="CODE", chapter=1,
                note=f"{len(objectives)} objectives and {len(questions)} questions." + ("" if same else " Each objective usually has its own question."),
            )
        )
    past = [(b, m.group(0)) for b in model.blocks if chapters.get(b.id) == 3 and b.kind in ("paragraph", "list_item") for m in evidence.PAST_TENSE.finditer(b.text)]
    items.append(
        ReadinessItem(
            id="N-tense", question="Planned work in the future tense", status="NEEDS_REVIEW" if past else "PASS", basis="CODE", chapter=3,
            note=(f"{len(past)} place(s) in the methodology describe the planned study as done, for example “{past[0][1]}”." if past else "No past-tense description of the planned study was found in the methodology."),
            where=past[0][0].id if past else "",
        )
    )
    checks = paper_checks.check(model)
    missing = [c for c in checks.items if c.kind == "CITED_NOT_LISTED" and c.certainty == "CONFIRMED"]
    unused = [c for c in checks.items if c.kind == "LISTED_NOT_CITED" and c.certainty == "CONFIRMED"]
    no_list = any(c.kind == "NO_REFERENCE_LIST" for c in checks.items)
    items.append(
        ReadinessItem(
            id="N-references", question="Every work cited is in the reference list", basis="CODE", chapter=0,
            status="MISSING" if no_list else "NEEDS_REVIEW" if missing or unused else "PASS",
            note="No reference list was found." if no_list else f"{len(missing)} citation(s) not in the list; {len(unused)} reference(s) never cited." if missing or unused else f"{checks.citations_found} citations match the {checks.references_found} references.",
        )
    )
    items.append(
        ReadinessItem(
            id="C2-REFS30", question="Are there a minimum of 30 quality references?", basis="CODE", chapter=2, status="PASS" if checks.references_found >= 30 else "NEEDS_REVIEW",
            note=f"{checks.references_found} references found. Many programmes expect at least 30; check yours.",
        )
    )
    body = sum(b.words for b in model.blocks if chapters.get(b.id) in (1, 2, 3) and b.kind in ("paragraph", "list_item", "heading", "quote", "caption", "table_cell"))
    pages = round(body / book["words_per_page"])
    low_p, high_p = book["levels"][level]["pages"]
    items.append(
        ReadinessItem(
            id="N-length", question=f"Length for a {book['levels'][level]['label']} proposal: {low_p}-{high_p} pages", basis="CODE", chapter=0,
            status="PASS" if low_p <= pages <= high_p else "NEEDS_REVIEW",
            note=f"About {pages} pages of main text ({body:,} words, at about 250 words a page).",
        )
    )
    return items, {"words": body, "references": checks.references_found}


def stage_analysing(ctx: "StageContext") -> None:
    job = ctx.job
    model = ctx.document()
    level = job.selection.level
    rulebook_id = rulebook.DEFAULT
    items, facts = code_checks(model, level, rulebook_id)
    runner = ctx.ai(ProposalRunner)
    assert isinstance(runner, ProposalRunner)
    chapters = _chapter_of(model)
    paragraphs = [
        {"id": b.id, "chapter": chapters.get(b.id, 1), "section": b.section or "", "text": b.text}
        for b in model.blocks
        if b.kind in ("paragraph", "list_item", "quote") and chapters.get(b.id) in (1, 2, 3) and b.text.strip()
    ]
    questions = {n: [q for q in rulebook.vetting(rulebook_id, n) if not q.get("deterministic")] for n in (1, 2, 3)}
    audit = runner.audit(
        {
            "level": level,
            "rules": rulebook.rules_for(rulebook_id),
            "vetting": {str(n): [{"id": q["id"], "question": q["question"]} for q in qs] for n, qs in questions.items()},
            "paragraphs": paragraphs,
            "paperaidChecks": [{"check": i.question, "status": i.status, "note": i.note} for i in items],
        }
    )
    blocks = model.by_id()
    answered = {a.id: a for a in audit.items}
    for n, qs in questions.items():
        for q in qs:
            a = answered.get(q["id"])
            items.append(
                ReadinessItem(
                    id=q["id"], question=q["question"], chapter=n, basis="AI",
                    status=a.status if a else "NEEDS_REVIEW", note=a.note if a else "Not assessed.", where=a.where if a and a.where in blocks else "",
                )
            )
    findings = [
        ReviewFinding(where=_where(blocks, f.where), kind=f.kind, severity=f.severity, issue=f.issue, suggestion=f.suggestion)
        for f in audit.findings[:25]
    ]
    items = [i.model_copy(update={"where": _where(blocks, i.where) if i.where else ""}) for i in items]
    result = ProposalReview(rulebook=rulebook_id, level=level, words=facts["words"], items=items, findings=findings, method=METHOD)

    def save(j: Job) -> Job:
        j.proposal_review = result
        return j

    ctx.update(save)


def _where(blocks: dict, block_id: str) -> str:
    block = blocks.get(block_id)
    if block is None:
        return ""
    start = " ".join(block.text.split()[:12])
    return f"{block.section} — “{start}…”" if block.section else f"“{start}…”"


# --- EXPORTING -----------------------------------------------------------------------------------

STATUS = {"PASS": "Pass", "NEEDS_REVIEW": "Needs review", "MISSING": "Missing", "NOT_APPLICABLE": "Not applicable", "BLOCKED": "Blocked"}
BASIS = {"CODE": "PaperAid check", "AI": "PaperAid review", "AUTHOR": "Your information"}
GREEN, MUTED = RGBColor(0x0F, 0x63, 0x3E), RGBColor(0x46, 0x55, 0x4D)


def report(name: str, when: datetime, review: ProposalReview) -> bytes:
    book = rulebook.load(review.rulebook)
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    brand = doc.add_paragraph().add_run("PaperAid")
    brand.bold, brand.font.color.rgb, brand.font.size = True, GREEN, Pt(14)
    doc.add_heading("Proposal review", level=1)
    meta = doc.add_paragraph().add_run(f"{name} · {book['levels'][review.level]['label']} · {when:%d %B %Y}")
    meta.font.color.rgb = MUTED
    doc.add_paragraph(
        "This is a readiness checklist of what examiners look for, not a mark, and "
        "faculties may set their own variations. PaperAid does not check plagiarism."
    )
    for n, title in ((0, "Whole proposal"), (1, "Chapter One: General Introduction"), (2, "Chapter Two: Literature Review"), (3, "Chapter Three: Methodology")):
        chapter = [i for i in review.items if i.chapter == n]
        if not chapter:
            continue
        doc.add_heading(title, level=2)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        for cell, text in zip(table.rows[0].cells, ("Requirement", "Result", "Why"), strict=True):
            cell.text = text
        for item in chapter:
            row = table.add_row().cells
            row[0].text = item.question
            row[1].text = f"{STATUS[item.status]} ({BASIS[item.basis]})"
            row[2].text = item.note + (f" See {item.where}." if item.where else "")
    if review.findings:
        doc.add_heading("What a supervisor is likely to raise", level=2)
        for f in review.findings:
            p = doc.add_paragraph(style="List Number")
            p.add_run(f"{f.severity.capitalize()} · {f.kind.lower()}: ").bold = True
            p.add_run(f"{f.issue} {f.suggestion}")
            if f.where:
                where = p.add_run(f" ({f.where})")
                where.font.color.rgb = MUTED
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def stage_exporting(ctx: "StageContext") -> None:
    job = ctx.job
    assert job.proposal_review is not None
    name = job.source.name if job.source else "proposal"
    data = report(name, utcnow(), job.proposal_review)
    path = f"{ctx.prefix}/output/proposal-review.docx"
    ctx.rt.files.put(path, data, DOCX_TYPE)
    stem = PurePosixPath(name).stem[:120]
    output = StoredOutput(id="proposal-review", label="Proposal review (Word)", name=f"{stem} – proposal review.docx", size_bytes=len(data), path=path, content_type=DOCX_TYPE)

    def save(j: Job) -> Job:
        j.outputs = [output]
        j.outcome = j.outcome or "FULL"
        return j

    ctx.update(save)


STAGES = {Stage.ANALYSING: stage_analysing, Stage.EXPORTING: stage_exporting}
