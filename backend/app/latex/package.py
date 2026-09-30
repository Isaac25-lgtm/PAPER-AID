"""The LaTeX download: main.tex, its figures, a README, and main.pdf when the server compiled it.

Compilation runs pdflatex offline in a temporary folder with shell escape disabled, a time
limit, and nothing but PaperAid's own escaped output as input. A compile error is reported, never
hidden: the .tex is still delivered and the job is marked as partial.
"""

import io
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from app.latex.convert import Result, convert

COMPILE_TIMEOUT_SEC = 90


def compile_pdf(result: Result) -> tuple[bytes | None, str]:
    """(PDF, "") when it compiled; (None, reason) when it did not or could not."""
    engine = shutil.which("pdflatex")
    if engine is None:
        return None, "no LaTeX engine is installed on this server"
    with tempfile.TemporaryDirectory(prefix="paperaid-tex-") as folder:
        root = Path(folder)
        (root / "main.tex").write_text(result.tex, encoding="utf-8")
        for name, blob in result.files.items():
            target = root / name
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(blob)
        try:
            run = subprocess.run(
                [engine, "-interaction=nonstopmode", "-halt-on-error", "-no-shell-escape", "main.tex"],
                cwd=root,
                capture_output=True,
                text=True,
                timeout=COMPILE_TIMEOUT_SEC,
                env={"PATH": str(Path(engine).parent), "HOME": folder, "TEXMFVAR": folder, "openout_any": "p", "openin_any": "p"},
            )
        except subprocess.TimeoutExpired:
            return None, "compiling took too long"
        pdf = root / "main.pdf"
        if run.returncode == 0 and pdf.exists():
            return pdf.read_bytes(), ""
        error = next((line[1:].strip() for line in run.stdout.splitlines() if line.startswith("!")), "unknown error")
        return None, error[:200]


README = """This is your paper as a LaTeX project, converted by PaperAid.

Files
  main.tex    the paper
  figures/    the images from your Word file
  main.pdf    the compiled paper (included when PaperAid could compile it)

Compile with pdfLaTeX, for example by uploading this zip to Overleaf (New project > Upload
project) or running:  pdflatex main.tex

Citations are kept exactly as they appear in your Word file, and the reference list exactly as
you wrote it. PaperAid does not convert references to BibTeX, because that would mean guessing
authors, titles and years.
"""


def package(docx: bytes) -> tuple[bytes, Result, bool, str]:
    """Convert and (when possible) compile. Returns (zip bytes, conversion result, compiled,
    compile problem)."""
    result = convert(docx)
    pdf, problem = compile_pdf(result)
    return _zip(result, pdf, problem), result, pdf is not None, problem


def project(docx: bytes) -> tuple[bytes, Result]:
    """The LaTeX project alone, not compiled (a proposal's PDF is its own download), with the conversion
    result, whose `omitted` says whether anything was left out."""
    result = convert(docx)
    return _zip(result, None, ""), result


def _zip(result: Result, pdf: bytes | None, problem: str) -> bytes:
    notes = "\n".join(f"- {w}" for w in result.warnings)
    readme = README + (f"\nThings to check\n{notes}\n" if notes else "") + (f"\nNot compiled: {problem}.\n" if pdf is None and problem else "")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("main.tex", result.tex)
        archive.writestr("README.txt", readme)
        for name, blob in sorted(result.files.items()):
            archive.writestr(name, blob)
        if pdf is not None:
            archive.writestr("main.pdf", pdf)
    return out.getvalue()


def self_test() -> None:
    """Run while the server image is built: every LaTeX construct the converter writes must
    compile with the installed engine, or the build fails."""
    from docx import Document

    doc = Document()
    doc.add_heading("A title", 0)
    doc.add_heading("1. Section", 1)
    p = doc.add_paragraph("Text with “quotes” — dashes, 50% & $5, x_y, a # b, 10°C, ≥ 3, α and café. ")
    p.add_run("Bold").bold = True
    p.add_run(" and ").italic = True
    doc.add_paragraph("An item", style="List Bullet")
    doc.add_paragraph("A step", style="List Number")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Campus", "n"
    table.cell(1, 0).text, table.cell(1, 1).text = "Main", "120"
    doc.add_heading("References", 1)
    doc.add_paragraph("Author, A. (2020). Title & subtitle. Journal, 1(2), 3–4.")
    buffer = io.BytesIO()
    doc.save(buffer)
    pdf, problem = compile_pdf(convert(buffer.getvalue()))
    if pdf is None:
        raise SystemExit(f"LaTeX self-test failed: {problem}")
    print("LaTeX self-test passed")


if __name__ == "__main__":
    self_test()
