"""Deterministic DOCX formatting. Models never touch layout: this code applies a FormattingSpec
through named styles and section properties, then proves the wording did not change."""

import io
import re
from copy import deepcopy

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from app.core.errors import PermanentStageError
from app.documents.docx_io import W, body_text_fingerprint, iter_paragraphs
from app.documents.model import DocumentModel
from app.formatting.presets import FormattingSpec
from app.jobs.models import FormattingResult, FormattingRule

PAPER = {"A4": (Cm(21.0), Cm(29.7)), "Letter": (Cm(21.59), Cm(27.94))}
_RUN_FONT_OVERRIDES = ("rFonts", "sz", "szCs", "color")
_PARA_OVERRIDES = ("spacing", "ind", "jc")
_CHAPTER_ONE = re.compile(r"^(chapter\s+(one|1|i)\b|1\.?\s+\w|introduction$)", re.I)
_CONTENTS = re.compile(r"^(table of contents|contents)$", re.I)


def _set_font(style, name: str, size: float) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        fonts.attrib.pop(qn(theme_attr), None)


def _style(doc, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        return style


def _strip_children(parent, local_names: tuple[str, ...]) -> None:
    if parent is None:
        return
    for name in local_names:
        for child in parent.findall(W + name):
            parent.remove(child)


def _page_field_paragraph(paragraph: Paragraph, align: WD_ALIGN_PARAGRAPH) -> None:
    paragraph.alignment = align
    paragraph._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"><w:r><w:t>1</w:t></w:r></w:fldSimple>'))


_AFTER_PGNUM = [W + t for t in ("cols", "formProt", "vAlign", "noEndnote", "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings", "sectPrChange")]


def _set_page_numbering(sectpr, fmt: str) -> None:
    """Replace <w:pgNumType>, inserted where the OOXML schema requires (Word rejects wrong order)."""
    for existing in sectpr.findall(W + "pgNumType"):
        sectpr.remove(existing)
    element = parse_xml(f'<w:pgNumType {nsdecls("w")} w:fmt="{fmt}" w:start="1"/>')
    successor = next((child for child in sectpr if child.tag in _AFTER_PGNUM), None)
    if successor is not None:
        successor.addprevious(element)
    else:
        sectpr.append(element)


def _page_break_only(p) -> bool:
    runs = p.findall(W + "r")
    return bool(runs) and all(
        all(child.tag in (W + "rPr", W + "br", W + "lastRenderedPageBreak") for child in r) for r in runs
    ) and any(br.get(W + "type") == "page" for r in runs for br in r.findall(W + "br"))


def apply_formatting(data: bytes, spec: FormattingSpec, model: DocumentModel) -> tuple[bytes, FormattingResult]:
    before = body_text_fingerprint(data)
    doc = Document(io.BytesIO(data))
    blocks = model.by_id()
    warnings: list[str] = []

    # 1. Named styles carry the rules, so the document stays editable in Word.
    normal = doc.styles["Normal"]
    _set_font(normal, spec.font, spec.size_pt)
    pf = normal.paragraph_format
    pf.line_spacing = spec.line_spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(spec.space_after_pt)
    pf.first_line_indent = Cm(spec.first_line_indent_cm) if spec.first_line_indent_cm else None
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if spec.alignment == "justify" else WD_ALIGN_PARAGRAPH.LEFT
    for level, hs in spec.headings.items():
        style = _style(doc, f"Heading {level}")
        _set_font(style, spec.font, hs.size_pt)
        style.font.bold, style.font.italic = hs.bold, hs.italic
        style.font.color.rgb = RGBColor(0, 0, 0)
        hpf = style.paragraph_format
        hpf.alignment = WD_ALIGN_PARAGRAPH.CENTER if hs.align == "center" else WD_ALIGN_PARAGRAPH.LEFT
        hpf.first_line_indent = Cm(0)
        hpf.space_before, hpf.space_after = Pt(spec.heading_space_before_pt), Pt(spec.heading_space_after_pt)
        hpf.line_spacing = spec.line_spacing
        hpf.keep_with_next = True
    for name in ("List Paragraph", "List Number", "List Bullet"):
        if name in [s.name for s in doc.styles]:
            doc.styles[name].paragraph_format.first_line_indent = Cm(0)
    for name in ("Title", "Caption"):
        style = _style(doc, name)
        _set_font(style, spec.font, spec.size_pt)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if name == "Title" else WD_ALIGN_PARAGRAPH.LEFT

    # 2. Paragraph-level cleanup and role-specific rules.
    converted: list[str] = []
    heading_count = reference_count = caption_count = 0
    paragraphs = list(iter_paragraphs(doc))
    for block_id, p, in_table in paragraphs:
        paragraph = Paragraph(p, doc._body)
        if in_table:  # tables stay compact: single spacing, no indent
            tpf = paragraph.paragraph_format
            tpf.line_spacing, tpf.first_line_indent, tpf.space_after = 1.0, Cm(0), Pt(0)
            continue
        block = blocks.get(block_id)
        if block is None:
            continue
        for run in p.iter(W + "r"):
            _strip_children(run.find(W + "rPr"), _RUN_FONT_OVERRIDES)
        if block.kind == "heading":
            heading_count += 1
            level = min(3, block.level or 1)
            if block.detected_heading:
                paragraph.style = _style(doc, f"Heading {level}")
                converted.append(block.text)
            _strip_children(p.pPr, ("jc", "ind"))
        elif block.kind == "paragraph":
            _strip_children(p.pPr, _PARA_OVERRIDES)
        elif block.kind == "reference":
            reference_count += 1
            _strip_children(p.pPr, _PARA_OVERRIDES)
            rpf = paragraph.paragraph_format
            rpf.left_indent = Cm(spec.references_hanging_cm)
            rpf.first_line_indent = Cm(-spec.references_hanging_cm)
            rpf.line_spacing = spec.references_line_spacing
            rpf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif block.kind == "caption":
            caption_count += 1
            paragraph.style = _style(doc, "Caption")

    # 3. Roman-numbered preliminary pages, when the preset asks and the paper has them.
    roman_applied = False
    if spec.roman_preliminary_pages:
        seen_level_one = False
        for index, (block_id, _p, in_table) in enumerate(paragraphs):
            block = blocks.get(block_id)
            if in_table or block is None or block.kind != "heading" or (block.level or 1) != 1:
                continue
            if seen_level_one and _CHAPTER_ONE.match(block.text) and index > 0:
                previous = paragraphs[index - 1][1]
                body_sectpr = doc.element.body.sectPr
                prelim_sectpr = deepcopy(body_sectpr)
                _set_page_numbering(prelim_sectpr, "lowerRoman")
                if _page_break_only(previous):
                    for r in previous.findall(W + "r"):
                        previous.remove(r)
                previous.get_or_add_pPr().append(prelim_sectpr)
                _set_page_numbering(body_sectpr, "decimal")
                roman_applied = True
                break
            seen_level_one = True

    # 4. Page setup on every section, keeping intentional landscape sections landscape.
    for section in doc.sections:
        landscape = section.orientation == WD_ORIENT.LANDSCAPE
        width, height = PAPER[spec.paper_size]
        section.page_width, section.page_height = (height, width) if landscape else (width, height)
        section.top_margin, section.bottom_margin, section.left_margin, section.right_margin = (Cm(v) for v in spec.margins_cm)
        section.header_distance = section.footer_distance = Cm(1.25)

    # 5. Page numbers in the first section; later sections inherit it.
    first = doc.sections[0]
    in_header = spec.page_numbers.startswith("top")
    container = first.header if in_header else first.footer
    existing = " ".join(par.text for par in container.paragraphs).strip()
    page_numbers_added = False
    if existing:
        warnings.append("Your document already has header or footer text, so we left it unchanged. Add page numbers there if needed.")
    else:
        container.is_linked_to_previous = False
        target = container.paragraphs[0] if container.paragraphs else container.add_paragraph()
        _page_field_paragraph(target, WD_ALIGN_PARAGRAPH.RIGHT if spec.page_numbers.endswith("right") else WD_ALIGN_PARAGRAPH.CENTER)
        for later in doc.sections[1:]:
            (later.header if in_header else later.footer).is_linked_to_previous = True
        page_numbers_added = True

    # 6. Table of contents field after a "Contents" heading, when the preset wants one.
    toc_inserted = False
    if spec.insert_toc:
        for block_id, p, _ in paragraphs:
            block = blocks.get(block_id)
            if block and _CONTENTS.match(block.text):
                toc = parse_xml(
                    f'<w:p {nsdecls("w")}><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
                    '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>'
                    '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
                )
                p.addnext(toc)
                toc_inserted = True
                break

    out = io.BytesIO()
    doc.save(out)
    result_bytes = out.getvalue()

    unchanged = body_text_fingerprint(result_bytes) == before
    if not unchanged:
        raise PermanentStageError("FORMAT_CHANGED_TEXT", "Formatting could not be applied without changing your text, so we stopped.")

    if converted:
        sample = ", ".join(f"“{t[:40]}”" for t in converted[:3])
        warnings.insert(0, f"{len(converted)} heading(s) were typed as bold or enlarged text rather than heading styles. We converted them — check {sample}.")
    warnings.append("Formatting covers layout only. Your citation and reference style were not changed.")

    top, bottom, left, right = spec.margins_cm
    margins = f"{top} cm" if len({top, bottom, left, right}) == 1 else f"{top} cm top/bottom, {left} cm left, {right} cm right"
    rules = [
        FormattingRule(label="Page", value=f"{spec.paper_size}, {margins}"),
        FormattingRule(label="Body text", value=f"{spec.font} {spec.size_pt:g} pt, {'double' if spec.line_spacing == 2 else f'{spec.line_spacing:g} line'} spacing, {spec.alignment}"),
        FormattingRule(
            label="Paragraphs",
            value=f"{spec.first_line_indent_cm:g} cm first-line indent" if spec.first_line_indent_cm else f"Block paragraphs, {spec.space_after_pt:g} pt after",
        ),
        FormattingRule(label="Headings", value=f"{heading_count} headings styled consistently ({len(converted)} converted from bold text)"),
    ]
    if page_numbers_added:
        positions = {"top-right": "Top right", "top-center": "Top centre", "bottom-center": "Bottom centre", "bottom-right": "Bottom right"}
        rules.append(FormattingRule(label="Page numbers", value=f"{positions[spec.page_numbers]} on every page"))
    if roman_applied:
        rules.append(FormattingRule(label="Preliminary pages", value="Numbered i, ii, iii…; main text restarts at 1"))
    if reference_count:
        rules.append(FormattingRule(label="References", value=f"{reference_count} entries, {spec.references_hanging_cm:g} cm hanging indent"))
    if caption_count:
        rules.append(FormattingRule(label="Captions", value=f"{caption_count} table/figure captions styled"))
    if toc_inserted:
        rules.append(FormattingRule(label="Contents", value="Table of contents inserted — right-click it in Word and choose Update Field"))
    return result_bytes, FormattingResult(preset=spec.label, rules=rules, body_text_unchanged=unchanged, warnings=warnings)
