"""DOCX reading and selective patching.

A paragraph is split into segments. Plain text runs that share the paragraph's dominant
formatting are editable. Everything else — fields (Zotero/Mendeley citations), hyperlinks,
footnote references, equations, images, tracked changes and runs with different formatting —
is an opaque ⟦Xn⟧ segment whose XML is moved back untouched when the paragraph is rebuilt.
Block IDs come from document order, so re-reading the unchanged source gives the same IDs.
"""

import io
import re
from collections import Counter
from collections.abc import Iterator
from copy import deepcopy
from dataclasses import dataclass, field

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

from app.documents.model import Block, BlockKind, DocumentModel
from app.documents.protect import TOKEN

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_TEXT_RUN_CHILDREN = {W + t for t in ("rPr", "t", "tab", "br", "lastRenderedPageBreak", "softHyphen", "noBreakHyphen")}
_MARKERS = {W + t for t in ("bookmarkStart", "bookmarkEnd", "proofErr", "commentRangeStart", "commentRangeEnd", "permStart", "permEnd")}
_TRACKED = {W + t for t in ("ins", "del", "moveFrom", "moveTo")}
_REFERENCE_HEADINGS = re.compile(r"^(references|reference list|bibliography|works cited|literature cited)$", re.I)
_CAPTION = re.compile(r"^(table|figure|fig\.)\s+\d+", re.I)
_NUMBERED_HEADING = re.compile(r"^(\d+(?:\.\d+)*)\.?\s+\S")


@dataclass
class Segment:
    kind: str  # "text" | "opaque" | "marker"
    text: str
    elements: list = field(default_factory=list)
    rpr_sig: str = ""


def _visible_text(el) -> str:
    parts = []
    for node in el.iter():
        if node.tag in (W + "t", "{http://schemas.openxmlformats.org/officeDocument/2006/math}t"):
            parts.append(node.text or "")
        elif node.tag == W + "tab":
            parts.append("\t")
    return "".join(parts)


def _rpr_sig(run) -> str:
    rpr = run.find(W + "rPr")
    return "" if rpr is None else etree.tostring(rpr).decode()


def _is_plain_text_run(run) -> bool:
    for child in run:
        if child.tag not in _TEXT_RUN_CHILDREN:
            return False
        if child.tag == W + "br" and child.get(W + "type") in ("page", "column"):
            return False
    return True


def segment_paragraph(p) -> tuple[list[Segment], bool]:
    """Split a <w:p> into segments. Returns (segments, contains_tracked_changes)."""
    segments: list[Segment] = []
    tracked = False
    field_depth = 0
    field_elements: list = []

    for child in p:
        tag = child.tag
        if tag == W + "pPr":
            continue
        if field_depth:
            field_elements.append(child)
            for fc in child.iter(W + "fldChar"):
                kind = fc.get(W + "fldCharType")
                field_depth += 1 if kind == "begin" else -1 if kind == "end" else 0
            if field_depth == 0:
                segments.append(Segment("opaque", _field_result(field_elements), field_elements))
                field_elements = []
            continue
        if tag == W + "r" and any(fc.get(W + "fldCharType") == "begin" for fc in child.iter(W + "fldChar")):
            kinds = [fc.get(W + "fldCharType") for fc in child.iter(W + "fldChar")]
            field_depth = kinds.count("begin") - kinds.count("end")
            field_elements = [child]
            if field_depth == 0:
                segments.append(Segment("opaque", _field_result(field_elements), field_elements))
                field_elements = []
            continue
        if tag in _MARKERS:
            segments.append(Segment("marker", "", [child]))
        elif tag == W + "r" and _is_plain_text_run(child):
            segments.append(Segment("text", _visible_text(child).replace("\r", ""), [child], _rpr_sig(child)))
        else:
            tracked = tracked or tag in _TRACKED or any(n.tag in _TRACKED for n in child.iter())
            segments.append(Segment("opaque", _visible_text(child), [child]))

    if field_elements:  # unterminated field: keep it opaque rather than guessing
        segments.append(Segment("opaque", _field_result(field_elements), field_elements))

    # Text runs whose formatting differs from the paragraph's dominant formatting stay untouched.
    weights: Counter[str] = Counter()
    for s in segments:
        if s.kind == "text":
            weights[s.rpr_sig] += len(s.text)
    if weights:
        dominant = weights.most_common(1)[0][0]
        for s in segments:
            if s.kind == "text" and s.rpr_sig != dominant:
                s.kind = "opaque"
    return segments, tracked


def _field_result(elements: list) -> str:
    """Visible text of a complex field: the runs between 'separate' and 'end'."""
    state, parts = "", []
    for el in elements:
        for node in el.iter():
            if node.tag == W + "fldChar":
                state = node.get(W + "fldCharType") or ""
            elif node.tag == W + "t" and state == "separate":
                parts.append(node.text or "")
    return "".join(parts)


def masked_text(segments: list[Segment]) -> str:
    out, n = [], 0
    for s in segments:
        if s.kind == "text":
            out.append(s.text)
        elif s.kind == "opaque":
            n += 1
            out.append(f"⟦X{n}⟧")
    return "".join(out)


def iter_paragraphs(doc: DocxDocument) -> Iterator[tuple[str, object, bool]]:
    """Yield (block_id, <w:p>, in_table) for every paragraph in body order, tables included."""
    counter = 0

    def walk(container, in_table: bool) -> Iterator[tuple[str, object, bool]]:
        nonlocal counter
        for child in container:
            if child.tag == W + "p":
                counter += 1
                yield f"b{counter:05d}", child, in_table
            elif child.tag == W + "tbl":
                for tr in child.findall(W + "tr"):
                    for tc in tr.findall(W + "tc"):
                        yield from walk(tc, True)
            elif child.tag == W + "sdt":
                content = child.find(W + "sdtContent")
                if content is not None:
                    yield from walk(content, in_table)

    yield from walk(doc.element.body, False)


def _style_name(paragraph: Paragraph) -> str:
    try:
        return paragraph.style.name if paragraph.style is not None else "Normal"
    except (KeyError, ValueError):
        return "Normal"


def detect_fake_heading(paragraph: Paragraph, text: str) -> int | None:
    """Level for a heading typed as bold or enlarged body text, else None."""
    words = text.split()
    if not words or len(words) > 12 or text.rstrip().endswith((".", ",", ";", ":")) or len(text) < 3:
        return None
    runs = [r for r in paragraph.runs if r.text.strip()]
    if not runs:
        return None
    all_bold = all(r.bold for r in runs)
    sizes = [r.font.size.pt for r in runs if r.font.size is not None]
    large = bool(sizes) and max(sizes) >= 13
    if not (all_bold or large):
        return None
    numbered = _NUMBERED_HEADING.match(text)
    if numbered:
        return min(3, numbered.group(1).count(".") + 1)
    return 1 if large and max(sizes) >= 14 else 2


def _classify(paragraph: Paragraph, text: str, in_table: bool, in_references: bool) -> tuple[BlockKind, int | None, bool]:
    style = _style_name(paragraph)
    if in_table:
        return "table_cell", None, False
    if style == "Title":
        return "title", None, False
    if style.startswith("Heading"):
        digits = "".join(c for c in style if c.isdigit())
        return "heading", int(digits) if digits else 1, False
    if style == "Caption" or _CAPTION.match(text):
        return "caption", None, False
    level = detect_fake_heading(paragraph, text)
    if level:
        return "heading", level, True
    if in_references:
        return "reference", None, False
    if "Quote" in style:
        return "quote", None, False
    if "List" in style or paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None:
        return "list_item", None, False
    return "paragraph", None, False


def read_docx(data: bytes) -> DocumentModel:
    doc = Document(io.BytesIO(data))
    blocks: list[Block] = []
    section, in_references, reference_level = "", False, 0
    tracked_found = False
    for block_id, p, in_table in iter_paragraphs(doc):
        paragraph = Paragraph(p, doc._body)
        segments, tracked = segment_paragraph(p)
        text = "".join(s.text for s in segments if s.kind != "marker").strip()
        if not text:
            continue
        kind, level, detected = _classify(paragraph, text, in_table, in_references)
        if kind == "heading":
            if in_references and (level or 1) <= reference_level:
                in_references = False
            if _REFERENCE_HEADINGS.match(text):
                in_references, reference_level = True, level or 1
            section = text
        tracked_found = tracked_found or tracked
        has_editable_text = any(s.kind == "text" and re.search(r"[A-Za-z]{3}", s.text) for s in segments)
        editable = kind in ("paragraph", "list_item") and has_editable_text and not tracked
        blocks.append(
            Block(
                id=block_id,
                kind=kind,
                level=level,
                section=section if kind != "heading" else text,
                text=text,
                masked=masked_text(segments) if editable else None,
                editable=editable,
                detected_heading=detected,
                locked=[s.text for s in segments if s.kind == "opaque"] if editable else [],
            )
        )
    warnings = []
    if tracked_found:
        warnings.append("This document contains someone else's tracked changes. Paragraphs with tracked changes were left untouched.")
    return DocumentModel(format="DOCX", blocks=blocks, warnings=warnings)


def _make_runs(text: str, rpr) -> list:
    runs = []
    for piece in re.split(r"(\t|\n)", text):
        if not piece:
            continue
        run = OxmlElement("w:r")
        if rpr is not None:
            run.append(deepcopy(rpr))
        if piece == "\t":
            run.append(OxmlElement("w:tab"))
        elif piece == "\n":
            run.append(OxmlElement("w:br"))
        else:
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = piece
            run.append(t)
        runs.append(run)
    return runs


def apply_revisions(data: bytes, revisions: dict[str, str]) -> bytes:
    """Patch revised masked text (⟦Xn⟧ tokens intact, ⟦Pn⟧ already restored) into a copy of the
    source. Paragraphs not in `revisions` are never touched."""
    doc = Document(io.BytesIO(data))
    for block_id, p, _ in iter_paragraphs(doc):
        if block_id not in revisions:
            continue
        segments, _tracked = segment_paragraph(p)
        dominant_rpr = next((s.elements[0].find(W + "rPr") for s in segments if s.kind == "text"), None)
        opaque = {}
        n = 0
        for s in segments:
            if s.kind == "opaque":
                n += 1
                opaque[f"X{n}"] = s.elements
        markers = [el for s in segments if s.kind == "marker" for el in s.elements]
        rpr_copy = deepcopy(dominant_rpr) if dominant_rpr is not None else None
        for child in list(p):
            if child.tag != W + "pPr":
                p.remove(child)
        for el in markers:
            p.append(el)
        for piece in re.split(r"(⟦X\d+⟧)", revisions[block_id]):
            token = TOKEN.fullmatch(piece)
            if token:
                for el in opaque[token.group(1)]:
                    p.append(el)
            elif piece:
                for run in _make_runs(piece, rpr_copy):
                    p.append(run)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def body_text_fingerprint(data: bytes) -> list[str]:
    """Normalised text of every paragraph, used to prove a formatting job changed no wording."""
    doc = Document(io.BytesIO(data))
    lines = []
    for _, p, _ in iter_paragraphs(doc):
        segments, _ = segment_paragraph(p)
        text = re.sub(r"\s+", " ", "".join(s.text for s in segments if s.kind != "marker")).strip()
        if text:
            lines.append(text)
    return lines
