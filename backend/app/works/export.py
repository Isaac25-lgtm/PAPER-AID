"""A work as a Word document: the render profile the page checks use (rulebook §17: the same layout
for measuring and for the download), headings and text with citations and figures filled in by
code, the funding tables rendered from the Results Model and budget (never written by a model),
the reference list in the required style, and, where it applies, the last-page note."""

import io
import re
import zipfile
from dataclasses import dataclass
from typing import Literal

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt

from app.core.errors import PermanentStageError
from app.formatting.apply import _set_font
from app.proposals import evidence as ev
from app.proposals.export import _footer_numbers
from app.proposals.models import EvidenceItem
from app.works import budget as budget_engine
from app.works import numbers
from app.works import results as results_engine
from app.works.models import Budget, ResolvedSpec, ResultsModel, WorkDocument
from app.works.results import TO_ADD

KIND_LABELS = {"FUNDING_CONCEPT": "Concept Note", "PROJECT_CONCEPT": "Project Concept Note", "NGO_PROJECT": "Project Proposal", "RESEARCH_GRANT": "Research Grant Proposal"}
# Where each table goes: after the section with this key (or at the end when the key is absent).
TABLES_AFTER = {"logframe": ("expected_results", "goal_objectives"), "workplan": ("implementation", "timeline", "timeline_budget"), "mel": ("mel",),
                "budget": ("budget_narrative", "timeline_budget")}


class Profile:
    """Font, size, spacing and margins: the call's or brief's instructions where it gives them."""

    def __init__(self, spec: ResolvedSpec):
        coursework = spec.kind == "COURSEWORK"
        self.font, self.size, self.spacing, self.margin = "Times New Roman", 12.0, 1.5 if coursework else 1.0, 1.0
        text = " ".join(r.value + " " + r.quote for r in spec.requirements if r.key == "format").lower()
        for name in ("times new roman", "arial", "calibri", "cambria", "garamond", "georgia", "helvetica", "trebuchet ms", "verdana"):
            if name in text:
                self.font = name.title().replace("Ms", "MS")
        size = re.search(r"\b(9|10|10\.5|11|11\.5|12|13|14)\s*(?:pt|point)", text)
        if size:
            self.size = float(size.group(1))
        if "double" in text:
            self.spacing = 2.0
        elif "1.5" in text or "one and a half" in text:
            self.spacing = 1.5
        elif "single" in text:
            self.spacing = 1.0
        margin = re.search(r"(\d(?:\.\d+)?)\s*(?:inch|in\b|\")", text) or re.search(r"(\d(?:\.\d+)?)\s*cm", text)
        if margin:
            value = float(margin.group(1))
            self.margin = value / 2.54 if "cm" in margin.group(0) else value


def _setup(doc, profile: Profile) -> None:
    normal = doc.styles["Normal"]
    _set_font(normal, profile.font, profile.size)
    normal.paragraph_format.line_spacing = profile.spacing
    normal.paragraph_format.space_after = Pt(6 if profile.spacing == 1.0 else 0)
    for level, size in ((1, profile.size + 2), (2, profile.size), (3, profile.size)):
        style = doc.styles[f"Heading {level}"]
        _set_font(style, profile.font, size)
        style.font.bold, style.font.italic = True, False
        style.font.color.rgb = None
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(12), Pt(6)
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(profile.margin)


def _table(doc, rows: list[list[str]], caption: str) -> None:
    if len(rows) < 2:
        return
    title = doc.add_paragraph()
    title.add_run(caption).bold = True
    width = max(len(row) for row in rows)  # uneven rows are padded, never cut (Codex audit 2026-09-30, second round)
    grid = doc.add_table(rows=len(rows), cols=width)
    grid.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, value in enumerate([*row, *[""] * (width - len(row))]):
            cell = grid.cell(r, c)
            cell.text = value
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.bold = r == 0
    doc.add_paragraph()


def _budget_tables(budget: Budget) -> tuple[list[list[str]], list[list[str]]]:
    """The budget's summary and detail. A cost the applicant has not entered yet prints as a marked
    gap, and so does every total it affects, never a misleading zero (owner decision 2026-10-01)."""
    t = budget_engine.totals(budget)
    cur = budget.currency
    complete = budget_engine.complete(budget)

    def money(value: float, known: bool = True) -> str:
        return budget_engine.money(value, cur) if known else TO_ADD

    summary = [["Category", "Amount"], *[[k, money(v, complete)] for k, v in sorted(t.by_category.items())], ["Total", money(t.total, complete)]]
    detail = [["Line", "Description", "Quantity", "Unit", "Unit cost", "Total", "Year", "Activities"]]
    for li in budget.lines:
        known = bool(li.quantity and li.unit_cost)
        detail.append([li.id, li.description, f"{li.quantity:g}" if li.quantity else TO_ADD, li.unit, money(li.unit_cost, bool(li.unit_cost)), money(budget_engine.line_total(li), known),
                       str(li.year), ", ".join(li.activity_ids) or ("Support" if li.support else "")])
    return summary, detail


COVER_FIELDS = (("name", "Name"), ("reg", "Registration number"), ("course", "Course"), ("lecturer", "Lecturer"), ("institution", "Institution"), ("due", "Date"))
EXPLORATORY = "Exploratory draft: not every eligibility criterion is met, so this is not ready to submit."
DRAFT_LABEL = "Draft: some requirements still need your attention (see PaperAid's checklist)."


@dataclass(frozen=True)
class Block:
    """One piece of the delivered document, in order. The Word file and PaperAid's final review are
    both built from the same list, so the reviewer judges exactly what prints (Codex audit 2026-10-01)."""

    kind: Literal["title", "label", "heading", "paragraph", "count", "table", "figure", "references_heading", "reference", "note"]
    text: str = ""
    rows: tuple[tuple[str, ...], ...] = ()
    section: str = ""  # the section key it belongs to ("" outside sections)


def layout(document: WorkDocument, spec: ResolvedSpec, results: ResultsModel | None, budget: Budget | None, library: dict[str, EvidenceItem],
           tokens: dict[str, tuple[str, str]], draft: bool) -> list[Block]:
    """The document as delivered: citations and figures filled in, tables that print (at least a header
    and one row), the reference list and the last-page note. `draft` adds the not-ready label, which
    only the stored file carries: it is decided after the final review, from its result."""
    citer = ev.Citer(library, spec.citation_style)

    def show(text: str) -> str:
        return numbers.render(citer.render(text), tokens)[0]

    blocks = [Block("title", document.title)]
    label = KIND_LABELS.get(document.variant, "")
    if label:
        blocks.append(Block("label", label))
    for key, name in COVER_FIELDS:  # the student's own details, as they gave them
        if document.cover.get(key, "").strip():
            blocks.append(Block("label", f"{name}: {document.cover[key].strip()}"))
    if document.exploratory:
        blocks.append(Block("label", EXPLORATORY))
    elif draft and document.status == "NOT_READY":
        blocks.append(Block("label", DRAFT_LABEL))
    generated: dict[str, list[tuple[list[list[str]], str]]] = {}
    for key, _, rows, caption in generated_tables(document, spec, results, budget):
        generated.setdefault(key, []).append((rows, caption))

    def table(rows: list[list[str]], caption: str, section: str) -> list[Block]:
        if len(rows) < 2:  # a header alone does not print
            return []
        width = max(len(r) for r in rows)
        return [Block("table", caption, tuple(tuple([*r, *[""] * (width - len(r))]) for r in rows), section)]

    for s in document.sections:
        blocks.append(Block("heading", s.heading, section=s.key))
        rendered = [show(paragraph) for paragraph in s.paragraphs]
        blocks += [Block("paragraph", text, section=s.key) for text in rendered]
        field = next((f for f in spec.fields if f.id == s.field_id), None) if s.field_id else None
        if field is not None:
            whole = " ".join(rendered)
            limit = f"{field.max_characters:,} characters" if field.max_characters else f"{field.max_words:,} words"
            count = f"{len(whole):,} characters" if field.max_characters else f"{len(whole.split()):,} words"
            blocks.append(Block("count", f"{count} (limit {limit})", section=s.key))
        if s.table:
            caption = show(s.table_caption) or s.heading
            if s.table_illustrative and ILLUSTRATIVE not in caption:
                caption = f"{caption} {ILLUSTRATIVE}"
            blocks += table([[show(c) for c in row] for row in s.table], caption, s.key)
        if s.figure is not None:
            figure_rows = [("Line", s.figure.x_axis, s.figure.y_axis), *((line.label, f"{p.x:g}", f"{p.y:g}") for line in s.figure.series for p in line.points)]
            blocks.append(Block("figure", f"{s.figure.caption} {ILLUSTRATIVE}", tuple(figure_rows), s.key))
        for rows, caption in generated.pop(s.key, []):
            blocks += table(rows, caption, s.key)
    for rows, caption in [t for group in generated.values() for t in group]:  # tables whose section is absent
        blocks += table(rows, caption, "")
    cited = [i for i in document.cited if i in library]
    if cited:
        blocks.append(Block("references_heading", "Reference List" if spec.citation_style == "HARVARD" else "References"))
        blocks += [Block("reference", entry) for entry in ev.reference_list([library[i].source for i in cited], spec.citation_style)]
    if document.ai_note:
        blocks.append(Block("note", document.ai_note))
    return blocks


def build(document: WorkDocument, spec: ResolvedSpec, results: ResultsModel | None, budget: Budget | None, library: dict[str, EvidenceItem],
          tokens: dict[str, tuple[str, str]], draft: bool) -> bytes:
    """The Word file, written block by block from `layout`."""
    profile = Profile(spec)
    doc = Document()
    _setup(doc, profile)
    _footer_numbers(doc.sections[0], "decimal")
    figures = 0
    for block in layout(document, spec, results, budget, library, tokens, draft):
        if block.kind == "title":
            title = doc.add_paragraph()
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title.add_run(block.text).bold = True
        elif block.kind == "label":
            label = doc.add_paragraph()
            label.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = label.add_run(block.text)
            run.italic = block.text in (EXPLORATORY, DRAFT_LABEL)
        elif block.kind == "heading":
            doc.add_heading(block.text, level=1 if not spec.fields else 2)
        elif block.kind == "paragraph":
            doc.add_paragraph(block.text).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif block.kind == "count":
            doc.add_paragraph().add_run(block.text).italic = True
        elif block.kind == "table":
            _table(doc, [list(r) for r in block.rows], block.text)
        elif block.kind == "figure":
            figures += 1
            doc.add_picture(io.BytesIO(draw_figure(block.rows)), width=Inches(5.8))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption = doc.add_paragraph()
            caption.add_run(f"Figure {figures}. ").bold = True
            caption.add_run(block.text).italic = True
        elif block.kind == "references_heading":
            doc.add_heading(block.text, level=1)
        elif block.kind == "reference":
            entry = doc.add_paragraph(block.text)
            entry.paragraph_format.left_indent = Inches(0.5)
            entry.paragraph_format.first_line_indent = Inches(-0.5)
        elif block.kind == "note":  # on a page of its own at the end (owner decision 2026-09-30)
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_paragraph(block.text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


ILLUSTRATIVE = "(illustrative values)"


def verify(data: bytes, document: WorkDocument, spec: ResolvedSpec, results: ResultsModel | None, budget: Budget | None, library: dict[str, EvidenceItem],
           tokens: dict[str, tuple[str, str]], draft: bool) -> None:
    """The built file opens and carries every heading and paragraph of the accepted text, in order (algorithm revision
    2026-10-09: nothing is published or charged for until the file itself has been read back)."""
    def plain(text: str) -> str:
        return " ".join(text.split())

    try:
        written = [plain(p.text) for p in Document(io.BytesIO(data)).paragraphs]
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise PermanentStageError("EXPORT_FAILED", "PaperAid could not build the Word file of this draft. Nothing was charged; please try again.",
                                  f"built file does not open: {type(exc).__name__}") from exc
    position = 0
    for block in layout(document, spec, results, budget, library, tokens, draft):
        if block.kind not in ("heading", "paragraph") or not plain(block.text):
            continue
        try:
            position = written.index(plain(block.text), position) + 1
        except ValueError:
            raise PermanentStageError("EXPORT_FAILED", "PaperAid could not build the Word file of this draft. Nothing was charged; please try again.",
                                      f"built file lacks a {block.kind} of section {block.section}") from None


def draw_figure(rows: tuple[tuple[str, ...], ...]) -> bytes:
    """A figure block drawn by code: its first row names the axes, then one row per point (line, x, y)."""
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    _, x_axis, y_axis = rows[0]
    lines: dict[str, list[tuple[float, float]]] = {}
    for label, x, y in rows[1:]:
        lines.setdefault(label, []).append((float(x), float(y)))
    fig, ax = plt.subplots(figsize=(6.4, 4.0), dpi=200)
    ax.spines[["top", "right"]].set_visible(False)
    for (label, points), colour in zip(lines.items(), ("#1f3a8a", "#c2410c", "#15803d", "#7c3aed", "#0e7490", "#b45309"), strict=False):
        xs, ys = zip(*sorted(points), strict=True)
        ax.plot(xs, ys, marker="o", markersize=3, linewidth=1.8, color=colour, label=label)
    ax.set_xlabel(x_axis)
    ax.set_ylabel(y_axis)
    if len(lines) > 1:
        ax.legend(frameon=False, fontsize=8)
    buffer = io.BytesIO()
    fig.tight_layout()
    fig.savefig(buffer, format="png", dpi=200)
    plt.close(fig)
    return buffer.getvalue()


def generated_tables(document: WorkDocument, spec: ResolvedSpec, results: ResultsModel | None, budget: Budget | None) -> list[tuple[str, str, list[list[str]], str]]:
    """The tables code renders from the Results Model and budget, as (the section they follow, kind,
    rows, caption): the Word file and the length checks use the same list."""
    out: list[tuple[str, str, list[list[str]], str]] = []
    if results is not None and spec.kind == "FUNDING_PROPOSAL":
        if document.variant == "NGO_PROJECT":
            out.append((_after("logframe", document), "logframe", results_engine.logframe(results), "Logframe"))
        out.append((_after("workplan", document), "workplan", results_engine.workplan(results, spec.duration_months), "Workplan"))
        out.append((_after("mel", document), "mel", results_engine.mel_table(results), "Monitoring and evaluation indicators"))
    if budget is not None and budget.lines and (spec.kind == "FUNDING_PROPOSAL" or any(li.unit_cost for li in budget.lines)):
        summary, detail = _budget_tables(budget)
        key = _after("budget", document)
        out.append((key, "budget", summary, "Budget summary"))
        if spec.kind == "FUNDING_PROPOSAL":
            out.append((key, "budget", detail, "Detailed budget"))
    return out


def _after(table: str, document: WorkDocument) -> str:
    keys = {s.key for s in document.sections}
    return next((k for k in TABLES_AFTER[table] if k in keys), "")
