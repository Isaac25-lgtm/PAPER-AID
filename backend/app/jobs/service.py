"""Job business rules. Independent of FastAPI: every function takes the runtime and the verified
user, so ownership, pricing and state rules can be unit-tested directly."""

import hashlib
import logging
import re
import secrets
from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import PurePosixPath
from typing import Literal

from app.ai.orchestration import current_engine
from app.analysis import signals
from app.core.config import Settings
from app.core.errors import AppError, Conflict, Forbidden, InvalidDocument, LimitExceeded, NotFound
from app.core.logging import log
from app.documents.docx_io import read_docx
from app.documents.intake import inspect_upload
from app.formatting.guideline import MAX_GUIDE_WORDS
from app.formatting.presets import PRESETS, PUBLIC_PRESETS
from app.integrations.store import cursor_of
from app.jobs import state
from app.jobs.models import (
    AdminAction,
    AdminJob,
    AdminSummary,
    BoundQuote,
    Camel,
    ChangedBlock,
    EstimateRun,
    FileMeta,
    Finding,
    ImageMeta,
    Job,
    JobEvent,
    JobStatus,
    JobView,
    LedgerEntry,
    Page,
    PaperCheck,
    PaymentStatus,
    Quote,
    QuoteResponse,
    ServiceSelection,
    Stage,
    StoredFile,
    StoredImage,
    Wallet,
    WalletSummary,
    WalletView,
    utcnow,
)
from app.pricing import credits
from app.pricing.billing import hold_for_job, is_bundled_document, refund_job, release_hold, reservation_key
from app.pricing.quote import ai_cap_usd, bound_quote, estimate_charged, estimate_scan_usd, needs_estimate, with_margin, work_prices_set
from app.proposals.models import Project, moved_path
from app.runtime import Runtime
from app.works.models import Work

logger = logging.getLogger("paperaid.jobs")

BUILT = ("AI_CHECK", "REFINE", "FORMAT", "TEMPLATE_FORMAT", "SOURCE_CHECK", "REDRAFT", "LATEX", "PROPOSAL", "CONCEPT_NOTE", "COURSEWORK", "FUNDING_PROPOSAL")
NEEDS_AI = ("AI_CHECK", "REFINE", "TEMPLATE_FORMAT", "SOURCE_CHECK", "REDRAFT", "PROPOSAL", "CONCEPT_NOTE", "COURSEWORK", "FUNDING_PROPOSAL")
WORKS = ("CONCEPT_NOTE", "COURSEWORK", "FUNDING_PROPOSAL")


def availability(settings: Settings, user: "User | None" = None) -> dict[str, str]:
    """"soon" = not built yet, or (a work service) not switched on or not priced yet; "not_configured"
    = built, but the AI keys are not set; "invite_only" = testing is limited to invited testers and
    this user isn't one."""
    result = {}
    for service in ("AI_CHECK", "REFINE", "FORMAT", "TEMPLATE_FORMAT", "REDRAFT", "LATEX", "SOURCE_CHECK", "PROPOSAL", *WORKS):
        if service not in BUILT or (service in WORKS and (service not in settings.works_enabled or not work_prices_set(settings, service))):
            result[service] = "soon"  # no invented prices: a work service without its token prices is not offered
        elif service in NEEDS_AI and not settings.ai_configured or service in WORKS and not settings.roles_configured:
            result[service] = "not_configured"
        elif service in NEEDS_AI and not may_use_ai(settings, user) or service in WORKS and not may_use_works(settings, user):
            result[service] = "invite_only"
        else:
            result[service] = "available"
    return result


def may_use_works(settings: Settings, user: "User | None") -> bool:
    """The work services' pilot fails closed: only listed testers and admins, even with an empty
    list, until `works_public` opens them (then the usual tester rule applies)."""
    if settings.works_public:
        return may_use_ai(settings, user)
    return user is not None and (user.is_admin or user.email.lower() in {e.strip().lower() for e in settings.tester_emails})


def may_use_ai(settings: Settings, user: "User | None") -> bool:
    if not settings.tester_emails:
        return True
    if user is None:
        return False
    return user.is_admin or user.email.lower() in {e.strip().lower() for e in settings.tester_emails}
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass(frozen=True)
class User:
    uid: str
    email: str
    is_admin: bool
    verified: bool = True


def public_config(rt: Runtime, user: "User | None" = None) -> dict:
    return {
        "paymentsEnabled": rt.settings.payments_enabled,
        "availability": availability(rt.settings, user),
        "creditsEnabled": rt.settings.credits_enabled,
        "minCredits": rt.settings.min_credits,
        "aiScore": rt.settings.show_ai_score,  # the AI-likeness percentage and band are shown only when on
        "minTopUpUgx": rt.settings.min_top_up_ugx,
        "ugxPerUsd": rt.settings.ugx_per_usd,
        "ugxPerToken": rt.settings.ugx_per_token,
        "pricing": {
            "mode": rt.settings.pricing_mode,
            "bandPages": rt.settings.band_pages,
            "bandStep": rt.settings.band_step,
            "tokens": rt.settings.fixed_tokens,
            "formatUgxPer300Words": rt.settings.format_ugx_per_300_words,
            "formatMinUgx": rt.settings.format_min_ugx,
            "latexUgxPer300Words": rt.settings.latex_ugx_per_300_words,
            "latexMinUgx": rt.settings.latex_min_ugx,
        },
        "retentionDays": rt.settings.retention_days,
        "presets": PUBLIC_PRESETS,
    }


def pipeline_for(selection: ServiceSelection) -> list[Stage]:
    if selection.work == "READ":
        return [Stage.ANALYSING, Stage.EXPORTING]
    if selection.work == "PLAN":
        return [Stage.RESEARCHING, Stage.PLANNING, Stage.EXPORTING]
    if selection.work == "DRAFT":
        return [Stage.RESEARCHING, Stage.DRAFTING, Stage.AUDITING, Stage.EXPORTING]
    if selection.work == "REVISE":
        return [Stage.AUDITING, Stage.EXPORTING]
    if selection.proposal == "REVIEW":
        return [Stage.EXTRACTING, Stage.ANALYSING, Stage.EXPORTING]
    if selection.proposal == "PLAN":
        return [Stage.RESEARCHING, Stage.PLANNING, Stage.EXPORTING]
    if selection.proposal.startswith("REVISE_"):
        return [Stage.AUDITING, Stage.EXPORTING]
    if selection.proposal == "PROFILE":
        return [Stage.PLANNING, Stage.EXPORTING]
    if selection.proposal != "NONE":
        return [Stage.RESEARCHING, Stage.PLANNING, Stage.DRAFTING, Stage.AUDITING, Stage.EXPORTING]
    stages = [Stage.EXTRACTING]
    if selection.writing != "NONE":
        stages.append(Stage.ANALYSING)
    if selection.source_check:
        stages.append(Stage.RESEARCHING)
    if selection.writing == "REFINE":
        stages += [Stage.PLANNING, Stage.REFINING, Stage.AUDITING]
    elif selection.writing == "REDRAFT":
        stages += [Stage.PLANNING, Stage.REDRAFTING, Stage.AUDITING]
    if selection.formatting != "NONE":
        stages.append(Stage.FORMATTING)
    if selection.latex:
        stages.append(Stage.CONVERTING)
    stages.append(Stage.EXPORTING)
    return stages


def _estimating(job: Job) -> bool:
    return job.estimate is not None and job.estimate.status == "RUNNING"


ESTIMATE_BUSY = "Your estimate is still running. Wait for it to finish (it takes a minute or two), then make changes."


def _owned(rt: Runtime, user: User, job_id: str) -> Job:
    job = rt.store.get(job_id) if _valid_id(job_id) else None
    if job is None or job.owner_uid != user.uid:
        raise NotFound("We couldn't find this job.")  # same answer whether it exists or belongs to someone else
    return job


def _valid_id(job_id: str) -> bool:
    return job_id.startswith("job_") and job_id[4:].isalnum() and len(job_id) <= 40


def _rate_limit(rt: Runtime, user: User, action: str, limit: int) -> None:
    if rt.store.hit(f"{user.uid}:{action}", 3600) > limit:
        log(logger, logging.WARNING, "rate limit hit", uid=user.uid, action=action)
        raise LimitExceeded("You've reached the limit for this action. Please try again in an hour.", code="RATE_LIMITED")


PAGE_SIZE = 500


def every_job(rt: Runtime, owner_uid: str | None, statuses: set[JobStatus] | None) -> list[Job]:
    """All matching jobs, following cursors page by page (maintenance must never stop at page one)."""
    jobs, cursor = rt.store.list(owner_uid, statuses, None, None, PAGE_SIZE)
    while cursor:
        page, cursor = rt.store.list(owner_uid, statuses, None, cursor, PAGE_SIZE)
        jobs.extend(page)
    return jobs


def task_name(job: Job, suffix: str = "") -> str:
    return f"{job.id}-g{job.generation}-s{len(job.completed_stages)}{suffix}"


# --- owner operations ------------------------------------------------------------------


ACCOUNT_CLOSING = "This account is being deleted, so nothing new can be started."
NOT_SCORABLE = "This paper has no passage long enough (25 words or more) for a writing check to assess, so it cannot be checked."


def create_draft(rt: Runtime, user: User) -> JobView:
    if not user.verified:
        raise Forbidden("Verify your email address before starting a job. Check your inbox for the link.", code="EMAIL_NOT_VERIFIED")
    _rate_limit(rt, user, "draft", rt.settings.quotes_per_hour)
    now = utcnow()
    job = Job(
        id=f"job_{secrets.token_hex(6)}",
        status=JobStatus.DRAFT,
        owner_uid=user.uid,
        owner_email=user.email,
        created_at=now,
        expires_at=now + timedelta(days=rt.settings.retention_days),
        events=[JobEvent(at=now, label="Draft created")],
    )
    if not rt.store.create_if_open(job):  # the account check and the write are one atomic step
        raise Conflict(ACCOUNT_CLOSING, code="ACCOUNT_CLOSING")
    return job.view()


FileRole = Literal["source", "guideline"]


def upload_file(rt: Runtime, user: User, job_id: str, role: FileRole, filename: str, data: bytes) -> FileMeta:
    """Attach the paper ("source") or the university formatting guide ("guideline") to a draft."""
    job = _owned(rt, user, job_id)
    if job.status not in (JobStatus.DRAFT, JobStatus.QUOTED):
        raise Conflict("This job has already been submitted. Start a new job to upload another file.")
    if _estimating(job):
        raise Conflict(ESTIMATE_BUSY, code="ESTIMATE_RUNNING")
    display_name = PurePosixPath(filename.replace("\\", "/")).name[:180] or role
    guide = role == "guideline"
    min_words = 20 if guide else 50  # guides can be short; papers cannot
    max_words = MAX_GUIDE_WORDS if guide else rt.settings.max_words
    try:
        model = inspect_upload(data, display_name, rt.settings.max_upload_bytes, max_words, rt.settings.max_pdf_pages, min_words=min_words)
    except InvalidDocument as exc:
        if guide and exc.code == "DOCUMENT_TOO_LONG":
            exc.message = f"We can't use this formatting guide. It is longer than {MAX_GUIDE_WORDS:,} words; upload only the section with the formatting rules."
        elif guide:
            exc.message = f"We can't use this formatting guide. {exc.message}"
        raise
    ext = "pdf" if model.format == "PDF" else "docx"
    digest = hashlib.sha256(data).hexdigest()
    # Every upload attempt gets its own immutable object (even for identical bytes), so a quote always
    # points at exactly what it priced and a replaced object can never become current again.
    path = f"{job.storage_prefix()}/input/{role}-{digest[:12]}-{secrets.token_hex(4)}.{ext}"
    rt.files.put(path, data, "application/pdf" if ext == "pdf" else DOCX_TYPE)
    stored = StoredFile(
        name=display_name,
        format=model.format,
        size_bytes=len(data),
        word_count=model.word_count,
        page_estimate=model.page_count or 1,
        heading_count=model.heading_count,
        path=path,
        sha256=digest,
        scorable_words=sum(b.words for b in signals.analysable(model)),
    )
    previous: list[str] = []

    def attach(j: Job) -> Job | None:
        if not _open_draft(j):
            return None  # submitted, being estimated or being deleted meanwhile: the new file must not touch it
        current = j.source if role == "source" else j.guideline
        if current and current.path != path:
            previous.append(current.path)
        if role == "source":
            j.source = stored
            j.fix_notes, j.scope_words = {}, None  # a "Fix selected" draft's passages belong to the file it came from
            if j.selection.only_blocks:
                j.selection = j.selection.model_copy(update={"only_blocks": []})
        else:
            j.guideline = stored
        j.quote = None  # a new file always needs a new quote
        if j.status == JobStatus.QUOTED:
            state.transition(j, JobStatus.DRAFT, "File replaced; quote cleared")
        return j

    if rt.store.update(job.id, attach) is None:
        rt.files.delete(path)  # unique to this attempt and never attached: always safe to remove
        raise Conflict("This job was submitted before the new file arrived. Start a new job to use another file.", code="ALREADY_SUBMITTED")
    for old in previous:
        rt.files.delete(old)
    return FileMeta.model_validate(stored.model_dump())



MAX_LOGO_BYTES = 2_000_000


def upload_logo(rt: Runtime, user: User, job_id: str, filename: str, data: bytes) -> ImageMeta:
    """An institution logo for the title page (PNG or JPEG, up to 2 MB). Replacing it clears the
    quote, like any file."""
    from docx.image.exceptions import InvalidImageStreamError, UnexpectedEndOfFileError, UnrecognizedImageError
    from docx.image.image import Image

    job = _owned(rt, user, job_id)
    if not _open_draft(job):
        raise Conflict("This job has already been submitted. Start a new job to use another logo.", code="ALREADY_SUBMITTED")
    if len(data) > MAX_LOGO_BYTES:
        raise InvalidDocument("This logo is larger than 2 MB. Upload a smaller PNG or JPEG.", code="LOGO_TOO_LARGE")
    kind = "PNG" if data[:8] == b"\x89PNG\r\n\x1a\n" else "JPEG" if data[:3] == b"\xff\xd8\xff" else None
    try:
        image = Image.from_blob(data) if kind else None
    except (UnrecognizedImageError, UnexpectedEndOfFileError, InvalidImageStreamError, ValueError, KeyError, IndexError) as exc:
        raise InvalidDocument("We could not read this image. Upload the logo as a PNG or JPEG.", code="LOGO_UNREADABLE") from exc
    if image is None or not image.px_width or not image.px_height:
        raise InvalidDocument("Upload the logo as a PNG or JPEG image.", code="LOGO_UNREADABLE")
    digest = hashlib.sha256(data).hexdigest()
    path = f"{job.storage_prefix()}/input/logo-{digest[:12]}-{secrets.token_hex(4)}.{'png' if kind == 'PNG' else 'jpg'}"
    rt.files.put(path, data, "image/png" if kind == "PNG" else "image/jpeg")
    name = PurePosixPath(filename.replace("\\", "/")).name[:180] or "logo"
    stored = StoredImage(name=name, format=kind, size_bytes=len(data), width_px=image.px_width, height_px=image.px_height, path=path, sha256=digest)  # type: ignore[arg-type]
    previous: list[str] = []

    def attach(j: Job) -> Job | None:
        if not _open_draft(j):
            return None
        if j.logo and j.logo.path != path:
            previous.append(j.logo.path)
        j.logo, j.quote = stored, None
        if j.status == JobStatus.QUOTED:
            state.transition(j, JobStatus.DRAFT, "Logo replaced; quote cleared")
        return j

    if rt.store.update(job.id, attach) is None:
        rt.files.delete(path)
        raise Conflict("This job was submitted before the logo arrived.", code="ALREADY_SUBMITTED")
    for old in previous:
        rt.files.delete(old)
    return ImageMeta.model_validate(stored.model_dump())


def remove_logo(rt: Runtime, user: User, job_id: str) -> None:
    job = _owned(rt, user, job_id)
    removed: list[str] = []

    def detach(j: Job) -> Job | None:
        if not _open_draft(j):
            return None
        if j.logo:
            removed.append(j.logo.path)
            j.logo, j.quote = None, None
            if j.status == JobStatus.QUOTED:
                state.transition(j, JobStatus.DRAFT, "Logo removed; quote cleared")
        return j

    if rt.store.update(job.id, detach) is None:
        raise Conflict("This job has already been submitted.", code="ALREADY_SUBMITTED")
    for path in removed:
        rt.files.delete(path)


def remove_guideline(rt: Runtime, user: User, job_id: str) -> None:
    """Detach the formatting guide from a draft. Idempotent; clears any quote that priced it."""
    job = _owned(rt, user, job_id)
    removed: list[str] = []

    def detach(j: Job) -> Job | None:
        if not _open_draft(j):
            return None
        if j.guideline is None:
            return j
        removed.append(j.guideline.path)
        j.guideline = None
        j.quote = None
        if j.status == JobStatus.QUOTED:
            state.transition(j, JobStatus.DRAFT, "Guide removed; quote cleared")
        return j

    if job.status not in (JobStatus.DRAFT, JobStatus.QUOTED) or rt.store.update(job.id, detach) is None:
        raise Conflict("This job has already been submitted.", code="ALREADY_SUBMITTED")
    for path in removed:
        rt.files.delete(path)


def request_quote(rt: Runtime, user: User, job_id: str, selection: ServiceSelection, start_estimate: bool = False) -> QuoteResponse:
    """Price a selection. Refinement first needs a paid AI estimate: without `start_estimate` this
    only reports what the estimate can cost; with it, the estimate is started (the worker finishes
    it). Everything else is priced from length at once. Either way the student needs credits: no
    estimate without a balance (owner decision 2026-09-24)."""
    job = _owned(rt, user, job_id)
    if job.status not in (JobStatus.DRAFT, JobStatus.QUOTED):
        raise Conflict("This job has already been submitted.")
    if job.source is None:
        raise AppError("Upload your paper first.", code="NO_FILE")
    services = selection.services()
    if not services:
        raise AppError("Choose at least one service.", code="NO_SERVICE")
    offered = availability(rt.settings, user)
    if any(offered.get(s.value) != "available" for s in services):
        raise AppError("That service is not available right now.", code="SERVICE_UNAVAILABLE")
    if selection.proposal not in ("NONE", "REVIEW") or job.project_id:
        raise AppError("Proposal steps are started from the proposal's page.", code="PROPOSAL_STEP")
    if selection.proposal == "REVIEW" and (len(services) > 1 or selection.writing != "NONE"):
        raise AppError("A proposal review runs on its own. Start another job for other services.", code="REVIEW_ALONE")
    if selection.writing == "AI_CHECK" and job.source.scorable_words == 0:  # refused before it is priced
        raise AppError(NOT_SCORABLE, code="NOT_SCORABLE")
    if selection.proposal == "REVIEW" and job.source.word_count > rt.settings.proposal_review_max_words:
        raise AppError(f"Proposal review accepts up to {rt.settings.proposal_review_max_words:,} words.", code="DOCUMENT_TOO_LONG")
    if selection.only_blocks and (selection.writing != "REFINE" or not all(re.fullmatch(r"b\d{5}", b) for b in selection.only_blocks)):
        raise AppError("Choose the passages to fix from your writing check.", code="INVALID_SELECTION")
    scope = _fix_scope(rt, job, selection.only_blocks) if selection.only_blocks else None
    # a source check may also run on its own: the "Check my sources" action on a paper's results
    if job.source.format == "PDF" and any(s.value not in ("AI_CHECK", "SOURCE_CHECK", "PROPOSAL") for s in services):
        raise AppError("A PDF can be checked (writing and sources) only. Upload the Word file to redraft or format it.", code="PDF_AI_CHECK_ONLY")
    if selection.formatting == "TEMPLATE_FORMAT" and job.guideline is None:
        raise AppError("Upload your university's formatting guide to use University templates.", code="NO_GUIDELINE")
    guideline_sha = job.guideline.sha256 if selection.formatting == "TEMPLATE_FORMAT" and job.guideline else None
    if selection.logo != "NONE" and (selection.formatting == "NONE" or job.logo is None):
        raise AppError("Upload your logo, and choose a formatting option, to place it on the title page.", code="NO_LOGO")
    if selection.formatting == "FORMAT" and selection.preset not in PRESETS:
        raise AppError("Choose an available formatting style.", code="INVALID_PRESET")
    if _estimating(job):
        if job.estimate and job.estimate.selection == selection:
            return QuoteResponse(estimate=job.estimate)
        raise Conflict(ESTIMATE_BUSY, code="ESTIMATE_RUNNING")
    now = utcnow()
    existing = job.quote
    if (
        existing
        and existing.selection == selection
        and existing.source_sha256 == job.source.sha256
        and existing.guideline_sha256 == guideline_sha
        and existing.expires_at > now + timedelta(minutes=2)
    ):
        return QuoteResponse(quote=Quote.model_validate(existing.model_dump()), estimate=job.estimate)
    settings = rt.settings
    source_sha, words = job.source.sha256, job.source.word_count
    finished = _finished_estimate(job, selection, guideline_sha, rt.settings)
    if needs_estimate(selection, settings) and finished is None:
        if not start_estimate:
            cap = _estimate_fee_cap(rt, job, selection) if estimate_charged(settings) else 0
            return QuoteResponse(estimate_fee_cap=cap, estimate=job.estimate)
        return _start_estimate(rt, user, job, selection, guideline_sha)
    if finished is None and settings.credits_enabled and settings.pricing_mode == "cost":
        wallet = rt.store.get_wallet(user.uid)
        available = wallet.available if wallet else 0
        if available <= 0:
            raise credits.InsufficientCredits(settings.min_top_up_ugx, available)
    _rate_limit(rt, user, "quote", settings.quotes_per_hour)
    issued: list[BoundQuote] = []

    def save(j: Job) -> Job | None:
        if not _open_draft(j) or j.source is None or j.source.sha256 != source_sha:
            return None
        if guideline_sha is not None and (j.guideline is None or j.guideline.sha256 != guideline_sha):
            return None
        run = _finished_estimate(j, selection, guideline_sha, settings)
        if needs_estimate(selection, settings) and run is None:
            return None
        # A finished estimate is repriced from its saved result: no new scan, no new fee. Any
        # estimate already charged on this job counts toward whichever quote follows it.
        quote = bound_quote(
            settings,
            selection,
            source_sha,
            words,
            (run.engine if run else None) or current_engine(settings),
            guideline_sha,
            run.guide_words if run else (j.guideline.word_count if guideline_sha and j.guideline else 0),
            run.passages if run else None,
            fee_paid=j.billing.fee_paid,
            estimate_id=run.id if run else None,
            scope_words=scope,
        )
        issued.append(quote)
        j.quote = quote
        j.selection = selection
        return state.transition(j, JobStatus.QUOTED, "Quote issued") if j.status == JobStatus.DRAFT else j

    if rt.store.update(job.id, save) is None:
        raise Conflict("Your files changed while we were pricing them. We'll price the new ones.", code="FILES_CHANGED")
    return QuoteResponse(quote=Quote.model_validate(issued[-1].model_dump()), estimate=job.estimate)


def _fix_scope(rt: Runtime, job: Job, blocks: list[str]) -> int:
    """"Fix selected" is priced by the words of the chosen passages, counted from the current file
    every time it is quoted: a selection can only name passages PaperAid may rewrite, and changing it
    changes the price (Codex audit 56c4f83 H08)."""
    assert job.source is not None
    if job.source.format != "DOCX":
        raise AppError("Fixing passages needs the Word file. Upload the .docx version of this paper.", code="PDF_AI_CHECK_ONLY")
    model = read_docx(rt.files.get(job.source.path))
    editable = {b.id: b for b in model.blocks if b.editable and b.kind in ("paragraph", "list_item")}
    if len(set(blocks)) != len(blocks) or any(b not in editable for b in blocks):
        raise AppError("Choose the passages to fix from your writing check.", code="INVALID_SELECTION")
    return sum(editable[b].words for b in blocks)


def _open_draft(j: Job) -> bool:
    """A draft that may still change: not submitted, not being estimated, not being deleted."""
    return j.status in (JobStatus.DRAFT, JobStatus.QUOTED) and not _estimating(j) and not j.deleting and not j.retiring and not j.files_deleted


def _finished_estimate(j: Job, selection: ServiceSelection, guideline_sha: str | None, settings: Settings) -> EstimateRun | None:
    """The job's finished estimate, if it sized exactly these files and this selection."""
    run = j.estimate
    if run is None or run.status != "READY" or j.source is None or not needs_estimate(selection, settings):
        return None
    if run.selection != selection or run.source_sha256 != j.source.sha256 or run.guideline_sha256 != guideline_sha:
        return None
    return run


def _estimate_fee_cap(rt: Runtime, job: Job, selection: ServiceSelection) -> int:
    assert job.source is not None
    return with_margin(estimate_scan_usd(rt.settings, job.source.word_count, selection), rt.settings)


def _start_estimate(rt: Runtime, user: User, job: Job, selection: ServiceSelection, guideline_sha: str | None) -> QuoteResponse:
    """Hold the most the scan can cost, then let the worker run it."""
    assert job.source is not None
    settings = rt.settings
    fee_cap = _estimate_fee_cap(rt, job, selection)
    _rate_limit(rt, user, "estimate", settings.quotes_per_hour)
    run = EstimateRun(
        id=f"est{secrets.token_hex(5)}",
        status="RUNNING",
        fee_cap=fee_cap,
        selection=selection,
        source_sha256=job.source.sha256,
        guideline_sha256=guideline_sha,
        budget_usd=ai_cap_usd(fee_cap, settings),
        engine=current_engine(settings),
    )

    def start(j: Job, w: Wallet) -> tuple[Job, Wallet] | None:
        if w.closing or not _open_draft(j) or j.source is None or j.source.sha256 != run.source_sha256:
            return None
        if guideline_sha is not None and (j.guideline is None or j.guideline.sha256 != guideline_sha):
            return None
        charged = estimate_charged(settings)  # under fixed prices the Deep Redraft preview is free
        if charged:
            credits.hold(w, fee_cap, j.id, "AI estimate (the most it can cost)")
        run.held = charged
        run.cost_base_usd = j.estimate_cost_usd
        j.estimate = run
        j.selection = selection
        j.quote = None  # earlier estimate fees stay on j.billing and count toward the next quote
        if j.status == JobStatus.QUOTED:
            state.transition(j, JobStatus.DRAFT, "New estimate requested; previous quote cleared")
        j.events.append(JobEvent(label=f"Estimate started (up to {credits.tokens(fee_cap)} held)" if charged else "Preview started (not charged)"))
        return j, w

    result = rt.store.update_job_and_wallet(job.id, start)
    if result is None:
        raise Conflict("Your files changed while we were pricing them. We'll price the new ones.", code="FILES_CHANGED")
    rt.queue.enqueue(job.id, f"{job.id}-e{run.id}")
    return QuoteResponse(estimate=run)


def _priced_input_intact(j: Job) -> bool:
    """The job still has exactly what its quote priced: the uploaded file, or for a proposal step
    the project input frozen when it was priced."""
    if j.quote is None:
        return False
    if j.project_id or j.work_id:
        return j.input_sha256 is not None and j.quote.source_sha256 == j.input_sha256
    return j.source is not None and j.quote.source_sha256 == j.source.sha256


ProjectGate = Callable[[Job, Project | None], Project | None]
WorkGate = Callable[[Job, Work | None], Work | None]


def submit(rt: Runtime, user: User, job_id: str, quote_id: str, project_gate: ProjectGate | None = None, work_gate: WorkGate | None = None) -> JobView:
    """Accept a quote. A proposal step is submitted through its project (app.proposals.service),
    whose `project_gate` runs in the same transaction as the hold: it refuses (None) or returns the
    project with the step claimed, so a project can never be deleted between its claim and the
    credits being held (Codex audit 2026-09-28 #4)."""
    job = _owned(rt, user, job_id)
    if job.project_id and project_gate is None:
        raise AppError("Start this step from the proposal's page.", code="PROPOSAL_STEP")
    if job.work_id and work_gate is None:
        raise AppError("Start this step from its page.", code="WORK_STEP")
    if job.status in state.SUBMITTED and job.quote and job.quote.id == quote_id:
        return job.view()  # double click or retried request: the same submission, nothing new
    if job.status != JobStatus.QUOTED or job.quote is None or job.quote.id != quote_id:
        raise Conflict("Your quote changed. Review the new price and submit again.", code="QUOTE_MISMATCH")
    if job.quote.expires_at < utcnow():
        raise Conflict("Your quote expired. Request a new one.", code="QUOTE_EXPIRED")
    if not _priced_input_intact(job):
        raise Conflict("Your file changed after it was priced. Request a new quote.", code="QUOTE_MISMATCH")
    if any(availability(rt.settings, user).get(service.value) != "available" for service in job.quote.selection.services()):
        raise AppError("That service is not available right now.", code="SERVICE_UNAVAILABLE")
    if rt.store.count_active(user.uid) >= rt.settings.max_active_jobs_per_user:
        raise LimitExceeded(
            f"You already have {rt.settings.max_active_jobs_per_user} jobs in progress. Submit this one when one of them finishes.",
            code="ACTIVE_JOB_LIMIT",
        )
    _rate_limit(rt, user, "submit", rt.settings.submits_per_hour)
    settings = rt.settings

    def accept(j: Job, w: Wallet) -> tuple[Job, Wallet] | None:
        # Re-checked inside the transaction: a concurrent upload or submit may have changed the job.
        if w.closing or j.status != JobStatus.QUOTED or j.deleting or j.retiring or j.files_deleted or j.quote is None or j.quote.id != quote_id or not _priced_input_intact(j):
            return None
        if any(availability(settings, user).get(service.value) != "available" for service in j.quote.selection.services()):
            return None
        if j.quote.guideline_sha256 is not None and (j.guideline is None or j.guideline.sha256 != j.quote.guideline_sha256):
            return None
        q = j.quote
        j.services = q.selection.services()
        j.pipeline = pipeline_for(q.selection)
        # The provider-spend cap keeps the margin: the quote's AI part at its frozen rate and multiplier.
        if q.budget_usd:  # fixed prices: the cap is the worst-case projection, not the price
            j.budget_usd = q.budget_usd
        else:
            j.budget_usd = max(0, q.amount - q.paid - q.fixed_ugx) / q.ugx_per_usd / q.multiplier if q.ugx_per_usd and q.multiplier else 0.0
        if settings.credits_enabled:
            key = reservation_key(j)
            if key and is_bundled_document(j):  # the credits Start reserved become this document's hold, in one transaction
                credits.release_reservation(w, key, "Reserved credits now held for your document")
            held = hold_for_job(j, w)  # raises InsufficientCredits, which aborts the whole transaction
            state.transition(j, JobStatus.QUEUED, f"Queued ({credits.tokens(held)} held)")
        else:
            j.payment_status = PaymentStatus.NOT_REQUIRED
            state.transition(j, JobStatus.QUEUED, "Queued (testing: not charged)")
        return j, w

    if job.project_id:
        assert project_gate is not None

        def accept_step(j: Job, w: Wallet, p: Project | None) -> tuple[Job, Wallet, Project] | None:
            claimed = project_gate(j, p)
            accepted = accept(j, w) if claimed is not None else None
            return (accepted[0], accepted[1], claimed) if accepted is not None and claimed is not None else None

        triple = rt.store.update_job_wallet_and_project(job.id, job.project_id, accept_step)
        result = (triple[0], triple[1]) if triple else None
    elif job.work_id:
        assert work_gate is not None

        def accept_work(j: Job, w: Wallet, k: Work | None) -> tuple[Job, Wallet, Work] | None:
            claimed = work_gate(j, k)
            accepted = accept(j, w) if claimed is not None else None
            return (accepted[0], accepted[1], claimed) if accepted is not None and claimed is not None else None

        triple_work = rt.store.update_job_wallet_and_work(job.id, job.work_id, accept_work)
        result = (triple_work[0], triple_work[1]) if triple_work else None
    else:
        result = rt.store.update_job_and_wallet(job.id, accept)
    updated = result[0] if result else None
    if updated is None:
        current = _owned(rt, user, job_id)
        if current.status in state.SUBMITTED and current.quote and current.quote.id == quote_id:
            return current.view()  # a concurrent identical submit won; same result
        raise Conflict("Your quote changed. Review the new price and submit again.", code="QUOTE_MISMATCH")
    if updated.status == JobStatus.QUEUED:
        rt.queue.enqueue(updated.id, task_name(updated))
    log(logger, logging.INFO, "job submitted", jobId=updated.id, services=[s.value for s in updated.services], words=updated.source.word_count if updated.source else 0)
    return updated.view()


def get_job(rt: Runtime, user: User, job_id: str) -> JobView:
    return _owned(rt, user, job_id).view()


def list_jobs(rt: Runtime, user: User, status: str | None, service: str | None, cursor: str | None, limit: int) -> Page[JobView]:
    """Submitted jobs by default. `status=DRAFT` lists unfinished drafts (priced or not) that have a
    paper, so a student can resume them."""
    if status == "DRAFT":
        statuses = {JobStatus.DRAFT, JobStatus.QUOTED}
    else:
        statuses = {JobStatus(status)} if status and status != "ALL" else state.SUBMITTED
    jobs, next_cursor = rt.store.list(user.uid, statuses, service if service and service != "ALL" else None, cursor, min(limit, 50))
    jobs = [j for j in jobs if not j.project_id and not j.work_id]  # project and work steps are shown on their own page
    if status == "DRAFT":
        jobs = [j for j in jobs if j.source is not None and not j.deleting and not j.files_deleted]
    return Page[JobView](items=[j.view() for j in jobs], next_cursor=next_cursor)


def cancel(rt: Runtime, user: User, job_id: str, actor: str | None = None) -> JobView:
    job = _owned(rt, user, job_id) if actor is None else _require(rt, job_id)
    if job.status == JobStatus.CANCELLED:
        return job.view()
    if job.status not in (JobStatus.DRAFT, JobStatus.QUOTED, JobStatus.AWAITING_PAYMENT, JobStatus.QUEUED):
        raise Conflict("This job has already started, so it can't be cancelled.", code="CANNOT_CANCEL")
    if _estimating(job):
        raise Conflict(ESTIMATE_BUSY, code="ESTIMATE_RUNNING")

    def do_cancel(j: Job, w: Wallet) -> tuple[Job, Wallet] | None:
        if j.status not in (JobStatus.DRAFT, JobStatus.QUOTED, JobStatus.AWAITING_PAYMENT, JobStatus.QUEUED) or _estimating(j):
            return None
        if actor:
            j.admin_actions.append(AdminAction(actor=actor, action="Cancelled"))
            refund_job(j, w, "Cancelled by PaperAid, so nothing was charged")  # our decision: everything back
        else:
            release_hold(j, w, "You cancelled before the job started")  # a finished estimate stays paid
        state.transition(j, JobStatus.CANCELLED, "Cancelled by admin" if actor else "Cancelled by user")
        return j, w

    result = rt.store.update_job_and_wallet(job.id, do_cancel)
    if result is None:
        raise Conflict("This job has already started, so it can't be cancelled.", code="CANNOT_CANCEL")
    return result[0].view()


ACTIVE = (JobStatus.QUEUED, JobStatus.PROCESSING, JobStatus.AWAITING_PAYMENT)


def _mark_deleting(rt: Runtime, job_id: str) -> Job | None:
    """Atomically claim a job for deletion. Refused (None) while work runs or credits are held;
    once claimed, nothing new can start on it (every operation checks `deleting` in its own
    transaction), so deleting the files and record afterwards cannot strand a hold."""

    def mark(j: Job) -> Job | None:
        if j.status in ACTIVE or _estimating(j) or j.billing.state == "HELD":
            return None
        j.deleting = True
        return j

    return rt.store.update(job_id, mark)


def _unmark_deleting(rt: Runtime, job_id: str) -> None:
    def unmark(j: Job) -> Job:
        j.deleting = False
        return j

    rt.store.update(job_id, unmark)


def _erase(rt: Runtime, job: Job) -> None:
    rt.files.delete_prefix(job.storage_prefix())
    rt.store.delete(job.id)


def delete(rt: Runtime, user: User, job_id: str) -> None:
    job = rt.store.get(job_id) if _valid_id(job_id) else None
    if job is None:
        return  # already gone: deletion is idempotent
    if job.owner_uid != user.uid:
        raise NotFound("We couldn't find this job.")
    claimed = _mark_deleting(rt, job.id)
    if claimed is None:
        current = rt.store.get(job.id)
        if current is None:
            return
        if _estimating(current):
            raise Conflict(ESTIMATE_BUSY, code="ESTIMATE_RUNNING")
        raise Conflict("Cancel the job or wait for it to finish before deleting it.", code="JOB_ACTIVE")
    _erase(rt, claimed)
    log(logger, logging.INFO, "job deleted", jobId=job.id)


def output_file(rt: Runtime, user: User, job_id: str, output_id: str) -> tuple[str, str, str]:
    """Returns (storage path, download name, content type) for an owned, unexpired output."""
    job = _owned(rt, user, job_id)
    if job.expires_at < utcnow():
        raise AppError("These files were deleted under our retention policy.", code="FILES_EXPIRED", status=410)
    output = next((o for o in job.outputs if o.id == output_id), None)
    if output is None or not rt.files.exists(output.path):
        raise NotFound("That file is not available.")
    return output.path, output.name, output.content_type


def delete_account(rt: Runtime, user: User) -> int:
    """Delete every job, file and the wallet of a user (Codex audit #2).

    1. The whole account is claimed first: the wallet is marked `closing` in a transaction, and
       every operation that starts work or moves credits checks that flag in its own transaction,
       so nothing new can begin while the account is being deleted.
    2. Every job is claimed; if one is still working, all claims are released and nothing changes.
    3. Jobs and files are erased, then a final sweep catches any draft created before the claim.
    4. The wallet becomes a closed tombstone (no email, no balance, no history) instead of being
       removed, so a request already past its checks can never recreate an open account: a late
       draft or credit grant finds it closed and is refused (Codex audit #2, second round).
    5. The complete credit history is erased last (Codex audit 56c4f83 H02): the tombstone is closed,
       so nothing can add to it meanwhile.
    An interrupted deletion leaves the account closing, and asking again finishes it.
    Credit balance (proposed default, awaiting the owner's decision): while credits are on, an
    account holding credit cannot be deleted until PaperAid has refunded it."""
    settings = rt.settings

    def close(w: Wallet) -> Wallet:
        if w.held:
            raise Conflict("One of your jobs is being processed. Delete your account when it finishes.", code="JOB_ACTIVE")
        if settings.credits_enabled and w.available > 0:
            raise Conflict(
                f"Your account still has {credits.tokens(w.available)}. Contact PaperAid to have them refunded before you delete your account.",
                code="ACCOUNT_HAS_CREDIT",
            )
        w.closing = True
        return w

    def reopen(w: Wallet) -> Wallet:
        w.closing = False
        return w

    rt.store.update_wallet(user.uid, user.email, close)
    claimed: list[Job] = []
    for job in every_job(rt, user.uid, None):
        marked = _mark_deleting(rt, job.id)
        if marked is None and rt.store.get(job.id) is not None:
            for done in claimed:
                _unmark_deleting(rt, done.id)
            rt.store.update_wallet(user.uid, user.email, reopen)
            raise Conflict("One of your jobs is being processed. Delete your account when it finishes.", code="JOB_ACTIVE")
        if marked is not None:
            claimed.append(marked)
    for job in claimed:
        _erase(rt, job)
    for _ in range(3):  # drafts created while the claim was being taken (they can hold no credits)
        leftover = [m for j in every_job(rt, user.uid, None) for m in [_mark_deleting(rt, j.id)] if m is not None]
        if not leftover:
            break
        for job in leftover:
            _erase(rt, job)
        claimed += leftover
    # Proposal projects: no step can be running (every job was claimed above) and none can be
    # created (the account is closing), so each is claimed and erased with its chapters and evidence.
    for project in rt.store.list_projects(user.uid):
        _erase_project(rt, project)
    for work in rt.store.list_works(user.uid):
        _erase_work(rt, work)

    def tombstone(w: Wallet) -> Wallet:
        return Wallet(uid=w.uid, email="", closing=True, grant_ops=w.grant_ops)

    rt.store.update_wallet(user.uid, "", tombstone)
    rt.store.delete_ledger(user.uid)
    log(logger, logging.WARNING, "account deleted", uid=user.uid, jobs=len(claimed))
    return len(claimed)


# --- credits ---------------------------------------------------------------------------------


def _wallet_view(rt: Runtime, wallet: Wallet) -> WalletView:
    return WalletView(
        available=wallet.available,
        held=wallet.held,
        entries=list(reversed(wallet.entries[-50:])),
        ugx_per_usd=rt.settings.ugx_per_usd,
        test_credits=rt.settings.env == "local",
    )


class LedgerPage(Camel):
    entries: list[LedgerEntry]
    next: str | None = None  # pass as `before` for the next (older) page


def wallet_history(rt: Runtime, user: User, before: str | None, limit: int = 50) -> LedgerPage:
    """The complete history, newest first, a page at a time (the wallet keeps only recent entries)."""
    limit = max(1, min(limit, 100))
    if before is None:
        rt.store.backfill_ledger(user.uid)  # earlier entries are in before the history is first shown (M12); once only
    entries = rt.store.ledger(user.uid, before, limit + 1)
    more = len(entries) > limit
    entries = entries[:limit]
    return LedgerPage(entries=entries, next=cursor_of(entries[-1]) if more else None)


def backfill_ledgers(rt: Runtime) -> int:
    """Copy the entries wallets kept from before the complete history existed into it, once per
    wallet (Codex audit 56c4f83 M12). Entries older than the wallet's display cap were not kept
    anywhere and cannot be recovered."""
    return sum(rt.store.backfill_ledger(uid) for uid in rt.store.all_wallet_ids())


def my_wallet(rt: Runtime, user: User) -> WalletView:
    """The student's balance. Opening it records their email, so an admin can find them."""
    wallet = rt.store.update_wallet(user.uid, user.email, lambda w: w if w.email or w.closing else w.model_copy(update={"email": user.email}))
    assert wallet is not None
    return _wallet_view(rt, wallet)


# --- admin operations ------------------------------------------------------------------


def admin_grant(rt: Runtime, admin: User, email: str, amount: int, note: str, op_id: str) -> WalletSummary:
    """Add credits by hand: test credits locally, refunds or goodwill in production. Payments will
    add credits through the aggregator instead. `op_id` identifies the request: repeating it (a
    retry after a lost response) adds nothing."""
    term = email.strip().lower()
    match = next((w for w in rt.store.find_wallets(term, 5) if w.email.lower() == term), None)
    if match is None:
        raise NotFound("No PaperAid account with that email has opened its credits yet. Ask them to sign in once, then try again.")
    default_note = "Test credits added by an admin" if rt.settings.env == "local" else "Credits added by PaperAid"
    text = (note.strip() or default_note)[:120]
    wallet = rt.store.update_wallet(match.uid, match.email, lambda w: credits.grant(w, amount, text, op_id, admin.email))
    assert wallet is not None
    log(logger, logging.WARNING, "credits granted", actor=admin.email, amount=amount, uid=wallet.uid, op=op_id)
    return WalletSummary(email=wallet.email, available=wallet.available, held=wallet.held, updated_at=wallet.updated_at)


def admin_wallets(rt: Runtime, search: str | None) -> list[WalletSummary]:
    return [WalletSummary(email=w.email, available=w.available, held=w.held, updated_at=w.updated_at) for w in rt.store.find_wallets(search, 50)]




def _require(rt: Runtime, job_id: str) -> Job:
    job = rt.store.get(job_id) if _valid_id(job_id) else None
    if job is None:
        raise NotFound("Job not found.")
    return job


# Every job field is classified (Codex audit 56c4f83 H01): PAPER_FIELDS can quote or paraphrase the
# student's paper and are emptied by `without_paper_text`; SUPPORT_FIELDS are metadata support may
# see. A test fails when a field is added to Job without being put in one of the two sets.
PAPER_FIELDS = frozenset({"analysis", "analysis_after", "refinement", "research", "references", "paper_checks", "proposal_review", "fix_notes"})
SUPPORT_FIELDS = frozenset(
    {
        "id", "status", "stage", "payment_status", "selection", "services", "pipeline", "source", "guideline", "logo", "quote", "estimate",
        "billing", "outcome", "warnings", "protected", "dismissed", "rejected_changes", "source_job", "latex", "formatting", "scope_words",
        "project_id", "work_id", "outputs", "failure", "created_at", "queued_at", "completed_at", "expires_at",
        "owner_uid", "owner_email", "completed_stages", "generation", "attempts", "lease_until", "cost_usd", "estimate_cost_usd",
        "refine_cost_usd", "budget_usd", "model_calls", "events", "admin_actions", "failure_detail", "files_deleted", "deleting", "retiring",
        "input_sha256", "delivery",
    }
)


def without_paper_text[T: JobView](job: T) -> T:
    """A copy that keeps only support metadata (an allowlist, Codex audit #5): ids, reason codes,
    severities, counts, outcomes and public sources. Every field that can quote or paraphrase the
    student's paper is emptied: excerpts, explanations and suggestions, section headings, before
    and after text, the plan's instructions, reviewer notes, checked claims and the citations quoted
    by paper checks. Used for admin views (support never sees papers) and for jobs past retention
    (the paper must not outlive its files)."""
    update: dict = {}
    for name in ("analysis", "analysis_after"):
        result = getattr(job, name)
        if result is not None:
            def bare(fs: list[Finding]) -> list[Finding]:
                return [f.model_copy(update={"section": "", "excerpt": "", "explanation": "", "suggestion": ""}) for f in fs]

            update[name] = result.model_copy(update={"findings": bare(result.findings), "review": bare(result.review)})
    if job.refinement is not None:
        changes = [ChangedBlock(block_id=c.block_id, kept=c.kept, section="", before="", after="") for c in job.refinement.changes]
        update["refinement"] = job.refinement.model_copy(update={"changes": changes})
    if job.research is not None:
        claims = [c.model_copy(update={"section": "", "claim": "", "note": ""}) for c in job.research.claims]
        update["research"] = job.research.model_copy(update={"claims": claims})
    if job.references is not None:
        items = [r.model_copy(update={"entry": "", "matched_title": "", "note": ""}) for r in job.references.items]
        update["references"] = job.references.model_copy(update={"items": items})
    if job.paper_checks is not None:
        items = [PaperCheck(kind=i.kind, certainty=i.certainty, item="", detail="") for i in job.paper_checks.items]
        update["paper_checks"] = job.paper_checks.model_copy(update={"items": items})
    if job.proposal_review is not None:  # notes, locations and findings can quote the proposal
        review = job.proposal_review
        update["proposal_review"] = review.model_copy(
            update={
                "items": [i.model_copy(update={"note": "", "where": ""}) for i in review.items],
                "findings": [f.model_copy(update={"where": "", "issue": "", "suggestion": ""}) for f in review.findings],
            }
        )
    update["fix_notes"] = {}  # the findings a "Fix selected" draft was made from (explanations, suggestions)
    return job.model_copy(update=update)


def _admin_view(job: Job) -> AdminJob:
    duration = int((job.completed_at - job.created_at).total_seconds()) if job.completed_at else None
    return AdminJob(
        job=without_paper_text(job.view()),
        owner_email=job.owner_email,
        cost_usd=round(job.cost_usd, 6),
        duration_sec=duration,
        model_calls=job.model_calls,
        events=job.events,
        admin_actions=job.admin_actions,
        failure_detail=job.failure_detail,
    )


def admin_summary(rt: Runtime) -> AdminSummary:
    jobs = every_job(rt, None, state.SUBMITTED)
    day_ago = utcnow() - timedelta(days=1)
    recent = [j for j in jobs if j.created_at > day_ago]
    finished = [j for j in jobs if j.status in (JobStatus.COMPLETED, JobStatus.FAILED)]
    return AdminSummary(
        jobs24h=len(recent),
        active=sum(1 for j in jobs if j.status == JobStatus.PROCESSING),
        queued=sum(1 for j in jobs if j.status == JobStatus.QUEUED),
        completion_rate=(sum(1 for j in finished if j.status == JobStatus.COMPLETED) / len(finished)) if finished else 1.0,
        failures24h=sum(1 for j in recent if j.status == JobStatus.FAILED),
        spend24h_usd=round(sum(j.cost_usd for j in recent), 4),
        processing_enabled=processing_enabled(rt),
    )


def admin_list(rt: Runtime, status: str | None, service: str | None, search: str | None, cursor: str | None, limit: int) -> Page[AdminJob]:
    if search:
        term = search.strip().lower()
        if _valid_id(term):
            job = rt.store.get(term)
            return Page[AdminJob](items=[_admin_view(job)] if job else [])
        jobs = every_job(rt, None, state.SUBMITTED)
        matches = [j for j in jobs if term in j.owner_email.lower() or (j.source and term in j.source.name.lower())]
        return Page[AdminJob](items=[_admin_view(j) for j in matches[:limit]])
    statuses = {JobStatus(status)} if status and status != "ALL" else state.SUBMITTED
    jobs, next_cursor = rt.store.list(None, statuses, service if service and service != "ALL" else None, cursor, min(limit, 50))
    return Page[AdminJob](items=[_admin_view(j) for j in jobs], next_cursor=next_cursor)


def admin_get(rt: Runtime, job_id: str) -> AdminJob:
    return _admin_view(_require(rt, job_id))


def admin_retry(rt: Runtime, admin: User, job_id: str) -> AdminJob:
    job = _require(rt, job_id)
    if job.status != JobStatus.FAILED or not (job.failure and job.failure.retryable):
        raise Conflict("Only retryable failures can be retried.", code="NOT_RETRYABLE")

    settings = rt.settings

    def retry(j: Job, w: Wallet) -> tuple[Job, Wallet] | None:
        if j.status != JobStatus.FAILED or j.deleting or j.retiring or j.files_deleted or w.closing:
            return None
        # A charged job's failure returned everything, so a retry needs its own hold, but only
        # while credits are on. A job accepted in testing mode is never charged by a retry.
        if j.billing.state != "NONE" and settings.credits_enabled:
            hold_for_job(j, w)
        elif j.billing.state != "NONE":
            j.billing.state = "NONE"
            j.payment_status = PaymentStatus.NOT_REQUIRED
        j.generation += 1
        j.attempts = 0
        j.failure = None
        j.failure_detail = None
        j.admin_actions.append(AdminAction(actor=admin.email, action="Retried"))
        state.transition(j, JobStatus.QUEUED, "Retried by admin — resumes from last completed stage")
        return j, w

    if job.project_id:
        # A proposal step is retried only onto its live project, claiming it in the same transaction
        # (Codex audit #4: admin retries too). A deleted project, or another step running, refuses.
        project = rt.store.get_project(job.project_id)
        if project is None or project.deleting:
            raise Conflict("This step's proposal was deleted, so it cannot be retried.", code="PROJECT_DELETED")
        if project.active_job != job.id and step_running(rt, project):
            raise Conflict("Another step of this proposal is running. Retry this one when it finishes.", code="STEP_RUNNING")
        seen_active = project.active_job

        def retry_step(j: Job, w: Wallet, p: Project | None) -> tuple[Job, Wallet, Project] | None:
            if p is None or p.deleting or p.active_job not in (seen_active, j.id):
                return None
            p.active_job = j.id
            done = retry(j, w)
            return (done[0], done[1], p) if done else None

        triple = rt.store.update_job_wallet_and_project(job.id, job.project_id, retry_step)
        result = (triple[0], triple[1]) if triple else None
    else:
        result = rt.store.update_job_and_wallet(job.id, retry)
    updated = result[0] if result else None
    if updated is None:
        return admin_get(rt, job_id)
    rt.queue.enqueue(updated.id, task_name(updated))
    return _admin_view(updated)


def set_processing_enabled(rt: Runtime, admin: User, enabled: bool) -> bool:
    rt.store.set_flag("processing_enabled", enabled)
    log(logger, logging.WARNING, "processing switch changed", enabled=enabled, actor=admin.email)
    return enabled


def processing_enabled(rt: Runtime) -> bool:
    return rt.store.get_flag("processing_enabled", rt.settings.processing_enabled)


# --- maintenance --------------------------------------------------------------------------


def reconcile(rt: Runtime) -> int:
    """Re-enqueue jobs that should be running but have no live worker (lost task, restart)."""
    now = utcnow()
    jobs = every_job(rt, None, {JobStatus.QUEUED, JobStatus.PROCESSING})
    stuck = [j for j in jobs if j.lease_until is None or j.lease_until < now]
    for job in stuck:
        rt.queue.enqueue(job.id, task_name(job, f"-r{int(now.timestamp()) // 300}"))
    estimates = [j for j in every_job(rt, None, {JobStatus.DRAFT}) if _estimating(j) and j.estimate and (j.estimate.lease_until is None or j.estimate.lease_until < now)]
    for job in estimates:
        rt.queue.enqueue(job.id, f"{job.id}-e{job.estimate.id}-r{int(now.timestamp()) // 300}" if job.estimate else job.id)
    return len(stuck) + len(estimates)


def cleanup_expired(rt: Runtime) -> int:
    """Delete files of jobs past retention, and every passage of the paper from their records. The
    record stays so history shows 'files expired'. Abandoned drafts and unsubmitted quotes are
    included (Codex audit #11); a draft whose estimate is still running waits for the next run. A
    paid estimate stays paid: the student chose not to go ahead."""
    now = utcnow()
    jobs = every_job(rt, None, {JobStatus.COMPLETED, JobStatus.FAILED, JobStatus.CANCELLED, JobStatus.DRAFT, JobStatus.QUOTED})
    cleaned = 0
    for job in jobs:
        if job.expires_at >= now or job.files_deleted:
            continue

        # Claim first, atomically (Codex audit #11, second round): nothing is deleted while work
        # runs or credits are held, and once claimed no submission, upload, estimate or retry can
        # start (each checks `retiring` in its own transaction). A claim left by an interrupted
        # run is picked up again here.
        def claim(j: Job) -> Job | None:
            if j.files_deleted or j.expires_at >= now or j.status in ACTIVE or _estimating(j) or j.billing.state == "HELD" or j.deleting:
                return None
            j.retiring = True
            return j

        if rt.store.update(job.id, claim) is None:
            continue
        rt.files.delete_prefix(job.storage_prefix())

        def strip_files(j: Job) -> Job:
            j = without_paper_text(j)
            j.files_deleted, j.retiring = True, False
            j.events.append(JobEvent(label="Files deleted under retention policy"))
            return j

        rt.store.update(job.id, strip_files)
        cleaned += 1
    return cleaned


def step_running(rt: Runtime, p: Project | Work) -> bool:
    """A step of this project (or work) is queued or running. Its claim (`active_job`) is set in the
    same transaction that holds the step's credits (app.jobs.service.submit)."""
    return bool(p.active_job) and (step := rt.store.get(p.active_job or "")) is not None and step.status in ACTIVE


def _erase_project(rt: Runtime, project: Project, expired_before: datetime | None = None) -> bool:
    """Delete a proposal project: its record, its files and every job run for it (their frozen
    inputs and saved AI answers hold the student's details; Codex audit 2026-09-28 #6).

    1. The project is claimed atomically. Refused while a step is queued or running, and, for
       expiry, if the student renewed it after it was listed (#5). A claim left by an interrupted
       deletion is resumed.
    2. Each child job, found by project id (not the project's own bounded list), is claimed and
       erased. One that cannot be claimed yet leaves the project claimed, to be finished next time.
    3. The files and the record go last."""

    def mark(p: Project) -> Project | None:
        if p.deleting:
            return p  # an interrupted deletion: finish it
        if expired_before is not None and p.expires_at >= expired_before:
            return None  # renewed since it was listed
        if step_running(rt, p):
            return None
        p.deleting = True
        return p

    if rt.store.update_project(project.id, mark) is None:
        return False
    for job in every_job(rt, project.owner_uid, None):
        if job.project_id != project.id:
            continue
        claimed = _mark_deleting(rt, job.id)
        if claimed is None:
            if rt.store.get(job.id) is not None:
                return False  # cannot be erased yet (a hold is still settling): resumed next time
            continue
        _erase(rt, claimed)
    rt.files.delete_prefix(project.storage_prefix())
    rt.files.delete_prefix(project.legacy_prefix())  # a project not yet migrated (Codex audit 56c4f83 H04)
    guide_prefix = f"users/{project.owner_uid}/guides/{project.id}"
    rt.files.delete_prefix(guide_prefix)
    rt.files.delete_prefix(guide_prefix + "-", flat=True)  # copies uploaded before guides had their own directory
    if project.guide is not None:
        rt.files.delete(project.guide.path)  # stored outside the project's prefix
    for profile_id in project.profiles:  # institution profiles built from this project's guide
        rt.files.delete(f"rulebooks/{profile_id}.json")
    rt.store.delete_project(project.id)
    return True


def migrate_legacy_project_files(rt: Runtime) -> int:
    """Move projects created before the storage change out of users/, where the bucket's fixed-age
    rule would delete their files (Codex audit 56c4f83 H04). Resumable and safe to repeat: files are
    copied first, the record's paths switch in one transaction, then the old files are removed.
    A project with a step running is left for the next run; steps priced earlier still find their
    files (reads fall back to the moved copy)."""
    migrated = 0
    for project in rt.store.all_projects():
        old = project.legacy_paths()
        if not old or project.deleting or step_running(rt, project):
            continue
        for path in old:
            target = moved_path(path)
            if rt.files.exists(path) and not rt.files.exists(target):
                rt.files.put(target, rt.files.get(path), "application/json")

        def switch(p: Project) -> Project | None:
            if p.deleting:
                return None
            for chapter in p.chapters:
                for version in chapter.versions:
                    version.path = moved_path(version.path)
            p.evidence_files = [moved_path(f) for f in p.evidence_files]
            if p.guide is not None:
                p.guide.path = moved_path(p.guide.path)
            return p

        if rt.store.update_project(project.id, switch) is not None:
            rt.files.delete_prefix(project.legacy_prefix())
            migrated += 1
    if migrated:
        log(logger, logging.INFO, "legacy projects migrated", projects=migrated)
    return migrated


def cleanup_expired_projects(rt: Runtime) -> int:
    """Delete proposal projects 30 days after the student's last action (owner decision
    2026-09-28). The expiry is shown on the project, so it is never a surprise. Pages through every
    expired project; one that is refused (renewed, or a step running) is skipped this time."""
    cutoff, erased = utcnow(), 0
    # Every expired id is listed first, so projects that cannot be erased yet never hide the ones
    # after them (Codex audit 56c4f83 M13).
    for project_id in rt.store.expired_project_ids(cutoff):
        project = rt.store.get_project(project_id)
        if project is not None:
            erased += _erase_project(rt, project, expired_before=cutoff)
    return erased


def _erase_work(rt: Runtime, work: Work, expired_before: datetime | None = None) -> bool:
    """Delete a work: its record, its files, and every job run for it (their frozen inputs and
    saved AI answers hold the student's text). Claimed first, like a proposal project."""

    def mark(k: Work) -> Work | None:
        if k.deleting:
            return k  # an interrupted deletion: finish it
        if expired_before is not None and k.expires_at >= expired_before:
            return None  # renewed since it was listed
        if step_running(rt, k):
            return None
        k.deleting = True
        return k

    if rt.store.update_work(work.id, mark) is None:
        return False
    for job in every_job(rt, work.owner_uid, None):
        if job.work_id != work.id:
            continue
        claimed = _mark_deleting(rt, job.id)
        if claimed is None:
            if rt.store.get(job.id) is not None:
                return False  # cannot be erased yet (a hold is still settling): resumed next time
            continue
        _erase(rt, claimed)
    rt.files.delete_prefix(work.storage_prefix())
    rt.store.delete_work(work.id)
    return True


def cleanup_expired_works(rt: Runtime) -> int:
    """Delete works 30 days after the student's last action, like proposal projects."""
    cutoff, erased = utcnow(), 0
    for work_id in rt.store.expired_work_ids(cutoff):
        work = rt.store.get_work(work_id)
        if work is not None:
            erased += _erase_work(rt, work, expired_before=cutoff)
    return erased


def ensure_valid_upload_name(filename: str) -> None:
    if not filename or len(filename) > 255:
        raise InvalidDocument("That file name isn't valid.", code="INVALID_NAME")
