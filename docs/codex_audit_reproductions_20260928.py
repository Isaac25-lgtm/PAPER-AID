"""Audit diagnostics: passing means the documented defect was reproduced on 655cda5.

Run from backend: python -m pytest ../docs/codex_audit_reproductions_20260928.py
  -q -s --basetemp=.codex_repro_tmp_20260928 -p no:cacheprovider
Uses the project's isolated, test-only models. Never calls a real AI provider.
These are NOT acceptance tests and are kept outside the application's test suite.
"""

import io
from datetime import timedelta

from docx import Document

from app.jobs.models import utcnow
from app.proposals import decisions, evidence, export, pipeline
from app.proposals.ai import ProposalRunner
from app.proposals.models import ChapterDocument, ChapterSection, ProposalInputs, ReadinessItem, StepInput
from app.runtime import get_runtime
from tests import conftest as audit_fixtures
from tests.fake_models import PLAN
from tests.test_api import STUDENT
from tests.test_proposals import DETAILS, LIB, _approved, _create, _run, plan

client = audit_fixtures.client  # expose the existing isolated pytest fixture


def chapter_setup(client):
    project = _create(client)
    _run(client, project["id"], "PLAN")
    _approved(client, project["id"], sampleSize=PLAN["sampleSize"])
    return project["id"]


def test_table_bypasses_checks_and_reaches_word(client):
    pid = chapter_setup(client)
    original = client.models.default

    def bad_table(payload):
        answer = original("p_draft", payload)
        for section in answer["sections"]:
            if section["key"] == "workplan":
                section["table"] = {"caption": "Smith (2019) found 73%", "rows": [["Activity", "Finding"], ["Interview", "73% ⟦E999999⟧"]]}
        return answer

    client.models.overrides["p_draft"] = bad_table
    job = _run(client, pid, "CHAPTER_3")
    downloaded = client.get(f"/api/projects/{pid}/export", headers=STUDENT)
    assert downloaded.status_code == 200
    doc = Document(io.BytesIO(downloaded.content))
    cells = " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "73% ⟦E999999⟧" in cells
    assert job["outcome"] == "FULL"
    print("TABLE: FULL output retains invented 73% and unresolved E999999 in Word.")


def test_final_fix_is_never_reviewed_and_three_fixes_run(client, monkeypatch):
    pid = chapter_setup(client)
    client.models.overrides["p_review"] = lambda p: {
        "results": [{"key": s["key"], "grade": "REPAIR", "issues": ["Correct the wording."], "note": ""} for s in p["sections"]]
    }
    sequence = []
    original_grade, original_fix = ProposalRunner.grade, ProposalRunner.fix

    def grade(self, *args, **kwargs):
        sequence.append("p_review")
        return original_grade(self, *args, **kwargs)

    def fix(self, *args, **kwargs):
        sequence.append("p_fix")
        return original_fix(self, *args, **kwargs)

    monkeypatch.setattr(ProposalRunner, "grade", grade)
    monkeypatch.setattr(ProposalRunner, "fix", fix)
    _run(client, pid, "CHAPTER_1")
    # Trace orchestration calls, including repeated inputs served from the paid-call cache.
    rounds = [t for i, t in enumerate(sequence) if i == 0 or t != sequence[i - 1]]
    assert rounds == ["p_review", "p_fix"] * 3, rounds
    print("REPAIR: review/fix runs three rounds, ends in an unreviewed fix:", rounds)


def test_unreviewed_budget_exhausted_chapter_is_full(client, monkeypatch):
    pid = chapter_setup(client)

    def no_review(self, *args, **kwargs):
        self.budget_reached = True
        return {}

    monkeypatch.setattr(ProposalRunner, "grade", no_review)
    job = _run(client, pid, "CHAPTER_1")
    assert job["outcome"] == "FULL", job
    view = client.get(f"/api/projects/{pid}/chapters/1", headers=STUDENT).json()
    assert view["warnings"] == []
    print("BUDGET: no section review, yet outcome FULL and no warnings.")


def test_reviewer_warning_is_not_delivered_to_the_student(client):
    pid = chapter_setup(client)
    client.models.overrides["p_review"] = lambda p: {
        "results": [{"key": s["key"], "grade": "PASS_WITH_WARNINGS", "issues": [], "note": "AUDIT: scope needs author confirmation."} for s in p["sections"]]
    }
    job = _run(client, pid, "CHAPTER_1")
    chapter = client.get(f"/api/projects/{pid}/chapters/1", headers=STUDENT).json()
    assert chapter["warnings"] == [] and job["warnings"] == []
    assert "scope needs author confirmation" not in str(chapter)
    print("REVIEW WARNING: PASS_WITH_WARNINGS note is discarded from both the chapter and job.")


def test_final_blockers_accept_empty_chapters_with_missing_readiness(client):
    pid = chapter_setup(client)
    project = get_runtime().store.get_project(pid)
    chapters = {}
    for n in (1, 2, 3):
        project.chapter(n).approved = True
        chapters[n] = ChapterDocument(
            number=n,
            title="Chapter",
            plan_version=project.plan_version,
            sections=[],
            cited=[],
            words=0,
            readiness=[ReadinessItem(id="MISSING", question="Required section", basis="CODE", status="MISSING", note="Absent.", chapter=n)],
        )
    assert export.final_blockers(project, chapters) == []
    data = export.build(project, chapters, {}, draft=False)
    doc = Document(io.BytesIO(data))
    assert not any("draft" in p.text.lower() for p in doc.paragraphs)
    print("FINAL: three empty, approved chapters with MISSING checks pass the complete-export gate.")


def test_project_delete_between_claim_and_wallet_submit(client, monkeypatch):
    pid = _create(client)["id"]
    rt = get_runtime()
    quoted = client.post(f"/api/projects/{pid}/steps", headers=STUDENT, json={"step": "PLAN"}).json()
    jobid = quoted["job"]["id"]
    original = rt.store.update_job_and_wallet

    def race(jid, mutate):
        # This is the exact boundary after the project claim and before paid acceptance.
        response = client.delete(f"/api/projects/{pid}", headers=STUDENT)
        assert response.status_code == 204, response.text
        return original(jid, mutate)

    monkeypatch.setattr(rt.store, "update_job_and_wallet", race)
    monkeypatch.setattr(rt.queue, "enqueue", lambda *a, **k: None)
    response = client.post(f"/api/projects/{pid}/steps/{jobid}/submit", headers=STUDENT, json={"quoteId": quoted["quote"]["id"]})
    assert response.status_code == 200, response.text
    assert rt.store.get_project(pid) is None
    job = rt.store.get(jobid)
    assert job.status.value == "QUEUED" and job.billing.state == "HELD"
    print("SUBMIT RACE: deleted project has a newly QUEUED job and HELD credits.")


def test_expiry_deletes_a_project_renewed_after_listing(client, monkeypatch):
    from app.jobs.service import cleanup_expired_projects

    pid = _create(client)["id"]
    rt = get_runtime()

    def expire(p):
        p.expires_at = utcnow() - timedelta(seconds=1)
        return p

    rt.store.update_project(pid, expire)
    original = rt.store.expired_projects

    def race(now, limit):
        listed = original(now, limit)
        response = client.post(f"/api/projects/{pid}/details", headers=STUDENT, json=DETAILS)
        assert response.status_code == 200
        assert rt.store.get_project(pid).expires_at > utcnow()
        return listed

    monkeypatch.setattr(rt.store, "expired_projects", race)
    assert cleanup_expired_projects(rt) == 1
    assert rt.store.get_project(pid) is None
    print("RETENTION: a freshly renewed project is erased from a stale expiry listing.")


def test_project_delete_leaves_private_child_job_input(client):
    pid = _create(client)["id"]
    job = _run(client, pid, "PLAN")
    rt = get_runtime()
    path = rt.store.get(job["id"]).storage_prefix() + "/internal/proposal_input.json"
    before = rt.files.get(path)
    assert b"Namukasa" in before
    assert client.delete(f"/api/projects/{pid}", headers=STUDENT).status_code == 204
    assert rt.files.get(path) == before
    assert rt.store.get(job["id"]) is not None
    print("DELETION: child job retains private frozen input after project deletion.")


def test_new_objective_does_not_mark_previous_objectives_section_stale():
    before = plan()
    doc = ChapterDocument(
        number=1,
        title="Introduction",
        plan_version=1,
        cited=[],
        words=1,
        sections=[
            ChapterSection(
                key="objectives", number="1.4", heading="Objectives", paragraphs=["Three objectives"], depends=decisions.stamp(1, "objectives", before)
            )
        ],
    )
    after = plan(
        specificObjectives=PLAN["specificObjectives"] + ["To examine vaccine affordability"],
        researchQuestions=PLAN["researchQuestions"] + ["How does affordability affect uptake?"],
        alignment=PLAN["alignment"] + [{"objective": 4, "data": "Costs", "collection": "Questionnaire", "analysis": "Regression"}],
    )
    assert "O4" in decisions.changed(before, after)
    assert decisions.stale(doc, after) == []
    print("DEPENDENCIES: adding objective 4 leaves the old three-objective section current.")


def test_apa6_first_narrative_citation_is_abbreviated_incorrectly():
    rendered = evidence.Citer(LIB, "APA6").render("⟦E00000b|n⟧ found low uptake. It remains low ⟦E00000b⟧.")
    assert rendered.startswith("Okello et al. (2022)"), rendered
    assert "(Okello, Namara, & Kato, 2022)" in rendered
    print("APA6:", rendered)


def test_a_study_year_is_accepted_as_author_population():
    inp = StepInput(
        project_id="prj_a",
        step="PLAN",
        chapter=0,
        rulebook="ucu-2018-v1",
        inputs=ProposalInputs.model_validate({**DETAILS["inputs"], "notes": "The study will be undertaken in 2026."}),
        plan_version=0,
    )
    generated = plan(sampleSize={**PLAN["sampleSize"], "population": 2026})
    guarded = pipeline._student_figures_only(generated, inp)
    assert guarded.sample_size.population == 2026
    print("AUTHOR FIGURES: a study year, 2026, is accepted as an author-supplied population N.")


def test_source_fetch_re_resolves_a_host_to_a_private_address(monkeypatch):
    """Real local socket reproduction: no external network or cloud endpoint is contacted."""
    import socket
    import threading
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

    from app.analysis import fetch

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            self.send_response(200)
            self.send_header("Content-Type", "text/plain")
            self.end_headers()
            self.wfile.write(b"private local audit endpoint")

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    original = socket.getaddrinfo
    resolutions = []

    def rebind(host, port, *args, **kwargs):
        if host in ("audit-rebind.invalid", b"audit-rebind.invalid"):
            resolutions.append(port)
            return original("93.184.216.34" if port is None else "127.0.0.1", port, *args, **kwargs)
        return original(host, port, *args, **kwargs)

    monkeypatch.setattr(socket, "getaddrinfo", rebind)
    for name in ("NO_PROXY", "no_proxy"):
        monkeypatch.setenv(name, "audit-rebind.invalid")
    try:
        got = fetch._get(f"http://audit-rebind.invalid:{server.server_port}/private")
        assert got is not None and got[0] == b"private local audit endpoint", got
        assert len(resolutions) >= 2
        print("SSRF: hostname passed the public-address check, then a second DNS lookup fetched loopback.")
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def test_formatting_erases_a_later_sections_distinct_header():
    from docx.enum.section import WD_SECTION

    from app.documents.docx_io import read_docx
    from app.formatting.apply import apply_formatting
    from app.formatting.presets import PRESETS

    doc = Document()
    doc.add_paragraph("The original body wording is unchanged.")
    doc.add_section(WD_SECTION.NEW_PAGE)
    later = doc.sections[1]
    later.header.is_linked_to_previous = False
    later.header.paragraphs[0].text = "Distinct chapter header that must survive"
    later.footer.is_linked_to_previous = False
    later.footer.paragraphs[0].text = "Distinct chapter footer that must survive"
    data = io.BytesIO()
    doc.save(data)
    source = data.getvalue()
    output, result = apply_formatting(source, PRESETS["apa7"], read_docx(source))
    produced = Document(io.BytesIO(output))
    position = PRESETS["apa7"].page_numbers
    container = produced.sections[1].header if position.startswith("top") else produced.sections[1].footer
    assert "Distinct chapter" not in " ".join(p.text for p in container.paragraphs)
    assert result.body_text_unchanged
    assert not any("left it unchanged" in w for w in result.warnings)
    print("FORMATTING: a later section loses its distinct header/footer while wording is reported unchanged.")
