"""Diagnostics for 56c4f83. PASS means the faulty behavior was reproduced, not fixed.

Run explicitly from backend with a NEW, short --basetemp and -p no:cacheprovider.
Isolated fixtures use tests.fake_models; no real keys or owner data are used.
"""

import hashlib
import io
import json
from datetime import timedelta
from pathlib import Path

import pytest
from docx import Document

from app.analysis.fetch import crossref_search as real_crossref_search
from app.documents.docx_io import read_docx
from app.integrations.store import LocalJobStore
from app.jobs.models import (
    AnalysisResult,
    Finding,
    Job,
    JobStatus,
    LedgerEntry,
    ReadinessItem,
    ServiceSelection,
    StoredFile,
    Wallet,
    utcnow,
)
from app.pricing import credits
from app.pricing.quote import price
from app.proposals import decisions, export, profile, rulebook, service
from app.proposals.models import ChapterDocument, ChapterSection, StoredChapterVersion
from app.runtime import get_runtime
from tests import conftest as fixtures
from tests.fake_models import profile_answer
from tests.test_api import ADMIN, STUDENT, get_quote, wait
from tests.test_formatting_options import _png
from tests.test_proposals import LIB, _create, _run, plan

client = fixtures.client
fixed_client = fixtures.fixed_client
UID = "u_" + hashlib.sha256(b"student@example.com").hexdigest()[:20]


def _paper(rt, *, words=None, scope=None, notes=None):
    data = fixtures.fixture_bytes("simple_essay.docx")
    model = read_docx(data)
    job = Job(
        id="job_auditpaper",
        status=JobStatus.DRAFT,
        owner_uid=UID,
        owner_email="student@example.com",
        expires_at=utcnow() + timedelta(days=30),
        scope_words=scope,
        fix_notes=notes or {},
    )
    path = f"{job.storage_prefix()}/input/source.docx"
    rt.files.put(path, data, "application/octet-stream")
    job.source = StoredFile(
        name="essay.docx",
        format="DOCX",
        size_bytes=len(data),
        word_count=words or model.word_count,
        page_estimate=1,
        heading_count=1,
        path=path,
        sha256=hashlib.sha256(data).hexdigest(),
    )
    rt.store.create(job)
    return job, model


def _project(client, number=1, *, unreviewed=False):
    rt = get_runtime()
    pid = _create(client)["id"]
    p = rt.store.get_project(pid)
    approved = plan()
    sections = [
        ChapterSection(
            key=s.key,
            number=s.number,
            heading=s.heading,
            paragraphs=[f"The proposed study will examine caregiver access in {s.heading}."],
            depends=decisions.stamp(number, s.key, approved),
        )
        for s in rulebook.sections(p.rulebook, number, p.inputs.level, approved)
    ]
    chapter = ChapterDocument(
        number=number,
        title=rulebook.chapter_spec(p.rulebook, number)["title"],
        plan_version=1,
        sections=sections,
        cited=[],
        words=200,
        warnings=["Background was not reviewed."] if unreviewed else [],
        readiness=[
            ReadinessItem(
                id=f"C{number}-REVIEWED",
                question="Every section reviewed",
                basis="CODE",
                status="NEEDS_REVIEW" if unreviewed else "PASS",
                note="Background unreviewed" if unreviewed else "",
                chapter=number,
            )
        ],
    )
    path = f"{p.storage_prefix()}/chapters/{number}/seed.json"
    rt.files.put(path, chapter.model_dump_json(by_alias=True).encode(), "application/json")

    def set_up(q):
        q.plan, q.plan_status, q.plan_version = approved, "APPROVED", 1
        state = q.chapter(number)
        state.current, state.approved = 1, True
        state.versions = [StoredChapterVersion(version=1, job_id="job_seed", words=200, plan_version=1, path=path)]
        return q

    rt.store.update_project(pid, set_up)
    return pid, chapter


def _comment(client, pid, chapter=1, key="problem"):
    response = client.post(f"/api/projects/{pid}/feedback", headers=STUDENT, json={"text": "Clarify the problem statement."})
    assert response.status_code == 200, response.text
    cid = response.json()["feedback"][-1]["id"]
    placed = client.post(f"/api/projects/{pid}/feedback/{cid}", headers=STUDENT, json={"chapter": chapter, "sections": [key]})
    assert placed.status_code == 200, placed.text
    return cid


def test_account_deletion_retains_durable_history(client):
    rt = get_runtime()
    assert rt.store.ledger(UID, None, 10)
    rt.settings.credits_enabled = False
    assert client.delete("/api/me", headers=STUDENT).status_code == 204
    assert rt.store.get_wallet(UID).closing
    assert rt.store.ledger(UID, None, 10)
    print("DELETION: account closed, but its durable ledger is still present.")


def test_admin_view_exposes_fix_notes(client):
    rt = get_runtime()
    job, _ = _paper(rt, notes={"b00001": ["PRIVATE author sentence and confidential fix instruction"]})
    response = client.get(f"/api/admin/jobs/{job.id}", headers=ADMIN)
    assert response.status_code == 200
    assert "PRIVATE author sentence" in json.dumps(response.json()["job"]["fixNotes"])
    from app.jobs.service import without_paper_text

    assert without_paper_text(job).fix_notes == job.fix_notes
    print("PRIVACY: admin/retention sanitizer preserves confidential fixNotes.")


def test_fix_scope_can_be_expanded_without_repricing_words(fixed_client):
    rt = get_runtime()
    job, _ = _paper(rt)
    doc = Document()
    for i in range(30):
        doc.add_paragraph(" ".join([f"passage{i}"] * 150))
    out = io.BytesIO()
    doc.save(out)
    model = read_docx(out.getvalue())
    rt.files.put(job.source.path, out.getvalue(), "application/octet-stream")
    job.source.word_count = model.word_count
    job.source.sha256 = hashlib.sha256(out.getvalue()).hexdigest()
    block_ids = [b.id for b in model.blocks if b.editable and b.kind == "paragraph"]
    job.status = JobStatus.COMPLETED
    job.analysis = AnalysisResult(
        band="LOW",
        confidence="HIGH",
        analysed_words=model.word_count,
        excluded_words=0,
        findings=[
            Finding(
                id="f-one",
                block_id=block_ids[0],
                section="Body",
                reason="GENERIC_PHRASING",
                severity="minor",
                excerpt="",
                explanation="Clarify the passage",
                suggestion="Be specific",
            )
        ],
        algorithm_version="audit",
    )
    rt.store.update(job.id, lambda j: job)
    rt.files.put(f"{job.storage_prefix()}/internal/document.json", model.model_dump_json(by_alias=True).encode(), "application/json")
    created = fixed_client.post(f"/api/jobs/{job.id}/fix", headers=STUDENT, json={"findingIds": ["f-one"], "safeOnly": False})
    assert created.status_code == 200, created.text
    job = rt.store.get(created.json()["id"])
    assert job.scope_words == 150 and model.word_count == 4500
    selection = {"writing": "REFINE", "academic": False, "onlyBlocks": [block_ids[0]]}
    first = fixed_client.post(f"/api/jobs/{job.id}/quote", headers=STUDENT, json={"selection": selection})
    assert first.status_code == 200
    expanded = {**selection, "onlyBlocks": block_ids}
    second = fixed_client.post(f"/api/jobs/{job.id}/quote", headers=STUDENT, json={"selection": expanded})
    assert second.status_code == 200 and second.json()["quote"]["amount"] == first.json()["quote"]["amount"] == 4000
    assert rt.store.get(job.id).selection.only_blocks == expanded["onlyBlocks"]
    assert price(rt.settings, ServiceSelection(writing="REFINE", academic=False), 4500).lines[0].amount == 7000
    print("SCOPE: an actual 150-word fix draft is expanded to 4500 valid words and stays 4000, instead of 7000.")


def test_workspace_rebuild_loses_uploaded_logo(fixed_client):
    rt = get_runtime()
    job, _ = _paper(rt)
    assert fixed_client.post(f"/api/jobs/{job.id}/files/logo", headers=STUDENT, files={"file": ("logo.png", _png(), "image/png")}).status_code == 200
    selection = {"writing": "REFINE", "academic": False, "formatting": "FORMAT", "preset": "apa7", "logo": "CENTER"}
    quote = get_quote(fixed_client, job.id, selection)
    assert fixed_client.post(f"/api/jobs/{job.id}/submit", headers=STUDENT, json={"quoteId": quote["id"]}).status_code == 200
    assert wait(fixed_client, job.id, timeout=120)["status"] == "COMPLETED"
    original = Document(io.BytesIO(fixed_client.get(f"/api/jobs/{job.id}/outputs/paper", headers=STUDENT).content))
    assert len(original.inline_shapes) == 1
    assert fixed_client.post(f"/api/jobs/{job.id}/rebuild", headers=STUDENT).status_code == 200
    rebuilt = Document(io.BytesIO(fixed_client.get(f"/api/jobs/{job.id}/outputs/paper-reviewed", headers=STUDENT).content))
    assert len(rebuilt.inline_shapes) == 0
    print("LOGO: ordinary output has the logo; download with choices has none.")


def test_browser_plan_accepts_unconfirmed_gap_evidence(client):
    pid = _create(client)["id"]
    edited = plan(researchGap={"known": "Established", "missing": "Unknown locally", "contribution": "Measure locally", "evidence": ["E999999", "not-an-id"]})
    saved = client.post(f"/api/projects/{pid}/plan", headers=STUDENT, json={"plan": edited.model_dump(by_alias=True), "baseVersion": 0})
    assert saved.status_code == 200 and saved.json()["plan"]["researchGap"]["evidence"] == ["E999999", "not-an-id"]
    assert client.post(f"/api/projects/{pid}/plan/approve", headers=STUDENT, json={"baseVersion": 1}).status_code == 200
    print("GAP: unknown evidence ids survive saving and plan approval.")


def test_empty_revision_marks_comment_applied_and_charges(fixed_client):
    pid, base = _project(fixed_client)
    cid = _comment(fixed_client, pid)
    fixed_client.models.overrides["p_fix"] = lambda p: {
        "sections": [{"key": s["key"], "paragraphs": [], "table": {"caption": "", "rows": []}} for s in p["sections"]]
    }
    job = _run(fixed_client, pid, "REVISE_1")
    assert job["status"] == "COMPLETED" and job["outcome"] == "PARTIAL"
    chapter = fixed_client.get(f"/api/projects/{pid}/chapters/1", headers=STUDENT).json()
    original = next(s.paragraphs for s in base.sections if s.key == "problem")
    assert next(s["paragraphs"] for s in chapter["sections"] if s["key"] == "problem") == original
    feedback = fixed_client.get(f"/api/projects/{pid}", headers=STUDENT).json()["feedback"]
    assert next(c for c in feedback if c["id"] == cid)["status"] == "APPLIED"
    assert job["billing"]["charged"] == 1500
    print("REVISION: earlier text kept, comment marked APPLIED, 1.5 of 2 tokens charged.")


def test_revision_erases_untouched_review_failure(client):
    pid, _ = _project(client, unreviewed=True)
    _comment(client, pid)
    job = _run(client, pid, "REVISE_1")
    assert job["status"] == "COMPLETED"
    chapter = client.get(f"/api/projects/{pid}/chapters/1", headers=STUDENT).json()
    assert next(r for r in chapter["readiness"] if r["id"] == "C1-REVIEWED")["status"] == "PASS"
    assert not any("Background" in w for w in chapter["warnings"])
    print("REVIEW: revising only the problem clears the untouched background's review failure.")


def test_table_only_revision_is_invisible_in_comparison(client):
    pid, base = _project(client, number=3)
    rt = get_runtime()
    table = next(s for s in base.sections if s.key == "workplan")
    table.table, table.table_caption = [["Task", "Month"], ["Pilot", "1"]], "Original schedule"
    project = rt.store.get_project(pid)
    first = project.chapter(3).versions[0].path
    rt.files.put(first, base.model_dump_json(by_alias=True).encode(), "application/json")
    newer = base.model_copy(deep=True)
    revised = next(s for s in newer.sections if s.key == "workplan")
    revised.table[1][1], revised.table_caption = "6", "Changed schedule"
    second = first.replace("seed.json", "second.json")
    rt.files.put(second, newer.model_dump_json(by_alias=True).encode(), "application/json")

    def append(q):
        q.chapter(3).versions.append(StoredChapterVersion(version=2, job_id="job_second", words=200, plan_version=1, path=second))
        return q

    rt.store.update_project(pid, append)
    response = client.get(f"/api/projects/{pid}/chapters/3/compare?older=1&newer=2", headers=STUDENT)
    assert response.status_code == 200 and response.json()["changed"] == 0
    print("COMPARISON: a changed schedule and caption report zero changes.")


def test_profile_numeric_error_is_not_classified_as_not_a_guide():
    answer = profile_answer({"guide": "research proposal", "reference": profile.reference()})
    answer["formatting"]["sizePt"] = "twelve"
    with pytest.raises(ValueError) as caught:
        profile.build(answer, "guide.docx")
    assert not isinstance(caught.value, profile.NotAGuide)
    print("PROFILE: invalid numeric answer raises ordinary ValueError, outside NOT_A_GUIDE handling.")


def test_missing_custom_rulebook_breaks_project_and_list(client):
    rt = get_runtime()
    pid = _create(client)["id"]
    book = profile.build(profile_answer({"guide": "research proposal", "reference": profile.reference()}), "guide.docx")
    path = rulebook.stored_path(book["id"])
    rt.files.put(path, json.dumps(book).encode(), "application/json")
    rt.store.update_project(pid, lambda p: p.model_copy(update={"rulebook": book["id"]}))
    assert client.get(f"/api/projects/{pid}", headers=STUDENT).status_code == 200
    rt.files.delete(path)
    rulebook.load.cache_clear()
    with pytest.raises(FileNotFoundError):
        client.get(f"/api/projects/{pid}", headers=STUDENT)
    with pytest.raises(FileNotFoundError):
        client.get("/api/projects", headers=STUDENT)
    print("MISSING PROFILE: one missing rulebook crashes its view and the owner's whole project list.")


def test_ledger_cursor_loses_entries_with_equal_times(client):
    rt = get_runtime()
    at = utcnow()

    def add(w):
        w.entries += [
            LedgerEntry(id=f"le_same{i:03}", kind="TOP_UP", amount=1000, at=at, note="same-time", available_after=1000, held_after=0) for i in range(60)
        ]
        return w

    rt.store.update_wallet(UID, "student@example.com", add)
    first = client.get("/api/wallet/history", headers=STUDENT).json()
    assert len([e for e in first["entries"] if e["note"] == "same-time"]) == 50 and first["next"]
    second = client.get("/api/wallet/history", headers=STUDENT, params={"before": first["next"]}).json()
    assert not any(e["note"] == "same-time" for e in second["entries"])
    print("CURSOR: ten entries sharing the first page's timestamp disappear from later pages.")


def test_existing_wallet_history_is_not_backfilled(tmp_path):
    rt = LocalJobStore(tmp_path)
    w = Wallet(uid="legacy", email="legacy@example.com")
    credits.top_up(w, 4000, "before-upgrade")
    rt._wallet_path(w.uid).write_text(w.model_dump_json(by_alias=True), encoding="utf-8")
    assert rt.get_wallet(w.uid).entries
    assert rt.ledger(w.uid, None, 50) == []
    rt.update_wallet(w.uid, w.email, lambda old: old)
    assert rt.ledger(w.uid, None, 50) == []
    print("MIGRATION: a pre-upgrade wallet has entries, but complete history stays empty.")


def test_local_wallet_failure_leaves_phantom_ledger_entry(tmp_path, monkeypatch):
    rt = LocalJobStore(tmp_path)
    rt.update_wallet("u", "u@example.com", lambda w: credits.top_up(w, 1000, "committed"))
    original = Path.replace

    def fail_wallet_replace(path, target):
        if Path(target) == rt._wallet_path("u"):
            raise OSError("synthetic interrupted wallet replace")
        return original(path, target)

    monkeypatch.setattr(Path, "replace", fail_wallet_replace)
    with pytest.raises(OSError):
        rt.update_wallet("u", "u@example.com", lambda w: credits.top_up(w, 5000, "never-committed"))
    assert rt.get_wallet("u").available == 1000
    assert any(e.note == "never-committed" for e in rt.ledger("u", None, 50))
    print("ATOMICITY: wallet stays at 1000 but history records the failed 5000 top-up.")


def test_list_of_tables_leaks_unrendered_evidence_tokens(client):
    pid, chapter = _project(client, number=3)
    rt = get_runtime()
    project = rt.store.get_project(pid)
    table = next(s for s in chapter.sections if s.key == "workplan")
    table.table, table.table_caption = [["Task", "Month"], ["Pilot", "1"]], "Schedule informed by ⟦E00000b⟧"
    data = export.build(project, {3: chapter}, LIB, draft=True)
    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert "Table 3.1: Schedule informed by ⟦E00000b⟧" in text
    assert "Schedule informed by (Okello, Namara, & Kato, 2022)" in text
    print("EXPORT: caption rendered in body, but raw evidence token appears in List of Tables.")


def test_truncated_png_is_an_unhandled_server_exception(client):
    from docx.image.exceptions import UnexpectedEndOfFileError

    job, _ = _paper(get_runtime())
    with pytest.raises(UnexpectedEndOfFileError):
        client.post(f"/api/jobs/{job.id}/files/logo", headers=STUDENT, files={"file": ("bad.png", b"\x89PNG\r\n\x1a\n", "image/png")})
    print("LOGO VALIDATION: a PNG signature without image data raises an unhandled exception.")


def test_feedback_add_has_no_per_user_quota(client):
    rt = get_runtime()
    rt.settings.quotes_per_hour = 1
    pid = _create(client)["id"]
    for _ in range(4):
        assert client.post(f"/api/projects/{pid}/feedback", headers=STUDENT, json={"text": "Clarify this."}).status_code == 200
    print("QUOTA: repeated feedback mutations bypass the configured per-user quota.")


def test_late_guide_upload_leaves_private_file_after_deletion(client, monkeypatch):
    rt = get_runtime()
    pid = _create(client)["id"]
    prefix = rt.store.get_project(pid).storage_prefix()
    original_put = rt.files.put
    paths = []

    def racing_put(path, data, content_type):
        if path.startswith(prefix + "/guide/"):
            from app.core.auth import User

            service.delete(rt, User(uid=UID, email="student@example.com", is_admin=False), pid)
            assert rt.store.get_project(pid) is None
            paths.append(path)
        return original_put(path, data, content_type)

    monkeypatch.setattr(rt.files, "put", racing_put)
    doc = Document()
    doc.add_paragraph("PRIVATE institution research guide " * 120)
    out = io.BytesIO()
    doc.save(out)
    response = client.post(f"/api/projects/{pid}/guide", headers=STUDENT, files={"file": ("guide.docx", out.getvalue(), "application/octet-stream")})
    assert response.status_code == 404, response.text
    assert rt.store.get_project(pid) is None
    assert paths and b"PRIVATE" in rt.files.get(paths[0])
    print("GUIDE RACE: the project is deleted, upload returns 404, but its private guide survives outside the TTL prefix.")


def test_legacy_project_prefix_is_not_erased(client):
    rt = get_runtime()
    pid, chapter = _project(client)
    legacy_path = f"users/{UID}/projects/{pid}/chapters/1/legacy.json"
    rt.files.put(legacy_path, chapter.model_dump_json(by_alias=True).encode(), "application/json")

    def legacy(p):
        p.chapter(1).versions[0].path = legacy_path
        return p

    rt.store.update_project(pid, legacy)
    assert client.get(f"/api/projects/{pid}/chapters/1", headers=STUDENT).status_code == 200
    assert client.delete(f"/api/projects/{pid}", headers=STUDENT).status_code == 204
    assert rt.files.get(legacy_path)
    print("LEGACY PREFIX: deletion succeeds but leaves the older project's chapter file behind.")


def test_cleanup_stops_behind_one_blocked_page(client, monkeypatch):
    from app.jobs import service as jobs

    rt = get_runtime()
    first, second = _create(client)["id"], _create(client)["id"]
    for pid, days in ((first, 3), (second, 2)):
        rt.store.update_project(pid, lambda p, days=days: p.model_copy(update={"expires_at": utcnow() - timedelta(days=days)}))
    now = utcnow()
    blocked = Job(
        id="job_blockcleanup", owner_uid=UID, owner_email="student@example.com", project_id=first, status=JobStatus.QUEUED, expires_at=now + timedelta(days=30)
    )
    rt.store.create(blocked)
    rt.store.update_project(first, lambda p: p.model_copy(update={"active_job": blocked.id}))
    monkeypatch.setattr(jobs, "PAGE_SIZE", 1)
    assert jobs.cleanup_expired_projects(rt) == 0
    assert rt.store.get_project(second) is not None
    print("CLEANUP: one blocked item fills page one, so an eligible later project is never visited.")


def _submit_quoted(client, pid, body):
    response = client.post(f"/api/projects/{pid}/steps/{body['job']['id']}/submit", headers=STUDENT, json={"quoteId": body["quote"]["id"]})
    assert response.status_code == 200, response.text
    return wait(client, body["job"]["id"])


def test_revision_from_stale_quote_replaces_newer_background(client):
    rt = get_runtime()
    pid, base = _project(client)
    _comment(client, pid)
    quoted = client.post(f"/api/projects/{pid}/steps", headers=STUDENT, json={"step": "REVISE_1"}).json()
    newer = base.model_copy(deep=True)
    next(s for s in newer.sections if s.key == "background").paragraphs = ["NEWER background that the student selected after the quote."]
    path = f"{rt.store.get_project(pid).storage_prefix()}/chapters/1/newer.json"
    rt.files.put(path, newer.model_dump_json(by_alias=True).encode(), "application/json")

    def change(p):
        p.chapter(1).versions.append(StoredChapterVersion(version=2, job_id="job_newer", words=200, plan_version=1, path=path))
        p.chapter(1).current = 2
        return p

    rt.store.update_project(pid, change)
    assert _submit_quoted(client, pid, quoted)["status"] == "COMPLETED"
    current = client.get(f"/api/projects/{pid}/chapters/1", headers=STUDENT).json()
    old_background = next(s.paragraphs for s in base.sections if s.key == "background")
    assert next(s["paragraphs"] for s in current["sections"] if s["key"] == "background") == old_background
    print("STALE REVISION: changing to version 2 after quoting still submits version 1; its older background becomes current again.")


def test_reassigned_feedback_marked_applied_to_unchanged_section(client):
    pid, base = _project(client)
    cid = _comment(client, pid)
    quoted = client.post(f"/api/projects/{pid}/steps", headers=STUDENT, json={"step": "REVISE_1"}).json()
    changed = client.post(f"/api/projects/{pid}/feedback/{cid}", headers=STUDENT, json={"chapter": 1, "sections": ["background"]})
    assert changed.status_code == 200
    assert _submit_quoted(client, pid, quoted)["status"] == "COMPLETED"
    current = client.get(f"/api/projects/{pid}/chapters/1", headers=STUDENT).json()
    assert next(s["paragraphs"] for s in current["sections"] if s["key"] == "background") == next(s.paragraphs for s in base.sections if s.key == "background")
    comment = next(c for c in client.get(f"/api/projects/{pid}", headers=STUDENT).json()["feedback"] if c["id"] == cid)
    assert comment["status"] == "APPLIED" and comment["sections"] == ["background"]
    print("FEEDBACK RACE: comment moved to the background after quoting; only the problem is revised, but the background comment says APPLIED.")


def test_reference_verification_accepts_second_author_as_first(monkeypatch):
    from app.analysis import fetch, references
    from tests.test_references import ENTRY, RECORD

    monkeypatch.setattr(fetch, "crossref_work", lambda doi: RECORD)
    monkeypatch.setattr(fetch, "openalex_retracted", lambda doi: False)
    wrong = ENTRY.replace("Ojakaa, D. I., & Jarvis, J. D.", "Jarvis, J. D.") + " https://doi.org/" + RECORD["doi"]
    assert references.verify(wrong).status == "VERIFIED"
    print("AUTHOR: a reference listing the second author as the only/first author is still VERIFIED.")


def test_reference_service_outage_says_doi_not_registered(monkeypatch):
    from app.analysis import fetch, references
    from tests.test_references import ENTRY, RECORD

    monkeypatch.setattr(fetch, "crossref_work", lambda doi: None)  # also the real timeout/error result
    result = references.verify(ENTRY + " https://doi.org/" + RECORD["doi"])
    assert result.status == "NOT_VERIFIED" and "DOI is not registered" in result.note
    print("OUTAGE: unavailable lookup is stated as proof the DOI is not registered.")


def test_reference_query_is_not_logged_by_real_fetch_client(client, monkeypatch, caplog):
    import logging

    from app.analysis import fetch

    real_client = fetch.httpx.Client
    transport = fetch.httpx.MockTransport(lambda request: fetch.httpx.Response(200, json={"message": {"items": []}}))
    monkeypatch.setattr(fetch, "_public_address", lambda host: "93.184.216.34")
    monkeypatch.setattr(fetch.httpx, "Client", lambda **kwargs: real_client(transport=transport, **kwargs))
    # Starlette's test client uses httpx2, but application fetch uses the silenced httpx client.
    assert real_client.__module__ == "httpx"
    assert logging.getLogger("httpx").getEffectiveLevel() == logging.WARNING
    logging.getLogger().addHandler(caplog.handler)
    with caplog.at_level(logging.INFO):
        real_crossref_search("PRIVATE_AUDIT_REFERENCE interview with a study participant")
    assert not any("PRIVATE_AUDIT_REFERENCE" in record.getMessage() for record in caplog.records)


def test_load_script_reports_success_when_student_future_crashes(monkeypatch, capsys):
    import sys

    from tests import load_test

    def crashing_student(base, number, rounds, submit, recorder):
        recorder.times["config"].append(0.01)
        raise ValueError("synthetic malformed successful response")

    monkeypatch.setattr(load_test, "student", crashing_student)
    monkeypatch.setattr(sys, "argv", ["load_test", "--students", "1", "--rounds", "1"])
    assert load_test.main() == 0
    output = capsys.readouterr().out
    assert "config" in output and "failure" not in output and "synthetic" not in output


def test_incomplete_academic_review_is_labelled_full(fixed_client, monkeypatch):
    from app.jobs import pipeline

    warning = "The academic review reached this job's spending limit before it covered the whole paper, so some sections were not reviewed."
    monkeypatch.setattr(pipeline, "_academic_review", lambda ctx, model: ([], warning))
    job, _ = _paper(get_runtime())
    quote = get_quote(fixed_client, job.id, {"writing": "AI_CHECK", "academic": True})
    assert fixed_client.post(f"/api/jobs/{job.id}/submit", headers=STUDENT, json={"quoteId": quote["id"]}).status_code == 200
    result = wait(fixed_client, job.id)
    assert result["outcome"] == "FULL" and warning in result["warnings"]
